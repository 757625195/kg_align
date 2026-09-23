from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips


SECTION_HEADING = "5.7 多数据集扩展验证与验证集选择后处理"
CONCLUSION_PREFIX = "补充五数据集实验进一步表明"

RESULT_ROWS = [
    ("DBP15K zh-en", "0.6877±0.0035", "0.6886±0.0036", "0.7002±0.0056", "+1.25", "0.7438 (+1.14)"),
    ("DBP15K ja-en", "0.7370±0.0013", "0.7368±0.0021", "0.7556±0.0021", "+1.86", "0.7954 (+1.65)"),
    ("DBP15K fr-en", "0.9025±0.0031", "0.9026±0.0028", "0.9080±0.0038", "+0.55", "0.9272 (+0.52)"),
    ("OpenEA en-fr", "0.5281±0.0050", "0.5315±0.0053", "0.5608±0.0052", "+3.26", "0.6280 (+2.83)"),
    ("OpenEA en-de", "0.3051±0.0059", "0.3069±0.0055", "0.3311±0.0045", "+2.61", "0.4124 (+2.36)"),
]


def set_run_fonts(run, east_asia: str = "STSong", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)


def style_added_paragraph(paragraph, heading: bool = False) -> None:
    for run in paragraph.runs:
        set_run_fonts(run)


def set_cell_width(cell, width_twips: int) -> None:
    cell.width = Twips(width_twips)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 80, left: int = 90, bottom: int = 80, right: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def configure_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_cols = table._tbl.tblGrid.gridCol_lst
    for grid_col, width in zip(grid_cols, widths):
        grid_col.w = Twips(width)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_table_text(table) -> None:
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                set_run_fonts(run)
                run.font.size = Pt(8.5)
                run.bold = row_index == 0 or col_index in (3, 4)
            if row_index == 0:
                shade_cell(cell, "F4F6F9")


def append_before(document, target_paragraph, element) -> None:
    target_paragraph._p.addprevious(element)


def add_extended_results(document: Document) -> None:
    discussion = next((p for p in document.paragraphs if p.text.strip() == "6 讨论"), None)
    references = next((p for p in document.paragraphs if p.text.strip() == "参考文献"), None)
    if discussion is None or references is None:
        raise RuntimeError("Could not locate the discussion or references heading.")

    heading = document.add_paragraph(SECTION_HEADING, style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    style_added_paragraph(heading, heading=True)

    intro = document.add_paragraph(
        "为检验最新简化配置与后处理方案的跨数据集稳定性，本节在既有 DBP15K zh-en 和 OpenEA en-fr 之外，"
        "新增 DBP15K ja-en、DBP15K fr-en 与 OpenEA en-de。每个数据集均运行随机种子 42、43 和 44，"
        "并在完整目标知识图谱实体集合上计算排名。实验依次比较基础模型、加入 top-20 MLP 重排序器以及在重排序分数上"
        "使用验证集选择参数的 CSLS。DBP15K 从原训练对中固定留出 10% 作为验证集，OpenEA 使用官方验证集；"
        "测试集仅用于最终报告。"
    )
    style_added_paragraph(intro)

    caption = document.add_paragraph(
        "表 9  五数据集扩展实验结果（均值±样本标准差，n=3）",
        style="Caption",
    )
    caption.paragraph_format.keep_with_next = True
    style_added_paragraph(caption)

    headers = ["数据集", "基础 Hits@1", "+ Reranker", "+ 验证选择 CSLS", "Hits@1 增益/百分点", "最终 MRR（增益/百分点）"]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in RESULT_ROWS:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value

    configure_table_geometry(table, [1720, 1370, 1470, 1760, 1000, 1800])
    repeat_table_header(table.rows[0])
    style_table_text(table)

    result_summary = document.add_paragraph(
        "最终组合在 15 次独立运行中均超过相应基础模型。五个数据集的平均 Hits@1 绝对增益为 1.91 个百分点。"
        "其中 OpenEA en-fr 和 en-de 分别提高 3.26 和 2.61 个百分点；DBP15K zh-en、ja-en 和 fr-en 分别提高 "
        "1.25、1.86 和 0.55 个百分点。fr-en 的基础 Hits@1 已达到 0.9025，因而可由后处理修正的剩余错误相对较少；"
        "这一现象说明收益大小会随基础性能和候选空间局部密度而变化。"
    )
    style_added_paragraph(result_summary)

    reranker_summary = document.add_paragraph(
        "Reranker 单独相对基础模型的 Hits@1 变化依次为 +0.10、-0.01、+0.01、+0.34 和 +0.18 个百分点，"
        "五数据集平均仅增加 0.12 个百分点。因此，当前证据不支持将重排序器单独表述为主要性能来源；稳定增量主要出现在"
        "验证集选择 CSLS 之后，但应将最终数字理解为 reranker 与 CSLS 的组合结果。"
    )
    style_added_paragraph(reranker_summary)

    parameter_summary = document.add_paragraph(
        "参数选择也呈现数据集差异：ja-en 的三个种子均选择 k=5，en-de 的三个种子均选择 blend=1.00；"
        "fr-en 的最优 k 在 5、10 和 20 之间变化。由此不宜根据单次测试结果为所有数据集写死同一参数，"
        "保留基于验证 MRR 的选择流程更为稳健。"
    )
    style_added_paragraph(parameter_summary)

    protocol_note = document.add_paragraph(
        "需要指出，表 9 的 DBP15K 实验使用训练对留出验证集，以便选择 CSLS 参数；表 3 的主实验则不设置额外验证集。"
        "因此，表 9 用于评价最新简化配置及后处理的跨数据集稳定性，不应与表 3 直接合并为同一训练协议，"
        "也不构成与外部文献方法的严格 SOTA 比较。"
    )
    style_added_paragraph(protocol_note)

    for element in (
        heading._p,
        intro._p,
        caption._p,
        table._tbl,
        result_summary._p,
        reranker_summary._p,
        parameter_summary._p,
        protocol_note._p,
    ):
        append_before(document, discussion, element)

    conclusion = document.add_paragraph(
        "补充五数据集实验进一步表明，在不改变模型主体结构的条件下，reranker 与验证集选择 CSLS 的组合能够稳定改善"
        "当前系统的最终检索结果：15 次独立运行均获得正向 Hits@1 变化，五数据集平均提升 1.91 个百分点。"
        "该结果扩大了实验覆盖范围，但不同数据集的增益和最优参数差异仍要求按验证集独立选择后处理配置。"
    )
    style_added_paragraph(conclusion)
    append_before(document, references, conclusion._p)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    if args.source.resolve() == args.output.resolve():
        raise ValueError("Output must differ from source to preserve the user's document.")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)
    if any(p.text.strip() == SECTION_HEADING for p in document.paragraphs):
        raise RuntimeError("The target document already contains the extended-results section.")
    if any(p.text.strip().startswith(CONCLUSION_PREFIX) for p in document.paragraphs):
        raise RuntimeError("The target document already contains the extended-results conclusion.")
    add_extended_results(document)
    document.save(args.output)


if __name__ == "__main__":
    main()
