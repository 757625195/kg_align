from __future__ import annotations

import argparse
import copy
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from mathify_all_document_symbols import set_markup  # noqa: E402


def combined_text(paragraph: Paragraph) -> str:
    return "".join(
        paragraph._p.xpath(".//w:t/text()|.//m:t/text()")
    ).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).startswith(prefix)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph beginning {prefix!r}, found {len(matches)}"
        )
    return matches[0]


def find_formula(document: Document, number: int) -> Paragraph:
    suffix = f"({number})"
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Formula" and combined_text(paragraph).endswith(suffix)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one formula ending {suffix}, found {len(matches)}")
    return matches[0]


def add_paragraph_after(paragraph: Paragraph, style: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    result = Paragraph(element, paragraph._parent)
    result.style = style
    return result


def replace_literal(paragraph: Paragraph, old: str, new: str) -> None:
    changed = False
    for node in paragraph._p.xpath(".//w:t"):
        if old in (node.text or ""):
            node.text = node.text.replace(old, new)
            changed = True
    if not changed:
        raise ValueError(f"Missing {old!r} in {combined_text(paragraph)!r}")


def renumber_existing_equations(document: Document) -> None:
    # Work backwards so that newly assigned numbers never collide with the
    # equation still being searched.
    for number in range(24, 4, -1):
        paragraph = find_formula(document, number)
        replace_literal(paragraph, f"({number})", f"({number + 1})")


def add_intermediate_transform(document: Document) -> None:
    equation4 = find_formula(document, 4)
    set_markup(
        equation4,
        r"[[\widetilde{\mathbf{h}}_i^{(\ell+1)}="
        r"\mathbf{s}_i^{(\ell)}+\mathbf{g}_i^{(\ell)}"
        r"\odot\bar{\mathbf{m}}_i^{(\ell)}]]." + "\t" + "(4)",
    )

    introduction = find_paragraph(document, "式（4）产生每一层的原始输出")
    set_markup(
        introduction,
        "式（4）产生第 [[\\ell]] 层的原始输出 "
        "[[\\widetilde{\\mathbf{h}}_i^{(\\ell+1)}]]。除最后一层外，"
        "该输出在作为下一层输入之前进行如下变换：",
    )

    equation5 = add_paragraph_after(introduction, "Formula")
    set_markup(
        equation5,
        r"[[\mathbf{h}_i^{(\ell+1)}="
        r"\operatorname{Dropout}\!\left("
        r"\operatorname{ReLU}\!\left("
        r"\operatorname{LayerNorm}\!\left("
        r"\widetilde{\mathbf{h}}_i^{(\ell+1)}"
        r"\right)\right)\right)]]." + "\t" + "(5)",
    )

    explanation = add_paragraph_after(equation5, "Normal")
    set_markup(
        explanation,
        "其中，LayerNorm 统一同一实体各维特征的尺度，ReLU 引入非线性，"
        "dropout 在训练时以 [[0.1]] 的概率随机屏蔽部分特征，并在验证和测试时关闭。"
        "式（5）只处理当前层输出，不改变图的连接关系；多跳信息仍来自图卷积层的连续堆叠。",
    )


def update_equation_references(document: Document) -> None:
    replacements = [
        ("层选择 MLP 的输入维度为", "式（6）", "式（7）"),
        ("式（13）—（20）定义主模型", "式（13）—（20）", "式（14）—（21）"),
        ("参数匹配晚期融合先分别归一化", "式（21）", "式（22）"),
        ("结构查询邻居基线与主模型共享", "式（13）—（20）", "式（14）—（21）"),
        ("结构查询邻居基线与主模型共享", "式（13）", "式（14）"),
        ("结构查询邻居基线与主模型共享", "式（15）", "式（16）"),
        ("结构查询邻居基线与主模型共享", "式（17）", "式（18）"),
        ("结构查询邻居基线与主模型共享", "式（19）—（20）", "式（20）—（21）"),
        ("式（23）即当前代码", "式（23）", "式（24）"),
    ]
    for prefix, old, new in replacements:
        replace_literal(find_paragraph(document, prefix), old, new)


def _math_char(properties, child_name: str) -> str | None:
    child = properties.find(qn(child_name))
    return None if child is None else child.get(qn("m:val"))


def replace_group_character_accents(document: Document) -> tuple[int, int]:
    bars = 0
    accents = 0
    groups = list(document.element.body.xpath(".//m:groupChr"))
    for group in groups:
        properties = group.find(qn("m:groupChrPr"))
        expression = group.find(qn("m:e"))
        if properties is None or expression is None:
            continue
        character = _math_char(properties, "m:chr")
        position = _math_char(properties, "m:pos")
        if position != "top":
            continue

        if character == "¯":
            replacement = OxmlElement("m:bar")
            replacement_properties = OxmlElement("m:barPr")
            position_element = OxmlElement("m:pos")
            position_element.set(qn("m:val"), "top")
            replacement_properties.append(position_element)
            replacement_properties.append(OxmlElement("m:ctrlPr"))
            replacement.append(replacement_properties)
            replacement.append(copy.deepcopy(expression))
            group.getparent().replace(group, replacement)
            bars += 1
        elif character in {"˜", "~", "ˆ", "^"}:
            replacement = OxmlElement("m:acc")
            replacement_properties = OxmlElement("m:accPr")
            character_element = OxmlElement("m:chr")
            character_element.set(qn("m:val"), character)
            replacement_properties.append(character_element)
            replacement_properties.append(OxmlElement("m:ctrlPr"))
            replacement.append(replacement_properties)
            replacement.append(copy.deepcopy(expression))
            group.getparent().replace(group, replacement)
            accents += 1
    return bars, accents


def hide_empty_nary_limits(document: Document) -> int:
    hidden = 0
    for tag, property_tag in (("m:sup", "m:supHide"), ("m:sub", "m:subHide")):
        for element in list(
            document.element.body.xpath(
                f".//{tag}[not(*) and not(normalize-space(text()))]"
            )
        ):
            parent = element.getparent()
            if parent.tag != qn("m:nary"):
                parent.remove(element)
                continue
            properties = parent.find(qn("m:naryPr"))
            if properties is None:
                properties = OxmlElement("m:naryPr")
                parent.insert(0, properties)
            hidden_property = properties.find(qn(property_tag))
            if hidden_property is None:
                hidden_property = OxmlElement(property_tag)
                hidden_property.set(qn("m:val"), "1")
                control = properties.find(qn("m:ctrlPr"))
                if control is None:
                    properties.append(hidden_property)
                else:
                    control.addprevious(hidden_property)
            hidden += 1
    return hidden


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)

    renumber_existing_equations(document)
    add_intermediate_transform(document)
    update_equation_references(document)
    bars, accents = replace_group_character_accents(document)
    hidden_limits = hide_empty_nary_limits(document)

    document.save(args.output)
    print(f"saved={args.output}")
    print(f"converted_bars={bars}")
    print(f"converted_accents={accents}")
    print(f"hidden_empty_nary_limits={hidden_limits}")


if __name__ == "__main__":
    main()
