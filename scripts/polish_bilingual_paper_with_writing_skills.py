from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
ZH_SOURCE = ROOT / "outputs" / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_Springer去防御性中文译稿_20260830_图片重绘版.docx"
EN_SOURCE = ROOT / "outputs" / "Relation_Aware_Neighbor_Context_Springer_Anti_Defensive_English_20260830_Figure_ReRendered.docx"
ZH_OUTPUT = ROOT / "outputs" / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_中文学术润色终稿_20260830.docx"
EN_OUTPUT = ROOT / "outputs" / "Relation_Aware_Neighbor_Context_Bilingual_Polished_English_20260830.docx"


# These revisions follow two constraints: make only necessary language changes,
# and state claims directly while retaining methodological scope conditions.
ZH_PLAIN = {
    3: (
        "跨语言知识图谱实体对齐用于识别不同语言图谱中指向同一对象的实体。本文提出一种语义引导的关系感知邻居上下文模型，"
        "使实体语义在联合编码前参与结构证据选择。结构编码器在消息传递中保留关系类型，并通过节点级层选择组合不同传播深度。"
        "语义编码器从同一实体文本序列提取 token、phrase 和 global 三种视图。融合模块使用语义表示为完整的一跳出邻域评分，"
        "以 1.5-entmax 抑制弱相关邻居，再将所得结构上下文与实体语义编码为联合表示。本文在五个数据集上采用三个随机种子进行评估。"
        "完整模型在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的 Hits@1 分别为 0.7285 和 0.6344。与仅语义检索相比，Hits@1 "
        "分别提高 9.09 和 9.66 个百分点；在邻域和可训练融合参数量固定时，语义查询相对结构查询分别提高 2.08 和 2.48 个百分点。"
        "移除关系类型后，Hits@1 分别下降 3.62 和 2.78 个百分点。完整模型相对参数匹配晚期融合分别提高 5.28 和 9.22 个百分点，"
        "该差异反映邻居访问与训练交互路径的共同作用。结果表明，关系类型、结构邻居上下文和语义引导选择能够稳定改善跨语言实体对齐。"
    ),
    11: (
        "本文围绕三个问题展开。第一，经过筛选的关系感知邻居上下文能否提供语义表示之外的对齐证据？第二，在邻域范围和可训练融合参数量保持一致时，"
        "使用语义表示为邻居评分是否优于使用结构表示？第三，关系类型、邻居上下文和多尺度语义编码分别对联合模型产生何种影响？"
    ),
    12: "本文的主要贡献包括：",
    13: (
        "提出轻量的关系感知结构编码方法。结构编码器通过共享变换控制参数规模，以关系嵌入保留边类型差异，并使用节点级层选择器组合不同传播深度的图状态。"
    ),
    14: (
        "设计语义引导的结构邻居选择机制。语义编码器从同一实体文本提取 token、phrase 和 global 三种视图，融合模块使用所得语义表示为完整的一跳出邻域评分。"
        "1.5-entmax 将弱相关邻居的权重置零，筛选后的邻居信息形成结构上下文，并与实体语义生成联合表示。"
    ),
    15: (
        "通过五个数据集和三个随机种子的实验确定模型增益的来源。匹配对照分别检验关系类型、多尺度语义、邻居访问和查询信号，并在固定邻域与融合参数量的条件下比较语义查询和结构查询。"
    ),
    18: (
        "MTransE 学习不同语言嵌入空间之间的映射[2]，JAPE 则在共享结构空间中加入属性关联[3]。这类方法以明确的几何映射连接不同图谱，但对关系类型及其局部邻域作用的表达较弱，后续研究因此转向关系感知结构编码。"
    ),
    20: (
        "GraphSAGE 通过聚合局部邻居学习节点表示[9]。R-GCN 在消息变换中保留关系类型，并通过共享基或分块结构控制参数规模[7]。这些研究表明，图编码器需要同时表示邻接实体与边类型。"
    ),
    21: (
        "RDGCN 通过实体图与关系对偶图引入关系证据[5]，RREA 使用关系反射变换保留关系差异[6]。RPR-RHGT 进一步使用图 Transformer 编码筛选后的多步关系路径[10]。"
        "这些方法分别从关系对偶图、几何变换和显式路径编码中利用关系结构。本文采用轻量的关系感知消息传递，并进一步考察一跳结构邻居在联合表示中的作用。"
    ),
    23: (
        "Transformer 能够建立序列中不同位置之间的联系[11]。层次注意力网络表明，不同文本粒度可以提供不同证据[12]。本文据此从同一实体文本中提取 token、phrase 和 global 三种视图，并将融合后的语义表示用于邻居选择。"
    ),
    24: (
        "文本增强实体对齐研究表明，语义信息可以补充图结构。RREA 分别报告文本增强和纯结构设置[6]，MCLEA 通过对比目标对齐多种数据类型[13]。另有研究指出，简单拼接或注意力融合可能混合尚未对齐的表示空间[14]。"
        "这些工作说明语义有助于实体对齐，也表明融合方式会影响异质表示的兼容性。基于此，本文在保持邻域和融合参数量不变的条件下替换邻居查询信号，以检验语义查询的作用。"
    ),
    25: (
        "近期方法还通过大语言模型[15]、域适应[16]和逐步推理代理[17]增强跨图监督或推理。本文关注不同层面的问题，即在固定实体对齐框架内组织局部结构邻居与实体语义的交互。"
    ),
    35: (
        "模型由结构编码、语义编码、邻居上下文融合和对比学习四部分构成。结构编码器保留关系类型并组合不同传播深度，语义编码器从同一实体文本构建 token、phrase 和 global 三种视图。"
        "融合模块使用语义表示为全部一跳出邻居评分，形成稀疏结构上下文，并与实体语义生成联合表示。双向 InfoNCE 同时约束联合表示和结构表示，检索阶段再以 CSLS 校正联合表示之间的相似度。图 1 给出完整信息流。"
    ),
    39: (
        "局部度与邻居统计常用于描述节点的结构角色[18–19]，结构身份方法也根据图中角色比较节点。受此启发，本文以八维局部拓扑向量初始化结构分支，具体定义如下。"
    ),
    46: (
        "结构编码与上下文选择使用不同方向的邻域。消息传递沿传入边更新目标节点，以聚合指向该实体的事实；融合模块读取当前实体的一跳出邻居，将该实体发出的事实作为候选上下文。两种方向分别服务于节点更新和上下文选择。"
    ),
    56: "每层传播后依次采用 LayerNorm、ReLU 和 dropout，以稳定不同深度的特征分布并降低过拟合风险。",
    57: "多层传播使实体逐步吸收更远的邻域证据。例如，在“姚明—出生地→上海—所在国家→中国”中，姚明的局部信息可以经上海传递至中国。",
    63: "三个语义视图使用同一实体文本，但编码范围不同。token 视图汇总词级线索，phrase 视图提取局部组合，global 视图建模跨位置依赖。",
    65: (
        "以实体“姚明”为例，共享序列可以包含名称“姚明”、关系短语“出生地”和“效力球队”，以及属性片段“身高 2.29 米”。token 视图汇总“姚明”“身高”等单项线索，phrase 视图识别“出生地 上海”或“效力球队 火箭队”等局部组合，"
        "global 视图则联系名称、关系和相距较远的属性。三种视图改变编码范围，但不改变数据来源。"
    ),
    68: "掩码排除填充位置。均值池化汇总全部有效 token，注意力池化提高区分性词项的贡献，二者取平均得到 token 视图。",
    79: (
        "融合模块从完整的一跳出邻域中构造与当前实体相关的结构上下文。语义查询先为结构邻居评分，1.5-entmax 将弱相关候选的权重置零；加权邻居摘要随后与实体自身结构状态结合，再通过联合门与语义表示合成为最终表示。"
    ),
    91: (
        "仍以“姚明”为例。若一跳出邻居包括“休斯敦火箭队”“上海”和一个低相关候选，语义查询可以依据“篮球运动员”“效力球队”等线索提高火箭队邻居的得分。1.5-entmax 随后将低相关候选的权重置零，式（19）再决定保留多少自身结构与加权邻居信息。"
    ),
    103: "语义查询先改变结构邻居的权重，最终门控再将所得结构上下文与实体语义结合。第 3.5.4 节通过匹配对照检验查询信号的作用。",
    105: (
        "结构查询对照用于检验查询信号。该对照保持邻域、动态补齐、1.5-entmax、门控和联合检索不变，仅以实体结构表示替代语义表示计算邻居权重。参数匹配晚期融合对照则在两个分支独立编码后形成联合表示，不读取邻居上下文，但保持可训练融合参数量与完整模型一致。"
    ),
    127: "本文以 Hits@1、Hits@10 和平均倒数排名（MRR）衡量完整候选集合上的检索质量。主实验与消融实验均报告三个随机种子的均值和样本标准差。",
    132: "表 2 报告五个数据集上的联合表示检索结果。每个数据集均使用三个随机种子。DBP15K 从训练对中留出 10% 作为验证集，OpenEA 和 EventEA 使用官方验证集。",
    135: "模型在五个数据集上的结果较为稳定，Hits@1 的最大样本标准差为 0.0053。该结果为后续匹配对照提供了稳定的实验基础。",
    137: "表 3 列出既有研究公开报告的结果。MTransE、JAPE、BootEA 和 RDGCN 的 DBP15K 数值取自 RDGCN[5]，RREA-text 数值取自 RREA[6]，OpenEA EN–FR-15K-V2 数值采用 OpenEA 官方五折平均结果[8]。",
    140: "表 3 将本文结果置于已有方法的公开结果中进行比较。各研究的数据划分、文本输入和后处理设置并不完全一致，因此该表用于说明相对位置；表 4–6 则采用统一协议检验各项方法设计。",
    143: "表 4 报告 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的组件控制结果。每个变体均运行三个随机种子，并在训练与检索阶段使用同一种表示。",
    145: (
        "关系类型是两个数据集上最稳定的结构因素。移除关系类型后，DBP15K 和 OpenEA 的 Hits@1 分别下降 3.62 和 2.78 个百分点。拓扑特征相对仅使用实体向量分别提高 3.46 和 8.03 个百分点。移除 phrase 视图后，两个数据集分别下降 0.81 和 3.95 个百分点。"
        "其余变体的影响较小或随数据集变化。"
    ),
    147: "表 5 比较完整模型、结构单分支、语义单分支和无可训练参数的平均融合。各变体在训练与检索阶段使用同一种表示，验证阶段不调整分支权重。",
    149: (
        "完整模型相对语义单分支在 DBP15K 和 OpenEA 上的 Hits@1 分别提高 9.09 和 9.66 个百分点，并且优于无可训练参数的平均融合。结果说明，学习型交互能够将结构上下文转化为语义表示之外的对齐证据。结构单分支的较低结果也表明，这部分增益依赖结构与语义的联合学习。"
    ),
    151: "表 6 在匹配条件下检验邻居查询信号。结构查询对照共享完整邻域、门控、融合参数和联合检索，仅使用结构表示为邻居评分；参数匹配晚期融合在两个分支独立编码后使用相同数量的融合参数。",
    153: (
        "语义查询相对结构查询在 DBP15K 和 OpenEA 上的 Hits@1 分别提高 2.08 和 2.48 个百分点。由于两者的邻域、融合参数量和检索表示相同，该差异反映语义引导邻居评分的贡献。完整模型相对参数匹配晚期融合分别提高 5.28 和 9.22 个百分点，"
        "该差异反映邻居访问与训练交互路径的共同作用。"
    ),
    156: (
        "组件控制显示，关系类型在两个数据集上均带来稳定增益。完整模型相对语义单分支的提升说明，结构邻居上下文能够补充文本证据；在邻域和融合参数量一致时，语义查询又优于结构查询，表明实体语义有助于从局部邻域中选择相关结构信息。"
        "这些结果共同支持关系感知邻居上下文与语义引导选择。"
    ),
    157: (
        "增益大小与输入数据有关。DBP15K 提供较有用的名称和局部图线索，EventEA 包含多样的事件关系和属性[25]，OpenEA 则使用编码后的实体标识以降低名称偏置[8]。在本文预处理中，GloVe[26] 对 OpenEA token 的覆盖率约为 52.18%。"
        "当文本能够引导邻居选择，且关系类型能够区分相似局部结构时，本方法的优势更明显。"
    ),
    159: (
        "实验覆盖五个数据集，组件控制集中在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2。三个随机种子用于估计均值与样本波动。固定 8 邻居加 softmax 的对照同时改变邻居数量和权重函数，因此其差异反映两者的共同作用。"
        "参数匹配晚期融合同时移除邻居访问并改变训练交互路径，其结果用于衡量整体设计差异。"
    ),
    161: (
        "后续研究可以从输入语义、图规模和任务设定三个方面扩展。多语言子词编码器可以替代未知 token 的随机向量[27–28]；更大规模图谱可以检验完整邻域选择的效率；开放世界数据集可以评估不存在对应实体的情况。"
        "这些实验将进一步检验语义引导邻居选择在弱文本线索和大候选空间下的稳健性。"
    ),
    163: (
        "本文研究关系感知结构邻居在结构—语义联合表示中的作用。关系类型在两个受控数据集上稳定增强结构证据，结构邻居上下文使完整模型超过语义单分支，语义查询也在匹配条件下优于结构查询。完整邻域访问与学习型交互路径的组合进一步超过参数匹配晚期融合。"
        "这些结果表明，关系感知邻居上下文与语义引导选择能够改善跨语言实体对齐。"
    ),
}


ZH_MARKUP = {
    86: (
        "其中，{{EQ0}} 保留两个输入向量，{{EQ1}} 描述对应维度的共同激活，{{EQ2}} 描述逐维绝对差。该组合借鉴自然语言推理中的匹配特征设计[20]，并为可学习门控提供判断邻居兼容性的输入。交互特征经可学习映射和 Sigmoid 函数得到邻居门值："
    ),
    109: (
        "模型在单一训练阶段采用双向 InfoNCE 目标进行优化[22]。每个批次从种子对齐集合采样 {{EQ0}} 个实体对，并使用同一批样本计算联合表示损失和结构分支损失；两项损失在一次前向与反向传播中共同优化。左图第 {{EQ1}} 个实体与右图第 {{EQ2}} 个实体之间的温度缩放余弦相似度为："
    ),
    129: (
        "模型使用 AdamW[24] 优化，batch size 为 512，表示维度为 128。结构编码器包含 3 层关系感知图神经网络，所有模块的 dropout 均为 0.1。语义输入由 300 维映射至 128 维；global 视图使用两层、四头 Transformer，phrase 视图采用宽度 3 和 5 的卷积。"
        "1.5-entmax 温度为 {{EQ0}}，InfoNCE 温度为 {{EQ1}}。实验使用随机种子 42、43 和 44。DBP15K 最多训练 36 个 epoch，OpenEA 与 EventEA 最多训练 50 个 epoch。模型每 5 个 epoch 依据验证集 MRR 保存检查点，连续 4 次验证未提升时提前停止。"
        "主检索直接使用训练得到的联合表示。每次运行按验证集 MRR 从 {{EQ2}} 中选择一个值。关系标识按规范化名称统一；DBP15K 另使用数据集提供的 sup_rel_ids 映射已知对应关系。"
    ),
}


EN_PLAIN = {
    24: "2.3 Structural-semantic interaction",
    5: (
        "Cross-lingual knowledge graph entity alignment identifies entities that refer to the same object in graphs written in different languages. This paper presents a relation-aware neighbor-context model in which entity semantics guides structural evidence selection before joint encoding. "
        "The structural encoder preserves relation types during message passing and combines graph states from different propagation depths. The semantic encoder extracts token, phrase, and global views from one entity-text sequence. The fusion module uses the semantic representation to score the complete one-hop outgoing neighborhood, applies 1.5-entmax to suppress weak neighbors, and combines the resulting structural context with entity semantics. "
        "Experiments use three random seeds on five datasets. The full model achieves Hits@1 scores of 0.7285 on DBP15K ZH–EN and 0.6344 on OpenEA EN–FR-15K-V2. It improves over semantic-only retrieval by 9.09 and 9.66 percentage points. With the neighborhood and trainable fusion parameter count fixed, semantic queries improve over structural queries by 2.08 and 2.48 points. "
        "Removing relation types reduces Hits@1 by 3.62 and 2.78 points. The full model also exceeds a parameter-matched late-fusion control by 5.28 and 9.22 points; this difference reflects the combined effect of neighbor access and the training interaction path. The results show that relation types, structural neighbor context, and semantics-guided selection consistently improve cross-lingual entity alignment."
    ),
    10: (
        "Existing methods establish the value of graph structure and text. MTransE and JAPE map graphs into shared or linked vector spaces [2, 3], BootEA adds high-confidence pairs to the training set [4], and RDGCN and RREA preserve relation and neighborhood information [5, 6]. These studies leave one design question open: when should structural neighbors enter a joint representation?"
    ),
    11: (
        "Neighbor selection must account for both relation type and relevance. A birthplace edge and a workplace edge can connect two people to the same city, yet they provide different evidence [5, 7]. In addition, not every neighbor helps identify the current entity. The proposed model therefore preserves relation types, uses entity semantics to weight structural neighbors, and combines the selected neighbor context with the semantic representation."
    ),
    13: (
        "This study addresses three questions. Does selected relation-aware neighbor context provide alignment evidence beyond the semantic representation? With the neighborhood and trainable fusion parameter count fixed, does a semantic representation score neighbors more effectively than a structural representation? How do relation types, neighbor context, and multi-scale semantic encoding affect the joint model?"
    ),
    14: "The main contributions are as follows.",
    15: (
        "The paper presents a lightweight relation-aware structural encoder. Shared transformations control the parameter count, relation embeddings preserve edge-type differences, and a node-level layer selector combines graph states from different propagation depths."
    ),
    16: (
        "The paper introduces semantics-guided structural neighbor selection. The semantic encoder extracts token, phrase, and global views from the same entity text, and the fusion module uses the resulting representation to score the complete one-hop outgoing neighborhood. The 1.5-entmax function assigns zero weight to weak neighbors before the selected context and entity semantics form the joint representation."
    ),
    17: (
        "Experiments with three random seeds on five datasets identify the sources of the gains. Matched controls test relation types, multi-scale semantics, neighbor access, and the query signal. The semantic-query comparison fixes both the neighborhood and the number of trainable fusion parameters."
    ),
    20: (
        "MTransE learns mappings between language-specific embedding spaces [2], while JAPE adds attribute links to a shared structural space [3]. These methods connect graphs through explicit geometric mappings, but they provide limited access to relation-specific local neighborhoods. Later work therefore turns to relation-aware structural encoders."
    ),
    22: (
        "GraphSAGE combines local neighbors to learn node representations [9]. R-GCN retains relation types during message transformation and limits parameter growth through shared bases or blocks [7]. These studies show that graph encoders should represent both neighboring entities and edge types."
    ),
    23: (
        "RDGCN introduces relation evidence through an entity graph and a relation dual graph [5], while RREA uses relation reflection to preserve relation differences [6]. RPR-RHGT encodes selected multi-step relation paths with a graph Transformer [10]. These methods use relation structure through dual-graph interaction, geometric transformations, or explicit path encoding. The present model instead uses lightweight relation-aware message passing and examines how one-hop structural neighbors enter the joint representation."
    ),
    25: (
        "Transformers connect positions across a sequence [11], and hierarchical attention networks show that different text levels provide distinct evidence [12]. The semantic encoder follows this multi-scale principle by extracting token, phrase, and global views from the same entity text and using the combined semantic representation for neighbor selection."
    ),
    26: (
        "Text-enhanced entity alignment shows that semantics can complement graph structure. RREA reports text-enhanced and structure-only settings [6], and MCLEA aligns several data types through contrastive objectives [13]. Other work finds that simple concatenation or attention may mix representation spaces that are not aligned [14]. These results motivate a controlled test of the neighbor-query signal. The present study changes the query source while keeping the neighborhood and trainable fusion parameter count fixed."
    ),
    27: (
        "Recent methods also use large language models [15], domain adaptation [16], and step-by-step reasoning agents [17] to improve cross-graph supervision or reasoning. This paper studies a different level of the problem: the interaction between local structural neighbors and entity semantics within a fixed alignment framework."
    ),
    37: (
        "The model contains four components: structural encoding, semantic encoding, neighbor-context fusion, and contrastive learning. The structural encoder preserves relation types and combines different propagation depths. The semantic encoder builds token, phrase, and global views from the same entity text. The fusion module uses semantics to score all one-hop outgoing neighbors, forms a sparse structural context, and combines this context with entity semantics. Bidirectional InfoNCE constrains both joint and structural representations, and CSLS adjusts similarities between trained joint representations during retrieval. Figure 1 shows the full information flow."
    ),
    41: (
        "Local degree and neighbor statistics describe structural roles [18, 19], and structural-identity methods compare nodes by their graph roles. Based on this idea, the model uses an eight-dimensional local topology vector to initialize the structural branch."
    ),
    48: (
        "Structural encoding and context selection use different edge directions. Message passing updates a target node through incoming edges, which aggregate facts that point to the entity. The fusion module reads one-hop outgoing neighbors and treats facts emitted by the entity as candidate context. These directions serve node updating and context selection, respectively."
    ),
    58: "Each propagation layer applies LayerNorm, ReLU, and dropout to stabilize feature distributions across depths and reduce overfitting.",
    59: "Stacked propagation lets an entity collect evidence from more distant neighbors. For example, information about Yao Ming can reach China along the path 'Yao Ming - birthplace -> Shanghai - country -> China'.",
    65: "The three semantic views use the same entity text but differ in their encoding range. The token view summarizes word-level clues, the phrase view extracts local combinations, and the global view models dependencies across positions.",
    67: (
        "For example, the shared sequence for Yao Ming may contain the name 'Yao Ming', the relation phrases 'birthplace' and 'team', and the attribute fragment 'height 2.29 m'. The token view summarizes single clues such as 'Yao' and 'height'. The phrase view captures local combinations such as 'birthplace Shanghai' and 'team Rockets'. The global view connects the name, relations, and distant attributes. The views change the encoding range, not the data source."
    ),
    70: "The mask excludes padded positions. Mean pooling summarizes all valid tokens, while attention pooling increases the contribution of discriminative terms. Their average forms the token view.",
    81: (
        "The fusion module constructs entity-specific structural context from the complete one-hop outgoing neighborhood. A semantic query first scores structural neighbors, and 1.5-entmax assigns zero weight to weak candidates. The model then combines the weighted neighbor summary with the entity's own structural state and uses a final gate to merge this context with entity semantics."
    ),
    80: "3.5 Semantics-guided structural context",
    82: "3.5.1 Semantics-guided neighbor weighting",
    93: (
        "The Yao Ming example also illustrates neighbor selection. If the outgoing neighbors include the Houston Rockets, Shanghai, and an unrelated entity, clues such as 'basketball player' and 'team' can increase the score of the Rockets. The 1.5-entmax function then assigns zero weight to the unrelated entity, and Eq. (19) balances self-structure against the weighted neighbor evidence."
    ),
    105: "The semantic query changes structural-neighbor weights before context construction. The final gate then combines the resulting context with entity semantics. Section 3.5.4 tests the role of the query signal with matched controls.",
    107: (
        "The structural-query control tests the query signal. It keeps the neighborhood, dynamic padding, 1.5-entmax, gates, and joint retrieval unchanged, but replaces the semantic query with the entity's structural representation. The parameter-matched late-fusion control forms a joint representation after separate branch encoding and does not read neighbor context. It has the same number of trainable fusion parameters as the full model."
    ),
    129: "The evaluation ranks each entity against the complete candidate set and reports Hits@1, Hits@10, and mean reciprocal rank (MRR). Main and ablation results give the mean and sample standard deviation over three random seeds.",
    134: "Table 2 reports retrieval with the joint representation on five datasets. Each dataset uses three random seeds. DBP15K reserves 10% of its training pairs for validation, while OpenEA and EventEA use their official validation sets.",
    137: "The results are stable across the five datasets: the largest sample standard deviation of Hits@1 is 0.0053. This stability supports the matched comparisons in the following sections.",
    139: "Table 3 lists results reported in earlier studies. The DBP15K values for MTransE, JAPE, BootEA, and RDGCN come from RDGCN [5], the RREA-text value comes from RREA [6], and the OpenEA EN–FR-15K-V2 values are the official five-fold averages from OpenEA [8].",
    142: "Table 3 places the proposed model among published results. Data splits, text inputs, and post-processing settings differ across studies, so the table indicates relative position. Tables 4–6 use a common protocol to test the model design.",
    145: "Table 4 reports component controls on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. Every variant uses three random seeds and the same representation during training and retrieval.",
    147: (
        "Relation types are the most stable structural factor across both datasets. Removing them reduces Hits@1 by 3.62 points on DBP15K and 2.78 points on OpenEA. Topology features improve Hits@1 by 3.46 and 8.03 points over learnable entity vectors alone. Removing the phrase view reduces Hits@1 by 0.81 and 3.95 points. The remaining variants have smaller or dataset-dependent effects."
    ),
    149: "Table 5 compares the full model, the two single branches, and parameter-free mean fusion. Each variant uses the same representation for training and retrieval, and validation does not adjust branch weights.",
    151: (
        "The full model improves Hits@1 over the semantic branch by 9.09 points on DBP15K and 9.66 points on OpenEA. It also outperforms parameter-free mean fusion. These results show that learned interaction turns structural context into alignment evidence beyond the semantic representation. The low structural-only scores further indicate that this gain depends on joint learning with semantics."
    ),
    153: "Table 6 tests the neighbor-query signal under matched conditions. The structural-query control shares the complete neighborhood, gates, fusion parameters, and joint retrieval but scores neighbors with structure. The parameter-matched late-fusion control uses the same number of fusion parameters after independent branch encoding.",
    155: (
        "Semantic queries improve Hits@1 over structural queries by 2.08 points on DBP15K and 2.48 points on OpenEA. Because the neighborhood, fusion capacity, and retrieval representation are fixed, this difference reflects semantics-guided neighbor scoring. The full model exceeds parameter-matched late fusion by 5.28 and 9.22 points. This difference reflects the combined effect of neighbor access and the training interaction path."
    ),
    158: (
        "The component controls identify three main sources of improvement. Relation types provide stable gains on both datasets. The advantage over the semantic branch shows that structural neighbor context adds evidence beyond text. Under matched neighborhood and fusion capacity, semantic queries also outperform structural queries, which shows that entity semantics helps select relevant local structure."
    ),
    159: (
        "The size of the gain depends on the input data. DBP15K provides useful names and local graph clues, EventEA contains varied event relations and attributes [25], and OpenEA uses encoded entity identifiers to reduce name bias [8]. In the preprocessing used here, GloVe [26] covers about 52.18% of the OpenEA tokens. The method works best when text can guide neighbor selection and relation types can distinguish similar local structures."
    ),
    161: (
        "The experiments cover five datasets, with component controls on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. Three random seeds estimate means and sample variation. The fixed 8-neighbor plus softmax control changes both the neighbor count and the weighting function, so its difference reflects both factors. The parameter-matched late-fusion control removes neighbor access and changes the training interaction path; its result measures the overall design difference."
    ),
    163: (
        "Future work can extend the input semantics, graph scale, and task setting. Multilingual subword encoders can replace random vectors for unknown tokens [27, 28]. Larger graphs can test the efficiency of complete-neighborhood selection, and open-world datasets can evaluate entities without matches. These studies can further test semantics-guided neighbor selection with weaker text and larger candidate spaces."
    ),
    165: (
        "This paper examines the role of relation-aware structural neighbors in a structural-semantic joint representation. Relation types consistently strengthen structural evidence on both controlled datasets. Structural neighbor context improves the full model over the semantic branch, and semantic queries outperform structural queries under matched conditions. Complete-neighborhood access with a learned interaction path also exceeds parameter-matched late fusion. These results show that relation-aware neighbor context and semantics-guided selection improve cross-lingual entity alignment."
    ),
}


EN_MARKUP = {
    88: (
        "The concatenation {{EQ0}} preserves both input vectors, {{EQ1}} captures co-activation in corresponding dimensions, and {{EQ2}} records the absolute difference in each dimension. This feature pattern follows matching features used in natural language inference [20] and provides input for a learned compatibility gate. A trainable mapping and a sigmoid function produce the neighbor-gate value."
    ),
    111: (
        "The model uses bidirectional InfoNCE in a single training stage [22]. Each mini-batch samples {{EQ0}} pairs from the seed alignment set and uses the same pairs for the joint and structural losses. Both losses are optimized in one forward and backward pass. The temperature-scaled cosine similarity between left entity {{EQ1}} and right entity {{EQ2}} is"
    ),
    131: (
        "The model uses AdamW [24], a batch size of 512, and 128-dimensional representations. The structural encoder contains three relation-aware graph neural network layers, and all components use a dropout rate of 0.1. The semantic input is projected from 300 to 128 dimensions. The global view uses a two-layer Transformer with four attention heads, and the phrase view uses convolution widths of 3 and 5. "
        "The 1.5-entmax temperature is {{EQ0}}, and the InfoNCE temperature is {{EQ1}}. Experiments use random seeds 42, 43, and 44. Training lasts at most 36 epochs for DBP15K and 50 epochs for OpenEA and EventEA. Validation MRR is checked every five epochs, and training stops after four checks without improvement. The main retrieval step directly uses the trained joint representation. "
        "For each run, validation MRR selects one value from {{EQ2}}. Relation identifiers are unified after relation-name normalization; DBP15K also uses the provided sup_rel_ids to map known relation pairs."
    ),
}


PLACEHOLDER_RE = re.compile(r"\{\{EQ(\d+)\}\}")


def equation_children(paragraph):
    return [
        copy.deepcopy(child)
        for child in paragraph._p.iterchildren()
        if child.tag in {qn("m:oMath"), qn("m:oMathPara")}
    ]


def set_run_font(run, ascii_font: str, east_asia_font: str) -> None:
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), ascii_font)
    fonts.set(qn("w:hAnsi"), ascii_font)
    fonts.set(qn("w:cs"), ascii_font)
    fonts.set(qn("w:eastAsia"), east_asia_font)
    fonts.set(qn("w:hint"), "eastAsia")


def append_text_run(paragraph, text: str, chinese: bool) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    set_run_font(run, "Times New Roman", "Songti SC" if chinese else "Times New Roman")


def set_plain_text(paragraph, text: str, chinese: bool) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    append_text_run(paragraph, text, chinese)


def set_paragraph_markup(paragraph, markup: str, chinese: bool) -> None:
    equations = equation_children(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    cursor = 0
    used: set[int] = set()
    for match in PLACEHOLDER_RE.finditer(markup):
        append_text_run(paragraph, markup[cursor:match.start()], chinese)
        index = int(match.group(1))
        if index >= len(equations):
            raise RuntimeError(f"Paragraph has {len(equations)} equations but markup requests EQ{index}")
        paragraph._p.append(copy.deepcopy(equations[index]))
        used.add(index)
        cursor = match.end()
    append_text_run(paragraph, markup[cursor:], chinese)
    if len(used) != len(equations):
        raise RuntimeError(f"Markup retained {len(used)} of {len(equations)} equations: {markup}")


def remove_word_joiners(document: Document) -> None:
    for node in document._element.xpath(".//w:t|.//m:t"):
        node.text = (node.text or "").replace("\u2060", "").replace("\ufeff", "")


def set_language(node, language: str) -> None:
    rpr = node.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), language)
    lang.set(qn("w:eastAsia"), "zh-CN" if language == "zh-CN" else language)


def normalize_fonts(document: Document, chinese: bool) -> None:
    east_asia_font = "Songti SC" if chinese else "Times New Roman"
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "Formula", "Reference", "List Number"]:
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        rpr = style._element.get_or_add_rPr()
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), east_asia_font)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_font(run, "Times New Roman", east_asia_font)
            set_language(run._element, "zh-CN" if chinese else "en-US")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = 1
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, "Times New Roman", east_asia_font)
                        set_language(run._element, "zh-CN" if chinese else "en-US")


def normalize_paragraph_layout(document: Document, chinese: bool) -> None:
    normal = document.styles["Normal"].paragraph_format
    normal.widow_control = True
    if chinese:
        normal.line_spacing = 1.5
        normal.space_after = Pt(0)
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name == "Normal" and paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if style_name == "Caption":
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if style_name == "Formula":
            paragraph.paragraph_format.keep_together = True


def apply_rewrites(document: Document, plain: dict[int, str], markup: dict[int, str], chinese: bool) -> None:
    expected = 202 if chinese else 204
    if len(document.paragraphs) != expected:
        raise RuntimeError(f"Expected {expected} paragraphs, found {len(document.paragraphs)}")
    for index, text in plain.items():
        if document.paragraphs[index]._p.xpath(".//m:oMath|.//m:oMathPara"):
            raise RuntimeError(f"Paragraph {index} contains math and needs markup replacement")
        set_plain_text(document.paragraphs[index], text, chinese)
    for index, text in markup.items():
        set_paragraph_markup(document.paragraphs[index], text, chinese)


def save_document(source: Path, output: Path, plain: dict[int, str], markup: dict[int, str], chinese: bool) -> None:
    document = Document(source)
    original_equations = len(document._element.xpath(".//m:oMath"))
    original_tables = len(document.tables)
    original_images = len(document.inline_shapes)
    apply_rewrites(document, plain, markup, chinese)
    remove_word_joiners(document)
    normalize_fonts(document, chinese)
    normalize_paragraph_layout(document, chinese)
    document.core_properties.subject = (
        "Cross-lingual entity alignment; academically polished bilingual manuscript"
        if not chinese
        else "跨语言实体对齐；中文学术润色终稿"
    )
    document.save(output)

    checked = Document(output)
    assert len(checked.tables) == original_tables
    assert len(checked.inline_shapes) == original_images
    assert len(checked._element.xpath(".//m:oMath")) == original_equations


def main() -> None:
    save_document(ZH_SOURCE, ZH_OUTPUT, ZH_PLAIN, ZH_MARKUP, chinese=True)
    save_document(EN_SOURCE, EN_OUTPUT, EN_PLAIN, EN_MARKUP, chinese=False)
    print(ZH_OUTPUT)
    print(EN_OUTPUT)


if __name__ == "__main__":
    main()
