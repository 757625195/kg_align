from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Dict, Iterable, List

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips


SECTION_HEADING = "5.7 多数据集扩展验证与验证集选择后处理"
TABLE_HEADERS = [
    "数据集",
    "基础 Hits@1",
    "+ Reranker",
    "+ 验证选择 CSLS",
    "Hits@1 增益/百分点",
    "最终 MRR（增益/百分点）",
]
NEW_DATASETS = {
    "openea_d_w": "OpenEA D-W",
    "openea_d_y": "OpenEA D-Y",
    "eventea_en_en": "EventEA en-en",
}
OLD_DATASET_GAINS = [1.25, 1.86, 0.55, 3.26, 2.61]
OLD_RERANKER_GAINS = [0.10, -0.01, 0.01, 0.34, 0.18]
EVENTEA_REFERENCE = (
    "[20] Tian X, Sun Z, Li G, Hu W (2022) EventEA: benchmarking entity alignment "
    "for event-centric knowledge graphs. arXiv:2211.02817"
)


def set_run_fonts(run, east_asia: str = "STSong", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)


def replace_paragraph_text(paragraph, text: str) -> None:
    paragraph.text = text
    for run in paragraph.runs:
        set_run_fonts(run)


def replace_text_preserving_embedded_objects(paragraph, old: str, new: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    combined = "".join(node.text or "" for node in text_nodes)
    start = combined.find(old)
    if start < 0:
        raise RuntimeError(f"Could not locate paragraph text to replace: {old}")
    end = start + len(old)

    positions = []
    cursor = 0
    for node in text_nodes:
        text = node.text or ""
        positions.append((node, cursor, cursor + len(text)))
        cursor += len(text)

    overlapping = [item for item in positions if item[1] < end and item[2] > start]
    if not overlapping:
        raise RuntimeError("Replacement span does not intersect any Word text node")

    first_node, first_start, _ = overlapping[0]
    last_node, last_start, _ = overlapping[-1]
    first_text = first_node.text or ""
    last_text = last_node.text or ""
    before = first_text[: start - first_start]
    after = last_text[end - last_start :]

    first_node.text = before + new + (after if first_node is last_node else "")
    for node, _, _ in overlapping[1:-1]:
        node.text = ""
    if last_node is not first_node:
        last_node.text = after


def set_cell_width(cell, width_twips: int) -> None:
    cell.width = Twips(width_twips)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int = 70, left: int = 80, bottom: int = 70, right: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def configure_table(table) -> None:
    widths = [1720, 1370, 1470, 1760, 1000, 1800]
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.w = Twips(width)

    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for col_index, (cell, width) in enumerate(zip(row.cells, widths)):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                set_run_fonts(run)
                run.font.size = Pt(8.2)
                run.bold = row_index == 0 or col_index in (3, 4)
            if row_index == 0:
                shade_cell(cell, "F4F6F9")


def load_results(paths: Iterable[Path]) -> Dict[str, Dict]:
    rows: Dict[str, Dict] = {}
    for path in paths:
        for row in json.loads(path.read_text(encoding="utf-8")):
            rows[row["dataset"]] = row
    missing = sorted(set(NEW_DATASETS) - set(rows))
    if missing:
        raise RuntimeError(f"Missing aggregate results for: {missing}")
    if any(int(rows[name]["runs"]) != 3 for name in NEW_DATASETS):
        raise RuntimeError("Each new dataset must contain exactly three completed runs")
    return rows


def format_mean_std(row: Dict, prefix: str) -> str:
    return f"{row[f'{prefix}_mean']:.4f}±{row[f'{prefix}_std']:.4f}"


def format_result_row(name: str, row: Dict) -> List[str]:
    hits_gain = 100.0 * (row["csls_Hits@1_mean"] - row["base_Hits@1_mean"])
    mrr_gain = 100.0 * (row["csls_MRR_mean"] - row["base_MRR_mean"])
    return [
        NEW_DATASETS[name],
        format_mean_std(row, "base_Hits@1"),
        format_mean_std(row, "reranker_Hits@1"),
        format_mean_std(row, "csls_Hits@1"),
        f"{hits_gain:+.2f}",
        f"{row['csls_MRR_mean']:.4f} ({mrr_gain:+.2f})",
    ]


def find_results_table(document: Document):
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers == TABLE_HEADERS:
            return table
    raise RuntimeError("Could not locate Table 9 by its headers")


def paragraph_after_heading(document: Document, heading: str, offset: int):
    heading_index = next(
        (index for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip() == heading),
        None,
    )
    if heading_index is None:
        raise RuntimeError(f"Could not locate heading: {heading}")
    return document.paragraphs[heading_index + offset]


def update_document(document: Document, results: Dict[str, Dict]) -> None:
    table = find_results_table(document)
    existing_names = {row.cells[0].text.strip() for row in table.rows[1:]}
    duplicates = sorted(existing_names.intersection(NEW_DATASETS.values()))
    if duplicates:
        raise RuntimeError(f"The source document already contains new result rows: {duplicates}")

    for dataset_name in NEW_DATASETS:
        cells = table.add_row().cells
        for cell, value in zip(cells, format_result_row(dataset_name, results[dataset_name])):
            cell.text = value
    configure_table(table)

    intro = paragraph_after_heading(document, SECTION_HEADING, 1)
    caption = paragraph_after_heading(document, SECTION_HEADING, 2)
    result_summary = paragraph_after_heading(document, SECTION_HEADING, 3)
    reranker_summary = paragraph_after_heading(document, SECTION_HEADING, 4)
    parameter_summary = paragraph_after_heading(document, SECTION_HEADING, 5)
    protocol_note = paragraph_after_heading(document, SECTION_HEADING, 6)

    replace_paragraph_text(
        intro,
        "为检验最新简化配置与后处理方案的跨数据集稳定性，本节在既有五个跨语言数据集之外，"
        "进一步加入同语种异源知识图谱 OpenEA D-W-15K-V2、D-Y-15K-V2[16]，以及关系和属性异质性更强的 "
        "EventEA EN-EN-20K[20]。OpenEA V2 对 DBpedia 和 YAGO 实体 URI 进行编码以降低名称偏差；EventEA 则读取其"
        "官方实体名称文件。八个数据集均运行随机种子 42、43 和 44，并在完整目标知识图谱实体集合上计算排名。"
        "实验依次比较基础模型、加入 top-20 MLP 重排序器以及在重排序分数上使用验证集选择参数的 CSLS。"
        "DBP15K 从原训练对中固定留出 10% 作为验证集，OpenEA 与 EventEA 使用官方验证集；测试集仅用于最终报告。",
    )
    replace_paragraph_text(caption, "表 9  八数据集扩展实验结果（均值±样本标准差，n=3）")

    new_gains = [
        100.0 * (results[name]["csls_Hits@1_mean"] - results[name]["base_Hits@1_mean"])
        for name in NEW_DATASETS
    ]
    all_gains = OLD_DATASET_GAINS + new_gains
    positive_new_datasets = sum(gain > 0.0 for gain in new_gains)

    replace_paragraph_text(
        result_summary,
        f"三个新增数据集的最终组合相对基础模型分别变化 {new_gains[0]:+.2f}、{new_gains[1]:+.2f} 和 "
        f"{new_gains[2]:+.2f} 个 Hits@1 百分点，其中 {positive_new_datasets}/3 个新增数据集的三随机种子平均增益为正。"
        f"纳入既有五数据集后，八个数据集的平均 Hits@1 绝对增益为 "
        f"{sum(all_gains) / len(all_gains):.2f} 个百分点。D-W 与 D-Y 排除了语言差异，因而其结果主要反映"
        "跨知识库关系、属性和邻域结构异质性；EventEA 则进一步检验事件实体及复杂属性条件下的鲁棒性。",
    )

    new_reranker_gains = [
        100.0 * (results[name]["reranker_Hits@1_mean"] - results[name]["base_Hits@1_mean"])
        for name in NEW_DATASETS
    ]
    all_reranker_gains = OLD_RERANKER_GAINS + new_reranker_gains
    replace_paragraph_text(
        reranker_summary,
        f"三个新增数据集上，Reranker 单独带来的 Hits@1 变化分别为 {new_reranker_gains[0]:+.2f}、"
        f"{new_reranker_gains[1]:+.2f} 和 {new_reranker_gains[2]:+.2f} 个百分点；八数据集平均变化为 "
        f"{sum(all_reranker_gains) / len(all_reranker_gains):+.2f} 个百分点。因此，仍不应将重排序器单独表述为"
        "主要性能来源；最终数字应解释为重排序器与验证集选择 CSLS 的组合结果。",
    )

    selected_labels = {
        "openea_d_w": "D-W",
        "openea_d_y": "D-Y",
        "eventea_en_en": "EventEA",
    }
    selected_parts = []
    for name in NEW_DATASETS:
        pairs = []
        for choice in results[name]["selected_csls"].split(","):
            k_value, blend_value = choice.split("/")
            pairs.append(
                f"({k_value.removeprefix('k=')}, {blend_value.removeprefix('b=')})"
            )
        selected_parts.append(f"{selected_labels[name]} " + "、".join(pairs))
    selected = "；".join(selected_parts)
    replace_paragraph_text(
        parameter_summary,
        "按随机种子 42、43、44 的顺序，三个新增数据集由验证集选出的 (k, blend) 分别为："
        + selected
        + "。不同来源和不同结构异质性下的最优参数并不固定，"
        "因此本文不根据测试集为所有数据集写死同一 CSLS 配置，而是继续以验证 MRR 独立选择 k 与 blend。",
    )
    replace_paragraph_text(
        protocol_note,
        "需要指出，表 9 的 DBP15K 实验使用训练对留出验证集，OpenEA 与 EventEA 使用官方训练、验证和测试划分。"
        "EventEA 排名还包含两侧知识图谱中的上下文实体，候选集合不局限于事件对齐真值，并使用官方名称文件；因此其"
        "绝对结果不宜与弱化名称信息的 OpenEA V2 直接比较。表 9 用于评价最新简化配置及后处理的跨数据集稳定性，"
        "不应与表 3 直接合并为同一训练协议，也不构成与外部文献方法的严格 SOTA 比较。",
    )

    limitation = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("内部有效性方面")
    )
    old_sentence = (
        "外部有效性方面，本文只覆盖 DBP15K ZH–EN 和 OpenEA EN–FR 15K V2，尚未验证更大规模、"
        "开放世界或无对应实体场景。"
    )
    new_sentence = (
        "外部有效性方面，补充实验已覆盖 DBP15K、OpenEA 的跨语言与同语种异源场景，以及 EventEA 的事件型异构场景，"
        "但尚未验证 100K 以上规模、开放世界或无对应实体场景。"
    )
    replace_text_preserving_embedded_objects(limitation, old_sentence, new_sentence)

    conclusion = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("补充五数据集实验进一步表明")
    )
    replace_paragraph_text(
        conclusion,
        f"补充八数据集实验进一步表明，在不改变模型主体结构的条件下，reranker 与验证集选择 CSLS 的组合能够"
        f"跨语言、同语种异源及事件型异构场景改善当前系统的最终检索结果；八数据集平均 Hits@1 提升 "
        f"{sum(all_gains) / len(all_gains):.2f} 个百分点。该结果扩大了实验覆盖范围，但不同数据集的增益和最优参数"
        "差异仍要求按验证集独立选择后处理配置。",
    )

    if not any(paragraph.text.strip().startswith("[20] Tian X") for paragraph in document.paragraphs):
        reference = document.add_paragraph(EVENTEA_REFERENCE, style="Reference")
        for run in reference.runs:
            set_run_fonts(run)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("results", type=Path, nargs="+")
    args = parser.parse_args()

    if args.source.resolve() == args.output.resolve():
        raise ValueError("Output must differ from source to preserve the user's document")

    results = load_results(args.results)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)
    update_document(document, results)
    document.save(args.output)


if __name__ == "__main__":
    main()
