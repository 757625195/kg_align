from __future__ import annotations

import copy
import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from revise_manuscript1_by_priority import (
    add_mixed,
    clear_paragraph,
    find_paragraph,
    insert_paragraph_before,
    paragraph_text,
    replace_display_equation,
    set_cell,
    set_cell_mixed,
    set_mixed,
    set_plain,
    set_run_font,
    set_table_geometry,
    set_table_pagination,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "1.docx"
RESULT_ROOT = ROOT / "outputs" / "reviewer_query_gate_2x2_20260831"
SUMMARY = RESULT_ROOT / "summary"
OUTPUT = ROOT / "outputs" / "1_二因素审稿实验标黄修订版_20260831.docx"

DATASET_LABELS = {
    "dbp15k_zh_en": "DBP15K ZH–EN",
    "openea_en_fr": "OpenEA EN–FR-15K-V2",
}
CELL_ORDER = ("SS", "TS", "ST", "TT")
CELL_LABELS = {
    "SS": "Semantic query + semantic gate",
    "TS": "Structural query + semantic gate",
    "ST": "Semantic query + structural gate",
    "TT": "Structural query + structural gate",
}


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def row_index(rows: list[dict[str, str]], **keys: str) -> dict[str, str]:
    for row in rows:
        if all(row[key] == value for key, value in keys.items()):
            return row
    raise RuntimeError(f"Missing summary row: {keys}")


def aggregate_pair(row: dict[str, str]) -> str:
    return (
        f"{float(row['Hits@1_mean']):.4f} ± {float(row['Hits@1_std']):.4f} / "
        f"{float(row['MRR_mean']):.4f} ± {float(row['MRR_std']):.4f}"
    )


def signed_points(value: float) -> str:
    if abs(value) < 0.005:
        value = 0.0
    return f"{value:+.2f}"


def add_run_highlight(run_element) -> None:
    r_pr = run_element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        run_element.insert(0, r_pr)
    highlight = r_pr.find(qn("w:highlight"))
    if highlight is None:
        highlight = OxmlElement("w:highlight")
        r_pr.append(highlight)
    highlight.set(qn("w:val"), "yellow")


def shade_paragraph(paragraph, fill: str = "FFF59D") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def highlight_paragraph(paragraph) -> None:
    for run in paragraph._p.xpath(".//w:r"):
        add_run_highlight(run)
    if paragraph._p.xpath(".//m:oMath"):
        shade_paragraph(paragraph)


def shade_cell(cell, fill: str = "FFF59D") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def build_table6(
    document: Document,
    aggregates: list[dict[str, str]],
    old_table,
):
    headers = (
        "Cell",
        "Query source",
        "Gate source",
        "Neighbor access",
        "DBP15K\nHits@1/MRR",
        "OpenEA\nHits@1/MRR",
    )
    rows: list[list[str]] = []
    for cell in CELL_ORDER:
        dbp = row_index(aggregates, dataset="dbp15k_zh_en", cell=cell)
        opn = row_index(aggregates, dataset="openea_en_fr", cell=cell)
        rows.append(
            [
                cell,
                "Semantic" if cell[0] == "S" else "Structure",
                "Semantic" if cell[1] == "S" else "Structure",
                "Complete one-hop",
                aggregate_pair(dbp),
                aggregate_pair(opn),
            ]
        )
    rows.append(
        [
            "Late fusion",
            "–",
            "–",
            "None",
            paragraph_text(old_table.cell(3, 3).paragraphs[0]),
            paragraph_text(old_table.cell(3, 4).paragraphs[0]),
        ]
    )

    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = old_table.style
    for column, header in enumerate(headers):
        set_cell(
            table.cell(0, column),
            header,
            bold=True,
            size=8.0,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_number, values in enumerate(rows, start=1):
        for column, value in enumerate(values):
            set_cell(
                table.cell(row_number, column),
                value,
                size=7.9,
                align=(
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column in {0, 3}
                    else WD_ALIGN_PARAGRAPH.CENTER
                ),
            )
            table.cell(row_number, column).vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )
    set_table_geometry(table, [900, 1200, 1200, 1400, 2150, 2150])
    set_table_pagination(table)
    for row in table.rows:
        for cell in row.cells:
            shade_cell(cell)
    old_table._tbl.addprevious(copy.deepcopy(table._tbl))
    table._tbl.getparent().remove(table._tbl)
    old_table._tbl.getparent().remove(old_table._tbl)


def add_per_seed_appendix(document: Document, runs: list[dict[str, str]]) -> None:
    references = find_paragraph(document, "References")
    heading = insert_paragraph_before(
        document, references, "Appendix A. Two-factor query–gate results", "Heading 1"
    )
    lead = insert_paragraph_before(
        document,
        references,
        "The first letter in each configuration denotes the dot-product query source; the second denotes the compatibility-gate source. S denotes the semantic state and T the structural state. Each cell reports Hits@1/MRR and the validation-selected CSLS k.",
    )
    caption = insert_paragraph_before(
        document,
        references,
        "Table A1 Per-seed results for the two-factor query–gate control",
        "Caption",
    )
    keep_with_next(caption)

    table = document.add_table(rows=1 + 2 * len(CELL_ORDER), cols=5)
    headers = ("Dataset", "Cell", "Seed 42", "Seed 43", "Seed 44")
    for column, header in enumerate(headers):
        set_cell(
            table.cell(0, column),
            header,
            bold=True,
            size=8.2,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    row_number = 1
    for dataset in DATASET_LABELS:
        for cell in CELL_ORDER:
            set_cell(
                table.cell(row_number, 0),
                DATASET_LABELS[dataset],
                size=7.7,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            set_cell(
                table.cell(row_number, 1),
                cell,
                size=8.0,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            for offset, seed in enumerate((42, 43, 44), start=2):
                row = row_index(runs, dataset=dataset, cell=cell, seed=str(seed))
                value = (
                    f"{float(row['Hits@1']):.4f}/{float(row['MRR']):.4f} "
                    f"(k={int(row['csls_k'])})"
                )
                set_cell(
                    table.cell(row_number, offset),
                    value,
                    size=7.9,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )
            row_number += 1
    set_table_geometry(table, [1700, 1000, 2100, 2100, 2100])
    set_table_pagination(table)
    for row in table.rows:
        for cell in row.cells:
            shade_cell(cell)
    references._p.addprevious(table._tbl)
    highlight_paragraph(heading)
    highlight_paragraph(lead)
    highlight_paragraph(caption)


def revise_method(document: Document) -> list:
    changed = []

    def plain(prefix: str, text: str):
        paragraph = find_paragraph(document, prefix)
        set_plain(paragraph, text)
        changed.append(paragraph)
        return paragraph

    def mixed(prefix: str, text: str):
        paragraph = find_paragraph(document, prefix)
        set_mixed(paragraph, text)
        changed.append(paragraph)
        return paragraph

    plain(
        "3.5.1 Semantics-guided neighbor weighting",
        "3.5.1 Query and compatibility-gate sources",
    )
    mixed(
        "For entity i, the fusion module receives four inputs",
        r"For entity $i$, let $s_i=z_i^{\mathrm{sem}}$ denote its semantic state, $t_i=z_i^{\mathrm{str}}$ its own structural state, and $t_{i,j}$ the structural state of outgoing neighbor $j$. The dot-product query uses source $u_i^{(q)}$, while the compatibility gate uses source $u_i^{(g)}$. Each source can be either $s_i$ or $t_i$. The complete model sets both sources to $s_i$. Keys and values are projected from $t_{i,j}$.",
    )
    equation14 = find_paragraph(document, "aij=d")
    replace_display_equation(
        equation14,
        r"a_{ij}=d^{-1/2}\left(W_Kt_{i,j}\right)^{\top}W_Qu_i^{(q)}",
        "14",
    )
    changed.append(equation14)
    mixed(
        "Equation (14) measures",
        r"Equation (14) measures compatibility between query source $u_i^{(q)}$ and structural neighbor $j$. The factor $d^{-1/2}$, with $d=128$, controls the numerical scale of the dot product.",
    )
    equation16 = find_paragraph(document, "pij=")
    replace_display_equation(
        equation16,
        r"p_{ij}=\sigma\!\left(W_P\phi(t_{i,j},u_i^{(g)})\right),\quad u_i^{(g)}\in\{s_i,t_i\}",
        "16",
    )
    changed.append(equation16)
    mixed(
        "The gate value",
        r"The scalar $p_{ij}\in(0,1)$ measures compatibility between neighbor $j$ and gate source $u_i^{(g)}$. The model combines this value with the dot-product score before normalization. Separating $u_i^{(q)}$ and $u_i^{(g)}$ makes the two semantic-conditioning paths independently testable.",
    )
    equation17 = find_paragraph(document, "αi=entmax")
    replace_display_equation(
        equation17,
        r"\boldsymbol{\alpha}_i=\operatorname{entmax}_{1.5}\!\left(\frac{\mathbf{a}_i+\log(\mathbf{p}_i+\epsilon)}{T}+\mathbf{m}_i\right)",
        "17",
    )
    changed.append(equation17)
    mixed(
        "In Eq. (17)",
        r"In Eq. (17), $\mathbf{a}_i$ and $\mathbf{p}_i$ collect the scores and gate values for all padded neighbor positions. Mask bias $m_{ij}$ equals zero for a valid neighbor and $-\infty$ for padding. The mask is applied after temperature scaling. Temperature $T$ controls concentration, and 1.5-entmax can assign exact zero weight to a valid low-scoring neighbor [21]. Valid weights sum to one; an entity without valid neighbors receives an all-zero summary.",
    )
    plain(
        "The model then uses the self structural representation",
        "The model then uses the self structural representation, neighbor summary, and semantic representation to compute a context gate for each dimension. Equation (18) keeps the semantic state fixed in every two-factor control, so the experiment changes only the query source in Eq. (14) and the compatibility-gate source in Eq. (16).",
    )
    plain(
        "The semantic query changes structural-neighbor weights",
        "In the complete model, semantics enters neighbor weighting through both the dot-product query and the compatibility gate. The resulting structural context is then combined with the original semantic state. Section 3.5.4 separates the two conditioning paths while keeping context construction and final fusion fixed.",
    )
    plain("3.5.4 Matched controls", "3.5.4 Two-factor matched controls")
    mixed(
        "The structural-query control tests the query signal",
        r"The two-factor control crosses query source $u_i^{(q)}\in\{s_i,t_i\}$ with compatibility-gate source $u_i^{(g)}\in\{s_i,t_i\}$. This produces semantic query + semantic gate (SS), structural query + semantic gate (TS), semantic query + structural gate (ST), and structural query + structural gate (TT). All four configurations use the same complete one-hop outgoing neighborhood, dynamic padding, 1.5-entmax, temperature, trainable layers, InfoNCE objective, Eq. (18) context gate, final joint gate, and joint retrieval. The parameter-matched late-fusion control retains the same fusion parameter count but omits neighbor context.",
    )
    return changed


def revise_claims(
    document: Document,
    effects: list[dict[str, str]],
    aggregates: list[dict[str, str]],
) -> list:
    changed = []

    def effect(dataset: str, name: str, metric: str = "Hits@1") -> dict[str, str]:
        return row_index(
            effects, dataset=dataset, effect=name, metric=metric
        )

    q_dbp = float(effect("dbp15k_zh_en", "query_main_effect")["points_mean"])
    q_open = float(effect("openea_en_fr", "query_main_effect")["points_mean"])
    g_dbp = float(effect("dbp15k_zh_en", "gate_main_effect")["points_mean"])
    g_open = float(effect("openea_en_fr", "gate_main_effect")["points_mean"])
    i_dbp = float(effect("dbp15k_zh_en", "query_gate_interaction")["points_mean"])
    i_open = float(effect("openea_en_fr", "query_gate_interaction")["points_mean"])

    def conditional(dataset: str, name: str) -> tuple[float, float]:
        row = effect(dataset, name)
        return float(row["points_mean"]), float(row["points_std"])

    dbp_qs, dbp_qs_sd = conditional("dbp15k_zh_en", "query_effect_sem_gate")
    dbp_qt, dbp_qt_sd = conditional("dbp15k_zh_en", "query_effect_struct_gate")
    open_qs, open_qs_sd = conditional("openea_en_fr", "query_effect_sem_gate")
    open_qt, open_qt_sd = conditional("openea_en_fr", "query_effect_struct_gate")
    dbp_gs, dbp_gs_sd = conditional("dbp15k_zh_en", "gate_effect_sem_query")
    dbp_gt, dbp_gt_sd = conditional("dbp15k_zh_en", "gate_effect_struct_query")
    open_gs, open_gs_sd = conditional("openea_en_fr", "gate_effect_sem_query")
    open_gt, open_gt_sd = conditional("openea_en_fr", "gate_effect_struct_query")

    abstract = find_paragraph(document, "Knowledge graph entity alignment identifies")
    set_plain(
        abstract,
        "Knowledge graph entity alignment identifies entities that refer to the same object across independently constructed graphs. This paper presents a relation-aware neighbor-context model. The structural encoder preserves relation types and combines graph states from different propagation depths. The semantic encoder extracts token, phrase, and global views from one entity-text sequence. The fusion module scores the complete one-hop outgoing neighborhood, applies 1.5-entmax to suppress weak neighbors, and combines the resulting structural context with entity semantics. Experiments use three random seeds on five datasets. The full model achieves Hits@1 scores of 0.7285 on DBP15K ZH–EN and 0.6344 on OpenEA EN–FR-15K-V2, improving over the semantic branch by 9.09 and 9.66 percentage points. A two-factor control independently changes the dot-product query and compatibility-gate sources. Across the two gate settings, a semantic query changes Hits@1 by "
        f"{signed_points(q_dbp)} and {signed_points(q_open)} points; across the two query settings, a semantic gate changes it by {signed_points(g_dbp)} and {signed_points(g_open)} points. "
        "The results identify relation types and structural neighbor context as the main sources of improvement, while query and gate sources have smaller effects.",
    )
    changed.append(abstract)

    research = find_paragraph(document, "This study addresses three questions")
    set_plain(
        research,
        "This study addresses three questions. Does relation-aware neighbor context provide alignment evidence beyond the semantic representation? When the neighborhood and fusion capacity are fixed, what effects come from the dot-product query source and the compatibility-gate source? How do relation types, neighbor context, and multi-scale semantic encoding affect the joint model?",
    )
    changed.append(research)

    contribution = find_paragraph(document, "Experiments with three random seeds on five datasets")
    set_plain(
        contribution,
        "Experiments with three random seeds on five datasets test the complete model across cross-lingual and event-centric settings. Component controls evaluate relation types, multi-scale semantics, and neighbor access. A two-factor experiment independently changes the dot-product query and compatibility-gate sources while fixing the neighborhood, trainable fusion capacity, context gate, and retrieval representation.",
    )
    changed.append(contribution)

    heading = find_paragraph(document, "5.5 Controlled test of neighbor-query signals")
    set_plain(heading, "5.5 Two-factor control of query and compatibility gate")
    changed.append(heading)

    intro = find_paragraph(document, "Table 6 tests the neighbor-query signal")
    set_plain(
        intro,
        "Table 6 reports a unified rerun of the four query–gate combinations. These rows share the complete one-hop outgoing neighborhood, 197,377 trainable fusion parameters, 1.5-entmax, temperature, context gate, final fusion, objective, and direct joint retrieval. The experiment changes only the source used by Eq. (14) and Eq. (16). The late-fusion row retains the same parameter count and provides the existing no-neighbor reference.",
    )
    changed.append(intro)

    caption = find_paragraph(document, "Table 6 Neighbor access, query signal")
    set_plain(
        caption,
        "Table 6 Two-factor query–gate control and late-fusion reference (Hits@1/MRR, mean ± sample standard deviation, n = 3)",
    )
    changed.append(caption)
    keep_with_next(caption)

    results = find_paragraph(document, "Semantic queries improve Hits@1")
    set_plain(
        results,
        "The semantic-query effect is measured as SS−TS under the semantic gate and ST−TT under the structural gate. The paired Hits@1 effects are "
        f"{signed_points(dbp_qs)} ± {abs(dbp_qs_sd):.2f} and {signed_points(dbp_qt)} ± {abs(dbp_qt_sd):.2f} points on DBP15K, and "
        f"{signed_points(open_qs)} ± {abs(open_qs_sd):.2f} and {signed_points(open_qt)} ± {abs(open_qt_sd):.2f} points on OpenEA. "
        "The semantic-gate effect is measured as SS−ST under the semantic query and TS−TT under the structural query. The corresponding effects are "
        f"{signed_points(dbp_gs)} ± {abs(dbp_gs_sd):.2f} and {signed_points(dbp_gt)} ± {abs(dbp_gt_sd):.2f} points on DBP15K, and "
        f"{signed_points(open_gs)} ± {abs(open_gs_sd):.2f} and {signed_points(open_gt)} ± {abs(open_gt_sd):.2f} points on OpenEA.",
    )
    changed.append(results)

    discussion_anchor = find_paragraph(document, "6 Discussion")
    interpretation = insert_paragraph_before(
        document,
        discussion_anchor,
        "Averaging over the other factor gives semantic-query main effects of "
        f"{signed_points(q_dbp)} and {signed_points(q_open)} Hits@1 points and semantic-gate main effects of {signed_points(g_dbp)} and {signed_points(g_open)} points. "
        f"The query–gate interactions are {signed_points(i_dbp)} and {signed_points(i_open)} points. These effects are smaller than the 5.28- and 9.22-point gaps to parameter-matched late fusion. The direct mechanism evidence therefore supports neighbor access and the learned interaction path more strongly than either semantic conditioning path alone.",
    )
    changed.append(interpretation)

    discussion = find_paragraph(document, "The component controls identify three main sources")
    set_plain(
        discussion,
        "The component controls identify relation types and structural neighbor context as the main sources of improvement. The full model exceeds the semantic branch on both datasets, so the selected structural context contributes evidence beyond text. The two-factor experiment further shows that the dot-product query and compatibility-gate sources account for smaller changes once the neighborhood, capacity, context gate, and retrieval representation are fixed.",
    )
    changed.append(discussion)

    scope = find_paragraph(document, "The main experiments cover four cross-lingual graph pairs")
    set_plain(
        scope,
        "The main experiments cover four cross-lingual graph pairs and the same-language EventEA EN–EN pair. Component controls focus on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. The two-factor rows in Table 6 form a self-contained three-seed rerun and support mechanism claims on these two datasets. The fixed eight-neighbor softmax control changes both neighbor count and weighting. Parameter-matched late fusion removes neighbor access and changes the training interaction path, so its difference represents the combined design effect.",
    )
    changed.append(scope)

    conclusion = find_paragraph(document, "This paper examines the role of relation-aware structural neighbors")
    set_plain(
        conclusion,
        "This paper examines relation-aware structural neighbors in a structural–semantic joint representation. Relation types strengthen structural evidence on both controlled datasets, and the complete model improves over the semantic branch. Complete-neighborhood access with a learned interaction path also exceeds parameter-matched late fusion. The two-factor experiment separates the dot-product query from the compatibility gate and shows that their source choices have smaller effects than neighbor access. The five-dataset results show that the complete configuration applies to cross-lingual graph pairs and a same-language event-centric graph pair.",
    )
    changed.append(conclusion)
    return changed


def validate_inputs(
    runs: list[dict[str, str]],
    aggregates: list[dict[str, str]],
    effects: list[dict[str, str]],
) -> None:
    if len(runs) != 24:
        raise RuntimeError(f"Expected 24 completed runs, found {len(runs)}")
    if len(aggregates) != 8:
        raise RuntimeError(f"Expected 8 aggregate rows, found {len(aggregates)}")
    if len(effects) != 28:
        raise RuntimeError(f"Expected 28 paired-effect rows, found {len(effects)}")
    for row in aggregates:
        if int(row["n"]) != 3:
            raise RuntimeError(f"Aggregate row does not contain three seeds: {row}")


def main() -> None:
    runs = read_tsv(SUMMARY / "runs.tsv")
    aggregates = read_tsv(SUMMARY / "aggregate.tsv")
    effects = read_tsv(SUMMARY / "paired_effects.tsv")
    validate_inputs(runs, aggregates, effects)

    document = Document(SOURCE)
    method_changes = revise_method(document)
    claim_changes = revise_claims(document, effects, aggregates)
    old_table6 = document.tables[5]
    build_table6(document, aggregates, old_table6)
    add_per_seed_appendix(document, runs)

    for paragraph in method_changes + claim_changes:
        highlight_paragraph(paragraph)
    for table in document.tables:
        set_table_pagination(table)

    document.core_properties.subject = (
        "Highlighted revision with independent query and compatibility-gate controls"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
