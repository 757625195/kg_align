from copy import deepcopy
from pathlib import Path
import re

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_教师要求Springer去重终稿_20260826.docx"
OUTPUT = ROOT / "outputs" / "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_Springer投稿删减版_20260826.docx"


def replace_paragraph(paragraph, text):
    """Replace prose while retaining paragraph properties and first-run styling."""
    first_rpr = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            first_rpr = deepcopy(run._r.rPr)
            break
    paragraph.clear()
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._r.insert(0, first_rpr)


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def find_paragraph(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"Paragraph not found: {prefix}")


def remove_section(document, heading_prefix, next_heading_prefix):
    heading = find_paragraph(document, heading_prefix)
    next_heading = find_paragraph(document, next_heading_prefix)
    current = heading._p
    while current is not None and current is not next_heading._p:
        following = current.getnext()
        remove_element(current)
        current = following


def replace_first_number_in_runs(paragraph, old_number, new_number):
    pattern = re.compile(rf"(?<!\d){old_number}(?!\d)")
    for run in paragraph.runs:
        replaced, count = pattern.subn(str(new_number), run.text, count=1)
        if count:
            run.text = replaced
            return
    raise RuntimeError(f"Table number {old_number} not found in: {paragraph.text}")


def main():
    document = Document(SOURCE)

    # Keep the author's abstract unchanged; this pass only reduces redundant
    # manuscript content and adjusts the positioning outside the abstract.
    replace_paragraph(
        find_paragraph(document, "关键词："),
        "关键词：知识图谱实体对齐；跨语言知识图谱；关系感知图神经网络；结构—语义融合；稀疏邻居选择；多尺度语义编码",
    )

    replace_paragraph(
        find_paragraph(document, "本文使用的 GraphSAGE"),
        "本文使用的 GraphSAGE 邻域聚合[6]、R-GCN 关系建模[7]、Transformer 自注意力[8]和 InfoNCE 对比目标[9]均来自已有研究，"
        "因而不把这些基础算子本身作为创新。本文关注的科学问题是：关系感知结构邻居应在何处进入结构—语义联合计算，以及在控制模型容量后，"
        "最终融合前可见的结构上下文是否优于两个分支独立编码后的晚期融合。局部拓扑统计仅作为结构初始状态的辅助设计，由消融实验评估，而不作为论文的核心定位。本文的主要贡献如下：",
    )
    replace_paragraph(
        find_paragraph(document, "实现一种拓扑初始化的轻量关系感知结构编码器"),
        "实现一种轻量关系感知结构编码器。共享关系投影在控制参数量的同时保留边类型差异，节点级层选择融合不同传播深度；"
        "局部拓扑统计与较小的可学习实体残差用于提供结构初始状态。",
    )

    replace_paragraph(
        find_paragraph(document, "模型由拓扑初始化的结构编码"),
        "模型由关系感知结构编码、多尺度语义编码、全一跳邻居上下文和对比学习四部分组成。结构编码器在两张知识图谱上执行关系感知消息传递，"
        "并融合不同传播深度；局部拓扑统计仅用于形成结构初始状态。语义编码器从名称、关系和属性序列构造 token、phrase 和 global 三个视图。"
        "对当前批次实体，模型保留全部一跳出邻居并按批次最大邻域宽度动态补齐，再以 1.5-entmax 形成稀疏结构上下文。训练同时约束联合表示与结构表示，"
        "验证阶段按数据集选择融合权重、有效传播深度和 CSLS 邻域。图 1 给出信息流与公式对应关系。",
    )
    figure_caption = find_paragraph(document, "图 1 拓扑初始化的关系感知结构上下文")
    for run in figure_caption.runs:
        run.text = run.text.replace(
            "图 1 拓扑初始化的关系感知结构上下文",
            "图 1 关系感知结构上下文",
        ).replace("八维图内拓扑统计", "局部拓扑统计")

    # Table 2 already contains every value visualized by Fig. 2. Remove the
    # redundant dataset-category section, Table 4, and Fig. 2 as one block.
    remove_section(document, "5.3 不同数据集类别的结果差异", "5.4 结构组件与语义视图消融")

    # Renumber the remaining result subsections and tables consecutively.
    heading_map = {
        "5.4 结构组件与语义视图消融": "5.3 结构组件与语义视图消融",
        "5.5 结构分支、语义分支与融合模块": "5.4 结构分支、语义分支与融合模块",
        "5.6 融合位置、邻居可见范围与容量控制": "5.5 融合位置、邻居可见范围与容量控制",
    }
    for old, new in heading_map.items():
        replace_paragraph(find_paragraph(document, old), new)

    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        match = re.match(r"表 ([5-8])", stripped)
        if match:
            old_number = int(match.group(1))
            replace_first_number_in_runs(paragraph, old_number, old_number - 1)
        for run in paragraph.runs:
            updated = run.text.replace("第 5.4 节", "第 5.3 节")
            if updated != run.text:
                run.text = updated

    replace_paragraph(
        find_paragraph(document, "DBP15K FR–EN 的 Hits@1"),
        "表 2 给出的五组结果均同时报告 Hits@1、Hits@10 和 MRR。三个随机种子下，各数据集 Hits@1 的样本标准差均不超过 0.0066，"
        "说明训练随机性造成的波动小于不同数据条件之间的差异；这些差异及其输入限制在第 6.2 节集中讨论。",
    )

    replace_paragraph(
        find_paragraph(document, "针对 RQ1"),
        "RQ1 的结果表明，模型能够在五个数据集上稳定训练，但绝对性能随语言、知识库来源和事件属性异构性而变化。"
        "RQ2 的消融显示，关系类型提供了跨两个代表性数据集的一致增益，其余结构和语义组件具有不同程度的数据集依赖性。"
        "RQ3 的控制实验支持在最终融合前保留结构邻居上下文，但没有证明语义查询本身优于结构查询。",
    )
    replace_paragraph(
        find_paragraph(document, "6.2 OpenEA 结果与语义输入限制"),
        "6.2 数据集差异与语义输入限制",
    )
    replace_paragraph(
        find_paragraph(document, "OpenEA EN–FR-15K-V2"),
        "DBP15K 的名称和局部结构线索相对充分，EventEA 则由 Wikidata 与 DBpedia 构成，并包含更强的事件关系和属性异构性[18]。"
        "OpenEA EN–FR-15K-V2 进一步对实体 URI 进行编码以降低名称偏置[14]，而当前输入使用英文 GloVe 覆盖词表内 token，并以基于 MD5 种子的确定性随机向量表示词表外 token[15]。"
        "本地统计显示，OpenEA 的 GloVe 覆盖率约为 52.18%，大量专名和编码标识缺少可迁移的跨语言语义，因此其检索结果低于 DBP15K。"
        "该现象与输入覆盖和跨图结构差异一致，但其因果贡献仍需通过替换语义编码器和控制输入覆盖率的实验验证。",
    )
    replace_paragraph(
        find_paragraph(document, "后续研究应优先解决三个问题"),
        "后续研究应优先解决三个问题。第一，以跨语言子词模型或可解释名称编码替代随机词表外向量，并在 OpenEA 全部官方划分上复现。"
        "第二，继续检验关系感知结构编码和稀疏邻居选择在更大图上的效率与稳定性。第三，在开放世界和存在无对应实体的场景中评估检索校准与泛化能力。",
    )
    replace_paragraph(
        find_paragraph(document, "本文实现了一种拓扑初始化的关系感知结构上下文"),
        "本文围绕关系感知结构上下文与多尺度语义表示的融合方式开展建模和受控实验。结果支持在最终融合前保留结构邻居上下文，"
        "同时表明查询信号和部分结构、语义组件的作用依赖具体数据条件。该结论限定了当前方法能够支持的主张，并为后续在更严格协议、"
        "更大规模和开放世界场景中的验证提供了实验基础。",
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
