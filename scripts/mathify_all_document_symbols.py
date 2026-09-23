from __future__ import annotations

import argparse
import copy
import html.entities
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn


ROOT = Path(__file__).resolve().parents[1]
MATH_DEPS = ROOT / "tmp" / "docx_math_deps"
sys.path.insert(0, str(MATH_DEPS))

import mathml2omml  # noqa: E402
from latex2mathml.converter import convert as latex_to_mathml  # noqa: E402


def latex_to_omml(expression: str):
    mathml = latex_to_mathml(expression.strip(), display="inline")
    omml = mathml2omml.convert(mathml, html.entities.name2codepoint)
    omml = re.sub(
        r"(<m:groupChr><m:groupChrPr>.*?)</m:groupChr>(<m:e>)",
        r"\1</m:groupChrPr>\2",
        omml,
    )
    omml = omml.replace("<m:oMath>", f"<m:oMath {nsdecls('m')}>", 1)
    return parse_xml(omml)


def has_math(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//m:oMath"))


def clear_paragraph(paragraph) -> None:
    properties = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)


def template_run_properties(paragraph):
    for run in paragraph._p.findall(qn("w:r")):
        if run.rPr is not None:
            return copy.deepcopy(run.rPr)
    return None


def set_markup(paragraph, markup: str) -> None:
    run_properties = template_run_properties(paragraph)
    clear_paragraph(paragraph)
    cursor = 0
    for match in re.finditer(r"\[\[(.+?)\]\]", markup):
        if match.start() > cursor:
            run = paragraph.add_run(markup[cursor : match.start()])
            if run_properties is not None:
                run._r.insert(0, copy.deepcopy(run_properties))
        paragraph._p.append(latex_to_omml(match.group(1)))
        cursor = match.end()
    if cursor < len(markup):
        run = paragraph.add_run(markup[cursor:])
        if run_properties is not None:
            run._r.insert(0, copy.deepcopy(run_properties))


def replace_literals(paragraph, replacements: list[tuple[str, str]]) -> None:
    if has_math(paragraph):
        raise ValueError(f"Paragraph already contains math: {paragraph.text[:80]}")
    markup = paragraph.text
    markers: list[tuple[str, str]] = []
    for index, (literal, latex) in enumerate(replacements):
        count = markup.count(literal)
        if count == 0:
            raise ValueError(f"Missing literal {literal!r} in {paragraph.text[:100]!r}")
        marker = f"@@MATH_{index}@@"
        markup = markup.replace(literal, marker)
        markers.append((marker, f"[[{latex}]]"))
    for marker, replacement in markers:
        markup = markup.replace(marker, replacement)
    set_markup(paragraph, markup)


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def mathify_numeric_statistics(document: Document) -> None:
    pattern = re.compile(r"(?<![\w.])([+−-]?\d+(?:\.\d+)?)±(\d+(?:\.\d+)?)")
    paragraphs = list(document.paragraphs)
    paragraphs.extend(
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    for paragraph in paragraphs:
        if has_math(paragraph) or not pattern.search(paragraph.text):
            continue
        text = paragraph.text
        cursor = 0
        parts: list[str] = []
        for match in pattern.finditer(text):
            parts.append(text[cursor : match.start()])
            value = match.group(1).replace("−", "-")
            parts.append(f"[[{value}\\pm {match.group(2)}]]")
            cursor = match.end()
        parts.append(text[cursor:])
        set_markup(paragraph, "".join(parts))


def mathify_body(document: Document) -> None:
    replacements: dict[str, list[tuple[str, str]]] = {
        "摘要": [],
        "跨语言知识图谱实体对齐需要": [("CSLS k", r"\operatorname{CSLS}\ k")],
        "图 1 关系感知": [
            ("CSLS k", r"\operatorname{CSLS}\ k"),
            ("α", r"\alpha"),
            ("L", "L"),
        ],
        "图 1 与公式": [
            ("z_i^str", r"z_i^{\mathrm{str}}"),
            ("X_i", r"X_i"),
            ("M_i", r"M_i"),
            ("z_i^sem", r"z_i^{\mathrm{sem}}"),
            ("c_i^str", r"c_i^{\mathrm{str}}"),
            ("z_i^joint", r"z_i^{\mathrm{joint}}"),
        ],
        "消息传递（message passing）": [
            ("h_i^(l)∈R^d", r"h_i^{(l)}\in\mathbb{R}^{d}"),
            (" i ", " i "),
            (" l ", " l "),
            ("e_r∈R^(d_r)", r"e_r\in\mathbb{R}^{d_r}"),
            (" r ", " r "),
            ("N_in(i)", r"\mathcal{N}_{\mathrm{in}}(i)"),
            ("(j,r,i)", r"(j,r,i)"),
        ],
        "其中 W_s^(l)": [
            ("W_s^(l)∈R^(d×d)", r"W_s^{(l)}\in\mathbb{R}^{d\times d}"),
            ("W_r^(l)∈R^(d×d_r)", r"W_r^{(l)}\in\mathbb{R}^{d\times d_r}"),
            ("e_r", r"e_r"),
            ("m_(j→i)^(l)∈R^d", r"m_{j\to i}^{(l)}\in\mathbb{R}^{d}"),
            ("(j,r,i)", r"(j,r,i)"),
            (" i ", " i "),
        ],
        "为避免只依赖最深层": [
            ("u_i^(0),…,u_i^(L)", r"u_i^{(0)},\ldots,u_i^{(L)}"),
            ("u_i^(0)", r"u_i^{(0)}"),
            ("u_i^(k)", r"u_i^{(k)}"),
            (" k ", " k "),
            ("u_bar_i", r"\bar{u}_i"),
            ("L+1", r"L+1"),
            ("a_i^(k)", r"a_i^{(k)}"),
            ("softmax", r"\operatorname{softmax}"),
        ],
        "整图结构编码只计算一次": [
            (" K ", " K "),
            ("K=8", "K=8"),
            ("K", "K"),
        ],
        "输入语义序列 X_i": [
            ("X_i∈R^(T×300)", r"X_i\in\mathbb{R}^{T\times300}"),
            ("M_i∈{0,1}^T", r"M_i\in\{0,1\}^{T}"),
            ("H_i∈R^(T×128)", r"H_i\in\mathbb{R}^{T\times128}"),
            ("H_i", "H_i"),
            ("M_i", "M_i"),
        ],
        "加权后的三个向量": [("L2", r"L_2")],
        "对于实体 i，融合模块": [
            (" i", " i"),
            ("s=z_i^sem∈R^d", r"s=z_i^{\mathrm{sem}}\in\mathbb{R}^{d}"),
            ("t_i=z_i^str∈R^d", r"t_i=z_i^{\mathrm{str}}\in\mathbb{R}^{d}"),
            ("T_i∈R^(K×d)", r"T_i\in\mathbb{R}^{K\times d}"),
            (" j ", " j "),
            ("t_(i,j)", r"t_{i,j}"),
            ("W_q、W_k、W_v∈R^(d×d)", r"W_q,W_k,W_v\in\mathbb{R}^{d\times d}"),
            ("q_i", "q_i"),
            ("k_j", "k_j"),
            ("v_j", "v_j"),
        ],
        "其中第一层将 2d=256": [
            ("2d=256", "2d=256"),
            ("d=128", "d=128"),
        ],
        "模型在联合表示上直接进行": [
            (" B ", " B "),
            (" i ", " i "),
            (" j ", " j "),
        ],
        "其中，τ 为温度参数": [("τ", r"\tau")],
        "模型采用 AdamW": [
            ("5×10⁻⁴", r"5\times10^{-4}"),
            ("3×10⁻⁴", r"3\times10^{-4}"),
        ],
        "训练阶段直接使用联合表示": [
            ("1−α", r"1-\alpha"),
            ("α", r"\alpha"),
            ("L2", r"L_2"),
        ],
        "其中，r_L 和 r_R": [
            ("r_L", r"r_L"),
            ("r_R", r"r_R"),
            (" k ", " k "),
            ("α", r"\alpha"),
            (" L ", " L "),
        ],
        "实验回答四个研究问题": [("CSLS k", r"\operatorname{CSLS}\ k")],
        "本文以 Hits@1": [("n=3", "n=3")],
        "公共设置为 batch size": [
            ("α∈{0.20,…,0.80}", r"\alpha\in\{0.20,\ldots,0.80\}"),
            ("L∈{1,2,3}", r"L\in\{1,2,3\}"),
            ("k∈{3,5,7,10,15,20}", r"k\in\{3,5,7,10,15,20\}"),
        ],
        "表 7 报告每个数据集": [
            ("CSLS", r"\operatorname{CSLS}"),
            ("α", r"\alpha"),
            ("L", "L"),
            ("k", "k"),
        ],
        "配置差异见表 7": [
            ("0.9038±0.0024", r"0.9038\pm0.0024"),
            ("0.5568±0.0096", r"0.5568\pm0.0096"),
            (" k ", " k "),
        ],
        "为回答结构分支和融合模块": [
            ("α=0", r"\alpha=0"),
            ("α=1", r"\alpha=1"),
            (" α ", r" \alpha "),
        ],
        "与完整模型相比，参数匹配": [
            ("+0.23", "+0.23"),
            ("+0.02", "+0.02"),
            ("+0.17", "+0.17"),
            ("−0.02", "-0.02"),
        ],
        "内部有效性方面": [("n=3", "n=3")],
        "本文构建了一套结合": [("CSLS k", r"\operatorname{CSLS}\ k")],
    }

    # Fix two missing inline objects in the indexing paragraph first.
    indexing = find_paragraph(document, "两张图首先映射到同一个全局节点编号空间")
    set_markup(
        indexing,
        "两张图首先映射到同一个全局节点编号空间，假设第一张图有 [[n_1]] 个实体，"
        "第二张图的节点编号整体增加 [[n_1]]。DBP15K 数据及其对齐划分源自 JAPE[2]，"
        "OpenEA 使用标准化实体对齐基准提供的划分[14]，EventEA 使用其官方 EN–EN "
        "事件实体划分[18]。DBP15K 的关系词表通过规范化 DBpedia URI 后缀和已知关系"
        "对齐文件合并为全局关系编号；OpenEA 与 EventEA 则按规范化后的关系名称建立"
        "共享关系编号。边保留原始方向，消息沿输入的方向聚合。",
    )

    symbol_convention = find_paragraph(document, "下标 i 和 j")
    set_markup(
        symbol_convention,
        "下标 [[i]] 和 [[j]] 分别表示待编码实体与其候选或邻居，[[l]] 表示消息传播层，"
        "[[k]] 在层选择公式中表示保留状态的深度。[[d=128]]，关系嵌入维度 "
        "[[d_r=128]]。[[\\sigma]] 表示 sigmoid，[[\\odot]] 表示逐元素乘积，"
        "[[\\lbrack\\cdot;\\cdot\\rbrack]] 表示向量拼接，[[\\operatorname{softmax}]] 在指定候选维度"
        "归一化，LayerNorm 进行特征层归一化，Norm 表示 [[L_2]] 归一化。",
    )

    protocol = find_paragraph(document, "表 3、表 5–8")
    set_markup(
        protocol,
        "表 3、表 5–8 和图 2 来自五个数据集、随机种子 42、43、44 的 15 次主模型训练。"
        "表 9–12 在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2 上复用同协议主模型结果，"
        "并新增 [[11]] 个变体、[[2]] 个数据集和 [[3]] 个种子的组合，共 66 次训练。"
        "所有运行采用相同数据划分、InfoNCE 目标、验证频率和提前停止规则；除单分支基线"
        "按定义固定 [[\\alpha=0]] 或 [[\\alpha=1]] 外，每个变体分别依据三个种子的平均"
        "验证 MRR 选择 [[\\alpha,L,k]]。表 4 的文献数值仅用于定位性能区间。",
    )

    selected_config = find_paragraph(document, "最优 α 位于")
    set_markup(
        selected_config,
        "最优 [[\\alpha]] 位于 0.50–0.70，[[k]] 覆盖 3、5、15 和 20，说明语义占比与 "
        "CSLS 局部尺度需要按数据集校准。五个主数据集均选择 [[L=3]]。新增消融也分别"
        "重新选择 [[\\alpha]]、[[L]]、[[k]]，避免把主模型的后处理参数强加给结构或语义"
        "已改变的变体。",
    )

    for prefix, items in replacements.items():
        if prefix == "摘要":
            continue
        replace_literals(find_paragraph(document, prefix), items)

    caption_rules = {
        "表 3 五数据集重新训练结果": [("±", r"\pm"), ("n=3", "n=3")],
        "表 5 验证选择前后": [("n=3", "n=3"), ("Δ", r"\Delta")],
        "图 2 五数据集原始检查点": [("n=3", "n=3")],
        "表 7 各数据集": [("±", r"\pm"), ("n=3", "n=3")],
        "表 8 五数据集最终": [("±", r"\pm"), ("n=3", "n=3")],
        "表 9 结构组件": [("±", r"\pm"), ("n=3", "n=3")],
        "表 10 分支与融合": [("±", r"\pm"), ("n=3", "n=3")],
        "表 12 参数匹配": [("±", r"\pm"), ("n=3", "n=3")],
    }
    for prefix, items in caption_rules.items():
        replace_literals(find_paragraph(document, prefix), items)


def mathify_symbol_table(document: Document) -> None:
    table = document.tables[0]
    rows = [
        (
            "[[G_q=(E_q,R_q,T_q)]]\n[[q\\in\\{L,R\\}]]",
            "集合",
            "左/右知识图谱及实体、关系、三元组集合",
        ),
        (
            r"[[A_{\mathrm{seed}},\ B]]",
            "集合；标量",
            "种子对齐集合；训练批次中的实体对数量",
        ),
        (
            r"[[h_i^{(l)},\ e_r]]",
            r"[[d;\ d_r]]",
            r"实体 [[i]] 的第 [[l]] 层结构状态；关系 [[r]] 的嵌入",
        ),
        (
            r"[[m_{j\to i}^{(l)},\ \bar{m}_i^{(l)}]]",
            r"[[d;\ d]]",
            r"单条关系消息；实体 [[i]] 的平均传入消息",
        ),
        (
            r"[[W_s^{(l)},\ W_r^{(l)},\ W_{\mathrm{self}}^{(l)}]]",
            "矩阵",
            "共享的源实体、关系和自身状态投影",
        ),
        (
            r"[[a_i^{(k)},\ z_i^{\mathrm{str}}]]",
            r"标量；[[d]]",
            r"第 [[k]] 个传播深度的选择权重；结构输出",
        ),
        (
            r"[[X_i,\ H_i,\ M_i]]",
            "[[T\\times300]]\n[[T\\times d;\\ T]]",
            "词向量输入、共享序列状态和有效 token 掩码",
        ),
        (
            r"[[z_i^{\mathrm{tok}},\ z_i^{\mathrm{phr}},\ z_i^{\mathrm{glo}},\ z_i^{\mathrm{sem}}]]",
            r"[[d]]",
            "三个语义视图及其门控融合结果",
        ),
        (
            r"[[\mathcal{N}_K(i),\ K,\ c_i^{\mathrm{str}}]]",
            r"集合；标量；[[d]]",
            "固定预算邻居、邻居数上限和结构上下文",
        ),
        (
            r"[[q_i,\ k_j,\ v_j,\ \alpha_{ij}]]",
            r"[[d;\ d;\ d]]；标量",
            "邻居注意力的查询、键、值和归一化权重",
        ),
        (
            r"[[g_i^c,\ g_i^j,\ z_i^{\mathrm{joint}}]]",
            r"[[d]]",
            "结构上下文门、联合表示门和最终实体表示",
        ),
        (
            r"[[S_{ij},\ \tau,\ \mathcal{L}_{\mathrm{align}}]]",
            "标量",
            "批内相似度、InfoNCE 温度和训练目标",
        ),
        (
            r"[[\alpha,\ L,\ k]]",
            "标量；整数；整数",
            "验证选择的语义权重、有效传播深度和 CSLS 邻域",
        ),
    ]
    if len(table.rows) != len(rows) + 1:
        raise ValueError("Unexpected symbol-table row count")
    for row, values in zip(table.rows[1:], rows):
        for cell, markup in zip(row.cells, values):
            set_markup(cell.paragraphs[0], markup)


def mathify_other_tables(document: Document) -> None:
    # Dataset tensor shapes.
    for row in document.tables[1].rows[1:]:
        value = row.cells[4].paragraphs[0].text
        latex = value.replace(",", "{,}").replace("×", r"\times")
        set_markup(row.cells[4].paragraphs[0], f"[[{latex}]]")

    # Signed deltas in Tables 5 and 6.
    for table_index, columns in ((4, (3, 6)), (5, (4,))):
        table = document.tables[table_index]
        for row in table.rows[1:]:
            for column in columns:
                value = row.cells[column].paragraphs[0].text.replace("−", "-")
                set_markup(row.cells[column].paragraphs[0], f"[[{value}]]")

    # Validation configuration table: variable headers, settings, and uncertainty.
    table = document.tables[6]
    for column, latex in ((1, r"\alpha"), (2, "L"), (3, "k")):
        set_markup(table.rows[0].cells[column].paragraphs[0], f"[[{latex}]]")
    for row in table.rows[1:]:
        for column in (1, 2, 3):
            value = row.cells[column].paragraphs[0].text
            set_markup(row.cells[column].paragraphs[0], f"[[{value}]]")

    # Ablation configuration table.
    table = document.tables[8]
    for column in (1, 3):
        text = table.rows[0].cells[column].paragraphs[0].text
        label = text.replace("α/L/k", "[[\\alpha/L/k]]")
        set_markup(table.rows[0].cells[column].paragraphs[0], label)
    for row in table.rows[1:]:
        for column in (1, 3):
            value = row.cells[column].paragraphs[0].text
            set_markup(row.cells[column].paragraphs[0], f"[[{value}]]")

    # Branch representations in Table 10.
    table = document.tables[9]
    representation_markup = {
        2: r"[[z^{\mathrm{str}},\ \alpha=0]]",
        3: r"[[z^{\mathrm{sem}},\ \alpha=1]]",
        4: r"[[\alpha z^{\mathrm{sem}}+(1-\alpha)z^{\mathrm{str}}]]",
    }
    for row_index, markup in representation_markup.items():
        set_markup(table.rows[row_index].cells[1].paragraphs[0], markup)

    # Delta symbol in table headers.
    for table_index in (4, 5):
        for cell in document.tables[table_index].rows[0].cells:
            if "Δ" in cell.paragraphs[0].text:
                replace_literals(cell.paragraphs[0], [("Δ", r"\Delta")])


def verify(document: Document) -> None:
    required = [
        r"h_i^{(l)}",
        r"\alpha",
        r"5\times10^{-4}",
        r"G_q=(E_q,R_q,T_q)",
    ]
    xml = document._element.xml
    if xml.count("<m:oMath") < 100:
        raise RuntimeError("Native-math count did not increase as expected")
    if any(token in "\n".join(p.text for p in document.paragraphs) for token in ("h_i^(l)", "d_r=128", "5×10⁻⁴")):
        raise RuntimeError("Plain-text mathematical notation remains in body paragraphs")
    if "有个实体" in "\n".join(p.text for p in document.paragraphs):
        raise RuntimeError("Missing indexing variable was not repaired")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    source = Path(args.source)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    document = Document(output)

    mathify_body(document)
    mathify_symbol_table(document)
    mathify_numeric_statistics(document)
    mathify_other_tables(document)
    document.core_properties.subject = "正文与表格数学符号统一为 Word 原生公式"
    verify(document)
    document.save(output)
    print(f"WROTE {output}")


if __name__ == "__main__":
    main()
