from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
PLAIN_SCRIPT = ROOT / "scripts" / "build_plain_academic_english_manuscript.py"
OUTPUT = ROOT / "outputs" / (
    "Relation_Aware_Neighbor_Context_"
    "Springer_English_Evidence_Focused_20260830.docx"
)
TITLE = (
    "Relation-Aware Neighbor Context and Semantic-Guided Selection "
    "for Cross-Lingual Entity Alignment"
)
RUNNING_TITLE = "Relation-Aware Neighbor Context for Entity Alignment"


def load_plain_builder():
    spec = importlib.util.spec_from_file_location("plain_builder", PLAIN_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load {PLAIN_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


FOCUS_OVERRIDES = {
    4: (
        "Cross-lingual knowledge graph entity alignment identifies entities that refer to the same object in graphs written in different languages. "
        "Many systems encode graph structure and entity text in separate branches. This design leaves a key question open. "
        "The model must decide which structural neighbors matter before it forms the joint representation. This study tests a relation-aware neighbor-context model with semantic-guided selection. "
        "The structural encoder preserves relation types and several propagation depths. The semantic encoder provides token, phrase, and global evidence. "
        "A semantic query scores every one-hop outgoing neighbor, and 1.5-entmax removes weak evidence. The model then combines the selected context with entity semantics. "
        "Under the same neighborhood and model size, semantic queries improve Hits@1 over structural queries by 2.08 points on DBP15K and 2.48 points on OpenEA. "
        "The full model also improves over semantic-only retrieval by 9.09 and 9.66 points. Removing relation types lowers Hits@1 by 3.62 and 2.78 points. "
        "The late-fusion control trails the full model by 5.28 and 9.22 points. This gap reflects both neighbor access and the training path. "
        "Three-seed experiments on five datasets show stable results. The evidence supports relation-aware neighbor context and semantic-guided neighbor selection."
    ),
    7: "1.1 Problem and evidence gap",
    8: (
        "Knowledge graphs usually represent facts as triples {{EQ0}}. In a triple, {{EQ1}}, {{EQ2}}, and {{EQ3}} denote the head entity, relation, and tail entity. "
        "The triple states that relation {{EQ5}} links head entity {{EQ4}} to tail entity {{EQ6}} [1]. Different systems often build knowledge graphs independently. "
        "They may also use different languages. The same object may therefore have different names, identifiers, attributes, and value formats. "
        "Entity alignment finds matching entities in sets {{EQ7}} and {{EQ8}}. These matches support graph integration."
    ),
    9: (
        "Existing methods provide strong evidence for both graph structure and text. MTransE and JAPE map graphs into shared or linked vector spaces [2, 3]. "
        "BootEA adds high-confidence pairs to the training set [4]. RDGCN and RREA preserve relation and neighborhood information [5, 6]. "
        "These methods establish the value of each information source. However, they provide less direct evidence about when structural neighbors should enter a joint representation."
    ),
    10: (
        "The open question matters for two reasons. First, an untyped edge may hide an important relation difference. "
        "A birthplace edge and a workplace edge can connect two people to the same city [5, 7]. Second, not every neighbor helps identify the current entity. "
        "A useful fusion method should keep relation meaning and reduce unrelated neighbor evidence. This study therefore focuses on relation-aware neighbor context. "
        "It also tests whether semantic evidence improves neighbor selection before the final joint representation."
    ),
    11: "1.2 Research questions and contributions",
    12: (
        "This study asks three questions. First, does relation-aware neighbor context add useful evidence beyond a semantic representation? "
        "Second, does a semantic query select neighbors better than a structural query under the same neighborhood and model size? "
        "Third, which structural and semantic components provide consistent support for the joint model?"
    ),
    13: "This study makes three contributions.",
    14: (
        "The model preserves relation types during structural message passing. It also uses a node-level selector to combine several propagation depths. "
        "These choices provide a compact structural state for later neighbor selection."
    ),
    15: (
        "The fusion module uses semantics to score the complete one-hop outgoing neighborhood. The 1.5-entmax function gives weak neighbors an exact zero weight. "
        "The model combines the selected neighbor context with the entity's own structure and semantics."
    ),
    16: (
        "The experiments use three seeds on five datasets. Matched controls compare semantic and structural neighbor queries under the same conditions. "
        "Other controls measure the value of relation types, semantic views, structural context, and late fusion without neighbor access."
    ),
    18: "2.1 Shared representation methods",
    19: (
        "MTransE learns mappings between language-specific embedding spaces [2]. JAPE adds attribute links to a shared structural space [3]. "
        "These methods give graph integration a clear geometric form. However, simple mappings may not describe complex multi-relation neighborhoods."
    ),
    20: "2.2 Relation-aware structural encoding",
    21: (
        "GraphSAGE combines local neighbors to learn node representations [9]. R-GCN also keeps relation types during message transformation [7]. "
        "It limits parameter growth through shared bases or blocks. These studies show that a graph encoder should preserve both neighboring entities and edge types."
    ),
    22: (
        "RDGCN adds relation evidence through an entity graph and a relation dual graph [5]. RREA uses relation reflection to keep relation differences [6]. "
        "RPR-RHGT encodes selected multi-step relation paths with a graph Transformer [10]. These models show the value of typed and multi-hop structure. "
        "The present study uses a lighter message-passing design. It focuses on how one-hop structural context enters the joint representation."
    ),
    23: "2.3 Structure-semantic interaction",
    24: (
        "Transformers connect positions across a sequence [11]. Hierarchical attention networks show that different levels of text may provide different evidence [12]. "
        "The present semantic encoder therefore uses token, phrase, and global views. These views support the main neighbor-selection mechanism."
    ),
    25: (
        "Text-based entity alignment studies show that semantics can support graph structure. RREA reports text-based and structure-only settings [6]. "
        "MCLEA aligns several data types through contrastive objectives [13]. Other work shows that simple fusion may mix spaces that are not aligned [14]. "
        "These findings motivate a more direct test. The present study holds the neighborhood and model size constant while changing the neighbor query signal."
    ),
    26: (
        "Recent studies use large language models [15], domain adaptation [16], or step-by-step agents [17]. These approaches change the supervision source or reasoning process. "
        "The present study instead tests the local interaction between structural neighbors and entity semantics."
    ),
    35: "3.2 Model overview",
    36: (
        "The model has four parts. The structural encoder preserves relation types and combines several propagation depths. "
        "The semantic encoder builds token, phrase, and global views from the same entity text. The fusion module uses semantics to score every one-hop outgoing neighbor. "
        "It forms a sparse structural context and combines that context with entity semantics. The training stage uses InfoNCE on the joint and structural representations. "
        "The retrieval stage uses the trained joint representation with CSLS. Figure 1 shows the full information flow."
    ),
    38: (
        "Fig. 1 Information flow of the proposed model. The structural branch preserves relation types and selects useful propagation depths. "
        "The semantic branch builds three views from one input. A semantic query scores every one-hop outgoing neighbor. "
        "The 1.5-entmax function removes weak evidence. The model then combines the selected structural context with semantics. "
        "Training and retrieval both use the joint representation. Validation selects only {{EQ0}}."
    ),
    78: "3.5 Semantic-guided structural context",
    79: (
        "This component implements the main interaction in the model. It reads the complete one-hop outgoing neighborhood and selects evidence for the current entity. "
        "The computation has three stages. A semantic query first scores each structural neighbor. The 1.5-entmax function removes weak candidates. "
        "The model then combines the selected neighbor summary with the entity's own structure. A final gate combines this structural context with entity semantics."
    ),
    98: "3.5.3 Final joint representation",
    103: (
        "The semantic query changes neighbor weights before the model builds the structural context. The final gate then combines this context with entity semantics. "
        "The architecture alone cannot show whether semantic querying is better. Section 3.5.4 therefore introduces a matched structural-query control."
    ),
    104: "3.5.4 Matched controls",
    105: (
        "The structural-query control isolates the query signal. It keeps the same neighborhood, dynamic padding, 1.5-entmax, gates, and joint retrieval. "
        "It replaces only the semantic query with the entity's structural representation. The late-fusion control tests a broader change. "
        "It keeps the same number of trainable fusion parameters but removes neighbor context. It joins the two branches only after separate encoding."
    ),
    108: "3.6 Training objective",
    130: "5.1 Performance across datasets",
    131: (
        "Table 2 reports joint-representation results on five datasets. Each dataset uses three random seeds. DBP15K reserves 10% of its training pairs for validation. "
        "OpenEA and EventEA use their official validation sets."
    ),
    134: (
        "The model produces stable results across the five datasets. The largest sample standard deviation for Hits@1 is 0.0053. "
        "This stability supports the controlled comparisons in the following sections."
    ),
    135: "5.2 Position relative to published results",
    136: (
        "Table 3 lists results reported in earlier studies. The DBP15K values for MTransE, JAPE, BootEA, and RDGCN come from RDGCN [5]. "
        "The RREA-text value comes from RREA [6]. The OpenEA EN-FR-V2 values are the official five-fold averages from OpenEA [8]."
    ),
    139: (
        "The methods in Table 3 use different splits, text inputs, and post-processing steps. The table therefore provides context rather than a same-protocol ranking. "
        "The controlled comparisons within the present model provide the main evidence for the paper's claims."
    ),
    141: "5.3 Evidence from component controls",
    142: (
        "Table 4 reports component controls on DBP15K ZH-EN and OpenEA EN-FR-15K-V2. Every variant uses three random seeds. "
        "Each variant also uses the same representation for training and retrieval."
    ),
    144: (
        "Relation types provide the clearest consistent structural gain. Their removal lowers Hits@1 by 3.62 points on DBP15K and 2.78 points on OpenEA. "
        "Topology features also support the structural state, with gains of 3.46 and 8.03 points over entity vectors alone. "
        "The phrase view improves both datasets by 0.81 and 3.95 points. The remaining controls show smaller or dataset-specific changes. "
        "The paper therefore treats them as supporting design choices rather than separate contributions."
    ),
    145: "5.4 Value of structural context in the joint representation",
    146: (
        "Table 5 compares the full model with the two single branches and mean fusion. Each result uses the same representation for training and retrieval. "
        "Validation does not change the weight of either branch."
    ),
    148: (
        "The full model improves Hits@1 over semantic-only retrieval by 9.09 points on DBP15K and 9.66 points on OpenEA. "
        "It also performs better than mean fusion without trainable parameters. These results show that structural evidence adds value after learned interaction with semantics. "
        "The structural branch alone is not the source of this gain. The gain appears when the model builds a trainable joint representation."
    ),
    149: "5.5 Controlled test of neighbor-query signals",
    150: (
        "Table 6 compares the full model with two controls. The structural-query control changes only the signal that scores neighbors. "
        "The late-fusion control keeps the same number of trainable fusion parameters but does not read neighbor context."
    ),
    152: (
        "The matched comparison shows that structural queries yield Hits@1 values 2.08 and 2.48 points below semantic queries. "
        "This result isolates the benefit of semantic neighbor scoring under the current protocol. Late fusion without neighbor context is 5.28 and 9.22 points below the full model. "
        "This second gap shows the value of a neighbor-visible interaction path. It does not measure fusion timing alone because two factors change together."
    ),
    154: "6.1 What the experiments establish",
    155: (
        "The experiments support three claims. First, relation types provide a consistent structural gain on both controlled datasets. "
        "Second, structural context adds clear value when the model learns it together with semantics. The full model improves over both semantic-only retrieval and simple mean fusion. "
        "Third, semantic queries select neighbors better than structural queries under the same neighborhood and model size. "
        "Together, these results support relation-aware neighbor context with semantic-guided selection."
    ),
    156: (
        "The size of the gain depends on the input data. DBP15K provides useful names and local graph clues. EventEA contains varied event relations and attributes [25]. "
        "OpenEA reduces name bias through encoded entity identifiers [8]. GloVe covers about 52.18% of the OpenEA tokens [26]. "
        "The method is most useful when text can guide neighbor selection and relation types can separate similar local structures."
    ),
    157: "6.2 Scope of the evidence",
    158: (
        "The study uses {{EQ0}} random seeds, so it reports effect size and sample variation without strong significance claims. "
        "The fixed-neighbor control changes both neighbor count and the weight function. The late-fusion control changes both neighbor access and the training path. "
        "The component controls cover two main datasets. The current evidence therefore supports the tested settings rather than every entity alignment setting."
    ),
    159: "6.3 Conditions and future work",
    160: (
        "Future work should test the method under three broader conditions. Multilingual subword encoders should first replace random vectors for unknown tokens [27, 28]. "
        "Larger graphs should then test the speed of complete-neighborhood selection. Open-world datasets should finally test entities without matches. "
        "These experiments will show whether semantic-guided neighbor selection remains useful under weaker text and larger candidate spaces."
    ),
    162: (
        "This study asks where structural neighbors should enter structure-semantic entity alignment. The results give four clear answers. "
        "Relation types improve structural evidence on both controlled datasets. A learned joint representation makes structural context useful beyond semantics alone. "
        "Semantic queries select neighbors better than structural queries under the same conditions. Complete neighbor access with a learned interaction path also outperforms late fusion without neighbor context. "
        "These findings support one main conclusion. Cross-lingual entity alignment benefits when relation-aware neighbor context is selected before the final joint representation."
    ),
}


def main() -> None:
    plain = load_plain_builder()
    plain.OVERRIDES.update(FOCUS_OVERRIDES)
    plain.OUTPUT = OUTPUT
    plain.main()

    base = plain.load_builder()
    document = Document(OUTPUT)
    base.set_plain_text(document.paragraphs[0], TITLE)
    document.paragraphs[0].alignment = base.WD_ALIGN_PARAGRAPH.CENTER
    document.paragraphs[0].paragraph_format.line_spacing = 1.0
    document.paragraphs[0].paragraph_format.space_after = base.Pt(10)
    document.paragraphs[0].paragraph_format.keep_with_next = True
    for run in document.paragraphs[0].runs:
        base.set_run_font(run, 17, bold=True)

    header = document.sections[0].header.paragraphs[0]
    base.set_plain_text(header, RUNNING_TITLE)
    header.alignment = base.WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.line_spacing = 1.0
    for run in header.runs:
        base.set_run_font(run, 9)

    document.core_properties.title = TITLE
    document.core_properties.subject = "Relation-aware neighbor context for cross-lingual entity alignment"
    document.core_properties.keywords = (
        "entity alignment; cross-lingual knowledge graphs; relation-aware GNN; "
        "semantic-guided neighbor selection; sparse structural context"
    )
    document.save(OUTPUT)

    checked = Document(OUTPUT)
    visible = "\n".join(base.combined_text(p) for p in checked.paragraphs)
    visible += "\n" + "\n".join(
        cell.text for table in checked.tables for row in table.rows for cell in row.cells
    )
    required = [
        TITLE,
        "0.7077 ± 0.0023 / 0.7497 ± 0.0016",
        "0.6096 ± 0.0043 / 0.6752 ± 0.0019",
        "structural queries yield Hits@1 values 2.08 and 2.48 points below semantic queries",
        "Relation types provide the clearest consistent structural gain",
    ]
    missing = [item for item in required if item not in visible]
    if missing:
        raise RuntimeError(f"Focused manuscript is missing required content: {missing}")
    if re.search(r"[\u3400-\u9fff]", visible):
        raise RuntimeError("Chinese text remains in the focused English manuscript")
    if len(checked.tables) != 6 or len(checked.inline_shapes) != 1:
        raise RuntimeError("Table or figure count changed")
    equations = sum(1 for node in checked.element.body.iter() if node.tag == qn("m:oMath"))
    if equations != 154:
        raise RuntimeError(f"Native equation count changed: {equations}")
    if "0.7277 ± 0.0023 / 0.7697 ± 0.0016" in visible:
        raise RuntimeError("Outdated structural-query value remains")
    print(OUTPUT)


if __name__ == "__main__":
    main()
