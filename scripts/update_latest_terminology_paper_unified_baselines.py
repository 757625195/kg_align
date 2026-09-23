from __future__ import annotations

import html.entities
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/Users/auroralin/Downloads") / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "术语与近期文献修订稿_20260827.docx"
)
OUTPUT = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "术语近期文献与统一协议经典模型重跑版_20260827.docx"
)

MATH_DEPS = ROOT / "tmp" / "docx_math_deps"
sys.path.insert(0, str(MATH_DEPS))

import mathml2omml  # noqa: E402
from latex2mathml.converter import convert as latex_to_mathml  # noqa: E402


def combined_text(paragraph) -> str:
    return "".join(
        paragraph._p.xpath(".//w:t/text()|.//m:t/text()")
    ).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning {prefix!r}, found {len(matches)}"
        )
    return matches[0]


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(text)


def latex_to_omml(expression: str):
    mathml = latex_to_mathml(expression.strip(), display="inline")
    omml = mathml2omml.convert(mathml, html.entities.name2codepoint)
    omml = re.sub(
        r"(<m:groupChr><m:groupChrPr>.*?)</m:groupChr>(<m:e>)",
        r"\1</m:groupChrPr>\2",
        omml,
    )
    omml = omml.replace("<m:oMath>", f"<m:oMath {nsdecls('m')}>", 1)
    return parse_xml(omml)


def set_paragraph_with_inline_math(
    paragraph, before: str, expression: str, after: str
) -> None:
    clear_paragraph(paragraph)
    paragraph.add_run(before)
    paragraph._p.append(latex_to_omml(expression))
    paragraph.add_run(after)


def insert_before_anchor(paragraph, anchor: str, insertion: str) -> None:
    matches = [
        node
        for node in paragraph._p.xpath(".//w:t")
        if node.text and anchor in node.text
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one text node containing {anchor!r}")
    matches[0].text = matches[0].text.replace(anchor, insertion + anchor, 1)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for run in paragraph.runs:
        run.bold = bold
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def update_comparison_table(document: Document) -> None:
    table = document.tables[2]
    if len(table.rows) != 14 or len(table.columns) != 5:
        raise RuntimeError(
            f"Unexpected comparison table shape: {len(table.rows)}x{len(table.columns)}"
        )

    table._tbl.remove(table.rows[-1]._tr)
    rows = [
        ("MTransE[2]", "0.1537", "0.5474", "0.2814"),
        ("JAPE[3]*", "0.1782", "0.5680", "0.3072"),
        ("BootEA[4]", "0.5439", "0.8041", "0.6324"),
        ("RDGCN[5]†", "0.0876", "0.3831", "0.1803"),
        ("RREA-basic[6]", "0.6512", "0.9095", "0.7451"),
        ("本文模型（三种子均值）", "0.7295", "0.8490", "0.7706"),
        ("MTransE[2]", "0.1923", "0.5448", "0.3089"),
        ("JAPE[3]*", "0.2170", "0.5752", "0.3350"),
        ("BootEA[4]", "0.6589", "0.9090", "0.7478"),
        ("RDGCN-random[5]†", "0.2039", "0.6851", "0.3625"),
        ("RREA-basic[6]", "0.7584", "0.9705", "0.8369"),
        ("本文模型（三种子均值）", "0.6324", "0.8196", "0.6978"),
    ]

    for column, value in enumerate(("数据集", "方法", "Hits@1", "Hits@10", "MRR")):
        set_cell_text(table.rows[0].cells[column], value, bold=True)

    set_cell_text(table.rows[1].cells[0], "DBP15K ZH–EN")
    set_cell_text(table.rows[7].cells[0], "OpenEA EN–FR-V2")
    best_rows = {
        5: {3},
        6: {2, 4},
        11: {2, 3, 4},
    }
    for row_index, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            set_cell_text(
                table.rows[row_index].cells[column],
                value,
                bold=column in best_rows.get(row_index, set()),
            )

    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))


def update_prose(document: Document) -> None:
    set_paragraph_text(
        find_paragraph(document, "5.2 与经典方法的定位性比较"),
        "5.2 经典方法源码重跑比较",
    )

    set_paragraph_with_inline_math(
        find_paragraph(document, "表 3 中 DBP15K"),
        "为减少不同论文实验设置造成的不可比性，表 3 在 DBP15K ZH–EN 和 OpenEA EN–FR-V2 上"
        "重新运行 MTransE、JAPE、BootEA、RDGCN 和 RREA-basic。五个经典方法均使用随机种子 42、"
        "与本文实验相同的训练/验证/测试划分以及完整目标图实体候选集，并依据验证集 MRR 独立选择 ",
        r"k_{\mathrm{CSLS}}\in\{3,5,7,10,15,20\}",
        "。本文模型保留表 2 的三随机种子均值，以保持最新主实验结果一致。"
        "因此，该表用于比较统一数据与检索协议下的实现结果，不用于随机种子配对显著性推断。",
    )

    caption = find_paragraph(document, "表 3 与经典实体对齐方法的定位性比较")
    set_paragraph_text(
        caption,
        "表 3 经典方法源码重跑与本文模型结果",
    )
    caption.paragraph_format.page_break_before = False
    caption.paragraph_format.keep_with_next = True

    paragraph_after_table = document.paragraphs[141]
    if combined_text(paragraph_after_table):
        raise RuntimeError("Expected the paragraph after Table 3 to be blank")
    set_paragraph_text(
        paragraph_after_table,
        "实现边界如下。MTransE、JAPE、BootEA 和 RDGCN 使用 OpenEA 官方源码；RREA-basic 使用原作者源码，"
        "且不启用半监督扩充。受计算预算限制，JAPE* 的属性子阶段训练 5 个 epoch，关系子阶段仍采用官方早停。"
        "RDGCN† 在 DBP15K 上的名称词向量覆盖为 626/34,460；OpenEA-V2 的实体标识经过匿名化，"
        "无法构造原始名称向量，故使用与实体身份无关的 Glorot 随机初始化，并记为 RDGCN-random。"
        "这些数值是统一本地协议下的源码重跑结果，不等同于原论文在各自设置下报告的数值。",
    )

    set_paragraph_text(
        find_paragraph(document, "在 DBP15K ZH–EN 上，本文模型的 Hits@1"),
        "在 DBP15K ZH–EN 上，本文模型的 Hits@1、Hits@10 和 MRR 分别为 0.7295、0.8490 和 0.7706。"
        "相较该数据集上表现最强的重跑基线 RREA-basic，Hits@1 和 MRR 分别高 7.83 和 2.55 个百分点，"
        "但 Hits@10 低 6.05 个百分点。OpenEA EN–FR-V2 上，本文模型的三项指标为 0.6324、0.8196 和 0.6978，"
        "低于 RREA-basic 的 0.7584、0.9705 和 0.8369，也低于 BootEA 的 0.6589、0.9090 和 0.7478。"
        "因此，源码重跑支持本文模型在 DBP15K 顶位检索上的优势，但不支持跨数据集一致领先的结论。",
    )

    limitations = find_paragraph(document, "内部有效性方面")
    insert_before_anchor(
        limitations,
        "外部有效性方面，",
        "表 3 的经典方法仅运行随机种子 42，而本文模型报告三随机种子均值，不能据此进行配对显著性推断；"
        "JAPE 和 RDGCN 还包含已明确标注的计算或输入适配，因此该表只支持本地统一协议下的实现比较。",
    )

    conclusion = find_paragraph(document, "本文围绕关系感知结构上下文")
    original = combined_text(conclusion)
    addition = (
        "经典方法源码重跑进一步表明，本文模型在 DBP15K 上具有较高的顶位检索结果，"
        "但在 OpenEA 上仍落后于 RREA-basic 和 BootEA。"
    )
    anchor = "由于该基线同时改变邻居可见范围"
    if anchor not in original:
        raise RuntimeError("Could not locate the conclusion anchor")
    set_paragraph_text(conclusion, original.replace(anchor, addition + anchor, 1))


def validate(document: Document) -> None:
    if len(document.tables) != 6:
        raise RuntimeError(f"Expected six tables, found {len(document.tables)}")
    if (len(document.tables[2].rows), len(document.tables[2].columns)) != (13, 5):
        raise RuntimeError("The comparison table does not have the expected 13x5 shape")

    body = "\n".join(combined_text(paragraph) for paragraph in document.paragraphs)
    required = [
        "5.2 经典方法源码重跑比较",
        "表 3 经典方法源码重跑与本文模型结果",
        "0.7295、0.8490 和 0.7706",
        "不支持跨数据集一致领先的结论",
        "RDGCN-random",
        "表 3 的经典方法仅运行随机种子 42",
    ]
    for text in required:
        if body.count(text) != 1:
            raise RuntimeError(f"Expected one occurrence of {text!r}")

    table_text = "\n".join(
        cell.text
        for row in document.tables[2].rows
        for cell in row.cells
    )
    for value in ("0.1537", "0.6512", "0.7295", "0.7584", "0.6324"):
        if value not in table_text:
            raise RuntimeError(f"Missing comparison value {value}")


def main() -> None:
    document = Document(SOURCE)
    update_comparison_table(document)
    update_prose(document)
    validate(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    validate(reopened)
    print(OUTPUT)


if __name__ == "__main__":
    main()
