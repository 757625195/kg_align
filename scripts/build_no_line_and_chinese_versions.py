from __future__ import annotations

import copy
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
ENGLISH_SOURCE = ROOT / "outputs" / (
    "Relation_Aware_Neighbor_Context_"
    "Springer_English_Evidence_Focused_20260830.docx"
)
CHINESE_SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "实验结论与逻辑一致性修订版_20260830.docx"
)
ENGLISH_OUTPUT = ROOT / "outputs" / (
    "Relation_Aware_Neighbor_Context_"
    "Springer_English_No_Line_Numbers_20260830.docx"
)
CHINESE_OUTPUT = ROOT / "outputs" / (
    "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_"
    "中文无行号版_20260830.docx"
)

CHINESE_TITLE = "跨语言实体对齐中的关系感知邻居上下文与语义引导选择"
CHINESE_RUNNING_TITLE = "关系感知邻居上下文实体对齐"


ZH_OVERRIDES = {
    0: CHINESE_TITLE,
    4: (
        "跨语言知识图谱实体对齐旨在识别不同语言图谱中指向同一对象的实体。许多方法分别编码图结构和实体文本，"
        "但模型仍需在形成联合表示之前判断哪些结构邻居真正有用。本文提出一种结合关系感知邻居上下文与语义引导选择的实体对齐方法。"
        "结构编码器保留关系类型和不同传播深度的信息，语义编码器提取 token、phrase 和 global 三种粒度的证据。"
        "模型使用语义查询为全部一跳出邻居评分，并通过 1.5-entmax 去除弱相关证据，随后将筛选后的结构上下文与实体语义结合。"
        "在邻域范围和模型规模一致的条件下，语义查询相对结构查询使 DBP15K 和 OpenEA 的 Hits@1 分别提高 2.08 和 2.48 个百分点。"
        "完整模型相对语义单分支分别提高 9.09 和 9.66 个百分点。移除关系类型后，Hits@1 分别下降 3.62 和 2.78 个百分点。"
        "不读取邻居上下文的晚期融合分别低 5.28 和 9.22 个百分点，但该差距同时包含邻居可见性和训练路径的变化。"
        "五个数据集上的三随机种子实验结果较为稳定。实验结果支持在跨语言实体对齐中使用关系感知邻居上下文和语义引导的邻居选择。"
    ),
    7: "1.1 问题与证据缺口",
    8: (
        "知识图谱通常使用三元组 {{EQ0}} 表示事实。在三元组中，{{EQ1}}、{{EQ2}} 和 {{EQ3}} 分别表示头实体、关系和尾实体。"
        "该三元组说明关系 {{EQ5}} 将头实体 {{EQ4}} 与尾实体 {{EQ6}} 相连[1]。不同系统通常独立构建知识图谱，也可能使用不同语言。"
        "因此，同一对象可能具有不同的名称、标识符、属性和值格式。实体对齐需要在实体集合 {{EQ7}} 和 {{EQ8}} 之间找到对应实体，以支持图谱整合。"
    ),
    9: (
        "现有研究已经证明图结构和文本信息均有价值。MTransE 和 JAPE 将不同图谱映射到共享或相互关联的向量空间[2–3]。"
        "BootEA 将高置信实体对加入训练集[4]。RDGCN 和 RREA 保留关系与邻域信息[5–6]。"
        "这些方法分别说明了不同信息源的作用，但没有充分回答结构邻居应在何时进入联合表示。"
    ),
    10: (
        "这一问题有两个实际原因。第一，无类型边可能掩盖重要的关系差异。例如，两个人都可以与同一城市相连，但一条边表示出生地，另一条边表示工作地点[5,7]。"
        "第二，并非每个邻居都有助于识别当前实体。有效的融合方法需要保留关系含义，并减少无关邻居的影响。"
        "因此，本文关注关系感知邻居上下文，并检验语义信息能否在形成最终联合表示之前改善邻居选择。"
    ),
    11: "1.2 研究问题与贡献",
    12: (
        "本文研究三个问题。第一，关系感知邻居上下文能否在语义表示之外提供有效证据？"
        "第二，在邻域和模型规模相同的条件下，语义查询能否比结构查询更有效地选择邻居？"
        "第三，哪些结构组件和语义组件能够稳定支持联合模型？"
    ),
    13: "本文的主要贡献如下。",
    14: (
        "模型在结构消息传递中保留关系类型，并使用节点级层选择器组合不同传播深度。"
        "这些设计为后续邻居选择提供紧凑的结构表示。"
    ),
    15: (
        "融合模块使用语义信息为全部一跳出邻居评分。1.5-entmax 将弱相关邻居的权重置为精确的零。"
        "模型随后将筛选后的邻居上下文与实体自身的结构和语义信息结合。"
    ),
    16: (
        "实验在五个数据集上使用三个随机种子。匹配控制在相同条件下比较语义查询和结构查询。"
        "其他对照进一步检验关系类型、语义视图、结构上下文以及不读取邻居的晚期融合。"
    ),
    18: "2.1 共享表示方法",
    19: (
        "MTransE 学习不同语言嵌入空间之间的映射[2]。JAPE 在共享结构空间中加入属性关联[3]。"
        "这类方法为图谱整合提供了清晰的几何解释，但简单映射难以充分描述复杂的多关系邻域。"
    ),
    20: "2.2 关系感知结构编码",
    21: (
        "GraphSAGE 通过聚合局部邻居学习节点表示[9]。R-GCN 在消息变换中保留关系类型，并通过共享基或分块结构控制参数规模[7]。"
        "这些研究说明，图编码器需要同时保留邻接实体和边类型。"
    ),
    22: (
        "RDGCN 通过实体图与关系对偶图引入关系证据[5]。RREA 使用关系反射变换保留关系差异[6]。"
        "RPR-RHGT 使用图 Transformer 编码筛选后的多步关系路径[10]。这些模型说明带类型和多跳结构具有价值。"
        "本文采用更轻量的消息传递设计，重点研究一跳结构上下文如何进入联合表示。"
    ),
    23: "2.3 结构—语义交互",
    24: (
        "Transformer 可以建立序列中不同位置之间的联系[11]。层次注意力网络表明，不同文本粒度可能提供不同证据[12]。"
        "因此，本文的语义编码器使用 token、phrase 和 global 三种视图，并将这些视图用于主要的邻居选择机制。"
    ),
    25: (
        "文本增强实体对齐研究表明，语义信息能够补充图结构。RREA 分别报告文本增强和纯结构设置[6]。"
        "MCLEA 通过对比目标对齐多种数据类型[13]。其他研究指出，简单融合可能混合尚未对齐的表示空间[14]。"
        "这些结果引出一个更直接的检验：本文保持邻域和模型规模不变，只改变邻居查询信号。"
    ),
    26: (
        "近期研究还使用大语言模型[15]、域适应[16]或逐步推理代理[17]。这些方法改变了监督来源或推理过程。"
        "本文则检验结构邻居与实体语义之间的局部交互。"
    ),
    35: "3.2 模型概述",
    36: (
        "模型包含四个部分。结构编码器保留关系类型，并组合不同传播深度。语义编码器从同一实体文本构造 token、phrase 和 global 三种视图。"
        "融合模块使用语义信息为全部一跳出邻居评分，形成稀疏结构上下文，并将该上下文与实体语义结合。"
        "训练阶段对联合表示和结构表示使用 InfoNCE，检索阶段对训练得到的联合表示使用 CSLS。图1给出了完整的信息流。"
    ),
    38: (
        "图1 所提出模型的信息流。结构分支保留关系类型并选择有效传播深度，语义分支从同一输入构造三种视图。"
        "语义查询为全部一跳出邻居评分，1.5-entmax 去除弱相关证据，模型随后将筛选后的结构上下文与语义信息结合。"
        "训练和检索均使用联合表示，验证集仅选择 {{EQ0}}。"
    ),
    78: "3.5 语义引导的结构上下文",
    79: (
        "该模块实现模型的主要交互过程。模块读取全部一跳出邻居，并为当前实体选择结构证据。"
        "计算分为三个阶段。语义查询首先为每个结构邻居评分，1.5-entmax 去除弱相关候选。"
        "模型随后将筛选后的邻居摘要与实体自身结构结合，最后使用门控将结构上下文与实体语义合成为联合表示。"
    ),
    98: "3.5.3 最终联合表示",
    103: (
        "语义查询在模型构造结构上下文之前改变邻居权重，最终门控再将该上下文与实体语义结合。"
        "模型结构本身不能证明语义查询更有效，因此第3.5.4节设置了条件相同的结构查询对照。"
    ),
    104: "3.5.4 匹配控制",
    105: (
        "结构查询对照用于分离查询信号的影响。该对照保留相同的邻域、动态补齐、1.5-entmax、门控和联合检索，只将语义查询替换为实体结构表示。"
        "晚期融合对照检验更广泛的变化。它保持相同的可训练融合参数量，但不读取邻居上下文，而是在两个分支独立编码后再形成联合表示。"
    ),
    108: "3.6 训练目标",
    130: "5.1 多数据集结果",
    131: (
        "表2报告五个数据集上的联合表示检索结果。每个数据集均使用三个随机种子。"
        "DBP15K 从训练对中留出10%作为验证集，OpenEA 和 EventEA 使用官方验证集。"
    ),
    134: (
        "模型在五个数据集上的结果较为稳定。Hits@1 的最大样本标准差为0.0053。"
        "这一稳定性为后续控制实验提供了基础。"
    ),
    135: "5.2 与文献结果的相对位置",
    136: (
        "表3列出既有研究公开报告的结果。MTransE、JAPE、BootEA 和 RDGCN 的 DBP15K 数值取自 RDGCN[5]。"
        "RREA-text 数值取自 RREA[6]，OpenEA EN–FR-V2 数值为 OpenEA 官方五折平均结果[8]。"
    ),
    139: (
        "表3中的方法采用不同的数据划分、文本输入和后处理步骤，因此该表只用于说明相对位置，而不是同一实验协议下的严格排名。"
        "本文结论主要依据模型内部的受控比较。"
    ),
    141: "5.3 组件控制提供的证据",
    142: (
        "表4报告 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的组件控制结果。每个变体均使用三个随机种子，"
        "并在训练和检索阶段使用同一种表示。"
    ),
    144: (
        "关系类型提供了最清楚且方向一致的结构增益。移除关系类型后，DBP15K 和 OpenEA 的 Hits@1 分别下降3.62和2.78个百分点。"
        "拓扑特征也支持结构初始状态，相对仅使用实体向量分别提高3.46和8.03个百分点。phrase 视图在两个数据集上分别提高0.81和3.95个百分点。"
        "其余控制的变化较小或依赖数据集，因此本文将其视为辅助设计，而不单独列为主要贡献。"
    ),
    145: "5.4 结构上下文在联合表示中的价值",
    146: (
        "表5比较完整模型、两个单分支模型和无可训练参数的平均融合。每项结果均直接检索训练时使用的表示，"
        "验证过程不改变任一分支的权重。"
    ),
    148: (
        "完整模型相对语义单分支在 DBP15K 和 OpenEA 上的 Hits@1 分别提高9.09和9.66个百分点，也优于无可训练参数的平均融合。"
        "这些结果说明，结构证据在与语义信息进行可训练交互后能够增加有效信息。结构单分支并不是该增益的直接来源；"
        "增益出现在模型形成可训练联合表示之后。"
    ),
    149: "5.5 邻居查询信号的受控检验",
    150: (
        "表6比较完整模型与两种对照。结构查询对照只改变用于为邻居评分的信号。"
        "晚期融合对照保持相同的可训练融合参数量，但不读取邻居上下文。"
    ),
    152: (
        "匹配控制结果显示，结构查询的 Hits@1 比语义查询分别低2.08和2.48个百分点。"
        "该比较在当前实验协议下分离出语义邻居评分的作用。不读取邻居上下文的晚期融合比完整模型分别低5.28和9.22个百分点。"
        "第二组差距说明能够读取邻居的交互路径具有价值，但它不能单独度量融合时机，因为邻居可见性和训练路径同时发生了变化。"
    ),
    154: "6.1 实验能够支持的结论",
    155: (
        "实验支持三项结论。第一，关系类型在两个受控数据集上均提供稳定的结构增益。"
        "第二，结构上下文与语义信息共同学习时能够增加有效信息，完整模型优于语义单分支和简单平均融合。"
        "第三，在邻域和模型规模相同的条件下，语义查询比结构查询更有效地选择邻居。"
        "这些结果共同支持关系感知邻居上下文与语义引导选择。"
    ),
    156: (
        "增益大小取决于输入数据。DBP15K 提供较有用的名称和局部图线索。EventEA 包含多样的事件关系和属性[25]。"
        "OpenEA 使用编码后的实体标识以降低名称偏置[8]，而 GloVe 仅覆盖约52.18%的 OpenEA token[26]。"
        "当文本能够引导邻居选择，并且关系类型能够区分相似局部结构时，本方法更容易发挥作用。"
    ),
    157: "6.2 证据范围",
    158: (
        "本文使用 {{EQ0}} 个随机种子，因此结果报告效应大小和样本波动，不作较强的显著性声明。"
        "固定邻居对照同时改变了邻居数量和权重函数，晚期融合对照同时改变了邻居访问和训练路径。"
        "组件控制仅覆盖两个主要数据集。因此，现有证据支持本文测试的实验条件，不能直接推广到所有实体对齐场景。"
    ),
    159: "6.3 适用条件与后续工作",
    160: (
        "后续研究需要检验三类更广泛的条件。首先，可使用多语言子词编码器替代未知 token 的随机向量[27–28]。"
        "其次，可在更大规模图谱上检验完整邻域选择的运行效率。最后，可在开放世界数据集上检验不存在对应实体的情况。"
        "这些实验将说明语义引导邻居选择在文本线索更弱和候选空间更大时是否仍然有效。"
    ),
    162: (
        "本文研究结构邻居应在何处进入结构—语义实体对齐。实验给出四项主要结果。关系类型在两个受控数据集上提高了结构证据的有效性。"
        "可训练联合表示使结构上下文在语义信息之外产生增益。在其他条件一致时，语义查询比结构查询更有效地选择邻居。"
        "能够读取完整邻居的交互路径也优于不读取邻居上下文的晚期融合。综上，跨语言实体对齐能够从关系感知邻居上下文中受益，"
        "而语义信息应在最终联合表示形成之前参与邻居选择。"
    ),
}


PLACEHOLDER_RE = re.compile(r"\{\{EQ(\d+)\}\}")


def combined_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).replace("\u00a0", " ")


def equation_children(paragraph: Paragraph):
    return [
        copy.deepcopy(child)
        for child in paragraph._p.iterchildren()
        if child.tag in {qn("m:oMath"), qn("m:oMathPara")}
    ]


def set_run_font(run, size: float, *, bold: bool | None = None, italic: bool | None = None) -> None:
    family = "STSong" if re.search(r"[\u3400-\u9fff]", run.text or "") else "Times New Roman"
    run.font.name = family
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), family)
    fonts.set(qn("w:hAnsi"), family)
    fonts.set(qn("w:eastAsia"), family)
    fonts.set(qn("w:cs"), family)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def append_text_run(paragraph: Paragraph, text: str) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    set_run_font(run, 12)


def set_plain_text(paragraph: Paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    append_text_run(paragraph, text)


def set_paragraph_markup(paragraph: Paragraph, markup: str) -> None:
    equations = equation_children(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    cursor = 0
    for match in PLACEHOLDER_RE.finditer(markup):
        append_text_run(paragraph, markup[cursor : match.start()])
        index = int(match.group(1))
        if index >= len(equations):
            raise RuntimeError(
                f"EQ{index} exceeds {len(equations)} equations in paragraph: {markup!r}"
            )
        paragraph._p.append(copy.deepcopy(equations[index]))
        cursor = match.end()
    append_text_run(paragraph, markup[cursor:])


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = style
    set_plain_text(inserted, text)
    return inserted


def set_style_font(
    style,
    size: float,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    family: str = "STSong",
) -> None:
    style.font.name = family
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), family)
    fonts.set(qn("w:hAnsi"), family)
    fonts.set(qn("w:eastAsia"), family)
    fonts.set(qn("w:cs"), family)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def remove_line_numbering(document: Document) -> None:
    for section in document.sections:
        section_pr = section._sectPr
        for node in list(section_pr.findall(qn("w:lnNumType"))):
            section_pr.remove(node)


def add_page_number(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instruction, separate, end])
    set_run_font(run, 10)


def apply_chinese_overrides(document: Document) -> None:
    if len(document.paragraphs) != 196:
        raise RuntimeError(f"Expected 196 source paragraphs, found {len(document.paragraphs)}")
    for index, markup in ZH_OVERRIDES.items():
        paragraph = document.paragraphs[index]
        if "{{EQ" in markup:
            set_paragraph_markup(paragraph, markup)
        else:
            set_plain_text(paragraph, markup)

    affiliation = insert_after(document.paragraphs[1], "作者单位：[院系，学校，城市，国家]")
    affiliation.paragraph_format.keep_with_next = True

    declarations = next(p for p in document.paragraphs if combined_text(p).strip() == "声明")
    anchor = insert_after(
        declarations,
        "基金项目  [请在投稿前填写准确的基金声明；如未获得外部资助，可写为：‘本研究未获得专项资助。’]",
    )
    insert_after(
        anchor,
        "作者贡献  林心怡负责研究设计、方法与软件实现、实验与分析、数据与图表整理，以及论文撰写和修改。",
    )

    conflict = next(p for p in document.paragraphs if combined_text(p).startswith("利益冲突"))
    anchor = insert_after(conflict, "伦理审批  不适用。")
    anchor = insert_after(anchor, "参与同意  不适用。")
    insert_after(anchor, "发表同意  不适用。")


def format_chinese_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    remove_line_numbering(document)

    styles = document.styles
    set_style_font(styles["Normal"], 12)
    normal = styles["Normal"].paragraph_format
    normal.line_spacing = 2.0
    normal.space_before = Pt(0)
    normal.space_after = Pt(0)
    normal.widow_control = True

    set_style_font(styles["Heading 1"], 14, bold=True)
    h1 = styles["Heading 1"].paragraph_format
    h1.line_spacing = 1.0
    h1.space_before = Pt(14)
    h1.space_after = Pt(6)
    h1.keep_with_next = True
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT

    set_style_font(styles["Heading 2"], 12, bold=True)
    h2 = styles["Heading 2"].paragraph_format
    h2.line_spacing = 1.0
    h2.space_before = Pt(12)
    h2.space_after = Pt(4)
    h2.keep_with_next = True
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT

    set_style_font(styles["Heading 3"], 12, bold=True)
    h3 = styles["Heading 3"].paragraph_format
    h3.line_spacing = 1.0
    h3.space_before = Pt(10)
    h3.space_after = Pt(3)
    h3.keep_with_next = True

    if "Caption" in styles:
        set_style_font(styles["Caption"], 10)
        caption = styles["Caption"].paragraph_format
        caption.line_spacing = 1.0
        caption.space_before = Pt(6)
        caption.space_after = Pt(6)
        caption.keep_together = True
        caption.keep_with_next = True
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "Formula" in styles:
        set_style_font(styles["Formula"], 11, family="Times New Roman")
        formula = styles["Formula"].paragraph_format
        formula.line_spacing = 1.0
        formula.space_before = Pt(4)
        formula.space_after = Pt(4)
        formula.keep_together = True
        formula.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "Reference" in styles:
        set_style_font(styles["Reference"], 11, family="Times New Roman")
        reference = styles["Reference"].paragraph_format
        reference.line_spacing = 1.5
        reference.left_indent = Mm(6)
        reference.first_line_indent = Mm(-6)
        reference.space_after = Pt(3)

    title = document.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    for run in title.runs:
        set_run_font(run, 17, bold=True)

    author = document.paragraphs[1]
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.line_spacing = 1.0
    author.paragraph_format.space_after = Pt(4)
    author.paragraph_format.keep_with_next = True
    for run in author.runs:
        set_run_font(run, 12)

    for paragraph in document.paragraphs[2:5]:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 10.5)

    abstract = next(p for p in document.paragraphs if combined_text(p).startswith("跨语言知识图谱实体对齐旨在识别"))
    abstract.paragraph_format.line_spacing = 1.15
    abstract.paragraph_format.space_after = Pt(6)
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in abstract.runs:
        set_run_font(run, 10.5)

    keywords = next(p for p in document.paragraphs if combined_text(p).startswith("关键词："))
    keywords.paragraph_format.line_spacing = 1.0
    keywords.paragraph_format.space_after = Pt(10)
    for run in keywords.runs:
        set_run_font(run, 10.5)

    excluded = {id(title), id(author), id(abstract), id(keywords)}
    for paragraph in document.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3", "Caption", "Formula", "Reference"}:
            continue
        if id(paragraph) in excluded or not combined_text(paragraph).strip():
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, 12)

    for table in document.tables:
        table.autofit = False
        header_props = table.rows[0]._tr.get_or_add_trPr()
        if header_props.find(qn("w:tblHeader")) is None:
            header_props.append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_mar = tc_pr.first_child_found_in("w:tcMar")
                if tc_mar is None:
                    tc_mar = OxmlElement("w:tcMar")
                    tc_pr.append(tc_mar)
                for side in ("top", "left", "bottom", "right"):
                    node = tc_mar.find(qn(f"w:{side}"))
                    if node is None:
                        node = OxmlElement(f"w:{side}")
                        tc_mar.append(node)
                    node.set(qn("w:w"), "90")
                    node.set(qn("w:type"), "dxa")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER if len(cell.text) < 30 else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        set_run_font(run, 8.5, bold=(row_index == 0))

    shape = document.inline_shapes[0]
    shape.width = Inches(6.35)
    shape.height = Inches(6.35 * 1500 / 2400)
    props = shape._inline.xpath(".//*[local-name()='docPr']")
    if props:
        props[0].set("name", "图1")
        props[0].set(
            "descr",
            "关系感知结构编码、多尺度语义编码、稀疏出邻居加权、训练目标与CSLS检索的信息流。",
        )

    header = section.header.paragraphs[0]
    set_plain_text(header, CHINESE_RUNNING_TITLE)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.line_spacing = 1.0
    for run in header.runs:
        set_run_font(run, 9)

    footer = section.footer.paragraphs[0]
    for child in list(footer._p):
        if child.tag != qn("w:pPr"):
            footer._p.remove(child)
    add_page_number(footer)


def equation_count(document: Document) -> int:
    return sum(1 for node in document.element.body.iter() if node.tag == qn("m:oMath"))


def has_line_numbering(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        return b"<w:lnNumType" in archive.read("word/document.xml")


def verify_english(path: Path) -> None:
    document = Document(path)
    text = "\n".join(combined_text(p) for p in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    required = [
        "Relation-Aware Neighbor Context and Semantic-Guided Selection",
        "0.7077 ± 0.0023 / 0.7497 ± 0.0016",
        "structural queries yield Hits@1 values 2.08 and 2.48 points below semantic queries",
    ]
    missing = [value for value in required if value not in text]
    if missing:
        raise RuntimeError(f"English output is missing required content: {missing}")
    if has_line_numbering(path):
        raise RuntimeError("English output still contains line numbering")
    if len(document.tables) != 6 or len(document.inline_shapes) != 1 or equation_count(document) != 154:
        raise RuntimeError("English table, figure, or equation count changed")


def verify_chinese(path: Path) -> None:
    document = Document(path)
    text = "\n".join(combined_text(p) for p in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    required = [
        CHINESE_TITLE,
        "语义查询相对结构查询使 DBP15K 和 OpenEA 的 Hits@1 分别提高 2.08 和 2.48 个百分点",
        "0.7077 ± 0.0023 / 0.7497 ± 0.0016",
        "关系类型提供了最清楚且方向一致的结构增益",
    ]
    missing = [value for value in required if value not in text]
    if missing:
        raise RuntimeError(f"Chinese output is missing required content: {missing}")
    if has_line_numbering(path):
        raise RuntimeError("Chinese output still contains line numbering")
    if len(document.tables) != 6 or len(document.inline_shapes) != 1 or equation_count(document) != 154:
        raise RuntimeError("Chinese table, figure, or equation count changed")
    if "0.7277 ± 0.0023 / 0.7697 ± 0.0016" in text:
        raise RuntimeError("Outdated structural-query result remains in Chinese output")


def build_english() -> None:
    shutil.copy2(ENGLISH_SOURCE, ENGLISH_OUTPUT)
    document = Document(ENGLISH_OUTPUT)
    remove_line_numbering(document)
    document.save(ENGLISH_OUTPUT)
    verify_english(ENGLISH_OUTPUT)


def build_chinese() -> None:
    shutil.copy2(CHINESE_SOURCE, CHINESE_OUTPUT)
    document = Document(CHINESE_OUTPUT)
    apply_chinese_overrides(document)
    format_chinese_document(document)
    document.core_properties.title = CHINESE_TITLE
    document.core_properties.author = "林心怡"
    document.core_properties.last_modified_by = "林心怡"
    document.core_properties.subject = "跨语言知识图谱实体对齐"
    document.core_properties.keywords = (
        "实体对齐；跨语言知识图谱；关系感知图神经网络；语义引导邻居选择；稀疏结构上下文"
    )
    document.save(CHINESE_OUTPUT)
    verify_chinese(CHINESE_OUTPUT)


def main() -> None:
    if not ENGLISH_SOURCE.exists() or not CHINESE_SOURCE.exists():
        raise FileNotFoundError("A required source manuscript is missing")
    build_english()
    build_chinese()
    print(ENGLISH_OUTPUT)
    print(CHINESE_OUTPUT)


if __name__ == "__main__":
    main()
