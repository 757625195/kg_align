from __future__ import annotations

import argparse
import html.entities
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


ROOT = Path(__file__).resolve().parents[1]
MATH_DEPS = ROOT / "tmp" / "docx_math_deps"
sys.path.insert(0, str(MATH_DEPS))

import mathml2omml  # noqa: E402
from latex2mathml.converter import convert as latex_to_mathml  # noqa: E402


def latex_to_omml(expression: str):
    mathml = latex_to_mathml(expression.strip(), display="inline")
    omml = mathml2omml.convert(mathml, html.entities.name2codepoint)
    omml = re.sub(
        r"(<m:groupChr><m:groupChrPr>.*?)</m:groupChr>(<m:e>)",
        r"\1</m:groupChrPr>\2",
        omml,
    )
    omml = omml.replace("<m:oMath>", f"<m:oMath {nsdecls('m')}>", 1)
    return parse_xml(omml)


def clear_paragraph(paragraph) -> None:
    properties = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)


def add_math(paragraph, expression: str) -> None:
    paragraph._p.append(latex_to_omml(expression))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    source = Path(args.source)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)

    document = Document(output)
    headings = [
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "1.1 研究背景与问题缺口"
    ]
    if len(headings) != 1:
        raise ValueError(f"Expected one section 1.1 heading, found {len(headings)}")

    paragraph = document.paragraphs[headings[0] + 1]
    remainder_marker = "不同机构"
    if remainder_marker not in paragraph.text:
        raise ValueError("Could not locate the unchanged remainder of section 1.1")
    remainder = paragraph.text[paragraph.text.index(remainder_marker) :]
    clear_paragraph(paragraph)
    paragraph.add_run("知识图谱通常使用三元组 ")
    add_math(paragraph, r"(h,r,t)")
    paragraph.add_run(" 表示关系事实，其中 ")
    add_math(paragraph, "h")
    paragraph.add_run("、")
    add_math(paragraph, "r")
    paragraph.add_run(" 和 ")
    add_math(paragraph, "t")
    paragraph.add_run(" 分别表示头实体、关系和尾实体；该三元组表示头实体 ")
    add_math(paragraph, "h")
    paragraph.add_run(" 通过关系 ")
    add_math(paragraph, "r")
    paragraph.add_run(" 指向尾实体 ")
    add_math(paragraph, "t")
    paragraph.add_run(f"[20]。{remainder}")

    if any(p.text.strip().startswith("[20]") for p in document.paragraphs):
        raise ValueError("Reference [20] already exists")
    reference = document.add_paragraph(style="Reference")
    reference.add_run(
        "[20] Bordes A, Usunier N, Garcia-Duran A, Weston J, Yakhnenko O "
        "(2013) Translating embeddings for modeling multi-relational data. "
        "In: Advances in Neural Information Processing Systems 26, pp 2787–2795"
    )

    document.core_properties.subject = (
        "教师要求逐项核验、统一协议补充消融与三元组基础来源版"
    )
    document.save(output)

    verified = Document(output)
    if not any(p.text.strip().startswith("[20] Bordes A") for p in verified.paragraphs):
        raise RuntimeError("Reference [20] was not saved")
    if not any(
        p.text.strip() == "给定两个知识图谱：" for p in verified.paragraphs
    ):
        raise RuntimeError("Section 3.1 task definition was unexpectedly modified")
    print(f"WROTE {output}")


if __name__ == "__main__":
    main()
