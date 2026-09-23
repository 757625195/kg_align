from __future__ import annotations

import copy
import hashlib
import os
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "实验结论与逻辑一致性修订版_20260830.docx"
)
OUTPUT = ROOT / "outputs" / (
    "Cross_Lingual_Knowledge_Graph_Entity_Alignment_"
    "Springer_English_Submission_20260830.docx"
)
FIGURE = ROOT / "outputs" / "paper_assets" / "figure_1_current_protocol_springer_en.png"
EXPECTED_SOURCE_SHA256 = "8c5b8e45d143b774d26f9081a2beb2d526d08d19aaf98ed56cbf8c69f166c5b3"


P = {
    0: "Cross-Lingual Knowledge Graph Entity Alignment with Relation-Aware Structural Context and Multi-Scale Semantic Fusion",
    1: "Xinyi Lin",
    2: "Corresponding author: Xinyi Lin, auroral020417@gmail.com",
    3: "Abstract",
    4: (
        "Cross-lingual knowledge graph entity alignment identifies entities that refer to the same real-world object in knowledge "
        "graphs expressed in different languages. Existing methods use graph structure and textual semantics, but relation types, "
        "propagation depths, and fragmented entity text are often modeled separately, and the role of structural neighbors in a joint "
        "representation remains insufficiently controlled. This study develops a model that combines relation-aware structural context "
        "with multi-scale semantics. The structural encoder uses shared relation projections for message passing and a node-wise selector "
        "to combine propagation depths. The semantic encoder derives token, phrase, and global views from the same entity text. The fusion "
        "module applies 1.5-entmax to select sparse evidence from the complete one-hop outgoing neighborhood and jointly encodes this context "
        "with the semantic representation. Training uses bidirectional InfoNCE, and cross-graph retrieval uses CSLS. Across five datasets "
        "and three random seeds, Hits@1 ranges from 0.6344 to 0.9239, with sample standard deviations no greater than 0.0053. On DBP15K "
        "ZH-EN and OpenEA EN-FR-15K-V2, removing relation types reduces Hits@1 by 3.62 and 2.78 percentage points, while semantic-only "
        "retrieval reduces it by 9.09 and 9.66 points. With the neighborhood and fusion capacity held constant, structural queries trail "
        "semantic queries by 2.08 and 2.48 points. Parameter-matched late fusion without neighbor context trails the full model by 5.28 "
        "and 9.22 points. These results support retaining relation-aware neighborhood evidence in the joint representation, but the late-"
        "fusion gap cannot be attributed to fusion timing alone."
    ),
    5: (
        "Keywords: knowledge graph entity alignment; cross-lingual knowledge graphs; relation-aware graph neural networks; "
        "structure-semantic fusion; sparse neighbor selection; multi-scale semantic encoding"
    ),
    6: "1 Introduction",
    7: "1.1 Background and research gap",
    8: (
        "Knowledge graphs commonly represent factual relations as triples {{EQ0}}, where {{EQ1}}, {{EQ2}}, and {{EQ3}} denote the head "
        "entity, relation, and tail entity, respectively. The triple states that head entity {{EQ4}} is linked to tail entity {{EQ6}} by "
        "relation {{EQ5}} [1]. Knowledge graphs are often constructed independently across institutions, languages, and information "
        "systems; consequently, the same real-world object may have different identifiers, names, attributes, and value domains. Entity "
        "alignment seeks equivalence relations between entity sets {{EQ7}} and {{EQ8}}, thereby providing a basis for graph integration."
    ),
    9: (
        "Existing methods broadly follow three lines. MTransE and JAPE embed different graphs into mappable or shared vector spaces [2, 3] "
        "and draw equivalent entities closer together. BootEA iteratively selects high-confidence candidates and adds them to the "
        "supervision set as pseudo-labels [4]. RDGCN and RREA explicitly use relations and neighborhood topology to strengthen structural "
        "representations [5, 6]. These studies established the main foundations of cross-lingual entity alignment, yet effective interaction "
        "between structural and textual evidence in a unified model still faces the following difficulties."
    ),
    10: (
        "First, structural similarity does not imply entity equivalence. For example, person A may be linked to Beijing by a birthplace "
        "relation, whereas person B is linked to Beijing by a workplace relation. If an encoder collapses both relations into untyped edges, "
        "it may overestimate the similarity between A and B [5, 7]. Second, semantic evidence is often fragmented. The token 'Washington' "
        "alone cannot distinguish a person, a state, and a city; relation or attribute context is required. Third, seed alignments available "
        "for training usually cover only part of the entity sets [4, 8], so the model must learn transferable structural and semantic cues "
        "from a limited set of known pairs."
    ),
    11: "1.2 Research objectives and completed work",
    12: (
        "The experiments address three questions: how accurately and consistently the model retrieves entities across five datasets; how "
        "relation-aware structural encoding, multi-scale semantic encoding, and neighborhood aggregation affect performance; and how semantic "
        "and structural neighbor queries differ when the neighborhood and parameter count are held constant, together with the change caused "
        "by removing neighbor context in a late-fusion baseline."
    ),
    13: "To address these questions, this study makes the following contributions.",
    14: (
        "A relation-aware structure-semantic entity alignment model is developed. The structural encoder distinguishes edge types and combines "
        "node representations from different propagation depths. The semantic encoder derives word-level, local phrase, and global contextual "
        "information from the same entity text. Both representations are learned under a common alignment objective."
    ),
    15: (
        "The complete one-hop structural neighborhood is introduced before the joint representation is formed. Sparse weights are assigned "
        "according to the relevance of each neighbor to the current entity, after which the neighbor summary, self-structure, and semantics are "
        "combined. Structural-query and parameter-matched late-fusion controls are included to identify the sources of performance differences."
    ),
    16: (
        "The model is evaluated in three independent runs on five entity alignment datasets. Component ablations and controlled comparisons on "
        "two representative datasets quantify the practical effects and empirical limits of the structural, semantic, and neighborhood mechanisms."
    ),
    17: "2 Related work",
    18: "2.1 Knowledge graph embeddings and shared representation spaces",
    19: (
        "MTransE represents entities and relations from language-specific knowledge graphs as vectors and learns mappings between their embedding "
        "spaces, bringing cross-lingual entities that refer to the same object closer in a common space [2]. JAPE extends a shared structural space "
        "with attribute correlations, allowing entity names, attributes, and graph structure to contribute jointly to matching; it also introduced "
        "the widely used DBP15K dataset [3]. Vector translations and linear mappings, however, have limited capacity to represent one-to-many, "
        "multi-relational, and structurally complex neighborhoods and may distort distances in the original spaces."
    ),
    20: "2.2 Relation modeling with graph neural networks",
    21: (
        "GraphSAGE learns node representations by sampling and aggregating local neighbors [9]. R-GCN additionally distinguishes relation types in "
        "message transformations and controls multi-relational parameter growth through basis or block decomposition [7]. These studies show that "
        "neighborhood aggregation must retain edge types as well as adjacent entities. Because relation frequencies and the structural scopes required "
        "by individual nodes vary, representing relation differences and multiple propagation depths under a limited parameter budget remains an open "
        "design problem."
    ),
    22: (
        "For entity alignment, RDGCN introduces relation evidence through attention between an entity graph and its relation dual graph [5]. RREA "
        "uses relational reflection transformations to retain both relation discrimination and the geometry of the representation space [6]. RPR-RHGT "
        "explicitly generates and filters multi-step relation paths and encodes them with a relation-aware heterogeneous graph Transformer [10]. Together, "
        "these studies demonstrate the value of relation and multi-hop structural information."
    ),
    23: "2.3 Semantic modeling and joint representations",
    24: (
        "Transformers use multi-head self-attention to model dependencies between sequence positions [11]. Hierarchical attention networks further show "
        "that semantic units at different granularities may contribute differently to a representation [12]. These findings motivate the preservation of "
        "fine-grained cues, local combinations, and global context, but they do not guarantee that every semantic view improves entity alignment. The "
        "present model therefore retains three granularities in a common encoder and tests their contributions through separate ablations."
    ),
    25: (
        "Text-enhanced entity alignment research shows that names, descriptions, and other modalities can complement structural evidence. RREA reports "
        "text-enhanced and structure-only settings separately [6]. MCLEA learns multimodal representations through intra-modal contrastive learning and "
        "cross-modal alignment [13]. A multimodal knowledge graph Transformer study further observes that direct concatenation or attention-based mixing "
        "can create heterogeneous but unaligned representation spaces [14]. These studies support combining structure and semantics, but simultaneous "
        "changes to architectures and objectives prevent the effect of the interaction position from being inferred in isolation."
    ),
    26: (
        "Recent work also explores supervision from large language models under noisy annotations [15], multi-source domain adaptation for real-world "
        "knowledge graphs [16], and entity alignment agents with structured multi-step reasoning [17]. These approaches alter the source of supervision "
        "or the reasoning procedure."
    ),
    27: "3 Problem formulation and method",
    28: "3.1 Entity alignment task",
    29: "Let the two knowledge graphs be",
    31: (
        "where {{EQ0}}, {{EQ1}}, and {{EQ2}} are the entity, relation, and triple sets on side {{EQ3}}, respectively. A small set of seed alignments is known:"
    ),
    33: (
        "Here, {{EQ0}} is the seed alignment set containing equivalent entity pairs confirmed by human annotation or dataset rules, and {{EQ1}} means "
        "that the two entities refer to the same real-world object."
    ),
    34: (
        "The objective is to learn a scoring function {{EQ0}} that ranks the true counterpart above other candidates. Evaluation follows a one-to-one, "
        "closed-world setting."
    ),
    35: "3.2 Overall architecture and interaction mechanism",
    36: (
        "The model comprises four components: relation-aware structural encoding, multi-scale semantic encoding, one-hop outgoing-neighborhood selection "
        "and structure-semantic fusion, and contrastive learning. The structural encoder performs relation-aware message passing on both knowledge graphs "
        "and combines different propagation depths. The semantic encoder constructs token, phrase, and global views from sequences of names, relations, "
        "and attributes. The fusion module sparsely selects relevant evidence from an untruncated one-hop outgoing neighborhood and combines it with the "
        "semantic representation to form a joint representation. Training constrains both the joint and structural representations, and retrieval uses "
        "CSLS for cross-graph matching. Figure 1 links the information flow to the equations."
    ),
    38: (
        "Fig. 1 Relation-aware structural context and multi-scale semantic fusion. The structural branch propagates messages along incoming edges and "
        "selects among propagation depths. The semantic branch constructs token, phrase, and global views from one shared input. The fusion module reads "
        "every one-hop outgoing neighbor along the original edge direction and assigns sparse weights with 1.5-entmax. The resulting structural context "
        "and semantic representation form the joint representation. Both training and primary retrieval use the joint representation; validation selects "
        "only {{EQ0}}."
    ),
    39: "3.3 Relation-aware structural encoder",
    40: (
        "Local degree and neighborhood statistics are commonly used to characterize structural roles, while structural-identity methods compare nodes by "
        "topological function [18, 19]. Motivated by these approaches, the structural branch is initialized from a set of local topology features; the "
        "particular feature combination used here is designed for this study."
    ),
    41: (
        "The local topology vector has eight dimensions. It is computed separately within each graph and does not require relation identifiers to agree "
        "across graphs. For entity {{EQ0}}, let {{EQ1}}, {{EQ2}}, and {{EQ3}} denote in-degree, out-degree, and total degree. The variables {{EQ4}} and {{EQ5}} "
        "are the numbers of distinct incoming and outgoing relation types. The quantities {{EQ6}} and {{EQ7}} are the means of {{EQ8}} over incoming source "
        "entities and outgoing target entities, respectively; a missing neighborhood is assigned a mean of zero. Directional balance is defined as {{EQ9}}. "
        "The raw feature vector is"
    ),
    43: (
        "The local topology features are standardized dimension by dimension within the left and right knowledge graphs, producing {{EQ0}}. The initial "
        "structural state combines a shared topology projection with a small entity-specific residual:"
    ),
    45: (
        "Here, {{EQ0}} is a linear projection shared by both knowledge graphs, and {{EQ1}} is a trainable entity vector. The topology term gives entities "
        "with similar local roles a comparable starting point, while the residual scaled by 0.1 prevents distinct entities with similar local statistics "
        "from collapsing to the same representation."
    ),
    46: (
        "At each forward pass, the structural encoder performs directed, relation-aware message passing over the complete knowledge graph. For triple "
        "{{EQ0}}, head entity {{EQ1}} sends a message to tail entity {{EQ2}} along the original edge direction. The encoder aggregates incoming edges only. "
        "Let {{EQ3}} be the state of entity {{EQ4}} at layer {{EQ5}}, {{EQ6}} the relation embedding, and {{EQ7}} the set of incoming edges to {{EQ8}}. Layer "
        "{{EQ9}} first computes"
    ),
    48: (
        "The matrices {{EQ0}} and {{EQ1}} project the source-entity state and relation embedding, respectively. All relations at a layer share these two "
        "projection matrices, while {{EQ2}} encodes relation-specific differences. This avoids assigning a separate full matrix to every relation. The vector "
        "{{EQ3}} is the message sent to target entity {{EQ5}} along edge {{EQ4}}. Messages arriving at entity {{EQ6}} are averaged:"
    ),
    50: (
        "The self-state is first projected as {{EQ0}}. An element-wise gate is then computed from the self-state and the mean neighbor message:"
    ),
    53: (
        "In Eq. (3), {{EQ0}} assigns a value between zero and one to each representation dimension. Equation (4) retains the projected self-state and adds "
        "neighbor evidence according to this gate. Consider the triple 'Yao Ming - birthplace -> Shanghai'. When Shanghai is updated, the message contains "
        "both Yao Ming's entity state and the embedding of the birthplace relation. The encoder can therefore distinguish this edge from another edge to "
        "Shanghai labeled workplace."
    ),
    54: "Equation (4) produces the raw output {{EQ1}} of layer {{EQ0}}. Except at the final layer, the output is transformed before entering the next layer:",
    56: (
        "Layer normalization, a nonlinear transformation, and dropout are applied after each propagation layer to stabilize feature distributions at "
        "different depths and reduce overfitting."
    ),
    57: (
        "Stacked propagation allows entities to accumulate evidence from progressively more distant neighborhoods. For example, in 'Yao Ming - birthplace "
        "-> Shanghai - country -> China', local evidence about Yao Ming can reach China through Shanghai."
    ),
    58: (
        "Using only the deepest layer may discard self-information or evidence from immediate neighbors. The encoder therefore retains the projected input "
        "state and each raw layer output, {{EQ0}}. The state {{EQ1}} contains no propagated neighbor information, whereas {{EQ2}} contains evidence from at "
        "most {{EQ3}} hops. The model first computes their mean context {{EQ4}}. A two-layer MLP then scores each candidate state after concatenation with "
        "the mean context, and a {{EQ5}} operation yields depth weights {{EQ6}}:"
    ),
    61: (
        "The layer-selection MLP has input dimension {{EQ0}} and hidden dimension {{EQ1}}, uses GELU and dropout, and outputs one scalar. Equation (7) "
        "weights states from depths {{EQ2}} through {{EQ3}} separately for each entity, followed by LayerNorm and L2 normalization to produce structural "
        "representation {{EQ4}}. Different entities can therefore use different mixtures of propagation depths."
    ),
    62: "3.4 Multi-scale semantic encoder",
    63: (
        "The three semantic views share the same entity text but use different receptive fields. The token view represents word-level cues, the phrase view "
        "captures local combinations, and the global view models dependencies across positions."
    ),
    64: (
        "For entity {{EQ0}}, its name, adjacent relation names, and attribute text are tokenized and concatenated into one sequence, and each token is mapped "
        "to a word vector. The phrase and global views use this same sequence and do not read additional descriptions, external sentences, or other corpora. "
        "A common projection, normalization, and positional encoding produce base representation {{EQ1}}; mask {{EQ2}} marks valid tokens."
    ),
    65: (
        "For example, the shared sequence for Yao Ming may contain the name 'Yao Ming', relation phrases such as 'birthplace' and 'team', and the attribute "
        "fragment 'height 2.29 m'. The token view summarizes individual cues such as 'Yao' and 'height'; the phrase view recognizes local combinations such "
        "as 'birthplace Shanghai' or 'team Rockets'; and the global view allows the name, relations, and distant attributes to refer to one another across "
        "the full sequence. The views differ in encoding range, not in data source."
    ),
    66: (
        "The token view directly summarizes the base sequence. Mean pooling retains information from all valid tokens, whereas attention pooling emphasizes "
        "tokens that are more discriminative for the current entity. Their average forms the word-level representation:"
    ),
    68: (
        "Attention weights are normalized over valid positions only, so padding tokens do not affect the entity representation. Combining mean and attention "
        "pooling retains broad lexical evidence while preventing a few high-weight tokens from dominating the representation."
    ),
    69: (
        "The phrase view aggregates adjacent tokens in the base sequence to capture local patterns in entity names, relation phrases, and attribute fragments:"
    ),
    72: (
        "The global view uses a Transformer encoder to model dependencies between distant tokens, allowing name, relation, and attribute cues to interact "
        "within the complete entity context:"
    ),
    75: "Because entities may depend on semantic evidence at different granularities, the model estimates view weights from the joint state of all three views:",
    77: (
        "The three weighted view vectors are concatenated and passed to an output MLP. Its output is added to a residual projection of the sum of the weighted "
        "views, followed by L2 normalization, yielding semantic representation {{EQ0}}."
    ),
    78: "3.5 Neighbor context and joint representation",
    79: (
        "This component extracts structural evidence relevant to the current entity from an untruncated one-hop outgoing neighborhood and combines it with "
        "the entity's semantic representation. Computation has three stages. A semantic query first evaluates the relevance of structural neighbors, and "
        "1.5-entmax assigns exact zero weights to low-relevance candidates. The weighted neighbor summary is then combined with the entity's self-structure "
        "to form structural context. Finally, an element-wise gate combines structural context with semantics to produce the joint representation."
    ),
    80: "3.5.1 Semantically guided neighbor weighting",
    81: (
        "For entity {{EQ0}}, the fusion module receives semantic representation {{EQ1}}, self structural representation {{EQ2}}, the structural states of "
        "all one-hop outgoing neighbors, and a validity mask. Projecting the semantic representation produces query {{EQ3}}. The {{EQ4}}th structural "
        "neighbor is projected into key {{EQ5}} and value {{EQ6}}. Their scaled dot product gives an initial relevance score:"
    ),
    83: (
        "Equation (14) measures the match between the entity semantics and structural neighbor {{EQ0}}. A larger dot product indicates greater semantic "
        "compatibility. The scale factor controls the numerical range of high-dimensional dot products, where {{EQ1}} is the query and key dimension."
    ),
    84: "A dot product alone may miss finer element-wise relations between vectors. The interaction feature is therefore defined as",
    86: (
        "The concatenation {{EQ0}} preserves both input vectors, {{EQ1}} captures co-activation in corresponding dimensions, and {{EQ2}} represents "
        "element-wise differences. A learnable mapping followed by a sigmoid produces a neighbor gate. This combination of original vectors, element-wise "
        "products, and absolute differences follows matching features used in natural language inference [20]. It is not itself the final similarity score; "
        "rather, it supplies features from which the model estimates neighbor compatibility."
    ),
    88: (
        "Gate value {{EQ0}} indicates the extent to which neighbor {{EQ1}} passes this feature-consistency check. The scaled dot-product score and gate value "
        "jointly determine normalized neighbor weights:"
    ),
    90: (
        "In Eq. (17), {{EQ0}} is the vector of scores for all valid neighbors, {{EQ1}} contains the gate values, and mask bias {{EQ2}} excludes padded "
        "positions. Unlike softmax, which assigns a positive weight to every valid position, 1.5-entmax [21] can set low-scoring neighbors exactly to zero. "
        "Temperature {{EQ3}} controls concentration. Valid weights still sum to one; an entity with no valid neighbor receives an all-zero summary."
    ),
    91: (
        "Again consider Yao Ming. If the one-hop outgoing neighbors include the Houston Rockets, Shanghai, and a low-relevance candidate, semantic cues such "
        "as 'basketball player' and 'team' may increase the Rockets score. The 1.5-entmax transformation can assign zero weight to the unrelated candidate, "
        "after which Eq. (19) decides how much self-structure and weighted neighbor evidence to retain."
    ),
    92: "3.5.2 Structural context construction",
    93: (
        "The sparse weights from Eq. (17) are used to average neighbor value vectors, producing structural neighbor summary {{EQ1}} for entity {{EQ0}}. "
        "Because the full model does not truncate the one-hop outgoing neighborhood before weighting, the summary can use the complete local neighborhood "
        "while sparse weights suppress irrelevant evidence."
    ),
    94: "An element-wise context gate is then computed from the self structural representation, neighbor summary, and semantic representation:",
    97: (
        "Gate vector {{EQ0}} determines, dimension by dimension, whether to retain more of the entity's self-structure or use the neighbor summary. A value "
        "near one makes Eq. (19) favor {{EQ1}}, whereas a value near zero favors {{EQ2}}. If no valid neighbor exists, the model sets {{EQ3}}, preventing "
        "padding from changing the entity representation."
    ),
    98: "3.5.3 Joint structural-semantic representation",
    99: (
        "After obtaining structural context {{EQ0}}, the model computes a joint gate that determines the dimension-wise contributions of structural and "
        "semantic evidence:"
    ),
    102: (
        "Equation (21) first yields gated vector {{EQ0}}. Dimensions with a joint-gate value near one use more structural context, whereas values near zero "
        "use more semantic evidence. Consistent with the implementation, the full model then computes the equal-weight base vector {{EQ1}} and outputs "
        "{{EQ2}}. This fixed residual limits the displacement introduced by the gate without adding trainable parameters."
    ),
    103: (
        "Thus, the semantic query affects neighbor weights before the neighbor summary is formed; the resulting structural context is then combined with the "
        "semantic representation. This mechanism is termed semantic neighbor querying. To avoid treating architectural position as proof of effectiveness, "
        "Sect. 3.5.4 introduces a structural-query control with the same neighborhood, in which semantics enters only at the final joint gate."
    ),
    104: "3.5.4 Controls for training-time fusion and neighbor visibility",
    105: (
        "Two controls compare the neighbor query signal and the effect of removing neighbor context. The structural-query control retains the same one-hop "
        "outgoing neighborhood, dynamic padding, 1.5-entmax, and gating parameters but computes neighbor weights from the entity's structural representation; "
        "semantics enters only after structural context has been formed. The parameter-matched late-fusion control does not access neighbor context. It forms "
        "a training representation after the structural and semantic branches have been encoded independently:"
    ),
    107: (
        "Equation (22) separately normalizes the structural and semantic representations and then uses concatenation and a nonlinear mapping to obtain "
        "{{EQ0}} for InfoNCE training and retrieval. This baseline has the same number of trainable fusion parameters as the full model."
    ),
    108: "3.6 Bidirectional InfoNCE for the joint and structural branches",
    109: (
        "The model uses single-stage bidirectional InfoNCE training [22] without a separate warm-up stage. Each mini-batch samples {{EQ0}} entity pairs from "
        "the seed alignment set. The same samples are used to compute joint-representation and structural-branch losses, which are optimized together in one "
        "forward and backward pass. The temperature-scaled cosine similarity between left-graph entity {{EQ1}} and right-graph entity {{EQ2}} is"
    ),
    111: (
        "Here, {{EQ0}} is the temperature. Correct pairs lie on the diagonal of the similarity matrix. Cross-entropy is computed for both left-to-right and "
        "right-to-left retrieval and then averaged to obtain joint-representation loss {{EQ1}}:"
    ),
    113: (
        "Equation (24) directly constrains the joint representation. To establish a cross-graph coordinate system for the topology-initialized structural "
        "branch early in training, the same bidirectional InfoNCE form is also applied to {{EQ0}}, giving structural loss {{EQ1}}. The total objective is"
    ),
    115: (
        "Structural supervision is used only early in training and decreases linearly with training progress. Let {{EQ0}} denote the current epoch and {{EQ1}} "
        "the total number of epochs. Normalized progress and the structural weight are defined as"
    ),
    117: "The structural auxiliary loss therefore has weight 0.1 in the first epoch and decays to zero in the final epoch.",
    118: "3.7 Retrieval with CSLS",
    119: (
        "Training optimizes the encoders with Eq. (25). Primary retrieval directly uses the trained joint representation {{EQ0}} and does not reweight the "
        "structural and semantic branches on validation data. Cosine similarities between joint representations are corrected with CSLS [23]:"
    ),
    121: (
        "The terms {{EQ0}} and {{EQ1}} are the mean cosine similarities from a query entity and a candidate entity to their respective {{EQ2}} nearest "
        "entities on the opposite side. For each training run, validation MRR alone selects {{EQ3}} from the candidate set, and the test set is evaluated once."
    ),
    122: "4 Experimental design",
    123: "4.1 Datasets and experimental splits",
    124: "Table 1 Statistics and experimental splits of the five datasets",
    125: "4.2 Evaluation metrics and statistical reporting",
    126: (
        "Retrieval quality over the complete candidate set is evaluated with Hits@1, Hits@10, and mean reciprocal rank (MRR). Main results on all five "
        "datasets and the ablation experiments report the mean and sample standard deviation over three random seeds."
    ),
    127: "4.3 Implementation details and hyperparameter selection",
    128: (
        "The model is optimized with AdamW [24] using a batch size of 512, a representation dimension of 128, three relation-aware GNN layers, and dropout "
        "of 0.1. Semantic inputs are projected from 300 to 128 dimensions. The global view uses a two-layer, four-head Transformer, and the phrase view uses "
        "convolutional widths of 3 and 5. The 1.5-entmax temperature is {{EQ0}}, and the InfoNCE temperature is {{EQ1}}. Random seeds are 42, 43, and 44. "
        "DBP15K is trained for at most 36 epochs; OpenEA and EventEA are trained for at most 50. Validation MRR is measured every five epochs, and training "
        "stops after four consecutive validations without improvement. Primary retrieval does not tune a branch-fusion weight. Instead, each run selects "
        "one value from {{EQ2}} by validation MRR. Relation identifiers are unified after relation names are normalized; DBP15K additionally uses the "
        "dataset-provided sup_rel_ids file to map known corresponding relations."
    ),
    129: "5 Results and analysis",
    130: "5.1 Overall test performance and training stability",
    131: (
        "Table 2 summarizes retrieval results obtained directly from the joint representation on five datasets. Each dataset is trained independently with "
        "three random seeds. DBP15K holds out a fixed 10% of the training pairs for validation, whereas OpenEA and EventEA use their official validation sets."
    ),
    132: "Table 2 Final test results on five datasets (mean ± sample standard deviation, n = 3)",
    134: "Across the five datasets, mean Hits@1 ranges from 0.6344 to 0.9239, and the sample standard deviation over three random seeds does not exceed 0.0053.",
    135: "5.2 Contextual comparison with published results",
    136: (
        "Table 3 uses values reported in the literature. For DBP15K ZH-EN, the MTransE, JAPE, BootEA, and RDGCN results are taken from the RDGCN summary "
        "[5], and RREA-text follows the text-enhanced setting reported by RREA [6]. OpenEA EN-FR-V2 values are the official five-fold averages reported by "
        "OpenEA [8]."
    ),
    137: "Table 3 Contextual comparison between published results and the proposed model",
    139: (
        "Because data splits, textual inputs, and post-processing differ across methods, Table 3 indicates only the position of the present results relative "
        "to prior work and is not a strict same-protocol ranking. On DBP15K ZH-EN, the proposed model outperforms the earlier structural methods listed in "
        "the table but remains below RREA-text. On OpenEA EN-FR-V2, it remains below BootEA, KDCoE, and RDGCN. These differences preclude a claim of state-"
        "of-the-art performance."
    ),
    141: "5.3 Ablation of structural components and semantic views",
    142: (
        "Table 4 presents ablations on DBP15K ZH-EN and OpenEA EN-FR-15K-V2. Every variant is retrained with three random seeds and is retrieved directly "
        "from the representation on which it was trained."
    ),
    143: "Table 4 Ablation of structural components, neighbor aggregation, and semantic views (Hits@1/MRR, mean {{EQ0}} sample standard deviation, {{EQ1}})",
    144: (
        "Table 4 shows that replacing topology initialization with learnable entity vectors alone reduces Hits@1 by 3.46 points on DBP15K and 8.03 points "
        "on OpenEA. Removing relation types reduces Hits@1 by 3.62 and 2.78 points. Removing the layer selector and jointly replacing the full neighborhood "
        "and 1.5-entmax with a fixed eight-neighbor budget and softmax produce smaller decreases in the same direction. Among the semantic views, the phrase "
        "view contributes positively on both datasets, whereas the marginal effect of the token view is no greater than 0.27 points. Removing the global "
        "view increases DBP15K by 0.49 points but decreases OpenEA by 1.15 points, indicating a dataset-dependent effect. Removing structural supervision "
        "reduces Hits@1 by only 0.07 and 0.17 points; this auxiliary objective is therefore not a principal source of performance."
    ),
    145: "5.4 Structural branch, semantic branch, and fusion module",
    146: (
        "Table 5 separates the training and primary retrieval representations for the full model, structural-only and semantic-only branches, and "
        "parameter-free mean fusion. Every result retrieves directly from the representation used for training; validation does not alter the structural "
        "or semantic weight."
    ),
    147: "Table 5 Controls for branches, training representations, and primary retrieval representations (Hits@1/MRR, mean {{EQ0}} sample standard deviation, {{EQ1}})",
    148: (
        "The semantic-only branch is substantially stronger than the structural-only branch but remains below the full model: Hits@1 is lower by 9.09 "
        "points on DBP15K and 9.66 points on OpenEA. Parameter-free mean fusion also underperforms the full model. Thus, structural information contributes "
        "primarily when it forms a trainable joint representation with semantic evidence, rather than when it is retrieved independently."
    ),
    149: "5.5 Neighbor context, query signal, and capacity control",
    150: (
        "Table 6 compares the full model, the structural-query control, and parameter-matched late fusion. The full model and structural-query control have "
        "the same one-hop outgoing neighborhood, fusion parameters, and direct joint retrieval; only the signal used to query neighbor weights differs. The "
        "parameter-matched late-fusion model retains the same number of trainable fusion parameters but does not read neighbor context."
    ),
    151: "Table 6 Neighbor visibility, query signal, and capacity controls (Hits@1/MRR, mean ± sample standard deviation, n = 3)",
    152: (
        "With the neighborhood, training capacity, and retrieval representation held constant, structural queries yield Hits@1 values 2.08 and 2.48 points "
        "below semantic queries on DBP15K and OpenEA, respectively. The direction is consistent across both datasets, indicating that semantic querying before "
        "neighbor aggregation improves retrieval under the current setting. Parameter-matched late fusion is lower by 5.28 and 9.22 points. Removing neighbor "
        "context while changing the training interaction path therefore causes a substantial decrease. Because both factors change together, this gap cannot "
        "be attributed to fusion timing alone."
    ),
    153: "6 Discussion",
    154: "6.1 Main findings and dataset differences",
    155: (
        "Three findings emerge. First, variation across random seeds is small on all five datasets, although absolute performance differs substantially across "
        "datasets. Second, semantic-only retrieval is markedly stronger than structural-only retrieval, yet the full model outperforms both the semantic branch "
        "and parameter-free mean fusion. Together with the topology-initialization and relation-type ablations, this result indicates that structural evidence "
        "is most useful within a trainable joint representation. Third, when neighborhood and fusion capacity are matched, semantic queries outperform structural "
        "queries on both representative datasets. Late fusion without neighbor context decreases further, showing that both the query signal and neighbor context "
        "affect the joint representation. The semantic views are not uniformly beneficial; in particular, the global view is dataset dependent."
    ),
    156: (
        "DBP15K provides comparatively informative names and local structural cues. EventEA combines Wikidata and DBpedia and contains greater heterogeneity in "
        "event relations and attributes [25]. OpenEA EN-FR-15K-V2 encodes entity URIs to reduce name bias [8]. The current input uses English GloVe vectors for "
        "in-vocabulary tokens and deterministic, MD5-seeded random vectors for out-of-vocabulary tokens [26]. Local statistics show approximately 52.18% GloVe "
        "coverage on OpenEA; many proper names and encoded identifiers therefore lack transferable cross-lingual semantics. This is consistent with its lower "
        "retrieval results relative to DBP15K, but the causal contribution requires experiments that replace the semantic encoder and control input coverage."
    ),
    157: "6.2 Limitations",
    158: (
        "Regarding internal validity, the main and ablation experiments share the same splits, validation rule, and three random seeds, but {{EQ0}} is still too "
        "small for high-powered significance tests. The fixed eight-neighbor plus softmax control changes both the neighbor budget and normalization function, "
        "whereas parameter-matched late fusion changes both neighbor visibility and the training interaction path; neither control identifies a single-factor "
        "causal effect. Cross-graph relation-identifier sharing and edge direction have not yet been controlled separately. Regarding external validity, component "
        "controls cover only two representative datasets, the OpenEA main result uses one official split, and open-world settings with unmatched entities have "
        "not been evaluated."
    ),
    159: "6.3 Implications and future work",
    160: (
        "Future work should address three priorities. First, random vectors for out-of-vocabulary tokens should be replaced with cross-lingual subword or multilingual "
        "retrieval encoders [27, 28], followed by replication on all official OpenEA splits. Second, the efficiency and stability of relation-aware structural "
        "encoding and sparse neighbor selection should be tested on larger graphs. Third, retrieval calibration and generalization should be evaluated in open-world "
        "settings that contain entities without counterparts."
    ),
    161: "7 Conclusion",
    162: (
        "This study presents a joint model of relation-aware structural context and multi-scale semantics and evaluates it with multiple random seeds on five entity "
        "alignment datasets. Topology initialization, relation types, and the phrase view provide useful information to the joint representation. Although the "
        "structural branch performs weakly in isolation, the full model outperforms the semantic-only branch and parameter-free mean fusion. With the same neighborhood, "
        "semantic queries outperform structural queries on two representative datasets. Parameter-matched late fusion without neighbor context decreases further, but "
        "this control changes both neighbor visibility and the training interaction path. The evidence therefore supports relation-aware neighbor context in a joint "
        "structure-semantic representation and semantic-guided neighbor weighting on the two controlled datasets. Their generality across datasets and the independent "
        "value of structural auxiliary supervision require further study."
    ),
    163: "Declarations",
    164: "Data availability  DBP15K, OpenEA, and EventEA are publicly available benchmarks; their sources are provided in the corresponding references.",
    165: "Code availability  Code, configurations, and result summaries required to reproduce the experiments are available from the corresponding author on reasonable request.",
    166: "Competing interests  The author declares no competing interests.",
    167: "References",
}


TABLE_TRANSLATIONS = {
    "数据集": "Dataset",
    "实体数（合计）": "Entities (total)",
    "三元组（合计）": "Triples (total)",
    "训练/验证/测试": "Train/validation/test",
    "语义序列形状": "Semantic sequence shape",
    "关系类型": "Relation types",
    "方法": "Method",
    "未报告": "Not reported",
    "本文模型": "Proposed model",
    "变体": "Variant",
    "完整模型": "Full model",
    "仅可学习实体初始化（无拓扑特征）": "Learnable entity initialization only (no topology features)",
    "固定 8 邻居 + softmax": "Fixed 8-neighbor budget + softmax",
    "去除关系类型": "Without relation types",
    "去除层选择器": "Without layer selector",
    "去除 token 视图": "Without token view",
    "去除 phrase 视图": "Without phrase view",
    "去除 global 视图": "Without global view",
    "去除结构监督": "Without structural supervision",
    "训练表示": "Training representation",
    "主检索表示": "Primary retrieval representation",
    "联合表示": "Joint representation",
    "仅结构分支": "Structural branch only",
    "结构表示": "Structural representation",
    "仅语义分支": "Semantic branch only",
    "语义表示": "Semantic representation",
    "无可训练平均融合": "Parameter-free mean fusion",
    "结构与语义等权平均": "Equal mean of structure and semantics",
    "同一等权平均": "Same equal-weight mean",
    "设计差异": "Design difference",
    "可训练融合参数": "Trainable fusion parameters",
    "完整一跳出邻域；语义查询；直接联合检索": "Complete one-hop outgoing neighborhood; semantic query; direct joint retrieval",
    "结构查询邻居": "Structural neighbor query",
    "相同邻域与门控；结构查询；直接联合检索": "Same neighborhood and gates; structural query; direct joint retrieval",
    "参数匹配晚期融合": "Parameter-matched late fusion",
    "不读取邻居；独立分支后形成并检索 z_late": "No neighbor context; form and retrieve z_late after independent branches",
}


PLACEHOLDER_RE = re.compile(r"\{\{EQ(\d+)\}\}")


def combined_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).replace("\u00a0", " ")


def equation_children(paragraph: Paragraph):
    return [
        copy.deepcopy(child)
        for child in paragraph._p.iterchildren()
        if child.tag in {qn("m:oMath"), qn("m:oMathPara")}
    ]


def append_text_run(paragraph: Paragraph, text: str) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_paragraph_markup(paragraph: Paragraph, markup: str) -> None:
    equations = equation_children(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    cursor = 0
    for match in PLACEHOLDER_RE.finditer(markup):
        append_text_run(paragraph, markup[cursor : match.start()])
        index = int(match.group(1))
        if index >= len(equations):
            raise RuntimeError(
                f"Equation placeholder EQ{index} exceeds {len(equations)} equations in paragraph: {markup!r}"
            )
        paragraph._p.append(copy.deepcopy(equations[index]))
        cursor = match.end()
    append_text_run(paragraph, markup[cursor:])


def set_plain_text(paragraph: Paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    append_text_run(paragraph, text)


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = style
    set_plain_text(inserted, text)
    return inserted


def set_run_font(run, size: float, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_style_font(style, size: float, *, bold: bool | None = None, italic: bool | None = None) -> None:
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def add_page_number(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instruction, separate, end])
    set_run_font(run, 10)


def add_line_numbering(section) -> None:
    section_pr = section._sectPr
    existing = section_pr.find(qn("w:lnNumType"))
    if existing is not None:
        section_pr.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    section_pr.append(line_numbers)


def image_member(document: Document) -> str:
    if len(document.inline_shapes) != 1:
        raise RuntimeError(f"Expected one figure, found {len(document.inline_shapes)}")
    shape = document.inline_shapes[0]
    blips = shape._inline.xpath(".//*[local-name()='blip']")
    rel_id = blips[0].get(qn("r:embed"))
    return f"word/{document.part.rels[rel_id].target_ref}"


def replace_zip_member(docx_path: Path, member_name: str, replacement_path: Path) -> None:
    temporary = docx_path.with_suffix(".replace.tmp.docx")
    replacement = replacement_path.read_bytes()
    with zipfile.ZipFile(docx_path, "r") as source_archive:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as output_archive:
            for item in source_archive.infolist():
                data = replacement if item.filename == member_name else source_archive.read(item.filename)
                output_archive.writestr(item, data)
    os.replace(temporary, docx_path)


def translate_document(document: Document) -> None:
    if len(document.paragraphs) != 196:
        raise RuntimeError(f"Expected 196 source paragraphs, found {len(document.paragraphs)}")
    for index, markup in P.items():
        paragraph = document.paragraphs[index]
        if "{{EQ" in markup:
            set_paragraph_markup(paragraph, markup)
        else:
            set_plain_text(paragraph, markup)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                source = cell.text.strip()
                if source in TABLE_TRANSLATIONS:
                    set_plain_text(cell.paragraphs[0], TABLE_TRANSLATIONS[source])
                    for extra in cell.paragraphs[1:]:
                        extra._element.getparent().remove(extra._element)

    affiliation = insert_after(document.paragraphs[1], "Affiliation: [Department, Institution, City, Country]")
    affiliation.paragraph_format.keep_with_next = True

    declarations = next(p for p in document.paragraphs if combined_text(p).strip() == "Declarations")
    anchor = declarations
    anchor = insert_after(
        anchor,
        "Funding  [Please provide the funding statement. If no external funding was received, replace this text with: "
        "'The author received no specific funding for this work.']",
    )
    anchor = insert_after(
        anchor,
        "Author contributions  Xinyi Lin: Conceptualization, methodology, software, validation, formal analysis, investigation, "
        "data curation, visualization, writing - original draft, and writing - review and editing.",
    )

    conflict = next(p for p in document.paragraphs if combined_text(p).startswith("Competing interests"))
    anchor = conflict
    anchor = insert_after(anchor, "Ethics approval  Not applicable.")
    anchor = insert_after(anchor, "Consent to participate  Not applicable.")
    insert_after(anchor, "Consent for publication  Not applicable.")


def format_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    add_line_numbering(section)

    styles = document.styles
    set_style_font(styles["Normal"], 12)
    normal = styles["Normal"].paragraph_format
    normal.line_spacing = 2.0
    normal.space_before = Pt(0)
    normal.space_after = Pt(0)
    normal.widow_control = True

    set_style_font(styles["Heading 1"], 14, bold=True)
    h1 = styles["Heading 1"].paragraph_format
    h1.line_spacing = 1.0
    h1.space_before = Pt(14)
    h1.space_after = Pt(6)
    h1.keep_with_next = True
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT

    set_style_font(styles["Heading 2"], 12, bold=True)
    h2 = styles["Heading 2"].paragraph_format
    h2.line_spacing = 1.0
    h2.space_before = Pt(12)
    h2.space_after = Pt(4)
    h2.keep_with_next = True
    h2.alignment = WD_ALIGN_PARAGRAPH.LEFT

    set_style_font(styles["Heading 3"], 12, bold=True, italic=True)
    h3 = styles["Heading 3"].paragraph_format
    h3.line_spacing = 1.0
    h3.space_before = Pt(10)
    h3.space_after = Pt(3)
    h3.keep_with_next = True

    if "Caption" in styles:
        set_style_font(styles["Caption"], 10)
        caption = styles["Caption"].paragraph_format
        caption.line_spacing = 1.0
        caption.space_before = Pt(6)
        caption.space_after = Pt(6)
        caption.keep_together = True
        caption.keep_with_next = True
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if "Formula" in styles:
        set_style_font(styles["Formula"], 11)
        formula = styles["Formula"].paragraph_format
        formula.line_spacing = 1.0
        formula.space_before = Pt(4)
        formula.space_after = Pt(4)
        formula.keep_together = True
        formula.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "Reference" in styles:
        set_style_font(styles["Reference"], 11)
        reference = styles["Reference"].paragraph_format
        reference.line_spacing = 1.5
        reference.left_indent = Mm(6)
        reference.first_line_indent = Mm(-6)
        reference.space_after = Pt(3)

    title = document.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    for run in title.runs:
        set_run_font(run, 17, bold=True)

    author = document.paragraphs[1]
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.line_spacing = 1.0
    author.paragraph_format.space_after = Pt(4)
    author.paragraph_format.keep_with_next = True
    for run in author.runs:
        set_run_font(run, 12)

    for paragraph in document.paragraphs[2:5]:
        if combined_text(paragraph).startswith("Affiliation:") or combined_text(paragraph).startswith("Corresponding author:"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, 10.5)

    abstract = next(p for p in document.paragraphs if combined_text(p).startswith("Cross-lingual knowledge graph entity alignment identifies"))
    abstract.paragraph_format.line_spacing = 1.15
    abstract.paragraph_format.space_after = Pt(6)
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in abstract.runs:
        set_run_font(run, 10.5)

    keywords = next(p for p in document.paragraphs if combined_text(p).startswith("Keywords:"))
    keywords.paragraph_format.line_spacing = 1.0
    keywords.paragraph_format.space_after = Pt(10)
    for run in keywords.runs:
        set_run_font(run, 10.5)

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name
        if style_name in {"Heading 1", "Heading 2", "Heading 3", "Caption", "Formula", "Reference"}:
            continue
        if paragraph in {title, author, abstract, keywords}:
            continue
        if not combined_text(paragraph).strip():
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_run_font(run, 12)

    for table in document.tables:
        table.autofit = False
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_mar = tc_pr.first_child_found_in("w:tcMar")
                if tc_mar is None:
                    tc_mar = OxmlElement("w:tcMar")
                    tc_pr.append(tc_mar)
                for side in ("top", "left", "bottom", "right"):
                    node = tc_mar.find(qn(f"w:{side}"))
                    if node is None:
                        node = OxmlElement(f"w:{side}")
                        tc_mar.append(node)
                    node.set(qn("w:w"), "90")
                    node.set(qn("w:type"), "dxa")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(cell.text) < 42 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_run_font(run, 8.5, bold=(row_index == 0))

    shape = document.inline_shapes[0]
    shape.width = Inches(6.35)
    shape.height = Inches(6.35 * 1500 / 2400)
    props = shape._inline.xpath(".//*[local-name()='docPr']")
    if props:
        props[0].set("name", "Figure 1")
        props[0].set(
            "descr",
            "Relation-aware structural context, multi-scale semantic encoding, sparse outgoing-neighbor weighting, training objective, and CSLS retrieval.",
        )

    header_paragraph = section.header.paragraphs[0]
    set_plain_text(header_paragraph, "Relation-Aware Structure-Semantic Entity Alignment")
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_paragraph.paragraph_format.line_spacing = 1.0
    for run in header_paragraph.runs:
        set_run_font(run, 9)

    footer_paragraph = section.footer.paragraphs[0]
    for child in list(footer_paragraph._p):
        if child.tag != qn("w:pPr"):
            footer_paragraph._p.remove(child)
    add_page_number(footer_paragraph)


def visible_text(document: Document) -> str:
    parts = [combined_text(p) for p in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def verify(path: Path, figure_member: str) -> None:
    document = Document(path)
    text = visible_text(document)
    required = [
        "Cross-Lingual Knowledge Graph Entity Alignment with Relation-Aware Structural Context",
        "0.7077 ± 0.0023 / 0.7497 ± 0.0016",
        "0.6096 ± 0.0043 / 0.6752 ± 0.0019",
        "structural queries yield Hits@1 values 2.08 and 2.48 points below semantic queries",
        "Funding  [Please provide the funding statement",
        "Author contributions  Xinyi Lin",
        "Affiliation: [Department, Institution, City, Country]",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise RuntimeError(f"Required English content is missing: {missing}")
    if re.search(r"[\u3400-\u9fff]", text):
        raise RuntimeError("Chinese text remains in visible manuscript content")
    if len(document.tables) != 6 or len(document.inline_shapes) != 1:
        raise RuntimeError("Table or figure count changed unexpectedly")
    equations = sum(1 for node in document.element.body.iter() if node.tag == qn("m:oMath"))
    if equations != 154:
        raise RuntimeError(f"Native equation count changed: {equations}")
    if "0.7277 ± 0.0023 / 0.7697 ± 0.0016" in text:
        raise RuntimeError("Outdated structural-query result remains")

    with zipfile.ZipFile(path) as archive:
        embedded = archive.read(figure_member)
        if hashlib.sha256(embedded).digest() != hashlib.sha256(FIGURE.read_bytes()).digest():
            raise RuntimeError("The current English architecture figure was not embedded")


def main() -> None:
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The approved Chinese source changed; refusing to translate a stale version. "
            f"Expected {EXPECTED_SOURCE_SHA256}, found {source_hash}."
        )
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)

    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    translate_document(document)
    figure_member = image_member(document)
    format_document(document)
    document.core_properties.title = P[0]
    document.core_properties.author = "Xinyi Lin"
    document.core_properties.last_modified_by = "Xinyi Lin"
    document.core_properties.subject = "Cross-lingual knowledge graph entity alignment"
    document.core_properties.keywords = "entity alignment; cross-lingual knowledge graphs; relation-aware GNN; multi-scale semantics; sparse neighbor selection"
    document.save(OUTPUT)
    replace_zip_member(OUTPUT, figure_member, FIGURE)
    verify(OUTPUT, figure_member)
    print(OUTPUT)


if __name__ == "__main__":
    main()
