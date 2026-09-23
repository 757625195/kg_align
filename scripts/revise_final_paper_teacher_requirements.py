# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from build_interim_new_main_paper import (
    add_mixed_text,
    clear_text_highlights,
    find_formula,
    find_starts,
    formula_text,
    insert_after,
    insert_before,
    replace_formula,
    replace_starts,
)


def find_contains(document: Document, fragment: str):
    matches = [paragraph for paragraph in document.paragraphs if fragment in paragraph.text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph containing {fragment!r}, found {len(matches)}")
    return matches[0]


def find_formula_starts(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Formula" and formula_text(paragraph).startswith(prefix)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one formula starting {prefix!r}, found {len(matches)}")
    return matches[0]


def append_text(paragraph, text: str) -> None:
    paragraph.add_run(text)


def append_mixed(paragraph, text: str) -> None:
    add_mixed_text(paragraph, text)


def remove_paragraph(paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def style_headings(document: Document) -> None:
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.color.rgb = RGBColor(43, 43, 43)
        style.font.name = "Heiti SC"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")


def keep_table_rows_intact(document: Document) -> None:
    for table in document.tables:
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            if not properties.xpath("./w:cantSplit"):
                properties.append(OxmlElement("w:cantSplit"))


def revise(document: Document) -> None:
    replace_starts(
        document,
        "第一，结构相似不等于实体等价",
        "第一，结构相似不等于实体等价。例如，人物 A 通过“出生地”连接北京，人物 B 通过“工作地点”连接北京；若模型把不同关系都退化为无类型连接，就会高估两人的相似性[4,7]。第二，语义证据往往是碎片化的。例如，单独出现的“Washington”无法区分人物、州和城市，需要结合关系或属性上下文。第三，人工核验跨图实体对需要领域知识和逐对检查，因此种子对齐通常只覆盖实体集合的一部分[3,14]。",
    )

    replace_starts(
        document,
        "本文利用的 GraphSAGE 邻域聚合",
        "本文使用的 GraphSAGE 邻域聚合[6]、R-GCN 关系建模[7]、Transformer 自注意力[8]和 InfoNCE 对比目标[9]均来自已有研究，因而不把这些基础算子本身作为创新。本文的科学贡献在于：以关系编号不变的拓扑初始化改善跨图结构起点，在不截断一跳邻居的条件下引入无新增参数的稀疏选择，并通过参数匹配与信息可见范围控制实验检验结构邻居在最终融合前进入计算是否有价值。该定位强调可检验机制和可复现实验证据，而非仅给出高层框架图或宣称达到最新最佳性能。本文的主要贡献如下：",
    )

    replace_starts(
        document,
        "BootEA 在少量已知实体对的基础上",
        "BootEA 在少量已知实体对的基础上迭代训练模型，并将高置信度预测实体对作为伪标签加入监督集合[3]。伪标签是模型生成但未经人工确认的训练标记。该策略可以利用未标注数据，但也会传播早期错误：若模型把行星“Mercury”错配为同名汽车品牌并将该实体对加入训练，后续表示学习会继续拉近两者，并影响新的候选排序。本文不使用伪标签扩充。",
    )

    for duplicate_prefix in (
        "实现并验证一种轻量级关系感知结构编码器",
        "构建语义编码与结构编码早期交互模块",
        "在五个数据集上完成三随机种子主实验，并在 DBP15K ZH–EN",
    ):
        remove_paragraph(find_starts(document, duplicate_prefix))
    empty_numbered = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Number" and not paragraph.text.strip()
    ]
    for paragraph in empty_numbered:
        remove_paragraph(paragraph)

    seed_formula = find_formula_starts(document, "Aseed")
    insert_after(
        document,
        seed_formula,
        "其中，$A_{\\mathrm{seed}}$ 称为种子对齐集合，包含少量已由人工或数据集构造规则确认的等价实体对；符号 $e_i\\equiv e_j$ 表示两者指向同一现实对象。种子对齐既是训练监督的来源，也是有限监督设定的具体含义：模型只能直接看到该集合中的对应关系，其余实体必须依靠结构、语义及已学得的跨图空间完成检索。",
    )

    architecture_caption = find_starts(document, "图 1 拓扑初始化的关系感知结构上下文")
    append_text(
        architecture_caption,
        " 结构编码、语义编码、上下文融合、训练目标和检索分别对应式（1）—（7）、式（8）—（13）、式（14）—（21）、式（23）—（26）和式（27）。",
    )

    replace_starts(
        document,
        "对于实体 i，本文将其名称、关系名称和属性文本组成 token 序列",
        "对于实体 $i$，本文将其名称、相邻关系名称和属性文本分词后拼接为同一 token 序列，并将每个 token 映射为词向量。phrase 与 global 视图中的词均来自这条共享序列，不额外读取实体描述、外部句子或其他语料。序列经统一维度映射、归一化和位置编码后得到基础表示 $H_i$，掩码 $M_i$ 标记有效 token，使后续池化和注意力忽略填充位置。",
    )

    interaction = find_contains(document, "描述逐维差异")
    append_text(
        interaction,
        " 这种“保留原向量、逐维乘积与绝对差”的组合借鉴了自然语言推理中的匹配特征设计[19]。它本身不是最终相似度，而是供后续可学习映射判断邻居兼容性的特征表示。",
    )

    replace_starts(
        document,
        "式（14）衡量实体语义与第",
        "式（14）衡量实体语义与第 $j$ 个结构邻居的匹配程度。点积越大，表示该邻居与当前实体的语义越一致；缩放因子用于控制高维点积的数值范围，其中 $d=128$ 为查询与键的表示维度。",
    )

    replace_formula(
        document,
        24,
        r"\mathcal{L}_{\mathrm{joint}}=\frac{1}{2}\left(\operatorname{CE}(\mathbf{S},[0,\ldots,B-1])+\operatorname{CE}(\mathbf{S}^{\mathsf{T}},[0,\ldots,B-1])\right)",
    )

    replace_starts(
        document,
        "模型采用单阶段训练，不设置独立预热阶段",
        "模型采用单阶段训练，不设置独立预热阶段。每个批次从种子对齐集合采样 $B$ 个实体对，并用同一批样本计算联合表示损失与结构分支损失；两者在一次前向和反向传播中共同优化。首先计算左图第 $i$ 个实体与右图第 $j$ 个实体之间的温度缩放余弦相似度：",
    )

    structure_decay = find_starts(document, "结构监督只用于训练前期")
    insert_before(
        document,
        structure_decay,
        "难负候选（hard negative candidate）是指在闭世界训练划分中并非真实配对、但被当前模型赋予较高相似度的实体。当前主实验不对这类候选进行额外挖掘，也不使用间隔排序损失；批内非配对实体仅通过式（24）的 InfoNCE 分母参与区分。",
    )
    remove_paragraph(find_starts(document, "模型采用 AdamW 优化器"))

    ablation_protocol = find_starts(document, "表 5 在 DBP15K ZH–EN")
    append_text(
        ablation_protocol,
        " 完整模型对应 token+phrase+global，去除 global 对应 token+phrase；此外分别去除 token、phrase 和 global 的留一消融，在保持其余模型与训练协议不变时估计每个视图的边际贡献。",
    )

    positioning = find_starts(document, "在 DBP15K ZH–EN 上，本文模型的 Hits@1")
    append_text(
        positioning,
        " 因此，本文对方法价值的论证主要来自五数据集稳定性、内部基线和受控消融，而不是把非同协议文献数值解释为最新最佳结果。",
    )

    replace_starts(
        document,
        "表 3 中 DBP15K ZH–EN 的 MTransE",
        "表 3 中 DBP15K ZH–EN 的 MTransE、JAPE、BootEA 和 RDGCN 数值取自 RDGCN 的汇总[4]，RREA-text 取自原论文[5]；OpenEA 数值取自官方基准的五折平均[14]，而本文 OpenEA 结果为官方第一折上的三随机种子均值。因此，该表用于方法定位，不构成完全同协议的严格排名。",
    )

    replace_starts(
        document,
        "公共设置为 batch size 512",
        "实现采用 AdamW[16]，batch size 为 512，表示维度为 128，关系感知 GNN 的最大深度为 3，dropout 为 0.1。语义输入由 300 维投影至 128 维；注意力池化评分网络为 $128\\rightarrow128\\rightarrow1$，phrase 卷积核宽度为 3 和 5，global 编码器包含两层、四个注意力头和宽度为 512 的前馈层；视图门控与输出 MLP 分别为 $384\\rightarrow128\\rightarrow3$ 和 $384\\rightarrow128\\rightarrow128$。DBP15K 的学习率为 $5\\times10^{-4}$，最多训练 36 个 epoch；OpenEA 与 EventEA 的学习率为 $3\\times10^{-4}$，最多训练 50 个 epoch。每 5 个 epoch 按验证集 MRR 保存检查点，连续 4 次验证未提升时提前停止。验证网格为 $\\alpha\\in\\{0.2,\\ldots,0.8\\}$、$L\\in\\{1,2,3\\}$ 和 $k_{\\mathrm{CSLS}}\\in\\{3,5,7,10,15,20\\}$；每个数据集和消融变体依据三个随机种子的平均验证 MRR 独立选择唯一配置。",
    )

    style_headings(document)
    find_starts(document, "5 实验结果与分析").paragraph_format.page_break_before = True
    keep_table_rows_intact(document)
    clear_text_highlights(document)


def verify(document: Document) -> None:
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required = [
        "种子对齐集合",
        "不设置独立预热阶段",
        "hard negative candidate",
        "注意力池化评分网络为",
        "不是最终相似度",
        "可检验机制和可复现实验证据",
    ]
    missing = [text for text in required if text not in body]
    if missing:
        raise RuntimeError(f"Missing required revisions: {missing}")
    if "Lalign" in formula_text(find_formula(document, 24)):
        raise RuntimeError("Equation (24) still uses the inconsistent L_align symbol")
    tags = []
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Formula":
            continue
        text = formula_text(paragraph)
        if text.endswith(")") and "(" in text:
            suffix = text.rsplit("(", 1)[-1][:-1]
            if suffix.isdigit():
                tags.append(int(suffix))
    if tags != list(range(1, 28)):
        raise RuntimeError(f"Unexpected equation numbering: {tags}")
    if body.count("伪标签是") != 1:
        raise RuntimeError("Pseudo-label definition must appear exactly once")
    if body.count("hard negative candidate") != 1:
        raise RuntimeError("Hard-negative definition must appear exactly once")
    duplicate_contributions = (
        "实现并验证一种轻量级关系感知结构编码器",
        "构建语义编码与结构编码早期交互模块",
    )
    if any(text in body for text in duplicate_contributions):
        raise RuntimeError("Duplicate contribution list remains")
    if "$d$" in body:
        raise RuntimeError("Literal LaTeX marker remains in body text")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)
    revise(document)
    verify(document)
    document.save(args.output)
    print(f"WROTE {args.output}")


if __name__ == "__main__":
    main()
