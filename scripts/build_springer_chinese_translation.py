from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ENGLISH = ROOT / "outputs" / "1_springer_proceedings.docx"
TRANSLATION_REFERENCE = ROOT / "outputs" / "1_二因素审稿实验标黄修订版_中文版_20260831.docx"
CHINESE_FIGURE = ROOT / "outputs" / "paper_assets" / "figure_1_springer_submission_zh_600dpi.png"
OUTPUT = ROOT / "outputs" / "1_springer_proceedings_中文版.docx"


EQ_RE = re.compile(r"\{\{EQ(\d+)(?::.*?)?\}\}")


def paragraph_markup(paragraph) -> str:
    parts: list[str] = []
    equation_index = 0
    for child in paragraph._p.iterchildren():
        if child.tag == qn("w:r"):
            parts.extend(node.text or "" for node in child.iter(qn("w:t")))
        elif child.tag == qn("w:hyperlink"):
            parts.extend(node.text or "" for node in child.iter(qn("w:t")))
        elif child.tag in {qn("m:oMath"), qn("m:oMathPara")}:
            # The formula content itself is retained from the English source.
            # Index-only placeholders avoid ambiguity when a formula ends in }.
            parts.append(f"{{{{EQ{equation_index}}}}}")
            equation_index += 1
        else:
            text = "".join(node.text or "" for node in child.iter(qn("w:t")))
            if text:
                parts.append(text)
    return "".join(parts).replace("\u00a0", " ")


def direct_math(paragraph):
    return [
        deepcopy(child)
        for child in paragraph._p.iterchildren()
        if child.tag in {qn("m:oMath"), qn("m:oMathPara")}
    ]


def rewrite_paragraph(paragraph, markup: str, bold_prefix: str | None = None) -> None:
    if paragraph._p.xpath(".//w:drawing|.//w:pict"):
        return
    math_nodes = direct_math(paragraph)
    requested = [int(match.group(1)) for match in EQ_RE.finditer(markup)]
    if requested and (max(requested) >= len(math_nodes) or sorted(requested) != list(range(len(math_nodes)))):
        raise RuntimeError(
            f"Equation placeholder mismatch: requested={requested}, available={len(math_nodes)}, text={markup!r}"
        )

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)

    position = 0
    prefix_pending = bold_prefix or ""
    for match in EQ_RE.finditer(markup):
        text = markup[position : match.start()]
        prefix_pending = append_text(paragraph, text, prefix_pending)
        paragraph._p.append(deepcopy(math_nodes[int(match.group(1))]))
        position = match.end()
    prefix_pending = append_text(paragraph, markup[position:], prefix_pending)
    if prefix_pending:
        raise RuntimeError(f"Bold prefix {bold_prefix!r} not found in {markup!r}")


def append_text(paragraph, text: str, bold_prefix: str) -> str:
    if not text:
        return bold_prefix
    if bold_prefix:
        if not text.startswith(bold_prefix):
            raise RuntimeError(f"Expected bold prefix {bold_prefix!r} at start of {text!r}")
        bold_run = paragraph.add_run(bold_prefix)
        bold_run.bold = True
        text = text[len(bold_prefix) :]
        bold_prefix = ""
    if text:
        paragraph.add_run(text)
    return bold_prefix


def translation_for_source(reference: Document, source_index: int, overrides: dict[int, str]) -> str:
    if source_index in overrides:
        return overrides[source_index]
    reference_index = source_index if source_index <= 150 else source_index + 1
    return paragraph_markup(reference.paragraphs[reference_index])


def bold_prefix_for(paragraph, translated: str) -> str | None:
    if paragraph.style.name == "Springer Abstract":
        return "摘要。"
    if paragraph.style.name == "Springer Keywords":
        return "关键词："
    if paragraph.style.name == "Heading 3":
        return translated.split("。", 1)[0] + "。"
    if paragraph.style.name == "Caption":
        match = re.match(r"^((?:图|表)\s*[A-Za-z]?\d+)", translated)
        return match.group(1) if match else None
    return None


def source_mapping(target_index: int):
    if target_index in (0, 1):
        return target_index
    if target_index == 2:
        return (2, 3)
    if 3 <= target_index <= 77:
        return target_index + 1
    if target_index == 78:
        return (79, 80)
    if 79 <= target_index <= 87:
        return target_index + 2
    if target_index == 88:
        return (90, 91)
    if 89 <= target_index <= 92:
        return target_index + 3
    if target_index == 93:
        return (96, 97)
    if 94 <= target_index <= 97:
        return target_index + 4
    if target_index == 98:
        return (102, 103)
    return target_index + 5


def set_east_asia_font(run, east_asia: str, size: float | None = None) -> None:
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), east_asia)


def configure_chinese_styles(document: Document) -> None:
    specs = {
        "Normal": (10, False, "Songti SC"),
        "Springer Paper Title": (14, True, "Heiti SC"),
        "Springer Author": (10, False, "Songti SC"),
        "Springer Abstract": (9, False, "Songti SC"),
        "Springer Keywords": (9, False, "Songti SC"),
        "Heading 1": (12, True, "Heiti SC"),
        "Heading 2": (10, True, "Heiti SC"),
        "Heading 3": (10, False, "Songti SC"),
        "Caption": (9, False, "Songti SC"),
        "Springer Reference": (9, False, "Times New Roman"),
        "List Number": (10, False, "Songti SC"),
    }
    for name, (size, bold, east_asia) in specs.items():
        try:
            style = document.styles[name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), east_asia)


def format_all_runs(document: Document) -> None:
    for paragraph in document.paragraphs:
        east_asia = "Heiti SC" if paragraph.style.name in {"Heading 1", "Heading 2", "Springer Paper Title"} else "Songti SC"
        if paragraph.style.name == "Springer Reference":
            east_asia = "Times New Roman"
        for run in paragraph.runs:
            set_east_asia_font(run, east_asia)

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_east_asia_font(run, "Heiti SC" if row_index == 0 else "Songti SC", 7.5)
                        if row_index == 0:
                            run.bold = True

    section = document.sections[0]
    for header in (section.header, section.even_page_header, section.first_page_header):
        for paragraph in header.paragraphs:
            for run in paragraph.runs:
                set_east_asia_font(run, "Songti SC", 9)


def set_cell_markup(cell, markup: str) -> None:
    rewrite_paragraph(cell.paragraphs[0], markup)
    for paragraph in cell.paragraphs[1:]:
        rewrite_paragraph(paragraph, "")


def translate_tables(document: Document, reference: Document) -> None:
    for table_index in range(5):
        target = document.tables[table_index]
        source = reference.tables[table_index]
        if len(target.rows) != len(source.rows) or len(target.columns) != len(source.columns):
            raise RuntimeError(f"Table {table_index + 1} shape mismatch")
        for row_index, row in enumerate(target.rows):
            for column_index, cell in enumerate(row.cells):
                markup = paragraph_markup(source.cell(row_index, column_index).paragraphs[0])
                set_cell_markup(cell, markup)

    table = document.tables[5]
    labels = {
        (0, 0): "变体",
        (0, 1): "设计差异",
        (0, 2): "可训练融合参数",
        (0, 3): "DBP15K\nHits@1/MRR",
        (0, 4): "OpenEA\nHits@1/MRR",
        (1, 0): "完整模型",
        (1, 1): "完整的一跳出邻域；语义查询；直接联合检索",
        (2, 0): "结构查询对照",
        (2, 1): "相同邻域与门控；结构查询；直接联合检索",
        (3, 0): "参数匹配晚期融合对照",
        (3, 1): "不使用邻居上下文；两个分支独立编码后形成并检索晚期融合表示",
    }
    for (row_index, column_index), text in labels.items():
        set_cell_markup(table.cell(row_index, column_index), text)


def replace_figure(document: Document) -> None:
    shape = document.inline_shapes[0]
    relation_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    document.part.related_parts[relation_id]._blob = CHINESE_FIGURE.read_bytes()
    shape._inline.docPr.set("name", "图1 所提出模型的信息流")
    shape._inline.docPr.set(
        "descr",
        "关系感知结构分支、多尺度语义分支、语义引导的一跳邻居选择、联合表示训练与检索流程图。",
    )
    # The translated caption is longer than the English caption.  A slightly
    # narrower figure allows the image and caption to remain together at the
    # foot of the preceding page instead of leaving a large blank area.
    aspect_ratio = shape.width / shape.height
    shape.width = Inches(4.0)
    shape.height = int(shape.width / aspect_ratio)


def translate_headers(document: Document) -> None:
    section = document.sections[0]
    for paragraph in section.header.paragraphs:
        text_nodes = paragraph._p.xpath(".//w:t")
        if text_nodes:
            text_nodes[0].text = "知识图谱实体对齐中的关系感知邻居上下文"
    for paragraph in section.even_page_header.paragraphs:
        text_nodes = paragraph._p.xpath(".//w:t")
        for node in text_nodes:
            if node.text == "X. Lin":
                node.text = "林心怡"


def set_language(document: Document) -> None:
    document.core_properties.language = "zh-CN"
    for part in document.part.package.parts:
        element = getattr(part, "_element", None)
        if element is None:
            continue
        for lang in element.xpath(".//w:lang"):
            lang.set(qn("w:val"), "zh-CN")
            lang.set(qn("w:eastAsia"), "zh-CN")


def format_reference_language(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Springer Reference":
            continue
        p_pr = paragraph._p.get_or_add_pPr()
        word_wrap = p_pr.find(qn("w:wordWrap"))
        if word_wrap is None:
            from docx.oxml import OxmlElement

            word_wrap = OxmlElement("w:wordWrap")
            p_pr.append(word_wrap)
        word_wrap.set(qn("w:val"), "0")
        for run in paragraph.runs:
            r_pr = run._element.get_or_add_rPr()
            lang = r_pr.find(qn("w:lang"))
            if lang is None:
                from docx.oxml import OxmlElement

                lang = OxmlElement("w:lang")
                r_pr.append(lang)
            lang.set(qn("w:val"), "en-US")
            lang.set(qn("w:eastAsia"), "en-US")


def pair_figure_and_caption(document: Document) -> None:
    for idx, paragraph in enumerate(document.paragraphs[:-1]):
        if paragraph._p.xpath(".//w:drawing|.//w:pict"):
            paragraph.paragraph_format.keep_with_next = True
            caption = document.paragraphs[idx + 1]
            if caption.style.name == "Caption":
                caption.paragraph_format.keep_with_next = False


def protect_latin_terms(document: Document) -> None:
    terms = (
        "OpenEA EN–FR-15K-V2",
        "EventEA EN–EN-20K",
        "DBP15K ZH–EN",
        "DBP15K JA–EN",
        "DBP15K FR–EN",
        "1.5-entmax",
        "RPR-RHGT",
        "RREA-text",
        "GraphSAGE",
        "Transformer",
        "LayerNorm",
        "InfoNCE",
        "Hits@10",
        "Hits@1",
        "softmax",
        "dropout",
        "phrase",
        "global",
        "token",
        "CSLS",
        "GELU",
        "ReLU",
        "MTransE",
        "BootEA",
        "RDGCN",
        "RREA",
        "MCLEA",
        "GloVe",
        "OpenEA",
        "EventEA",
        "DBP15K",
    )
    for node in document._element.xpath(".//w:t"):
        value = node.text or ""
        for term in terms:
            if term in value:
                value = value.replace(term, "\u2060".join(term))
        node.text = value


def build() -> Path:
    document = Document(ENGLISH)
    reference = Document(TRANSLATION_REFERENCE)

    overrides = {
        0: "知识图谱实体对齐中的关系感知\n邻居上下文与语义引导选择",
        3: "知识图谱实体对齐旨在识别独立构建图谱中指向同一对象的实体。本文提出一种关系感知邻居上下文模型，其中实体语义在联合编码前引导结构证据选择。结构编码器在消息传递过程中保留关系类型，并组合不同传播深度的图状态。语义编码器从同一实体文本序列提取 token、phrase 和 global 三种视图。融合模块使用语义表示为完整的一跳出邻域评分，利用 1.5-entmax 抑制弱相关邻居，并将所得结构上下文与实体语义结合。实验在五个数据集上使用三个随机种子。完整模型在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的 Hits@1 分别达到 0.7285 和 0.6344，较语义分支分别提高 9.09 和 9.66 个百分点。在邻域和可训练融合参数量固定时，语义查询较结构查询分别提高 2.08 和 2.48 个百分点。结果表明，关系类型、结构邻居上下文和语义引导选择能够在不同图谱设置下改善实体对齐。",
        11: "本文研究三个问题。筛选后的关系感知邻居上下文能否提供语义表示之外的对齐证据？在邻域和可训练融合参数量固定时，语义表示是否比结构表示更有效地为邻居评分？关系类型、邻居上下文和多尺度语义编码如何影响联合模型？",
        15: "本文在五个数据集上使用三个随机种子，检验完整模型在不同图谱设置下的表现。匹配对照考察关系类型、多尺度语义、邻居访问和查询信号。",
        79: "3.5.1 语义引导的邻居加权",
        80: "对于实体 {{EQ0}}，融合模块接收四类输入：语义表示 {{EQ1}}、自身结构表示 {{EQ2}}、全部一跳出邻居状态以及有效性掩码。投影层将语义表示转换为查询 {{EQ3}}。模型将结构邻居 {{EQ4}} 投影为键 {{EQ5}} 和值 {{EQ6}}，其缩放点积给出第一个相关性分数。",
        82: "式（14）衡量实体语义与结构邻居 {{EQ0}} 的匹配程度。点积越大，表示语义匹配越强。缩放因子使高维点积保持在稳定范围内。变量 {{EQ1}} 表示查询向量和键向量的维度。",
        87: "门值 {{EQ0}} 表示邻居 {{EQ1}} 通过该特征检查的程度。模型结合点积分数与门值，计算归一化邻居权重。",
        89: "式（17）中，{{EQ0}} 包含全部有效邻居的分数，{{EQ1}} 包含门值，掩码值 {{EQ2}} 用于去除填充位置。softmax 会为每个有效位置赋予正权重；相比之下，1.5-entmax[21] 可以将弱相关邻居的权重精确置零。温度 {{EQ3}} 控制权重的集中程度。有效权重之和仍为 1；当实体不存在有效邻居时，其邻居摘要为零向量。",
        91: "模型使用式（17）的稀疏权重对邻居值向量加权平均，得到结构邻居摘要 {{EQ0}}，对应实体 {{EQ1}}。完整模型在加权前保留全部一跳出邻居，因此该摘要可以利用完整局部邻域，稀疏权重则降低无关邻居的影响。",
        92: "模型随后使用自身结构表示、邻居摘要和语义表示，为每个维度计算上下文门控。",
        101: "语义查询在构造上下文之前改变结构邻居的权重，最终门控再将所得上下文与实体语义结合。第 3.5.4 节通过匹配对照检验查询信号的作用。",
        102: "3.5.4 匹配对照",
        103: "结构查询对照用于检验查询信号。它保持 1.5-entmax、邻域、动态填充、门控和联合检索不变，只将语义查询替换为实体结构表示。参数匹配的晚期融合对照在两个分支独立编码后形成联合表示，且不读取邻居上下文。它与完整模型具有相同数量的可训练融合参数。",
        146: "完整模型相对语义分支的 Hits@1 在 DBP15K 上提高 9.09 个百分点，在 OpenEA 上提高 9.66 个百分点，并且优于无可训练参数的平均融合。这些结果表明，学习型交互能够将结构上下文转化为语义表示之外的对齐证据。结构单分支的较低结果进一步说明，这部分增益依赖结构与语义的联合学习。",
        147: "5.5 邻居查询信号的受控检验",
        148: "表 6 在匹配条件下检验邻居查询信号。结构查询对照共享完整邻域、门控、融合参数和联合检索流程，但使用结构表示为邻居评分。参数匹配晚期融合对照在两个分支独立编码后使用相同数量的融合参数。",
        149: "表 6 邻居访问、查询信号与容量控制结果（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
        150: "语义查询相对结构查询使 DBP15K 和 OpenEA 的 Hits@1 分别提高 2.08 和 2.48 个百分点。由于邻域、融合容量和检索表示保持不变，该差异反映语义引导邻居评分的作用。完整模型相对参数匹配晚期融合分别提高 5.28 和 9.22 个百分点；这一差异反映邻居访问与训练交互路径的共同作用。",
        153: "组件对照识别出三项主要改进来源。关系类型在两个数据集上均带来稳定增益。完整模型优于语义分支，说明结构邻居上下文能够提供文本之外的证据。在邻域和融合容量匹配时，语义查询也优于结构查询，表明实体语义有助于选择相关的局部结构。",
        156: "主实验覆盖四个跨语言图谱对和同语言 EventEA EN–EN 图谱对。组件对照集中于 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2。因此，组件结论来自两个有代表性的跨语言设置，而 EventEA 用于在事件中心图谱对上检验完整配置。三个随机种子用于估计均值和样本波动。固定 8 邻居加 softmax 对照同时改变邻居数量和加权函数，因此其差异反映两个因素。参数匹配晚期融合对照移除邻居访问并改变训练交互路径，其结果衡量整体设计差异。",
        160: "本文考察关系感知结构邻居在结构—语义联合表示中的作用。关系类型在两个受控数据集上增强结构证据。结构邻居上下文使完整模型优于语义分支，语义查询在匹配条件下优于结构查询。具有学习型交互路径的完整邻域访问也优于参数匹配晚期融合。这些结果支持将关系感知邻居上下文和语义引导选择用于实体对齐。五个数据集的实验还表明，完整配置适用于跨语言图谱对和同语言事件中心图谱对。",
        161: "参考文献",
    }

    for target_index, paragraph in enumerate(document.paragraphs):
        mapping = source_mapping(target_index)
        if isinstance(mapping, tuple):
            first = translation_for_source(reference, mapping[0], overrides).rstrip("。")
            second = translation_for_source(reference, mapping[1], overrides)
            translated = first + "。" + second
        else:
            if mapping >= 162:
                continue
            translated = translation_for_source(reference, mapping, overrides)
        rewrite_paragraph(paragraph, translated, bold_prefix_for(paragraph, translated))

    translate_tables(document, reference)
    replace_figure(document)
    translate_headers(document)
    pair_figure_and_caption(document)
    protect_latin_terms(document)
    configure_chinese_styles(document)
    format_all_runs(document)
    set_language(document)
    format_reference_language(document)

    document.core_properties.title = "知识图谱实体对齐中的关系感知邻居上下文与语义引导选择"
    document.core_properties.subject = "关系感知结构编码与语义引导邻居选择"
    document.core_properties.keywords = "知识图谱实体对齐；跨语言知识图谱；事件中心知识图谱；关系感知图神经网络；结构—语义融合；稀疏邻居选择"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
