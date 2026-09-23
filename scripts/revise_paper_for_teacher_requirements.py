import argparse
import json
import os
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from rewrite_chinese_paper_controlled_ablation import (
    add_caption_before,
    add_paragraph_before,
    add_table_before,
    clear_text_highlights,
    find_paragraph,
    format_table,
    insert_after,
    metric,
    replace_heading,
    replace_paragraph,
)


def remove_table(table) -> None:
    table._tbl.getparent().remove(table._tbl)


def remove_body_range(start, end) -> None:
    current = start._p
    end_element = end._p
    parent = current.getparent()
    while current is not end_element:
        next_element = current.getnext()
        parent.remove(current)
        current = next_element


def replace_zip_member(docx_path: Path, member_name: str, replacement_path: Path) -> None:
    temp_path = docx_path.with_suffix(".tmp.docx")
    replacement_data = replacement_path.read_bytes()
    with zipfile.ZipFile(docx_path, "r") as source_zip:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as output_zip:
            for item in source_zip.infolist():
                data = replacement_data if item.filename == member_name else source_zip.read(item.filename)
                output_zip.writestr(item, data)
    os.replace(temp_path, docx_path)


def replace_first_figure(document: Document, output: Path, figure: Path) -> None:
    shape = document.inline_shapes[0]
    blips = shape._inline.xpath(".//*[local-name()='blip']")
    relationship_id = blips[0].get(qn("r:embed"))
    relationship = document.part.rels[relationship_id]
    member_name = f"word/{relationship.target_ref}"
    doc_properties = shape._inline.xpath(".//*[local-name()='docPr']")
    if doc_properties:
        doc_properties[0].set(
            "descr",
            "关系感知结构上下文、多尺度语义编码、固定预算邻居融合、InfoNCE训练和验证驱动检索流程",
        )
    document.save(output)
    replace_zip_member(output, member_name, figure)


def signed_points(delta: float) -> str:
    return f"{delta * 100:+.2f}"


def result_row(indexed, variant: str):
    return indexed[("dbp15k_zh_en", variant)], indexed[("openea_en_fr", variant)]


def table_metric_rows(indexed, variants):
    rows = []
    for variant, label in variants:
        dbp, open_ea = result_row(indexed, variant)
        rows.append([
            label,
            f"{dbp['semantic_weight']:.2f}/{dbp['depth']}/{dbp['csls_k']}",
            f"{metric(dbp, 'Hits@1')} / {metric(dbp, 'MRR')}",
            f"{open_ea['semantic_weight']:.2f}/{open_ea['depth']}/{open_ea['csls_k']}",
            f"{metric(open_ea, 'Hits@1')} / {metric(open_ea, 'MRR')}",
        ])
    return rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument(
        "--results",
        default="outputs/controlled_ablation_20260820/ablation_summary.json",
    )
    parser.add_argument("--figure", required=True)
    args = parser.parse_args()

    source = Path(args.source)
    output = Path(args.output)
    results_path = Path(args.results)
    figure = Path(args.figure)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    document = Document(output)

    results = json.loads(results_path.read_text(encoding="utf-8"))
    indexed = {(row["dataset"], row["variant"]): row for row in results}
    required = {
        (dataset, variant)
        for dataset in ("dbp15k_zh_en", "openea_en_fr")
        for variant in (
            "main",
            "no_relation_types",
            "no_layer_selector",
            "no_token_view",
            "no_phrase_view",
            "no_global_view",
            "token_only",
            "structure_only",
            "semantic_only",
            "mean_fusion",
            "late_concat_matched",
            "structural_neighbor_query",
        )
    }
    missing = sorted(required - set(indexed))
    if missing:
        raise ValueError(f"Missing experiment summaries: {missing}")

    token_dbp, token_open = result_row(indexed, "token_only")
    structure_dbp, structure_open = result_row(indexed, "structure_only")
    semantic_dbp, semantic_open = result_row(indexed, "semantic_only")
    mean_dbp, mean_open = result_row(indexed, "mean_fusion")

    replace_paragraph(
        document,
        "跨语言知识图谱实体对齐需要同时处理",
        "跨语言知识图谱实体对齐需要同时处理关系结构不一致、实体文本碎片化和监督样本有限等问题。本文构建一种关系感知结构上下文与多尺度语义融合方法：关系嵌入和共享投影形成轻量关系消息，节点级层选择器融合不同传播深度，token、phrase 和 global 三个视图编码实体文本，固定预算结构邻居形成局部上下文。模型采用单阶段双向 InfoNCE 训练，并仅依据验证集按数据集选择语义权重、有效传播深度和 CSLS k。五个数据集的三随机种子主结果达到平均 Hits@1 0.7262、MRR 0.7721。在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上进一步完成 66 次同协议控制训练。关系类型和层选择在两个数据集上均有稳定贡献；仅 token 视图的 Hits@1 分别为 "
        f"{token_dbp['Hits@1_mean']:.4f} 和 {token_open['Hits@1_mean']:.4f}。结构单分支、语义单分支和无可训练融合基线进一步区分两个信息源与融合模块的作用。参数匹配晚期融合明显弱于保留邻居上下文的模型，但结构查询与语义查询结果接近，因此现有证据支持关系感知、多深度选择和结构邻居可见性，尚不能证明语义必须在邻居聚合前参与查询。",
    )

    research = replace_paragraph(
        document,
        "本研究围绕四个可检验问题展开",
        "本研究围绕五个可检验问题展开：关系类型和节点级多深度选择是否提高结构区分性；token、token+phrase 和 token+phrase+global 的累积语义视图是否逐步改善结果；结构与语义分支是否分别提供不可替代的信息；在容量相同的条件下，性能差异主要来自查询时机还是结构邻居可见范围；按数据集进行验证选择是否稳定改善检索。由此，本文不预设每个组件都有效，而以多随机种子控制实验区分得到支持、数据集依赖和未获支持的假设。",
    )
    naming = insert_after(
        research,
        "方法名称直接对应代码中的两个可观察机制。“关系感知结构上下文”表示关系类型进入边消息、不同传播深度被节点级选择，并为每个实体提取固定预算邻居；“多尺度语义融合”表示 token、phrase 和 global 视图先形成语义表示，再与结构上下文组合。名称不再使用缺少独立技术定义的“conceptual”或“collaborative”，且“融合”只描述信息流，不预设一定产生增益。",
    )
    insert_after(
        naming,
        "本文的科学新意不在于重新提出 GraphSAGE、Transformer 或 InfoNCE，而在于把轻量关系消息、节点级传播深度选择和固定预算结构邻居置于同一可复现实验协议中，并用参数匹配、信息可见范围和查询信号三个控制变量检验性能来源。这种设计既能确认关系结构组件的作用，也能识别语义查询和部分语义视图未获支持的情形。",
    )

    replace_paragraph(
        document,
        "在五个数据集上完成三随机种子主实验",
        "在五个数据集上完成三随机种子主实验，并在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2 上新增 66 次同协议消融训练。消融覆盖关系类型、层选择、三个语义视图、累积视图、结构/语义单分支、无可训练融合、参数匹配晚期融合和查询信号；所有配置只依据验证集平均 MRR 选择。",
    )

    symbol_table = document.tables[0]
    remove_table(symbol_table)
    replace_paragraph(document, "表 1  主要符号说明", "表 1 主要符号、维度与来源")
    section34 = find_paragraph(document, "3.4 总体架构与交互机制")
    add_table_before(
        document,
        section34,
        ["符号", "维度", "定义与来源"],
        [
            ["G_q=(E_q,R_q,T_q), q∈{L,R}", "集合", "左/右知识图谱及实体、关系、三元组集合"],
            ["A_seed, B", "集合；标量", "种子对齐集合；训练批次中的实体对数量"],
            ["h_i^(l), e_r", "d；d_r", "实体 i 的第 l 层结构状态；关系 r 的嵌入"],
            ["m_(j→i)^(l), m_bar_i^(l)", "d；d", "单条关系消息；实体 i 的平均传入消息"],
            ["W_s^(l), W_r^(l), W_self^(l)", "矩阵", "共享的源实体、关系和自身状态投影"],
            ["a_i^(k), z_i^str", "标量；d", "第 k 个传播深度的选择权重；结构输出"],
            ["X_i, H_i, M_i", "T×300；T×d；T", "词向量输入、共享序列状态和有效 token 掩码"],
            ["z_i^tok, z_i^phr, z_i^glo, z_i^sem", "d", "三个语义视图及其门控融合结果"],
            ["N_K(i), K, c_i^str", "集合；标量；d", "固定预算邻居、邻居数上限和结构上下文"],
            ["q_i, k_j, v_j, α_ij", "d；d；d；标量", "邻居注意力的查询、键、值和归一化权重"],
            ["g_i^c, g_i^j, z_i^joint", "d", "结构上下文门、联合表示门和最终实体表示"],
            ["S_ij, τ, L_align", "标量", "批内相似度、InfoNCE 温度和训练目标"],
            ["α, L, k", "标量；整数；整数", "验证选择的语义权重、有效传播深度和 CSLS 邻域"],
        ],
        [2550, 1500, 5270],
    )
    add_paragraph_before(
        section34,
        "下标 i 和 j 分别表示待编码实体与其候选或邻居，l 表示消息传播层，k 在层选择公式中表示保留状态的深度。d=128，关系嵌入维度 d_r=128。σ 表示 sigmoid，⊙ 表示逐元素乘积，[·;·] 表示向量拼接，softmax 在指定候选维度归一化，LayerNorm 进行特征层归一化，Norm 表示 L2 归一化。",
    )

    replace_paragraph(
        document,
        "每个实体首先具有可学习初始向量",
        "消息传递（message passing）指沿知识图谱有向边把邻居及关系证据逐层送入目标实体。设 h_i^(l)∈R^d 为实体 i 在第 l 层的状态，e_r∈R^(d_r) 为关系 r 的嵌入，N_in(i) 为所有指向 i 的传入边。对于边 (j,r,i)，第 l 层首先计算：",
    )
    replace_paragraph(
        document,
        "其中  和  是共享线性投影",
        "其中 W_s^(l)∈R^(d×d) 与 W_r^(l)∈R^(d×d_r) 是同一层所有关系共享的线性投影；关系差异来自 e_r，而不是每种关系各自拥有一套大矩阵。m_(j→i)^(l)∈R^d 是边 (j,r,i) 传递给目标实体 i 的消息。把所有传入消息取均值得到：",
    )
    replace_paragraph(
        document,
        "为避免只依赖最深层",
        "为避免只依赖最深层，编码器保留投影后的输入状态及全部层输出 u_i^(0),…,u_i^(L)，其中 u_i^(0) 表示未传播的自身信息，u_i^(k) 表示吸收至多 k 跳证据后的状态。平均上下文 u_bar_i 由 L+1 个状态求均值得到，两层 MLP 再为每个深度输出标量分数并通过 softmax 得到 a_i^(k)：",
    )
    replace_paragraph(
        document,
        "输入语义序列为由名称、关系和属性 token 组成的矩阵",
        "输入语义序列 X_i∈R^(T×300) 由名称、关系和属性 token 构成，每一行是一个 300 维词向量。非零行形成掩码 M_i∈{0,1}^T。输入经过 300→128 的线性投影、LayerNorm、正弦位置编码和 dropout 得到 H_i∈R^(T×128)；token、phrase 和 global 三个视图都从同一 H_i 和 M_i 构造，不额外引入外部句子。",
    )
    replace_paragraph(
        document,
        "对于实体 ，输入为语义表示",
        "对于实体 i，融合模块的输入为语义表示 s=z_i^sem∈R^d、自身结构表示 t_i=z_i^str∈R^d、邻居结构矩阵 T_i∈R^(K×d) 和有效掩码。对第 j 个邻居 t_(i,j)，W_q、W_k、W_v∈R^(d×d) 分别产生查询 q_i、键 k_j 和值 v_j；缩放点积注意力首先计算：",
    )

    replace_paragraph(
        document,
        "图 1  关系感知结构—语义早期交互模型",
        "图 1 关系感知结构上下文与多尺度语义融合方法及训练—检索协议。（a）关系类型进入三层全图消息传播，节点级层选择得到结构表示；token、phrase 和 global 视图形成语义表示；主模型使用语义查询固定预算结构邻居，并经结构上下文门和联合表示门输出实体表示。结构查询、单分支、无可训练融合和参数匹配晚期融合仅作为第 5 节控制组。（b）种子对齐用于单阶段双向 InfoNCE 训练，验证集选择 α、L 和 CSLS k，测试集只执行一次最终评估。",
    )
    section35 = find_paragraph(document, "3.5 数据预处理与统一索引")
    add_paragraph_before(
        section35,
        "图 1 与公式的对应关系如下：结构分支接收关系图并由式（1）—（6）输出 z_i^str；语义分支接收 X_i 和 M_i 并由式（7）—（12）输出 z_i^sem；邻居上下文模块由式（13）—（18）输出 c_i^str；联合表示模块由式（19）—（20）输出 z_i^joint；式（22）—（23）对应图 1(b) 的训练目标，式（24）对应验证和测试阶段的 CSLS 检索。",
    )
    replace_paragraph(
        document,
        "模型采用 AdamW 优化器",
        "模型采用 AdamW 优化器[16]。DBP15K 的学习率为 5×10⁻⁴，最多训练 36 个 epoch；OpenEA 与 EventEA 的学习率为 3×10⁻⁴，最多训练 50 个 epoch。所有数据集每 5 个 epoch 在验证集上评估 MRR 并保存最佳检查点，连续 4 次验证未提升时允许提前停止。主实验与新增 66 次消融均不进行结构预热、联合阶段切换或末轮权重平均。",
    )
    replace_paragraph(
        document,
        "DBP15K 的语义序列使用已有预处理特征",
        "DBP15K 的语义序列使用已有预处理特征，其长度与形状见表 2。OpenEA 与 EventEA 的序列按“实体名称—关系名称—属性名称—属性值”的顺序分词，最大长度为 32。OpenEA EN–FR-15K-V2 对实体 URI 进行编码以降低名称偏置[14]，因此名称 token 往往不能直接提供跨语言同名线索。词表内 token 直接使用 GloVe 300 维向量[15]；词表外 token 使用基于 MD5 种子的确定性随机向量。该处理保证复现性，却不能为编码 URI 或专名提供真正的跨语言语义，这构成 OpenEA 结果低于 DBP15K 的重要可能原因。",
    )

    comparison_table = document.tables[3]
    remove_table(comparison_table)
    replace_paragraph(
        document,
        "表 4  与经典方法的非严格对照",
        "表 4 与经典方法的定位性比较（文献协议不完全一致）",
    )
    comparison_anchor = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("本文重新训练模型在 DBP15K")
    )
    add_table_before(
        document,
        comparison_anchor,
        ["数据集", "方法", "来源与设置", "Hits@1", "Hits@10", "MRR"],
        [
            ["DBP15K ZH–EN", "MTransE[1]", "原论文", "0.3083", "0.6141", "未报告"],
            ["", "JAPE[2]", "原论文", "0.4118", "0.7446", "未报告"],
            ["", "BootEA[3]", "原论文", "0.6294", "0.8475", "未报告"],
            ["", "RDGCN[4]", "原论文", "0.7075", "0.8455", "未报告"],
            ["", "RREA-text[5]", "文本增强设置", "0.8220", "0.9640", "未可靠报告"],
            ["", "本文模型", "三种子、当前划分", "0.6809", "0.8119", "0.7289"],
            ["OpenEA EN–FR-V2", "MTransE[14]", "官方五折平均", "0.2404", "0.5251", "0.336"],
            ["", "JAPE[14]", "官方五折平均", "0.2918", "0.6244", "0.402"],
            ["", "GCN-Align[14]", "官方五折平均", "0.4141", "0.7955", "0.542"],
            ["", "BootEA[14]", "官方五折平均", "0.6599", "0.9055", "0.745"],
            ["", "KDCoE[14]", "官方五折平均", "0.7303", "0.8690", "0.778"],
            ["", "RDGCN[14]", "官方五折平均", "0.8469", "0.9339", "0.880"],
            ["", "本文模型", "第1折、三种子", "0.5568", "0.7368", "0.6196"],
        ],
        [1450, 1450, 2800, 1200, 1200, 1220],
    )
    replace_paragraph(
        document,
        "本文重新训练模型在 DBP15K",
        "在 DBP15K ZH–EN 上，本文模型的 Hits@1 高于 MTransE、JAPE 和 BootEA，但低于 RDGCN 与 RREA-text。在 OpenEA EN–FR-15K-V2 上，本文结果高于官方 MTransE、JAPE 和 GCN-Align，低于 BootEA、KDCoE 和 RDGCN[14]。因此本文不宣称达到最新最佳性能，其价值主要来自统一协议下对关系消息、传播深度、语义尺度、分支必要性和邻居上下文的受控证据。OpenEA 文献值为官方五折平均，而本文只使用第 1 折并报告三随机种子，两组数值只能定位性能区间，不能作严格排名。",
    )

    replace_paragraph(
        document,
        "表 3、表 5–8 和图 2 来自",
        "表 3、表 5–8 和图 2 来自五个数据集、随机种子 42、43、44 的 15 次主模型训练。表 9–12 在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2 上复用同协议主模型结果，并新增 11 个变体×2 个数据集×3 个种子，共 66 次训练。所有运行采用相同数据划分、InfoNCE 目标、验证频率和提前停止规则；除单分支基线按定义固定 α=0 或 1 外，每个变体分别依据三个种子的平均验证 MRR 选择 α、L、k。表 4 的文献数值仅用于定位性能区间。",
    )
    replace_paragraph(
        document,
        "本文以 Hits@1、Hits@10",
        "本文以 Hits@1、Hits@10 和平均倒数排名（MRR）衡量完整候选集合上的检索质量。五数据集主结果及十一项消融均报告随机种子 42、43、44 的均值与样本标准差。每个数据集—变体组合只使用训练集或验证集选择检查点和检索配置，测试集仅用于唯一配置确定后的最终报告。由于 n=3 不足以支持可靠的正态性与显著性检验，本文报告绝对差值和跨种子离散程度，不使用夸大的显著性措辞。",
    )
    replace_paragraph(
        document,
        "五数据集主模型始终保留 token",
        "五数据集主模型始终保留 token、phrase、global 三个语义视图，global 分支及其两层 Transformer 未作结构改动。新增实验只在两个代表性数据集上构造 token、token+phrase 和完整三视图的累积组合，并设置结构单分支、语义单分支及无可训练融合基线，因此不会改变五数据集主配置。",
    )
    replace_paragraph(
        document,
        "DBP15K 的选择后平均 Hits@1",
        "DBP15K 的选择后平均 Hits@1 为 0.7726，EventEA 为 0.7563，OpenEA EN–FR-15K-V2 为 0.5568。OpenEA V2 通过编码 URI 降低名称偏置[14]，而当前词表外 token 只获得确定性随机向量；同时本文采用单阶段 InfoNCE，不使用 BootEA 式自举或 KDCoE 式描述协同。因此该数据集上的绝对结果低于强文献方法并不意外。验证选择把 Hits@1 从 0.5230 提高到 0.5568，说明检索校准有效，但不能弥补输入语义和训练策略的差距。",
    )

    section57 = find_paragraph(document, "5.7 关系、层选择与语义视图消融")
    section6 = find_paragraph(document, "6 讨论")
    remove_body_range(section57, section6)

    add_paragraph_before(section6, "5.7 结构组件与累积语义视图消融", "Heading 2")
    add_paragraph_before(
        section6,
        "表 9 同时报告关系类型、层选择、单视图删除和老师要求的累积语义组合。完整模型即 token+phrase+global；token+phrase 复用去除 global 的同协议运行。每个配置训练三个随机种子并独立进行验证选择。",
    )
    add_caption_before(section6, "表 9 结构组件与语义视图消融（Hits@1/MRR，均值±样本标准差，n=3）")
    add_table_before(
        document,
        section6,
        ["变体", "DBP α/L/k", "DBP H@1 / MRR", "OpenEA α/L/k", "OpenEA H@1 / MRR"],
        table_metric_rows(
            indexed,
            [
                ("main", "完整模型：token+phrase+global"),
                ("no_relation_types", "去除关系类型"),
                ("no_layer_selector", "去除层选择器"),
                ("token_only", "仅 token"),
                ("no_global_view", "token+phrase"),
                ("no_token_view", "phrase+global（去除 token）"),
                ("no_phrase_view", "token+global（去除 phrase）"),
            ],
        ),
        [2000, 1150, 2260, 1150, 2760],
    )
    add_paragraph_before(
        section6,
        "去除关系类型后，DBP15K 与 OpenEA 的 Hits@1 分别下降 3.40 和 2.35 个百分点；去除层选择器分别下降 2.15 和 1.90 个百分点，两项结构组件在两个数据集上的方向一致。累积视图方面，仅 token 的 Hits@1 分别为 "
        f"{token_dbp['Hits@1_mean']:.4f} 和 {token_open['Hits@1_mean']:.4f}；加入 phrase 后对应 token+phrase 结果为 {indexed[('dbp15k_zh_en', 'no_global_view')]['Hits@1_mean']:.4f} 和 {indexed[('openea_en_fr', 'no_global_view')]['Hits@1_mean']:.4f}；完整三视图为 0.6809 和 0.5568。OpenEA 的主要增益出现在加入 phrase 时，global 未继续稳定改善；DBP15K 三种组合差异较小。单项删除结果与此一致：去除 phrase 在 OpenEA 上下降 4.24 个百分点，而去除 token 或 global 未产生跨数据集一致下降。",
    )

    add_paragraph_before(section6, "5.8 结构分支、语义分支与融合模块", "Heading 2")
    add_paragraph_before(
        section6,
        "为回答结构分支和融合模块是否必要，表 10 比较完整模型、结构单分支、语义单分支和无可训练融合。单分支训练目标只作用于保留的表示；验证检索分别固定 α=0 和 α=1，防止被删除分支通过后处理重新进入。无可训练融合使用验证选择的 α 对两个分支加权平均，不含融合 MLP、邻居注意力或门控参数。",
    )
    add_caption_before(section6, "表 10 分支与融合必要性实验（均值±样本标准差，n=3）")
    add_table_before(
        document,
        section6,
        ["变体", "训练/检索表示", "DBP H@1 / MRR", "OpenEA H@1 / MRR"],
        [
            ["完整模型", "结构邻居上下文 + 语义 + 门控", f"{metric(indexed[('dbp15k_zh_en', 'main')], 'Hits@1')} / {metric(indexed[('dbp15k_zh_en', 'main')], 'MRR')}", f"{metric(indexed[('openea_en_fr', 'main')], 'Hits@1')} / {metric(indexed[('openea_en_fr', 'main')], 'MRR')}"],
            ["仅结构分支", "z^str，α=0", f"{metric(structure_dbp, 'Hits@1')} / {metric(structure_dbp, 'MRR')}", f"{metric(structure_open, 'Hits@1')} / {metric(structure_open, 'MRR')}"],
            ["仅语义分支", "z^sem，α=1", f"{metric(semantic_dbp, 'Hits@1')} / {metric(semantic_dbp, 'MRR')}", f"{metric(semantic_open, 'Hits@1')} / {metric(semantic_open, 'MRR')}"],
            ["无可训练融合", "αz^sem+(1−α)z^str", f"{metric(mean_dbp, 'Hits@1')} / {metric(mean_dbp, 'MRR')}", f"{metric(mean_open, 'Hits@1')} / {metric(mean_open, 'MRR')}"],
        ],
        [1800, 2800, 2360, 2360],
    )
    add_paragraph_before(
        section6,
        "相对完整模型，仅结构分支的 Hits@1 变化为 DBP "
        f"{signed_points(structure_dbp['delta_Hits@1'])}、OpenEA {signed_points(structure_open['delta_Hits@1'])} 个百分点，表明随机初始化的结构表示在有限种子监督下不能独立建立可靠的跨图坐标系；仅语义分支分别变化 {signed_points(semantic_dbp['delta_Hits@1'])} 和 {signed_points(semantic_open['delta_Hits@1'])} 个百分点，说明语义提供主要跨图锚点，而结构上下文仍带来额外消歧。无可训练融合在 DBP 的 Hits@1 仅变化 {signed_points(mean_dbp['delta_Hits@1'])} 个百分点，但 MRR 下降 {abs(mean_dbp['delta_MRR']) * 100:.2f} 个百分点；在 OpenEA 上 Hits@1 和 MRR 分别下降 {abs(mean_open['delta_Hits@1']) * 100:.2f} 和 {abs(mean_open['delta_MRR']) * 100:.2f} 个百分点。因此，可训练融合的证据在 OpenEA 上更清楚，在 DBP15K 上主要体现为排序质量而非首位命中，不能概括为跨数据集同等幅度的增益。",
    )

    add_paragraph_before(section6, "5.9 交互位置、邻居可见范围与容量控制", "Heading 2")
    add_paragraph_before(
        section6,
        "表 11 固定结构与语义编码器，并使三种融合路径具有相同的 197,377 个有效参数，以分离容量、结构邻居可见范围和查询信号。",
    )
    add_caption_before(section6, "表 11 参数匹配融合控制的设计差异")
    add_table_before(
        document,
        section6,
        ["变体", "结构邻居可见", "邻居查询信号", "语义进入位置 / 融合参数"],
        [
            ["完整模型", "是", "语义表示", "邻居聚合前 + 最终门控 / 197,377"],
            ["无语义查询邻居", "是", "自身结构表示", "仅最终门控 / 197,377"],
            ["参数匹配晚期融合", "否", "无邻居查询", "独立编码后末端融合 / 197,377"],
        ],
        [1800, 1600, 2200, 3520],
    )
    add_caption_before(section6, "表 12 参数匹配融合控制的测试结果（均值±样本标准差，n=3）")
    add_table_before(
        document,
        section6,
        ["变体", "DBP H@1", "DBP MRR", "OpenEA H@1", "OpenEA MRR"],
        [
            [label, metric(dbp, "Hits@1"), metric(dbp, "MRR"), metric(open_ea, "Hits@1"), metric(open_ea, "MRR")]
            for variant, label in (
                ("main", "完整模型"),
                ("structural_neighbor_query", "无语义查询邻居"),
                ("late_concat_matched", "参数匹配晚期融合"),
            )
            for dbp, open_ea in [result_row(indexed, variant)]
        ],
        [1800, 1830, 1830, 1830, 1830],
    )
    add_paragraph_before(
        section6,
        "与完整模型相比，参数匹配晚期融合在 DBP15K 和 OpenEA 上的 Hits@1 分别下降 1.32 和 7.52 个百分点，说明相同末端参数容量不能替代结构邻居上下文。无语义查询邻居基线的 Hits@1 仅变化 +0.23 和 +0.02 个百分点，MRR 仅变化 +0.17 和 −0.02 个百分点，均小于跨种子标准差。因此，结构邻居可见性得到支持，语义必须在聚合前参与查询的假设没有得到支持。",
    )

    replace_paragraph(
        document,
        "主实验与受控消融共同支持四点",
        "主实验与 66 次受控消融共同区分了四类结论。第一，关系类型和节点级层选择在两个代表性数据集上均产生方向一致的增益。第二，累积语义实验表明 phrase 对 OpenEA 重要，而 global 未在当前协议下继续带来稳定增益；DBP15K 对语义视图组合不敏感。第三，结构/语义单分支和无可训练融合明确显示两个信息源及融合机制的相对作用。第四，参数匹配晚期融合显著弱于保留邻居上下文的模型，但结构查询与语义查询几乎相同，因此主要证据落在结构邻居可见性而非查询时机。",
    )
    replace_heading(document, "6.2 结果解释与未获支持的假设", "6.2 OpenEA 低结果与未获支持的假设")
    replace_paragraph(
        document,
        "消融结果对最初假设形成了区分",
        "OpenEA EN–FR-15K-V2 的 Hits@1 为 0.5568，低于官方 BootEA、KDCoE 和 RDGCN 五折均值[14]。一方面，V2 对 URI 编码以降低名称偏置，当前 MD5 随机 OOV 向量无法恢复专名的跨语言语义；另一方面，本文坚持单阶段 InfoNCE，没有使用伪标签扩充、描述协同或更强负样本策略。内部基线显示，加入 phrase 能明显改善 token-only，但 global 和语义查询没有稳定独立增益。由此，较低结果既反映输入与训练策略的限制，也说明不能把所有多尺度或早期交互组件都写成已验证创新。",
    )
    replace_paragraph(
        document,
        "内部有效性方面",
        "内部有效性方面，所有新增消融使用相同数据划分、训练目标和三随机种子，但 n=3 仍不足以支持高功效显著性检验；每个变体独立选择配置虽保证公平，却也引入验证网格方差。单分支与无融合实验只覆盖 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2，不能直接推广到其余三个主数据集。外部有效性方面，尚未验证 100K 以上规模、开放世界和无对应实体场景。OpenEA 官方结果为五折平均，而本文只运行第 1 折；表 4 因输入、划分和后处理不同不能替代统一复现。",
    )
    replace_paragraph(
        document,
        "后续研究应优先解决三个问题",
        "后续研究应优先解决三个问题。第一，以跨语言子词或编码名称表示替代随机 OOV 向量，并在 OpenEA 五个折上复现。第二，重新设计语义查询或采用对比约束，使语义在邻居选择阶段产生可测量而非冗余的作用。第三，在更大规模、开放世界、无对应实体和跨知识库来源场景中评估效率、校准稳定性与泛化性。",
    )
    replace_paragraph(
        document,
        "两个代表性数据集上的 42 次新增训练",
        "两个代表性数据集上的 66 次新增训练进一步限定了方法贡献：关系类型和层选择得到一致支持；累积视图显示 phrase 在 OpenEA 上承担主要语义增益，global 未显示稳定附加贡献；结构单分支、语义单分支和无可训练融合明确了信息源与融合模块的作用；参数匹配晚期融合弱于结构邻居上下文，而语义提前查询邻居没有稳定独立增益。",
    )
    replace_paragraph(
        document,
        "因此，本文的科学结论不是",
        "因此，本文的科学结论不是“所有模块都有效”或“达到最新最佳性能”，而是通过累积视图、单分支、无融合、容量匹配和查询信号控制，将可靠证据定位到关系感知、多深度结构选择和结构邻居上下文，并明确 OpenEA 结果、global 视图和语义查询的适用边界。",
    )

    document.core_properties.subject = "教师要求逐项核验与统一协议补充消融版"
    clear_text_highlights(document)
    replace_first_figure(document, output, figure)

    verified = Document(output)
    if len(verified.tables) != 12:
        raise RuntimeError(f"Expected 12 tables, found {len(verified.tables)}")
    if len(verified.inline_shapes) != 2:
        raise RuntimeError(f"Expected 2 figures, found {len(verified.inline_shapes)}")
    print(f"WROTE {output}")


if __name__ == "__main__":
    main()
