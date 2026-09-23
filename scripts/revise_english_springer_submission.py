from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "Relation_Aware_Neighbor_Context_Springer_English_No_Line_Numbers_20260830.docx"
OUTPUT = ROOT / "outputs" / "Relation_Aware_Neighbor_Context_Springer_English_Formatted_20260830.docx"
FIGURE = ROOT / "outputs" / "paper_assets" / "figure_1_springer_submission_600dpi.png"


def combined_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).strip()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if combined_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def find_caption(document: Document, prefix: str) -> Paragraph:
    matches = [
        p for p in document.paragraphs
        if p.style.name == "Caption" and combined_text(p).startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one caption beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def set_plain_text(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_text_nodes(root, replacements: list[tuple[str, str]]) -> None:
    for node in root.xpath(".//w:t"):
        value = node.text or ""
        for source, target in replacements:
            value = value.replace(source, target)
        node.text = value


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = style
    new_paragraph.add_run(text)
    return new_paragraph


def set_caption_label(paragraph: Paragraph, label: str) -> None:
    prefix = f"{label} "
    first_text = next((node for node in paragraph._p.xpath(".//w:t") if (node.text or "").startswith(prefix)), None)
    if first_text is None:
        raise RuntimeError(f"Caption does not begin with {label!r}: {combined_text(paragraph)!r}")
    first_text.text = (first_text.text or "")[len(prefix):]
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_properties.append(OxmlElement("w:b"))
    run.append(run_properties)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = prefix
    run.append(text)
    paragraph._p.insert(1, run)


def rewrite_plain_caption(paragraph: Paragraph, label: str, body: str) -> None:
    paragraph.clear()
    label_run = paragraph.add_run(f"{label} ")
    label_run.bold = True
    paragraph.add_run(body.rstrip("."))
    paragraph.paragraph_format.keep_with_next = True


def remove_caption_final_period(paragraph: Paragraph) -> None:
    for node in reversed(paragraph._p.xpath(".//w:t")):
        value = node.text or ""
        if value:
            if value.endswith("."):
                node.text = value[:-1]
            return


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0


def remove_table_row_by_label(table, label: str) -> None:
    matches = [row for row in table.rows if row.cells and row.cells[0].text.strip() == label]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table row labelled {label!r}, found {len(matches)}")
    row = matches[0]
    row._tr.getparent().remove(row._tr)


def set_table_widths(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for grid_col, width in zip(list(grid), widths):
        grid_col.set(qn("w:w"), str(width))
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(widths)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))


def format_tables(document: Document) -> None:
    widths = {
        0: [1700, 1250, 1350, 1800, 1900, 1072],
        1: [2450, 2200, 2200, 2222],
        2: [1800, 1900, 1400, 1400, 2572],
        3: [2400, 3336, 3336],
        4: [1500, 1650, 1650, 2136, 2136],
        5: [1900, 2500, 1200, 1736, 1736],
    }
    left_columns = {
        0: set(),
        1: set(),
        2: set(),
        3: {0},
        4: {0, 1, 2},
        5: {0, 1},
    }
    for table_index, table in enumerate(document.tables):
        set_table_widths(table, widths[table_index])
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT
                        if row_index > 0 and column_index in left_columns[table_index]
                        else WD_ALIGN_PARAGRAPH.CENTER
                    )
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(8.5)
                        if row_index == 0:
                            run.bold = True


def bold_label(paragraph: Paragraph, label: str) -> None:
    text = paragraph.text
    if not text.startswith(label):
        raise RuntimeError(f"Expected label {label!r} in {text!r}")
    remainder = text[len(label):].lstrip()
    paragraph.clear()
    label_run = paragraph.add_run(label)
    label_run.bold = True
    if remainder:
        paragraph.add_run(f"  {remainder}")


def replace_figure(document: Document) -> None:
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)
    shape = document.inline_shapes[0]
    relation_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    document.part.related_parts[relation_id]._blob = FIGURE.read_bytes()
    with Image.open(FIGURE) as image:
        width_px, height_px = image.size
    shape.width = Inches(6.35)
    shape.height = Inches(6.35 * height_px / width_px)
    doc_properties = shape._inline.docPr
    doc_properties.set("name", "Figure 1. Proposed model information flow")
    doc_properties.set(
        "descr",
        "Two-panel diagram of the relation-aware structural branch, multi-scale semantic branch, "
        "semantics-guided one-hop context module, training objective, and joint-representation retrieval.",
    )


def set_math_vectors(document: Document) -> None:
    vector_symbols = {"W", "H", "M", "h", "e", "m", "s", "t", "u", "z", "f", "g", "c", "q", "k", "v", "x", "y", "α"}
    active = False
    for paragraph in document.paragraphs:
        text = combined_text(paragraph)
        if text.startswith("3.3 "):
            active = True
        if text.startswith("3.6 "):
            active = False
        if not active:
            continue
        for math_run in paragraph._p.xpath(".//m:r"):
            token = "".join((node.text or "") for node in math_run.iter(qn("m:t")))
            if token not in vector_symbols:
                continue
            run_properties = math_run.find(qn("m:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("m:rPr")
                math_run.insert(0, run_properties)
            style = run_properties.find(qn("m:sty"))
            if style is None:
                style = OxmlElement("m:sty")
                run_properties.append(style)
            style.set(qn("m:val"), "bi")


def set_document_language(document: Document) -> None:
    document.core_properties.language = "en-US"
    for part in document.part.package.parts:
        element = getattr(part, "_element", None)
        if element is None:
            continue
        for lang in element.xpath(".//w:lang"):
            lang.set(qn("w:val"), "en-US")
            lang.set(qn("w:eastAsia"), "en-US")
            lang.set(qn("w:bidi"), "en-US")


def add_table_spacing(document: Document) -> None:
    children = list(document._body._element)
    for index, child in enumerate(children[:-1]):
        if child.tag != qn("w:tbl"):
            continue
        next_child = children[index + 1]
        if next_child.tag == qn("w:p"):
            paragraph = Paragraph(next_child, document._body)
            paragraph.paragraph_format.space_before = Pt(8)


def remove_line_numbering(document: Document) -> None:
    for section_properties in document._element.xpath(".//w:sectPr"):
        line_numbering = section_properties.find(qn("w:lnNumType"))
        if line_numbering is not None:
            section_properties.remove(line_numbering)


def revise() -> Path:
    document = Document(SOURCE)

    document.core_properties.title = (
        "Relation-Aware Neighbor Context and Semantics-Guided Selection for Cross-Lingual Entity Alignment"
    )
    document.core_properties.subject = "Cross-lingual knowledge graph entity alignment"
    document.core_properties.keywords = (
        "knowledge graph entity alignment; cross-lingual knowledge graphs; relation-aware graph neural networks; "
        "structural–semantic fusion; sparse neighbor selection; multi-scale semantic encoding"
    )

    set_plain_text(
        document.paragraphs[0],
        "Relation-Aware Neighbor Context and Semantics-Guided Selection for Cross-Lingual Entity Alignment",
    )
    set_plain_text(
        document.paragraphs[2],
        "Affiliation: [Department, Institution, City, Country — complete before submission]",
    )
    set_plain_text(
        document.paragraphs[5],
        "Cross-lingual knowledge graph entity alignment identifies matching entities across graphs written in "
        "different languages. Most systems encode graph structure and entity text in separate branches. This design "
        "leaves an open question: how should structural neighbors contribute before the model forms a joint "
        "representation? This study evaluates a relation-aware neighbor-context model with semantics-guided selection. "
        "Its structural encoder preserves relation types and several propagation depths. Its semantic encoder builds "
        "token, phrase, and global views from one entity-text sequence. A semantic query scores the complete one-hop "
        "outgoing neighborhood. The 1.5-entmax function removes weak evidence. The model then combines the selected "
        "structural context with entity semantics. Experiments use three random seeds on five datasets. On DBP15K "
        "ZH–EN and OpenEA EN–FR-15K-V2, the full model obtains hits at rank 1 (Hits@1) scores of 0.7285 and 0.6344. "
        "Under the same neighborhood and trainable fusion parameter count, semantic queries improve Hits@1 over "
        "structural queries by 2.08 and 2.48 percentage points. The full model also improves over semantic-only "
        "retrieval by 9.09 and 9.66 points. Removing relation types lowers Hits@1 by 3.62 and 2.78 points. A "
        "parameter-matched late-fusion control falls by 5.28 and 9.22 points. This difference reflects both neighbor "
        "access and the training interaction path. The results support relation-aware neighbor context and "
        "semantics-guided selection under the tested settings.",
    )
    set_plain_text(
        document.paragraphs[6],
        "Keywords: knowledge graph entity alignment; cross-lingual knowledge graphs; relation-aware graph neural "
        "networks; structural–semantic fusion; sparse neighbor selection; multi-scale semantic encoding",
    )

    set_plain_text(
        find_paragraph(document, "This study asks three questions"),
        "This study asks three questions. First, does relation-aware neighbor context add useful evidence beyond a "
        "semantic representation? Second, does a semantic query select neighbors better than a structural query when "
        "the neighborhood and trainable fusion parameter count are fixed? Third, which structural and semantic "
        "components provide consistent support for the joint model?",
    )
    set_plain_text(
        find_paragraph(document, "The model preserves relation types"),
        "The model preserves relation types during structural message passing. A node-level layer selector lets each "
        "entity combine shallow and deep graph states. This design provides a compact relation-aware state without a "
        "separate full transformation matrix for every relation.",
    )
    set_plain_text(
        find_paragraph(document, "The fusion module uses semantics"),
        "The model builds token, phrase, and global views from one entity-text sequence. It uses the resulting semantic "
        "state to score the complete one-hop outgoing neighborhood. The 1.5-entmax function removes weak neighbors "
        "before the model forms the joint representation.",
    )
    set_plain_text(
        find_paragraph(document, "The experiments use three seeds"),
        "The experiments cover five datasets and three random seeds. Matched controls isolate relation types, semantic "
        "views, neighbor access, and the query signal. The query comparison "
        "keeps the neighborhood and trainable fusion parameter count fixed.",
    )

    set_plain_text(
        find_paragraph(document, "The model has four parts"),
        "The model has four parts. The structural encoder preserves relation types and combines several propagation "
        "depths. The semantic encoder builds token, phrase, and global views from the same entity text. The fusion "
        "module uses semantics to score every one-hop outgoing neighbor. It forms a sparse structural context and "
        "combines that context with entity semantics. The training stage uses information noise-contrastive estimation "
        "(InfoNCE) on the joint and structural representations. The retrieval stage uses the trained joint "
        "representation with cross-domain similarity local scaling (CSLS). Figure 1 shows the full information flow.",
    )

    direction_anchor = find_paragraph(document, "During each forward pass")
    insert_after(
        direction_anchor,
        "The structural encoder and the context module use different edge roles. Incoming edges update the current node "
        "during graph propagation. Outgoing neighbors represent facts emitted by the current entity and serve as "
        "candidate context during fusion. This distinction explains why the two modules use opposite edge directions.",
    )

    overview_caption = find_paragraph(document, "Fig. 1")
    replace_text_nodes(
        overview_caption._p,
        [
            ("A semantic query", "A semantics-guided query"),
            ("Validation selects only", "Validation selects only"),
        ],
    )
    set_caption_label(overview_caption, "Fig. 1")
    remove_caption_final_period(overview_caption)
    overview_caption.paragraph_format.keep_with_next = False

    set_plain_text(
        find_paragraph(document, "Table 4 reports component controls"),
        "Table 4 reports component controls on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. Every variant uses three random "
        "seeds. Each variant also uses the same representation for training and retrieval.",
    )
    set_plain_text(
        find_paragraph(document, "Relation types provide the clearest"),
        "Relation types provide the clearest consistent structural gain. Their removal lowers Hits@1 by 3.62 points on "
        "DBP15K and 2.78 points on OpenEA. Topology features also support the structural state, with gains of 3.46 and "
        "8.03 points over entity vectors alone. The phrase view improves both datasets by 0.81 and 3.95 points. The "
        "remaining controls show smaller or dataset-specific changes. They are therefore treated as supporting design "
        "choices rather than separate contributions.",
    )
    set_plain_text(
        find_paragraph(document, "Table 6 compares the full model"),
        "Table 6 compares the full model with two controls. The structural-query control changes only the signal that "
        "scores neighbors. The parameter-matched late-fusion control keeps the same number of trainable fusion parameters "
        "but does not read neighbor context.",
    )
    set_plain_text(
        find_paragraph(document, "The matched comparison shows"),
        "The matched comparison shows that structural queries yield Hits@1 values 2.08 and 2.48 points below semantic "
        "queries. This result isolates the benefit of semantics-guided neighbor scoring under the current protocol. The "
        "parameter-matched late-fusion control is 5.28 and 9.22 points below the full model. This second gap reflects both "
        "neighbor access and the training interaction path. It does not measure fusion timing alone because two factors "
        "change together.",
    )
    set_plain_text(
        find_paragraph(document, "The experiments support three claims"),
        "The experiments support three claims. First, relation types provide a consistent structural gain on both "
        "controlled datasets. Second, structural context "
        "adds value when the model learns it together with semantics. Third, semantic queries outperform structural "
        "queries when the neighborhood and trainable fusion parameter count are fixed. Together, these findings support "
        "relation-aware neighbor context with semantics-guided selection under the tested protocol.",
    )
    set_plain_text(
        find_paragraph(document, "The study uses"),
        "The study uses three random seeds, so it reports effect size and sample variation without strong significance "
        "claims. The fixed-neighbor control changes both neighbor count and the weight function. The parameter-matched "
        "late-fusion control changes both neighbor access and "
        "the training interaction path. The component controls cover two main datasets. The evidence therefore supports "
        "the tested settings rather than every entity-alignment setting.",
    )
    set_plain_text(
        find_paragraph(document, "This study asks where structural neighbors"),
        "This study examines how structural neighbors should enter structural–semantic entity alignment. Relation types "
        "improve structural evidence on both controlled datasets. A learned joint representation makes structural "
        "context useful beyond semantics alone. Semantic "
        "queries also outperform structural queries under matched conditions. Complete neighbor access with a learned "
        "interaction path performs better than late fusion without neighbor context, but this difference does not isolate "
        "fusion timing. The results support relation-aware neighbor context and semantics-guided selection under the "
        "evaluated protocol.",
    )

    global_replacements = [
        ("Semantic-Guided", "Semantics-Guided"),
        ("Semantic-guided", "Semantics-Guided"),
        ("semantic-guided", "semantics-guided"),
        ("Semantically guided", "Semantics-Guided"),
        ("Structure-semantic", "Structural–Semantic"),
        ("structure-semantic", "structural–semantic"),
        ("DBP15K zh-en", "DBP15K ZH–EN"),
        ("DBP15K ja-en", "DBP15K JA–EN"),
        ("DBP15K fr-en", "DBP15K FR–EN"),
        ("DBP15K ZH-EN", "DBP15K ZH–EN"),
        ("OpenEA en-fr", "OpenEA EN–FR-15K-V2"),
        ("OpenEA EN-FR-V2", "OpenEA EN–FR-15K-V2"),
        ("OpenEA EN–FR-V2", "OpenEA EN–FR-15K-V2"),
        ("OpenEA EN-FR-15K-V2", "OpenEA EN–FR-15K-V2"),
        ("EventEA en-en", "EventEA EN–EN-20K"),
        ("under the same neighborhood and model size", "under the same neighborhood and trainable fusion parameter count"),
        ("same neighborhood and model size", "same neighborhood and trainable fusion parameter count"),
        ("three seeds", "three random seeds"),
        ("training path", "training interaction path"),
        ("node-level selector", "node-level layer selector"),
        ("Structural neighbor query", "Structural-query control"),
        ("Parameter-matched late fusion", "Parameter-matched late-fusion control"),
        ("simple mean fusion", "parameter-free mean fusion"),
        ("mean fusion without trainable parameters", "parameter-free mean fusion"),
        ("Neighbor visibility", "Neighbor access"),
    ]
    replace_text_nodes(document._element, global_replacements)

    set_plain_text(find_paragraph(document, "2.3 "), "2.3 Structural–Semantic Interaction")
    set_plain_text(find_paragraph(document, "3.5 Semantics-Guided"), "3.5 Semantics-Guided Structural Context")
    set_plain_text(
        find_paragraph(document, "3.5.1 Semantics-Guided"),
        "3.5.1 Semantics-Guided Neighbor Weighting",
    )

    replace_text_nodes(
        find_paragraph(document, "The deepest layer may lose")._p,
        [("A two-layer MLP then", "A two-layer multilayer perceptron (MLP) then")],
    )
    replace_text_nodes(
        find_paragraph(document, "The layer-selection MLP")._p,
        [
            ("It uses GELU", "It uses the Gaussian error linear unit (GELU)"),
        ],
    )
    replace_text_nodes(
        find_paragraph(document, "Text-based entity alignment studies")._p,
        [
            (
                "holds the neighborhood and model size constant",
                "holds the neighborhood and trainable fusion parameter count constant",
            )
        ],
    )
    replace_text_nodes(
        find_paragraph(document, "The terms")._p,
        [("validation MRR selects", "validation mean reciprocal rank (MRR) selects")],
    )
    replace_text_nodes(
        find_paragraph(document, "The model uses AdamW")._p,
        [
            ("GNN layers", "graph neural network layers"),
            (
                "DBP15K training lasts at most 36 epochs. OpenEA and EventEA training lasts at most 50 epochs.",
                "Training lasts at most 36 epochs for DBP15K and 50 epochs for OpenEA and EventEA.",
            ),
        ],
    )

    table1_caption = find_paragraph(document, "Table 1")
    intro = table1_caption.insert_paragraph_before(
        "Table 1 summarizes the dataset statistics and experimental splits.", style="Normal"
    )
    intro.paragraph_format.keep_with_next = True

    captions = {
        "Table 1": "Statistics and experimental splits of the five datasets",
        "Table 2": "Final test results on five datasets (mean ± sample standard deviation, n = 3)",
        "Table 3": "Contextual comparison between published results and the proposed model",
        "Table 4": "Ablation of structural components, neighbor aggregation, and semantic views (Hits@1/MRR, mean ± sample standard deviation, n = 3)",
        "Table 5": "Controls for branches, training representations, and primary retrieval representations (Hits@1/MRR, mean ± sample standard deviation, n = 3)",
        "Table 6": "Neighbor access, query signal, and capacity controls (Hits@1/MRR, mean ± sample standard deviation, n = 3)",
    }
    for label, body in captions.items():
        rewrite_plain_caption(find_caption(document, label), label, body)

    dataset_names = [
        "DBP15K ZH–EN",
        "DBP15K JA–EN",
        "DBP15K FR–EN",
        "OpenEA EN–FR-15K-V2",
        "EventEA EN–EN-20K",
    ]
    for table_index in (0, 1):
        for row_index, dataset_name in enumerate(dataset_names, start=1):
            set_cell_text(document.tables[table_index].cell(row_index, 0), dataset_name)
    set_cell_text(document.tables[2].cell(7, 0), "OpenEA EN–FR-15K-V2")

    set_cell_text(document.tables[3].cell(0, 1), "DBP15K Hits@1/MRR")
    set_cell_text(document.tables[3].cell(0, 2), "OpenEA Hits@1/MRR")
    set_cell_text(document.tables[4].cell(0, 3), "DBP15K Hits@1/MRR")
    set_cell_text(document.tables[4].cell(0, 4), "OpenEA Hits@1/MRR")
    set_cell_text(document.tables[5].cell(0, 3), "DBP15K Hits@1/MRR")
    set_cell_text(document.tables[5].cell(0, 4), "OpenEA Hits@1/MRR")
    set_cell_text(document.tables[5].cell(1, 0), "Full model")
    set_cell_text(document.tables[5].cell(2, 0), "Structural-query control")
    set_cell_text(document.tables[5].cell(3, 0), "Parameter-matched late-fusion control")
    set_cell_text(
        document.tables[5].cell(3, 1),
        "No neighbor context; form and retrieve the late-fusion representation after independent branches",
    )
    remove_table_row_by_label(document.tables[3], "Without structural supervision")
    format_tables(document)

    declarations = find_paragraph(document, "Declarations")
    set_plain_text(declarations, "Statements and Declarations")
    declarations.style = "Heading 1"
    set_plain_text(
        find_paragraph(document, "Funding"),
        "Funding  The author received no specific funding for this work.",
    )
    bold_label(document.paragraphs[6], "Keywords:")
    for label in (
        "Funding",
        "Author contributions",
        "Data availability",
        "Code availability",
        "Competing interests",
        "Ethics approval",
        "Consent to participate",
        "Consent for publication",
    ):
        bold_label(find_paragraph(document, label), label)
    reference_eight = find_paragraph(document, "[8]")
    if "10.14778/3407790.3407828" not in combined_text(reference_eight):
        for run in reference_eight.runs:
            if run.text:
                run.text = run.text.rstrip() + " https://doi.org/10.14778/3407790.3407828"
                break

    replace_figure(document)
    set_math_vectors(document)
    add_table_spacing(document)
    set_document_language(document)
    remove_line_numbering(document)

    for paragraph in document.paragraphs:
        if paragraph.style.name == "Caption":
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
        if paragraph.style.name == "Formula":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(revise())
