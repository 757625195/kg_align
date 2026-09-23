from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "Springer引用与逻辑修订版_20260826.docx"
)
OUTPUT = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "完整模型7295版_20260827.docx"
)

sys.path.insert(0, str(ROOT / "scripts"))
from mathify_all_document_symbols import set_markup  # noqa: E402


def combined_text(element) -> str:
    return "".join(
        element._element.xpath(".//w:t/text()|.//m:t/text()")
    ).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if combined_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def rewrite(document: Document, prefix: str, text: str) -> None:
    set_markup(find_paragraph(document, prefix), text)


def set_stat(cell, latex: str) -> None:
    set_markup(cell.paragraphs[0], f"[[{latex}]]")
    set_math_size(cell, 17)


def set_pair(cell, left: str, right: str) -> None:
    set_markup(cell.paragraphs[0], f"[[{left}]] / [[{right}]]")
    set_math_size(cell, 17)


def set_math_size(cell, half_points: int) -> None:
    for math_run in cell._tc.xpath(".//m:r"):
        run_properties = math_run.find(qn("w:rPr"))
        if run_properties is None:
            run_properties = OxmlElement("w:rPr")
            math_properties = math_run.find(qn("m:rPr"))
            if math_properties is None:
                math_run.insert(0, run_properties)
            else:
                math_run.insert(math_run.index(math_properties) + 1, run_properties)
        size = run_properties.find(qn("w:sz"))
        if size is None:
            size = OxmlElement("w:sz")
            run_properties.append(size)
        size.set(qn("w:val"), str(half_points))
        size_cs = run_properties.find(qn("w:szCs"))
        if size_cs is None:
            size_cs = OxmlElement("w:szCs")
            run_properties.append(size_cs)
        size_cs.set(qn("w:val"), str(half_points))


def set_plain_like(cell, template_cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    properties = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    template_runs = template_cell._tc.xpath(".//w:r[w:rPr][1]")
    if template_runs:
        run._r.insert(0, deepcopy(template_runs[0].find(qn("w:rPr"))))


def remove_row_by_label(table, label: str) -> None:
    matches = [row for row in table.rows if combined_text(row.cells[0]) == label]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one row labelled {label!r}, found {len(matches)}")
    row = matches[0]
    row._tr.getparent().remove(row._tr)


def update_narrative(document: Document) -> None:
    rewrite(
        document,
        "跨语言知识图谱实体对齐需要",
        "跨语言知识图谱实体对齐需要同时处理关系结构不一致、实体文本碎片化和监督样本有限等问题。"
        "本文构建一种关系感知结构上下文与多尺度语义融合方法。结构分支以八维局部拓扑统计和小比例可学习实体残差初始化，"
        "通过关系感知消息传播与节点级层选择融合不同结构深度；语义分支从共享输入构造 token、phrase 和 global 三种视图。"
        "融合阶段保留全部一跳出邻居，采用批次内动态补齐和温度为 0.25 的 1.5-entmax 稀疏选择。"
        "五个数据集的三随机种子测试平均 Hits@1、Hits@10 和 MRR 分别为 0.7640、0.8862 和 0.8082。"
        "统一协议消融表明，共享拓扑初始化和关系类型在两个代表性数据集上均产生稳定增益；phrase 与 global 视图呈现数据集依赖性；"
        "参数匹配晚期融合的 Hits@1 分别低 3.15 和 9.44 个百分点，而将语义查询替换为结构查询分别低 1.41 和 0.22 个百分点。",
    )
    rewrite(
        document,
        "开展多数据集实验和受控消融分析",
        "开展多数据集实验和受控消融分析。 本文在五个数据集上进行三随机种子主实验，并在两个代表性数据集上分别检验"
        "拓扑初始化、邻居选择、关系类型、层选择及三种语义视图。参数匹配的晚期训练融合基线控制训练容量，"
        "结构查询基线则在邻居可见范围不变时替换查询信号，从而分别提供关于邻居上下文和查询来源的受控证据。",
    )
    rewrite(
        document,
        "因此，结构分支在第一个 epoch",
        "因此，结构分支辅助目标在第一个 epoch 的权重为 0.1，并在最后一个 epoch 衰减为 0。"
        "该目标仅在训练前期约束结构表示空间，后期优化由联合表示上的对齐目标主导。",
    )
    rewrite(
        document,
        "实验围绕三个研究问题展开",
        "实验围绕三个研究问题展开。RQ1：新主配置在跨语言、跨知识库来源与事件知识图谱上能否获得稳定结果？"
        "RQ2：拓扑初始化、全部邻居稀疏选择、关系类型、节点级层选择以及三种语义视图分别产生何种贡献？"
        "RQ3：在训练阶段融合参数量相同的条件下，最终检索表示包含一跳结构邻居时与仅使用独立分支时有何差异，且邻居查询是否必须使用语义信息？",
    )
    rewrite(
        document,
        "在 DBP15K ZH–EN 上",
        "在 DBP15K ZH–EN 上，本文模型的 Hits@1 为 0.7295，高于 MTransE、JAPE、BootEA 和 RDGCN 表中数值，但低于 RREA-text。"
        "OpenEA EN–FR 的 Hits@1 为 0.6324，高于官方 MTransE、JAPE 和 GCN-Align，低于 BootEA、KDCoE 和 RDGCN。 ",
    )
    rewrite(
        document,
        "将共享拓扑初始化替换为",
        "将共享拓扑初始化替换为独立可学习实体向量后，Hits@1 在 DBP15K 和 OpenEA 上分别下降 4.48 与 7.45 个百分点；"
        "去除关系类型分别下降 4.00 与 3.10 个百分点，这两项在两个数据集上方向一致。"
        "固定 8 个邻居并恢复 softmax 分别下降 1.59 与 1.02 个百分点，表明全部一跳邻居与稀疏选择具有一致收益。"
        "去除层选择器在 DBP15K 和 OpenEA 上分别下降 1.08 与 1.71 个百分点。语义视图方面，去除 token 分别下降 0.97 与 0.02 个百分点；"
        "去除 phrase 分别下降 1.26 与 4.87 个百分点；去除 global 后 DBP15K 提高 0.33 个百分点而 OpenEA 下降 0.84 个百分点，"
        "表明不同语义视图的作用具有数据集依赖性。",
    )
    rewrite(
        document,
        "结构单分支的 Hits@1",
        "结构单分支的 Hits@1 在 DBP15K 和 OpenEA 上分别为 0.1684 和 0.1014，语义单分支分别为 0.6424 和 0.5379。"
        "这说明语义是当前输入条件下的主要独立检索信号，但结构分支并非没有边际作用：完整模型相对语义单分支分别提高 8.71 和 9.46 个百分点。"
        "无可训练平均融合达到 0.7019 和 0.5832，仍比完整模型低 2.76 和 4.93 个百分点，表明可训练的邻居上下文与门控优于结构和语义的固定等权组合。",
    )
    rewrite(
        document,
        "相对完整模型，结构查询邻居",
        "相对完整模型，结构查询邻居的 Hits@1 在 DBP15K 和 OpenEA 上分别下降 1.41 与 0.22 个百分点。"
        "OpenEA 上的差值较小，而 DBP15K 上的差值更明显，因此现有结果表明查询来源的影响具有数据集依赖性，不能把语义查询概括为普遍的主要增益来源。"
        "参数匹配晚期训练融合分别下降 3.15 与 9.44 个百分点，表明在当前统一检索规则下，保留一跳结构邻居上下文与更高结果相关。"
        "由于晚期基线同时移除了邻居上下文并改变训练交互路径，该差值是两项变化的联合效应，不能单独证明提前融合优于晚期融合。",
    )
    rewrite(
        document,
        "RQ1 的结果表明",
        "RQ1 的结果表明，模型能够在五个数据集上稳定训练，但绝对性能随语言、知识库来源和事件属性异构性而变化。"
        "RQ2 的消融显示，拓扑初始化、关系类型和层选择在两个代表性数据集上均提供方向一致的增益；"
        "邻居选择和语义视图的贡献大小则随数据集变化。RQ3 的控制实验表明，查询来源的影响具有数据集依赖性；"
        "包含一跳结构邻居的主模型优于不读取邻居的参数匹配基线，但现有设计不能把邻居可见性与训练交互位置完全分离。",
    )


def update_tables(document: Document) -> None:
    main = document.tables[1]
    set_plain_like(main.rows[1].cells[1], main.rows[2].cells[1], "0.7295 ± 0.0037")
    set_plain_like(main.rows[1].cells[2], main.rows[2].cells[2], "0.8490 ± 0.0017")
    set_plain_like(main.rows[1].cells[3], main.rows[2].cells[3], "0.7706 ± 0.0034")

    comparison = document.tables[2]
    comparison.rows[6].cells[2].text = "0.7295"
    comparison.rows[6].cells[3].text = "0.8490"
    comparison.rows[6].cells[4].text = "0.7706"

    ablation = document.tables[3]
    remove_row_by_label(ablation, "去除早期结构监督")
    set_plain_like(
        ablation.rows[1].cells[1],
        ablation.rows[2].cells[1],
        "0.7295 ± 0.0037 / 0.7706 ± 0.0034",
    )

    branches = document.tables[4]
    set_plain_like(
        branches.rows[1].cells[2],
        branches.rows[2].cells[2],
        "0.7295 ± 0.0037 / 0.7706 ± 0.0034",
    )

    controls = document.tables[6]
    set_plain_like(controls.rows[1].cells[1], controls.rows[2].cells[1], "0.7295 ± 0.0037")
    set_plain_like(controls.rows[1].cells[2], controls.rows[2].cells[2], "0.7706 ± 0.0034")


def citation_order(document: Document) -> list[int]:
    citation_re = re.compile(r"\[(\d+(?:\s*(?:,|–|-)\s*\d+)*)\]")
    seen: list[int] = []
    for paragraph in document.paragraphs:
        if paragraph.style is not None and paragraph.style.name == "Reference":
            continue
        for match in citation_re.finditer(combined_text(paragraph)):
            for raw in re.findall(r"\d+", match.group(1)):
                number = int(raw)
                if number not in seen:
                    seen.append(number)
    return seen


def validate(document: Document) -> None:
    paragraphs = "\n".join(combined_text(p) for p in document.paragraphs)
    table_text = "\n".join(
        combined_text(cell)
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    body = paragraphs + "\n" + table_text

    if "去除早期结构监督" in body:
        raise RuntimeError("The removed structural-supervision ablation is still present")
    if len(document.tables[3].rows) != 9:
        raise RuntimeError(f"Expected eight ablations plus header, found {len(document.tables[3].rows)} rows")

    expected_rows = [
        combined_text(document.tables[1].rows[1].cells[1]),
        combined_text(document.tables[2].rows[6].cells[2]),
        combined_text(document.tables[3].rows[1].cells[1]),
        combined_text(document.tables[4].rows[1].cells[2]),
        combined_text(document.tables[6].rows[1].cells[1]),
    ]
    if any("0.7295" not in value for value in expected_rows):
        raise RuntimeError(f"Complete-model Hits@1 is inconsistent: {expected_rows}")
    if "0.7640、0.8862 和 0.8082" not in paragraphs:
        raise RuntimeError("Abstract averages were not updated")
    if "3.15 和 9.44" not in paragraphs or "1.41 和 0.22" not in paragraphs:
        raise RuntimeError("Dependent comparison deltas were not updated")

    expected_citations = list(range(1, 21))
    if citation_order(document) != expected_citations:
        raise RuntimeError(f"Citation order changed: {citation_order(document)}")
    references = [combined_text(p) for p in document.paragraphs if p.style.name == "Reference"]
    reference_numbers = [int(re.match(r"\[(\d+)\]", text).group(1)) for text in references]
    if reference_numbers != expected_citations:
        raise RuntimeError(f"Reference list changed: {reference_numbers}")


def main() -> None:
    document = Document(SOURCE)
    update_narrative(document)
    update_tables(document)
    validate(document)
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    validate(reopened)
    print(f"WROTE {OUTPUT}")
    print(f"ablation_rows={len(reopened.tables[3].rows) - 1} citations={len(citation_order(reopened))}")


if __name__ == "__main__":
    main()
