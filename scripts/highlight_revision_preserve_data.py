from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from revise_manuscript1_by_priority import (
    find_paragraph,
    paragraph_text,
    set_cell,
    set_mixed,
    set_plain,
    set_table_pagination,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "1.docx"
REVISED = ROOT / "outputs" / "1_按修稿优先级修订版_20260831.docx"
OUTPUT = ROOT / "outputs" / "1_修订标黄版_保留原实验数据_20260831.docx"


def remove_paragraph(paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def add_run_highlight(run_element) -> None:
    r_pr = run_element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        run_element.insert(0, r_pr)
    highlight = r_pr.find(qn("w:highlight"))
    if highlight is None:
        highlight = OxmlElement("w:highlight")
        r_pr.append(highlight)
    highlight.set(qn("w:val"), "yellow")


def shade_paragraph(paragraph, fill: str = "FFF59D") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def highlight_paragraph(paragraph) -> None:
    for run in paragraph._p.xpath(".//w:r"):
        add_run_highlight(run)
    if paragraph._p.xpath(".//m:oMath"):
        shade_paragraph(paragraph)


def shade_cell(cell, fill: str = "FFF59D") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def replace_table(target, replacement) -> None:
    target._tbl.addprevious(copy.deepcopy(replacement._tbl))
    target._tbl.getparent().remove(target._tbl)


def restore_original_experimental_data(document: Document, source: Document) -> None:
    # The source leaves semantic sequence shapes blank; the highlighted revision preserves that table.
    for row in document.tables[0].rows[1:]:
        set_cell(row.cells[4], "", align=WD_ALIGN_PARAGRAPH.CENTER)

    # Preserve the source manuscript's literature-reported comparison table.
    replace_table(document.tables[2], source.tables[2])
    set_plain(
        find_paragraph(document, "Table 3 reports a seed-42 RREA-basic"),
        paragraph_text(find_paragraph(source, "Table 3 lists results reported")),
    )
    set_plain(
        find_paragraph(document, "Table 3 Same-protocol comparison"),
        paragraph_text(find_paragraph(source, "Table 3 Contextual comparison")),
    )
    remove_paragraph(find_paragraph(document, "The three-seed mean of the proposed model"))

    # Preserve the structural-query aggregate values recorded in the source manuscript.
    table6 = document.tables[5]
    source6 = source.tables[5]
    set_cell(table6.cell(2, 3), paragraph_text(source6.cell(2, 3).paragraphs[0]), size=8.7)
    set_cell(table6.cell(2, 4), paragraph_text(source6.cell(2, 4).paragraphs[0]), size=8.7)

    set_mixed(
        find_paragraph(document, "Semantic conditioning exceeds structural conditioning"),
        "Semantic conditioning exceeds structural conditioning by 2.08 and 2.48 Hits@1 points on DBP15K and OpenEA. Because $u_i$ enters Eqs. (14), (16), and (18), "
        "this difference measures the complete conditioning path rather than the dot-product query alone. The full model exceeds parameter-matched late fusion by 5.28 "
        "and 9.22 points. That gap measures the combined effect of neighbor access and the training interaction path.",
    )
    set_plain(
        find_paragraph(document, "The paired Hits@1 differences"),
        "The paired Hits@1 differences (control minus complete model; mean ± sample standard deviation) are Complete-neighborhood softmax: DBP15K -0.35 ± 0.42 "
        "points and OpenEA -0.51 ± 0.60 points; Parameter-matched late fusion: DBP15K -5.28 ± 0.67 points and OpenEA -9.22 ± 1.25 points. Table 6 retains "
        "the manuscript's structural-query aggregate. Appendix A omits structural-query per-seed rows because the available run logs correspond to a different control implementation.",
    )
    set_plain(
        find_paragraph(document, "The controlled experiments support two main conclusions"),
        "The controlled experiments support two main conclusions on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. Relation types and topology initialization contribute "
        "consistent structural evidence. The complete model also exceeds the semantic branch, which shows that trainable structural context contributes beyond the text "
        "representation. Replacing semantic conditioning with the structural-query control lowers Hits@1 by 2.08 and 2.48 points under the aggregate results reported in Table 6.",
    )
    set_plain(
        find_paragraph(document, "This paper examines relation-aware structural context"),
        "This paper examines relation-aware structural context in a joint entity-alignment representation. On DBP15K ZH–EN and OpenEA EN–FR-15K-V2, relation types "
        "and topology initialization provide consistent structural gains, and the complete model exceeds both the semantic branch and parameter-free mean fusion. Replacing "
        "semantic conditioning with the structural-query control lowers Hits@1 by 2.08 and 2.48 points. The complete-neighborhood softmax control separates the normalization "
        "function from neighbor truncation, while the late-fusion comparison measures the joint effect of neighbor access and the training interaction path. Results on five "
        "datasets show that the complete configuration can be trained in cross-lingual and same-language event-centric settings. Component conclusions apply to the two datasets "
        "used for controlled experiments.",
    )
    set_plain(
        find_paragraph(document, "The experiments cover five datasets and three random seeds"),
        "The experiments cover five datasets and three random seeds. Controlled component studies use two representative cross-lingual pairs. Published-method comparisons "
        "retain the values and protocols reported by the cited studies. The complete-neighborhood softmax control and the reported component analyses follow the model's unified "
        "experimental protocol. Per-seed results and paired differences are reported when compatible run records are available.",
    )

    # Remove incompatible structural-conditioning per-seed rows from the added appendix.
    appendix_controls = document.tables[7]
    for row in list(appendix_controls.rows[1:]):
        if "Structural conditioning" in paragraph_text(row.cells[1].paragraphs[0]):
            appendix_controls._tbl.remove(row._tr)
    set_table_pagination(appendix_controls)


def apply_revision_highlighting(document: Document, source: Document) -> None:
    source_paragraphs = {
        paragraph_text(paragraph)
        for paragraph in source.paragraphs
        if paragraph_text(paragraph)
    }
    for paragraph in document.paragraphs:
        text = paragraph_text(paragraph)
        if text and text not in source_paragraphs:
            highlight_paragraph(paragraph)

    # Added direct control in Table 4.
    table4 = document.tables[3]
    for row in table4.rows:
        if paragraph_text(row.cells[0].paragraphs[0]) == "Complete neighborhood + softmax":
            for cell in row.cells:
                shade_cell(cell)

    # Expanded structural-control definition, while its numerical cells remain unchanged.
    table6 = document.tables[5]
    shade_cell(table6.cell(2, 0))
    shade_cell(table6.cell(2, 1))

    # Both appendix tables are additions to the source manuscript.
    for table in document.tables[6:]:
        for row in table.rows:
            for cell in row.cells:
                shade_cell(cell)

    # Equations (16)-(18) were revised and need a visible marker even inside OMML.
    for tag in ("(16)", "(17)", "(18)"):
        for paragraph in document.paragraphs:
            if tag in paragraph_text(paragraph):
                shade_paragraph(paragraph)


def main() -> None:
    source = Document(SOURCE)
    document = Document(REVISED)
    restore_original_experimental_data(document, source)
    apply_revision_highlighting(document, source)
    document.core_properties.subject = "Highlighted revision preserving source experimental data"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
