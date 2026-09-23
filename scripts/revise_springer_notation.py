from __future__ import annotations

import argparse
import copy
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from mathify_all_document_symbols import set_markup  # noqa: E402
from rewrite_chinese_paper_controlled_ablation import format_table  # noqa: E402


def combined_text(paragraph) -> str:
    return "".join(
        paragraph._p.xpath(".//w:t/text()|.//m:t/text()")
    ).replace("\u00a0", " ").strip()


def find_combined(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).startswith(prefix)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def find_formula(document: Document, number: int):
    suffix = f"({number})"
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).endswith(suffix)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one equation {suffix}, found {len(matches)}")
    return matches[0]


def replace_formula(document: Document, number: int, latex: str) -> None:
    set_markup(find_formula(document, number), f"[[{latex}]].\t({number})")


def set_cell_markup(cell, markup: str) -> None:
    set_markup(cell.paragraphs[0], markup.replace(r"\n", "\n"))


def replace_symbol_table(document: Document) -> None:
    old_table = document.tables[0]
    table = document.add_table(rows=1, cols=3)
    old_table._tbl.addprevious(table._tbl)

    headers = ["符号", "类型或维度", "定义"]
    rows = [
        (
            r"[[\mathcal{G}_s=(\mathcal{E}_s,\mathcal{R}_s,\mathcal{T}_s)]]\n[[s\in\{L,R\}]]",
            "集合",
            "左侧或右侧知识图谱，以及对应的实体、关系和三元组集合",
        ),
        (
            r"[[\mathcal{A}_{\mathrm{seed}},\ B]]",
            r"集合；[[\mathbb{N}]]",
            "种子对齐集合；训练批次中的实体对数量",
        ),
        (
            r"[[\mathbf{h}_i^{(\ell)},\ \mathbf{e}_r]]",
            r"[[\mathbb{R}^{d};\ \mathbb{R}^{d_r}]]",
            r"实体 [[i]] 在第 [[\ell]] 层的结构状态；关系 [[r]] 的嵌入",
        ),
        (
            r"[[\mathbf{m}_{j\to i}^{(\ell)},\ \bar{\mathbf{m}}_i^{(\ell)}]]",
            r"[[\mathbb{R}^{d}]]",
            r"邻居 [[j]] 到实体 [[i]] 的关系消息；实体 [[i]] 的平均传入消息",
        ),
        (
            r"[[\mathbf{W}_s^{(\ell)},\ \mathbf{W}_r^{(\ell)},\ \mathbf{W}_{\mathrm{self}}^{(\ell)}]]",
            r"[[\mathbb{R}^{d\times d}]]\n[[\mathbb{R}^{d\times d_r};\ \mathbb{R}^{d\times d}]]",
            "共享的源实体、关系和自身状态投影矩阵",
        ),
        (
            r"[[a_i^{(p)},\ \mathbf{z}_i^{\mathrm{str}}]]",
            r"[[\mathbb{R};\ \mathbb{R}^{d}]]",
            r"第 [[p]] 个保留深度的选择权重；最终结构表示",
        ),
        (
            r"[[\mathbf{X}_i,\ \mathbf{H}_i,\ \mathbf{M}_i]]",
            r"[[\mathbb{R}^{n_i\times300}]]\n[[\mathbb{R}^{n_i\times d};\ \{0,1\}^{n_i}]]",
            "词向量输入、共享序列状态和有效 token 掩码",
        ),
        (
            r"[[\mathbf{z}_i^{\mathrm{tok}},\ \mathbf{z}_i^{\mathrm{phr}},\ \mathbf{z}_i^{\mathrm{glo}},\ \mathbf{z}_i^{\mathrm{sem}}]]",
            r"[[\mathbb{R}^{d}]]",
            "token、phrase、global 三个语义视图及其融合结果",
        ),
        (
            r"[[\mathcal{N}_{K_{\mathrm{nbr}}}(i),\ K_{\mathrm{nbr}},\ \mathbf{T}_i]]",
            r"集合；[[\mathbb{N}]]\n[[\mathbb{R}^{K_{\mathrm{nbr}}\times d}]]",
            "固定预算邻居集合、邻居数上限和邻居结构表示矩阵",
        ),
        (
            r"[[\mathbf{q}_i,\ \mathbf{k}_j,\ \mathbf{v}_j,\ \alpha_{ij}]]",
            r"[[\mathbb{R}^{d};\ \mathbb{R}^{d};\ \mathbb{R}^{d};\ \mathbb{R}]]",
            "邻居注意力的查询、键、值和归一化权重",
        ),
        (
            r"[[\bar{\mathbf{c}}_i,\ \mathbf{g}_i^c,\ \mathbf{c}_i^{\mathrm{str}}]]\n"
            r"[[\mathbf{g}_i^{\mathrm{joint}},\ \mathbf{z}_i^{\mathrm{joint}}]]",
            r"[[\mathbb{R}^{d}]]",
            "邻居摘要、上下文门、结构上下文、联合门和最终实体表示",
        ),
        (
            r"[[S_{ij},\ \tau,\ \mathcal{L}_{\mathrm{align}}]]",
            r"[[\mathbb{R};\ \mathbb{R}_{>0};\ \mathbb{R}]]",
            "批内相似度、InfoNCE 温度和训练目标",
        ),
        (
            r"[[\alpha,\ L,\ k_{\mathrm{CSLS}}]]",
            r"[[[0,1];\ \mathbb{N};\ \mathbb{N}]]",
            "验证集选择的语义权重、有效传播深度和 CSLS 邻域大小",
        ),
    ]

    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for cell, markup in zip(cells, values):
            set_cell_markup(cell, markup)
    format_table(table, [2950, 2350, 4020])
    for row in table.rows[1:]:
        row.cells[2].paragraphs[0].alignment = 0
    old_table._tbl.getparent().remove(old_table._tbl)


def revise_task_and_conventions(document: Document) -> None:
    set_markup(find_combined(document, "3.2 符号约定"), "3.2 符号与记号")
    set_markup(find_combined(document, "表 1 主要符号"), "表 1 主要符号及其定义")
    replace_symbol_table(document)

    set_markup(
        find_combined(document, "下标 i 和 j"),
        "为避免后续公式产生歧义，表 1 仅汇总全文反复使用的主要记号。本文以花体大写字母表示集合，"
        "粗体小写字母表示向量，粗体大写字母表示矩阵，普通斜体字母表示标量或索引。"
        "[[s\\in\\{L,R\\}]] 表示知识图谱侧别，[[i]] 和 [[j]] 表示实体索引，"
        "[[\\ell]] 表示消息传播层，[[p]] 表示层选择器中的保留状态深度。[[n_i]] 是实体 [[i]] 的有效 token 数，"
        "[[K_{\\mathrm{nbr}}]] 是结构邻居预算，[[k_{\\mathrm{CSLS}}]] 是 CSLS 邻域大小。"
        "标签性上下标（如 [[\\mathrm{str}]]、[[\\mathrm{sem}]] 和 [[\\mathrm{joint}]]) 采用正体；"
        "[[\\operatorname{softmax}]]、[[\\operatorname{LayerNorm}]] 和 [[\\operatorname{Norm}]] 等算子采用正体。"
        "[[\\sigma]] 表示 sigmoid，[[\\odot]] 表示逐元素乘积，[[[\\cdot;\\cdot]]] 表示向量拼接。"
        "模型维度和邻居预算的具体取值见第 4.3 节。",
    )

    set_markup(
        find_combined(document, "G1=(E1,R1,T1)"),
        r"[[\mathcal{G}_L=(\mathcal{E}_L,\mathcal{R}_L,\mathcal{T}_L),\quad "
        r"\mathcal{G}_R=(\mathcal{E}_R,\mathcal{R}_R,\mathcal{T}_R)]]",
    )
    set_markup(
        find_combined(document, "其中 E、R 和 T"),
        "其中 [[\\mathcal{E}_s]]、[[\\mathcal{R}_s]] 和 [[\\mathcal{T}_s]] 分别表示侧别 "
        "[[s\\in\\{L,R\\}]] 的实体集合、关系集合和三元组集合。已知少量种子对齐：",
    )
    set_markup(
        find_combined(document, "Aseed="),
        r"[[\mathcal{A}_{\mathrm{seed}}=\{(e_i,e_j)\mid e_i\in\mathcal{E}_L,"
        r"\ e_j\in\mathcal{E}_R,\ e_i\equiv e_j\}]]",
    )


def revise_method_notation(document: Document) -> None:
    set_markup(
        find_combined(document, "图 1 关系感知结构上下文"),
        "图 1 关系感知结构上下文与多尺度语义融合方法及训练—检索协议。（a）关系类型进入三层全图消息传播，"
        "节点级层选择得到结构表示；token、phrase 和 global 视图形成语义表示；主模型使用语义查询固定预算结构邻居，"
        "并经结构上下文门和联合表示门输出实体表示。结构查询、单分支、无可训练融合和参数匹配晚期融合仅作为第 5 节控制组。"
        "（b）种子对齐用于单阶段双向 InfoNCE 训练，验证集选择 [[\\alpha]]、[[L]] 和 "
        "[[k_{\\mathrm{CSLS}}]]，测试集只执行一次最终评估。",
    )
    set_markup(
        find_combined(document, "图 1 与公式的对应关系"),
        "图 1 与公式的对应关系如下：结构分支接收关系图并由式（1）—（6）输出 "
        "[[\\mathbf{z}_i^{\\mathrm{str}}]]；语义分支接收 [[\\mathbf{X}_i]] 和 [[\\mathbf{M}_i]] 并由式（7）—（12）输出 "
        "[[\\mathbf{z}_i^{\\mathrm{sem}}]]；邻居上下文模块由式（13）—（18）输出 "
        "[[\\mathbf{c}_i^{\\mathrm{str}}]]；联合表示模块由式（19）—（20）输出 "
        "[[\\mathbf{z}_i^{\\mathrm{joint}}]]；式（22）—（23）对应图 1(b) 的训练目标，"
        "式（24）对应验证和测试阶段的 CSLS 检索。",
    )

    set_markup(
        find_combined(document, "消息传递（message passing）"),
        "消息传递（message passing）指沿知识图谱有向边把邻居及关系证据逐层送入目标实体。设 "
        "[[\\mathbf{h}_i^{(\\ell)}\\in\\mathbb{R}^{d}]] 为实体 [[i]] 在第 [[\\ell]] 层的状态，"
        "[[\\mathbf{e}_r\\in\\mathbb{R}^{d_r}]] 为关系 [[r]] 的嵌入，"
        "[[\\mathcal{N}_{\\mathrm{in}}(i)]] 为所有指向 [[i]] 的传入边。"
        "对于边 [[(j,r,i)]]，第 [[\\ell]] 层首先计算：",
    )
    replace_formula(
        document,
        1,
        r"\mathbf{m}_{j\to i}^{(\ell)}=\operatorname{LayerNorm}\!\left("
        r"\mathbf{W}_s^{(\ell)}\mathbf{h}_j^{(\ell)}+\mathbf{W}_r^{(\ell)}\mathbf{e}_r\right)",
    )
    set_markup(
        find_combined(document, "其中 Ws"),
        "其中 [[\\mathbf{W}_s^{(\\ell)}\\in\\mathbb{R}^{d\\times d}]] 与 "
        "[[\\mathbf{W}_r^{(\\ell)}\\in\\mathbb{R}^{d\\times d_r}]] 是同一层所有关系共享的线性投影；"
        "关系差异来自 [[\\mathbf{e}_r]]，而不是每种关系各自拥有一套大矩阵。"
        "[[\\mathbf{m}_{j\\to i}^{(\\ell)}\\in\\mathbb{R}^{d}]] 是边 [[(j,r,i)]] 传递给实体 [[i]] 的消息。"
        "把所有传入消息取均值得到：",
    )
    replace_formula(
        document,
        2,
        r"\bar{\mathbf{m}}_i^{(\ell)}=\frac{1}{\max(1,|\mathcal{N}_{\mathrm{in}}(i)|)}"
        r"\sum_{(j,r,i)\in\mathcal{N}_{\mathrm{in}}(i)}\mathbf{m}_{j\to i}^{(\ell)}",
    )
    set_markup(
        find_combined(document, "自身状态投影为"),
        "自身状态投影为 [[\\mathbf{s}_i^{(\\ell)}=\\mathbf{W}_{\\mathrm{self}}^{(\\ell)}"
        "\\mathbf{h}_i^{(\\ell)}]]。模型根据自身与邻居证据计算逐维门控：",
    )
    replace_formula(
        document,
        3,
        r"\mathbf{g}_i^{(\ell)}=\sigma\!\left(\mathbf{W}_g^{(\ell)}"
        r"[\mathbf{s}_i^{(\ell)};\bar{\mathbf{m}}_i^{(\ell)}]\right)",
    )
    replace_formula(
        document,
        4,
        r"\mathbf{h}_i^{(\ell+1)}=\mathbf{s}_i^{(\ell)}+\mathbf{g}_i^{(\ell)}"
        r"\odot\bar{\mathbf{m}}_i^{(\ell)}",
    )
    set_markup(
        find_combined(document, "为避免只依赖最深层"),
        "为避免只依赖最深层，编码器保留投影后的输入状态及全部层输出 "
        "[[\\mathbf{u}_i^{(0)},\\ldots,\\mathbf{u}_i^{(L)}]]，其中 [[\\mathbf{u}_i^{(0)}]] 表示未传播的自身信息，"
        "[[\\mathbf{u}_i^{(p)}]] 表示吸收至多 [[p]] 跳证据后的状态。平均上下文 "
        "[[\\bar{\\mathbf{u}}_i]] 由 [[L+1]] 个状态求均值得到，两层 MLP 再为每个深度输出标量分数并通过 "
        "[[\\operatorname{softmax}]] 得到 [[a_i^{(p)}]]：",
    )
    replace_formula(
        document,
        5,
        r"a_i^{(p)}=\operatorname{softmax}_{p}\!\left(\operatorname{MLP}"
        r"([\mathbf{u}_i^{(p)};\bar{\mathbf{u}}_i])\right)",
    )
    replace_formula(
        document,
        6,
        r"\mathbf{z}_i^{\mathrm{str}}=\operatorname{Norm}\!\left(\operatorname{LayerNorm}\!\left("
        r"\sum_{p=0}^{L}a_i^{(p)}\mathbf{u}_i^{(p)}\right)\right)",
    )

    set_markup(
        find_combined(document, "整图结构编码只计算一次"),
        "整图结构编码只计算一次，批次中的实体及其邻居随后从结构表示矩阵中索引。对于每个实体，"
        "系统按照邻接表中的确定性顺序提取至多 [[K_{\\mathrm{nbr}}]] 个邻居，主配置取 "
        "[[K_{\\mathrm{nbr}}=8]]。若实体没有邻居，则使用实体自身编号填充并将掩码置为 0；"
        "若邻居少于 [[K_{\\mathrm{nbr}}]]，则循环填充已有邻居并保持有效掩码。"
        "该步骤仅控制批次张量的固定形状，不引入额外打分函数或可训练参数。",
    )

    set_markup(
        find_combined(document, "输入语义序列"),
        "输入语义序列 [[\\mathbf{X}_i\\in\\mathbb{R}^{n_i\\times300}]] 由名称、关系和属性 token 构成，"
        "其中 [[n_i]] 是实体 [[i]] 的有效序列长度，每一行是一个 300 维词向量。非零行形成掩码 "
        "[[\\mathbf{M}_i\\in\\{0,1\\}^{n_i}]]。输入经过 300→128 的线性投影、LayerNorm、正弦位置编码和 dropout 得到 "
        "[[\\mathbf{H}_i\\in\\mathbb{R}^{n_i\\times d}]]；token、phrase 和 global 三个视图都从同一 "
        "[[\\mathbf{H}_i]] 和 [[\\mathbf{M}_i]] 构造，不额外引入外部句子。",
    )
    replace_formula(
        document,
        7,
        r"\mathbf{z}_i^{\mathrm{tok}}=\frac{1}{2}\!\left(\operatorname{MeanPool}"
        r"(\mathbf{H}_i,\mathbf{M}_i)+\operatorname{AttnPool}(\mathbf{H}_i,\mathbf{M}_i)\right)",
    )
    replace_formula(
        document,
        8,
        r"\mathbf{H}_i^{\mathrm{phr}}=\operatorname{LayerNorm}\!\left(\mathbf{H}_i+\frac{1}{2}"
        r"(\operatorname{Conv}_3(\mathbf{H}_i)+\operatorname{Conv}_5(\mathbf{H}_i))\right)",
    )
    replace_formula(
        document,
        9,
        r"\mathbf{z}_i^{\mathrm{phr}}=\frac{1}{2}\!\left(\operatorname{MeanPool}"
        r"(\mathbf{H}_i^{\mathrm{phr}},\mathbf{M}_i)+\operatorname{AttnPool}"
        r"(\mathbf{H}_i^{\mathrm{phr}},\mathbf{M}_i)\right)",
    )
    replace_formula(
        document,
        10,
        r"\mathbf{H}_i^{\mathrm{glo}}=\operatorname{Transformer}(\mathbf{H}_i,\mathbf{M}_i)",
    )
    replace_formula(
        document,
        11,
        r"\mathbf{z}_i^{\mathrm{glo}}=\frac{1}{2}\!\left(\operatorname{MeanPool}"
        r"(\mathbf{H}_i^{\mathrm{glo}},\mathbf{M}_i)+\operatorname{AttnPool}"
        r"(\mathbf{H}_i^{\mathrm{glo}},\mathbf{M}_i)\right)",
    )
    replace_formula(
        document,
        12,
        r"[\pi_i^{\mathrm{tok}},\pi_i^{\mathrm{phr}},\pi_i^{\mathrm{glo}}]="
        r"\operatorname{softmax}\!\left(\operatorname{MLP}"
        r"([\mathbf{z}_i^{\mathrm{tok}};\mathbf{z}_i^{\mathrm{phr}};\mathbf{z}_i^{\mathrm{glo}}])\right)",
    )

    set_markup(
        find_combined(document, "对于实体i，融合模块"),
        "对于实体 [[i]]，融合模块的输入为语义表示 "
        "[[\\mathbf{s}_i=\\mathbf{z}_i^{\\mathrm{sem}}\\in\\mathbb{R}^{d}]]、自身结构表示 "
        "[[\\mathbf{t}_i=\\mathbf{z}_i^{\\mathrm{str}}\\in\\mathbb{R}^{d}]]、邻居结构矩阵 "
        "[[\\mathbf{T}_i\\in\\mathbb{R}^{K_{\\mathrm{nbr}}\\times d}]] 和有效掩码。"
        "对第 [[j]] 个邻居 [[\\mathbf{t}_{i,j}]]，[[\\mathbf{W}_Q,\\mathbf{W}_K,\\mathbf{W}_V"
        "\\in\\mathbb{R}^{d\\times d}]] 分别产生查询 [[\\mathbf{q}_i=\\mathbf{W}_Q\\mathbf{s}_i]]、"
        "键 [[\\mathbf{k}_j=\\mathbf{W}_K\\mathbf{t}_{i,j}]] 和值 "
        "[[\\mathbf{v}_j=\\mathbf{W}_V\\mathbf{t}_{i,j}]]；缩放点积注意力首先计算：",
    )
    replace_formula(
        document,
        13,
        r"a_{ij}=d^{-1/2}\mathbf{k}_j^{\mathsf{T}}\mathbf{q}_i",
    )
    replace_formula(
        document,
        14,
        r"\phi(\mathbf{x},\mathbf{y})=[\mathbf{x};\mathbf{y};\mathbf{x}\odot\mathbf{y};"
        r"|\mathbf{x}-\mathbf{y}|]",
    )
    replace_formula(
        document,
        15,
        r"p_{ij}=\sigma\!\left(\mathbf{W}_P\phi(\mathbf{t}_{i,j},\mathbf{s}_i)\right)",
    )
    replace_formula(
        document,
        16,
        r"\alpha_{ij}=\operatorname{softmax}_{j}\!\left(a_{ij}+\log(p_{ij}+\epsilon)+"
        r"\mathrm{mask}_{ij}\right)",
    )
    set_markup(
        find_combined(document, "结构邻居摘要为"),
        "结构邻居摘要为 [[\\bar{\\mathbf{c}}_i=\\sum_{j=1}^{K_{\\mathrm{nbr}}}"
        "\\alpha_{ij}\\mathbf{v}_j]]。"
        "随后使用另一个逐维门控决定信任自身还是邻居：",
    )
    replace_formula(
        document,
        17,
        r"\mathbf{g}_i^c=\sigma\!\left(\mathbf{W}_C[\phi(\mathbf{t}_i,\bar{\mathbf{c}}_i);"
        r"\mathbf{s}_i]\right)",
    )
    replace_formula(
        document,
        18,
        r"\mathbf{c}_i^{\mathrm{str}}=\mathbf{g}_i^c\odot\mathbf{t}_i+"
        r"(\mathbf{1}-\mathbf{g}_i^c)\odot\bar{\mathbf{c}}_i",
    )
    set_markup(
        find_combined(document, "没有有效邻居时"),
        "没有有效邻居时，[[\\mathbf{c}_i^{\\mathrm{str}}=\\mathbf{t}_i]]。"
        "最后分别归一化 [[\\mathbf{s}_i]] 与 [[\\mathbf{c}_i^{\\mathrm{str}}]]，并用联合门控",
    )
    replace_formula(
        document,
        19,
        r"\mathbf{g}_i^{\mathrm{joint}}=\sigma\!\left(\mathbf{W}_J"
        r"\phi(\mathbf{s}_i,\mathbf{c}_i^{\mathrm{str}})\right)",
    )
    replace_formula(
        document,
        20,
        r"\mathbf{z}_i^{\mathrm{joint}}=\operatorname{Norm}\!\left("
        r"\mathbf{g}_i^{\mathrm{joint}}\odot\mathbf{c}_i^{\mathrm{str}}+"
        r"(\mathbf{1}-\mathbf{g}_i^{\mathrm{joint}})\odot\mathbf{s}_i\right)",
    )
    replace_formula(
        document,
        21,
        r"\mathbf{z}_i^{\mathrm{late}}=\operatorname{Norm}\!\left(\mathbf{W}_2"
        r"\operatorname{Dropout}\!\left(\operatorname{GELU}\!\left(\mathbf{W}_1"
        r"[\operatorname{Norm}(\mathbf{t}_i);\operatorname{Norm}(\mathbf{s}_i)]\right)\right)\right)",
    )
    replace_formula(
        document,
        22,
        r"S_{ij}=\frac{\cos(\mathbf{z}_{L_i}^{\mathrm{joint}},"
        r"\mathbf{z}_{R_j}^{\mathrm{joint}})}{\tau}",
    )
    replace_formula(
        document,
        23,
        r"\mathcal{L}_{\mathrm{align}}=\frac{1}{2}\!\left("
        r"\operatorname{CE}(\mathbf{S},[0,\ldots,B-1])+"
        r"\operatorname{CE}(\mathbf{S}^{\mathsf{T}},[0,\ldots,B-1])\right)",
    )
    replace_formula(
        document,
        24,
        r"\operatorname{CSLS}(\mathbf{x},\mathbf{y})=2\cos(\mathbf{x},\mathbf{y})-r_L-r_R",
    )
    set_markup(
        find_combined(document, "其中，rL 和 rR"),
        "其中，[[r_L]] 和 [[r_R]] 分别表示实体到另一侧最近 [[k_{\\mathrm{CSLS}}]] 个候选的平均相似度。"
        "[[\\alpha]]、有效传播深度 [[L]] 和 [[k_{\\mathrm{CSLS}}]] 仅由三个随机种子的平均验证 MRR 选择；"
        "CSLS 校正强度固定为 1.0。该步骤不增加训练参数，测试集不参与配置选择。",
    )


def replace_csls_mentions(document: Document) -> None:
    replacements = {
        "跨语言知识图谱实体对齐需要": (
            "CSLS k。",
            "[[k_{\\mathrm{CSLS}}]]。",
        ),
        "实验回答四个研究问题": (
            "CSLS k 是否改善检索",
            "CSLS 邻域 [[k_{\\mathrm{CSLS}}]] 是否改善检索",
        ),
        "表 3、表 5–8": (
            "选择 α,L,k。",
            "选择 [[\\alpha,L,k_{\\mathrm{CSLS}}]]。",
        ),
        "公共设置为 batch size": (
            "k∈{3,5,7,10,15,20}",
            "[[k_{\\mathrm{CSLS}}\\in\\{3,5,7,10,15,20\\}]]",
        ),
        "表 7 报告每个数据集": (
            "k 从 3、5、7、10、15、20 中搜索",
            "[[k_{\\mathrm{CSLS}}]] 从 3、5、7、10、15、20 中搜索",
        ),
        "最优 α 位于": (
            "k 覆盖 3、5、15 和 20",
            "[[k_{\\mathrm{CSLS}}]] 覆盖 3、5、15 和 20",
        ),
        "本文构建了一套结合": (
            "CSLS k。",
            "[[k_{\\mathrm{CSLS}}]]。",
        ),
    }
    for prefix, (old, new) in replacements.items():
        paragraph = find_combined(document, prefix)
        text = combined_text(paragraph)
        if old not in text:
            raise ValueError(f"Could not replace {old!r} in {prefix!r}")
        markup = text.replace(old, new)
        if prefix == "最优 α 位于":
            markup = markup.replace("新增消融也分别重新选择 α、L、k", "新增消融也分别重新选择 [[\\alpha]]、[[L]]、[[k_{\\mathrm{CSLS}}]]")
            markup = markup.replace("最优 α", "最优 [[\\alpha]]")
            markup = markup.replace("选择 L=3", "选择 [[L=3]]")
        elif prefix == "公共设置为 batch size":
            markup = markup.replace("α∈{0.20,…,0.80}", "[[\\alpha\\in\\{0.20,\\ldots,0.80\\}]]")
            markup = markup.replace("L∈{1,2,3}", "[[L\\in\\{1,2,3\\}]]")
        elif prefix == "表 7 报告每个数据集":
            markup = markup.replace("α 从 0.20–0.80", "[[\\alpha]] 从 0.20–0.80")
            markup = markup.replace("L 从 1–3", "[[L]] 从 1–3")
        set_markup(paragraph, markup)

    table = document.tables[6]
    set_cell_markup(table.rows[0].cells[1], r"[[\alpha]]")
    set_cell_markup(table.rows[0].cells[2], r"[[L]]")
    set_cell_markup(table.rows[0].cells[3], r"[[k_{\mathrm{CSLS}}]]")
    table = document.tables[8]
    set_cell_markup(table.rows[0].cells[1], r"DBP [[\alpha/L/k_{\mathrm{CSLS}}]]")
    set_cell_markup(table.rows[0].cells[3], r"OpenEA [[\alpha/L/k_{\mathrm{CSLS}}]]")


def replace_first_figure(document: Document, figure_path: Path) -> None:
    if not figure_path.exists():
        raise FileNotFoundError(figure_path)
    shape = document.inline_shapes[0]
    blip = shape._inline.xpath(".//*[local-name()='blip']")[0]
    relationship_id = blip.get(qn("r:embed"))
    relationship = document.part.rels[relationship_id]
    relationship.target_part._blob = figure_path.read_bytes()
    doc_properties = shape._inline.xpath(".//*[local-name()='docPr']")
    if doc_properties:
        doc_properties[0].set(
            "descr",
            "关系感知结构上下文、多尺度语义编码、固定预算邻居融合、InfoNCE 训练和验证驱动检索流程",
        )


def verify(document: Document) -> None:
    table = document.tables[0]
    if len(table.columns) != 3 or len(table.rows) != 14:
        raise RuntimeError("The rebuilt notation table has an unexpected shape")
    xml = document._element.xml
    required = (
        "k_{\\mathrm{CSLS}}",
        "K_{\\mathrm{nbr}}",
        "\\mathcal{G}_s",
        "\\mathbf{g}_i^{\\mathrm{joint}}",
    )
    # The XML stores OMML rather than LaTeX; verify through reconstructed text too.
    all_text = "\n".join(combined_text(p) for p in document.paragraphs)
    if "3.2 符号与记号" not in all_text:
        raise RuntimeError("Section 3.2 heading was not updated")
    if "第 k 个传播深度" in all_text or "最近k个候选" in all_text:
        raise RuntimeError("Ambiguous legacy k notation remains")
    if "定义与来源" in "".join(
        "".join(cell._tc.xpath(".//w:t/text()|.//m:t/text()"))
        for cell in table.rows[0].cells
    ):
        raise RuntimeError("Legacy notation-table header remains")
    if xml.count("<m:oMath") < 100:
        raise RuntimeError("Native Word equations were unexpectedly lost")
    del required


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--figure")
    args = parser.parse_args()

    source = Path(args.source)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    document = Document(output)

    revise_task_and_conventions(document)
    revise_method_notation(document)
    replace_csls_mentions(document)
    if args.figure:
        replace_first_figure(document, Path(args.figure))

    document.core_properties.subject = "Springer 风格符号约定与全文数学记号统一"
    verify(document)
    document.save(output)
    print(f"WROTE {output}")


if __name__ == "__main__":
    main()
