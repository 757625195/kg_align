from __future__ import annotations

import argparse
import copy
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from mathify_all_document_symbols import set_markup  # noqa: E402


def combined_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).strip()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if combined_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def find_containing(document: Document, needle: str) -> Paragraph:
    matches = [p for p in document.paragraphs if needle in combined_text(p)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph containing {needle!r}, found {len(matches)}")
    return matches[0]


def replace_literal(paragraph: Paragraph, old: str, new: str) -> None:
    changed = False
    for node in paragraph._p.xpath(".//w:t"):
        if old in (node.text or ""):
            node.text = node.text.replace(old, new)
            changed = True
    if not changed:
        raise ValueError(f"Text {old!r} not found in paragraph: {combined_text(paragraph)!r}")


def add_paragraph_after(paragraph: Paragraph, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    result = Paragraph(element, paragraph._parent)
    result.style = style
    return result


def remove_block(first: Paragraph, last: Paragraph) -> None:
    parent = first._p.getparent()
    current = first._p
    stop = last._p.getnext()
    while current is not stop:
        following = current.getnext()
        parent.remove(current)
        current = following


def merge_results_sections(document: Document) -> None:
    heading = find_paragraph(document, "5.1 整体性能与训练稳定性")
    caption = find_paragraph(document, "表 3 训练结果")
    old_summary = find_paragraph(document, "表 3 汇总五个数据集的结果")
    duplicate_heading = find_paragraph(document, "5.6 五个数据集最终测试结果")
    duplicate_last = find_paragraph(document, "最终结果中，DBP15K FR–EN")

    # Replace the abbreviated Table 3 with the complete final-results table from
    # the duplicate section, preserving its native Word equations and styling.
    short_table = document.tables[2]
    final_table = document.tables[6]
    short_table._tbl.addprevious(copy.deepcopy(final_table._tbl))
    short_table._tbl.getparent().remove(short_table._tbl)

    set_markup(heading, "5.1 总体测试结果与训练稳定性")
    set_markup(
        caption,
        "表 3 五个数据集的最终测试结果（均值[[\\pm]]样本标准差，[[n=3]]）",
    )

    intro = add_paragraph_after(heading)
    set_markup(
        intro,
        "表 3 汇总完整模型在五个数据集上的最终检索结果。每个数据集均使用随机种子 42、43 和 44 "
        "独立训练，并报告 Hits@1、Hits@10 和 MRR 的均值与样本标准差。DBP15K 从训练对中固定留出 "
        "10% 作为验证集；OpenEA 与 EventEA 使用官方验证集。OpenEA EN–FR-15K-V2 的结果对应官方第 "
        "1 折，而非五折平均；全部模型选择和 CSLS 参数选择只使用验证集。",
    )

    set_markup(
        old_summary,
        "完整模型在 DBP15K FR–EN 上取得最高结果，Hits@1、Hits@10 和 MRR 分别为 "
        "[[0.9038\\pm0.0024]]、[[0.9572\\pm0.0018]] 和 [[0.9239\\pm0.0017]]。"
        "EventEA EN–EN 的 Hits@1 为 [[0.7563\\pm0.0029]]，表明该模型能够迁移到事件和属性异构场景。"
        "OpenEA EN–FR-15K-V2 的 Hits@1 最低，为 [[0.5568\\pm0.0096]]；该结果应结合其名称去偏设置、"
        "当前跨语言词向量覆盖不足以及第一折评估范围理解，不能与采用五折平均或额外文本资源的方法作严格同条件比较。",
    )

    stability = add_paragraph_after(old_summary)
    set_markup(
        stability,
        "五个数据集的 Hits@1 样本标准差介于 [[0.0024]] 和 [[0.0108]] 之间，MRR 样本标准差介于 "
        "[[0.0015]] 和 [[0.0087]] 之间，均明显小于不同数据集之间的绝对性能差异。"
        "因此，表 3 所示数据集差异不是由某一次随机初始化单独造成的；但每个数据集仅运行三个随机种子，"
        "这些统计仍不足以代替更大样本的显著性检验。",
    )

    # Remove the later duplicate section, including its caption and table.
    remove_block(duplicate_heading, duplicate_last)

    # Close the numbering gaps created by removing the duplicate section/table.
    replacements = [
        ("5.7 结构组件与累积语义视图消融", "5.5 结构组件与累积语义视图消融"),
        ("5.8 结构分支、语义分支与融合模块", "5.6 结构分支、语义分支与融合模块"),
        ("5.9 交互位置、邻居可见范围与容量控制", "5.7 交互位置、邻居可见范围与容量控制"),
        ("表 9 同时报告", "表 7 同时报告"),
        ("表 9 结构组件与语义视图消融", "表 7 结构组件与语义视图消融"),
        ("表 10 比较", "表 8 比较"),
        ("表 10 分支与融合必要性实验", "表 8 分支与融合必要性实验"),
        ("表 11 固定结构与语义编码器", "表 9 固定结构与语义编码器"),
        ("表 11 参数匹配融合控制的设计差异", "表 9 参数匹配融合控制的设计差异"),
        ("表 12 参数匹配融合控制的测试结果", "表 10 参数匹配融合控制的测试结果"),
    ]
    for old, new in replacements:
        paragraph = find_containing(document, old)
        replace_literal(paragraph, old, new)

    cross_references = [
        ("前向计算包含两个交互点", "第 5.8 节", "第 5.7 节"),
        ("加权后的三个向量以拼接形式", "第 5.7 节", "第 5.5 节"),
        ("公式（13）—（20）给出", "第 5.8 节", "第 5.7 节"),
    ]
    for prefix, old, new in cross_references:
        replace_literal(find_paragraph(document, prefix), old, new)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)
    merge_results_sections(document)
    document.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
