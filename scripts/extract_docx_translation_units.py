from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_markup(paragraph) -> str:
    parts: list[str] = []
    equation_index = 0
    for child in paragraph._p.iterchildren():
        if child.tag == qn("w:r"):
            parts.extend(node.text or "" for node in child.iter(qn("w:t")))
        elif child.tag == qn("w:hyperlink"):
            parts.extend(node.text or "" for node in child.iter(qn("w:t")))
        elif child.tag in {qn("m:oMath"), qn("m:oMathPara")}:
            equation_text = "".join(node.text or "" for node in child.iter(qn("m:t")))
            parts.append(f"{{{{EQ{equation_index}:{equation_text}}}}}")
            equation_index += 1
        else:
            text = "".join(node.text or "" for node in child.iter(qn("w:t")))
            if text:
                parts.append(text)
    return "".join(parts).replace("\u00a0", " ")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--end", type=int)
    parser.add_argument("--tables", action="store_true")
    args = parser.parse_args()

    document = Document(args.docx)
    end = len(document.paragraphs) if args.end is None else args.end
    for index, paragraph in enumerate(document.paragraphs[args.start:end], args.start):
        print(f"P{index:03d}\t{paragraph.style.name}\t{paragraph_markup(paragraph)}")

    if args.tables:
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    text = " / ".join(paragraph_markup(p) for p in cell.paragraphs)
                    print(f"T{table_index:02d}R{row_index:02d}C{cell_index:02d}\t{text}")


if __name__ == "__main__":
    main()
