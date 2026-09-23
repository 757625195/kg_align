from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_中文无行号版_20260830.docx"
OUTPUT = ROOT / "outputs" / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_Springer中文修订稿_20260830.docx"
FIGURE = ROOT / "outputs" / "paper_assets" / "figure_1_springer_submission_zh_600dpi.png"


def combined_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).strip()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if combined_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def find_caption(document: Document, prefix: str) -> Paragraph:
    matches = [
        p for p in document.paragraphs
        if p.style.name == "Caption" and combined_text(p).startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one caption beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def set_plain_text(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_text_nodes(root, replacements: list[tuple[str, str]]) -> None:
    for node in root.xpath(".//w:t"):
        value = node.text or ""
        for source, target in replacements:
            value = value.replace(source, target)
        node.text = value


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = style
    new_paragraph.add_run(text)
    return new_paragraph


def rewrite_plain_caption(paragraph: Paragraph, label: str, body: str) -> None:
    paragraph.clear()
    label_run = paragraph.add_run(f"{label} ")
    label_run.bold = True
    paragraph.add_run(body.rstrip("。"))
    paragraph.paragraph_format.keep_with_next = True


def set_caption_label(paragraph: Paragraph, source_label: str, target_label: str) -> None:
    prefix = f"{source_label} "
    first_text = next(
        (node for node in paragraph._p.xpath(".//w:t") if (node.text or "").startswith(prefix)),
        None,
    )
    if first_text is None:
        raise RuntimeError(f"Caption does not begin with {source_label!r}: {combined_text(paragraph)!r}")
    first_text.text = (first_text.text or "")[len(prefix) :]
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:b"))
    run.append(properties)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = f"{target_label} "
    run.append(text)
    paragraph._p.insert(1, run)


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0


def remove_table_row_by_label(table, label: str) -> None:
    matches = [row for row in table.rows if row.cells and row.cells[0].text.strip() == label]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one table row labelled {label!r}, found {len(matches)}")
    row = matches[0]
    row._tr.getparent().remove(row._tr)


def set_table_widths(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for grid_col, width in zip(list(grid), widths):
        grid_col.set(qn("w:w"), str(width))
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(widths)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))


def apply_run_font(run, size: float | None = None, east_asia: str = "Songti SC") -> None:
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), east_asia)


def format_tables(document: Document) -> None:
    widths = {
        0: [1700, 1250, 1350, 1800, 1900, 1072],
        1: [2450, 2200, 2200, 2222],
        2: [1800, 1900, 1400, 1400, 2572],
        3: [2400, 3336, 3336],
        4: [1500, 1650, 1650, 2136, 2136],
        5: [1900, 2500, 1200, 1736, 1736],
    }
    left_columns = {
        0: set(),
        1: set(),
        2: set(),
        3: {0},
        4: {0, 1, 2},
        5: {0, 1},
    }
    for table_index, table in enumerate(document.tables):
        set_table_widths(table, widths[table_index])
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT
                        if row_index > 0 and column_index in left_columns[table_index]
                        else WD_ALIGN_PARAGRAPH.CENTER
                    )
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        apply_run_font(run, 8.5)
                        if row_index == 0:
                            run.bold = True


def bold_label(paragraph: Paragraph, label: str) -> None:
    text = paragraph.text
    if not text.startswith(label):
        raise RuntimeError(f"Expected label {label!r} in {text!r}")
    remainder = text[len(label):].lstrip()
    paragraph.clear()
    label_run = paragraph.add_run(label)
    label_run.bold = True
    if remainder:
        paragraph.add_run(f"  {remainder}")


def replace_figure(document: Document) -> None:
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)
    shape = document.inline_shapes[0]
    relation_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    document.part.related_parts[relation_id]._blob = FIGURE.read_bytes()
    with Image.open(FIGURE) as image:
        width_px, height_px = image.size
    shape.width = Inches(6.35)
    shape.height = Inches(6.35 * height_px / width_px)
    properties = shape._inline.docPr
    properties.set("name", "图1 所提出模型的信息流")
    properties.set(
        "descr",
        "关系感知结构分支、多尺度语义分支、语义引导的一跳邻居上下文模块、训练目标与联合表示检索流程图。",
    )


def set_math_vectors(document: Document) -> None:
    vector_symbols = {"W", "H", "M", "h", "e", "m", "s", "t", "u", "z", "f", "g", "c", "q", "k", "v", "x", "y", "α"}
    active = False
    for paragraph in document.paragraphs:
        text = combined_text(paragraph)
        if text.startswith("3.3 "):
            active = True
        if text.startswith("3.6 "):
            active = False
        if not active:
            continue
        for math_run in paragraph._p.xpath(".//m:r"):
            token = "".join((node.text or "") for node in math_run.iter(qn("m:t")))
            if token not in vector_symbols:
                continue
            properties = math_run.find(qn("m:rPr"))
            if properties is None:
                properties = OxmlElement("m:rPr")
                math_run.insert(0, properties)
            style = properties.find(qn("m:sty"))
            if style is None:
                style = OxmlElement("m:sty")
                properties.append(style)
            style.set(qn("m:val"), "bi")


def set_document_language(document: Document) -> None:
    document.core_properties.language = "zh-CN"
    for part in document.part.package.parts:
        element = getattr(part, "_element", None)
        if element is None:
            continue
        for lang in element.xpath(".//w:lang"):
            lang.set(qn("w:val"), "zh-CN")
            lang.set(qn("w:eastAsia"), "zh-CN")
            lang.set(qn("w:bidi"), "zh-CN")


def format_reference_language(document: Document) -> None:
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Reference":
            continue
        properties = paragraph._p.get_or_add_pPr()
        word_wrap = properties.find(qn("w:wordWrap"))
        if word_wrap is None:
            word_wrap = OxmlElement("w:wordWrap")
            properties.append(word_wrap)
        word_wrap.set(qn("w:val"), "0")
        for run in paragraph.runs:
            run_properties = run._r.get_or_add_rPr()
            lang = run_properties.find(qn("w:lang"))
            if lang is None:
                lang = OxmlElement("w:lang")
                run_properties.append(lang)
            lang.set(qn("w:val"), "en-US")
            lang.set(qn("w:eastAsia"), "en-US")
            lang.set(qn("w:bidi"), "en-US")


def add_table_spacing(document: Document) -> None:
    children = list(document._body._element)
    for index, child in enumerate(children[:-1]):
        if child.tag != qn("w:tbl"):
            continue
        next_child = children[index + 1]
        if next_child.tag == qn("w:p"):
            paragraph = Paragraph(next_child, document._body)
            paragraph.paragraph_format.space_before = Pt(8)


def remove_line_numbering(document: Document) -> None:
    for section_properties in document._element.xpath(".//w:sectPr"):
        numbering = section_properties.find(qn("w:lnNumType"))
        if numbering is not None:
            section_properties.remove(numbering)


def protect_latin_terms(document: Document) -> None:
    terms = (
        "OpenEA EN–FR-15K-V2",
        "EventEA EN–EN-20K",
        "DBP15K ZH–EN",
        "DBP15K JA–EN",
        "DBP15K FR–EN",
        "parameter-matched",
        "1.5-entmax",
        "RPR-RHGT",
        "RREA-text",
        "GCN-Align",
        "Transformer",
        "InfoNCE",
        "DBP15K",
        "OpenEA",
        "EventEA",
        "Hits@10",
        "Hits@1",
    )
    replacements = []
    for term in terms:
        protected = "\ufeff".join(term)
        replacements.append((term, protected))
    replace_text_nodes(document._element, replacements)


def format_chinese_typography(document: Document) -> None:
    style_specs = {
        "Normal": (12, False, "Songti SC"),
        "List Number": (12, False, "Songti SC"),
        "Heading 1": (14, True, "Heiti SC"),
        "Heading 2": (12, True, "Heiti SC"),
        "Heading 3": (12, True, "Heiti SC"),
        "Caption": (10, False, "Songti SC"),
        "Reference": (10.5, False, "Songti SC"),
    }
    for name, (size, bold, east_asia) in style_specs.items():
        if name not in document.styles:
            continue
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        rpr = style._element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), east_asia)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            east_asia = "Heiti SC" if paragraph.style.name.startswith("Heading") else "Songti SC"
            apply_run_font(run, east_asia=east_asia)

    document.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in document.paragraphs[0].runs:
        apply_run_font(run, 16, "Heiti SC")
        run.bold = True
    for index, size in ((1, 12), (2, 10.5), (3, 10.5)):
        document.paragraphs[index].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in document.paragraphs[index].runs:
            apply_run_font(run, size)


def revise() -> Path:
    document = Document(SOURCE)

    document.core_properties.title = "跨语言实体对齐中的关系感知邻居上下文与语义引导选择"
    document.core_properties.subject = "跨语言知识图谱实体对齐"
    document.core_properties.keywords = (
        "知识图谱实体对齐；跨语言知识图谱；关系感知图神经网络；结构—语义融合；稀疏邻居选择；多尺度语义编码"
    )

    set_plain_text(document.paragraphs[0], "跨语言实体对齐中的关系感知邻居上下文与语义引导选择")
    set_plain_text(document.paragraphs[2], "作者单位：[院系，学校，城市，国家——请在投稿前补全]")
    set_plain_text(
        document.paragraphs[5],
        "跨语言知识图谱实体对齐旨在识别使用不同语言构建的图谱中的对应实体。大多数系统使用独立分支编码图结构和实体文本。这种设计留下了一个问题：模型形成联合表示之前，结构邻居应如何发挥作用？本文评估一种采用语义引导选择的关系感知邻居上下文模型。结构编码器保留关系类型和多个传播深度。语义编码器从同一实体文本序列构建 token、phrase 和 global 三种视图。语义查询为完整的一跳出邻域评分，1.5-entmax 去除弱相关证据。模型随后将筛选后的结构上下文与实体语义结合。实验在五个数据集上使用三个随机种子。在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上，完整模型的排名第一命中率（Hits@1）分别为 0.7285 和 0.6344。在邻域和可训练融合参数量相同的条件下，语义查询相对结构查询的 Hits@1 分别提高 2.08 和 2.48 个百分点。完整模型相对仅语义检索也分别提高 9.09 和 9.66 个百分点。去除关系类型后，Hits@1 分别下降 3.62 和 2.78 个百分点。参数匹配晚期融合对照分别下降 5.28 和 9.22 个百分点。该差异同时反映邻居访问和训练交互路径的变化。实验结果支持在当前测试条件下使用关系感知邻居上下文与语义引导选择。",
    )
    set_plain_text(
        document.paragraphs[6],
        "关键词：知识图谱实体对齐；跨语言知识图谱；关系感知图神经网络；结构—语义融合；稀疏邻居选择；多尺度语义编码",
    )

    set_plain_text(
        find_paragraph(document, "本文研究三个问题"),
        "本文围绕三个问题展开。第一，关系感知邻居上下文能否在语义表示之外提供有效证据？第二，在邻域和可训练融合参数量固定时，语义查询是否比结构查询更有效地选择邻居？第三，哪些结构和语义组件能够稳定支持联合模型？",
    )
    set_plain_text(
        find_paragraph(document, "模型在结构消息传递中保留关系类型"),
        "模型在结构消息传递中保留关系类型。节点级层选择器使每个实体能够组合浅层和深层图状态。该设计不需要为每种关系配置独立的完整变换矩阵，同时能够形成紧凑的关系感知状态。",
    )
    set_plain_text(
        find_paragraph(document, "融合模块使用语义信息"),
        "模型从同一实体文本序列构建 token、phrase 和 global 三种视图，并使用所得语义状态为完整的一跳出邻域评分。1.5-entmax 在模型形成联合表示之前去除弱相关邻居。",
    )
    set_plain_text(
        find_paragraph(document, "实验在五个数据集上使用三个随机种子"),
        "实验覆盖五个数据集，并使用三个随机种子。匹配控制分别检验关系类型、语义视图、邻居访问和查询信号。查询对照保持邻域与可训练融合参数量不变。",
    )
    set_plain_text(
        find_paragraph(document, "模型包含四个部分"),
        "模型包含四个部分。结构编码器保留关系类型并组合多个传播深度。语义编码器从同一实体文本构建 token、phrase 和 global 三种视图。融合模块使用语义信息为全部一跳出邻居评分，形成稀疏结构上下文，并将该上下文与实体语义结合。训练阶段在联合表示和结构表示上使用信息噪声对比估计（InfoNCE）。检索阶段使用训练后的联合表示，并采用跨域相似度局部缩放（CSLS）进行校正。图 1 展示完整的信息流。",
    )

    direction_anchor = find_paragraph(document, "接着结构编码器在每次前向计算中")
    insert_after(
        direction_anchor,
        "结构编码器与上下文模块对边的使用方式不同。图传播使用传入边更新当前节点；融合阶段则把当前实体发出的事实所指向的出邻居视为候选上下文。这一区分解释了两个模块为何采用相反的边方向。",
    )

    figure_caption = find_caption(document, "图1")
    figure_text_nodes = figure_caption._p.xpath(".//w:t")
    figure_text_nodes[0].text = (
        "图1 所提出模型的信息流。结构分支保留关系类型并选择传播深度，语义分支从同一输入构建三种视图。"
        "语义引导查询为完整的一跳出邻域评分，1.5-entmax 去除弱相关证据，模型随后形成联合表示。"
        "训练和主检索均使用该联合表示，验证集只选择 "
    )
    if figure_text_nodes[-1].text == "。":
        figure_text_nodes[-1].text = ""
    set_caption_label(figure_caption, "图1", "图 1")
    figure_caption.paragraph_format.keep_with_next = False

    replace_text_nodes(
        find_paragraph(document, "但仅使用最深层")._p,
        [("由两层 MLP 输出标量分数", "由两层多层感知机（MLP）输出标量分数")],
    )
    replace_text_nodes(
        find_paragraph(document, "层选择 MLP")._p,
        [("采用 GELU", "采用高斯误差线性单元（GELU）")],
    )

    set_plain_text(
        find_paragraph(document, "表4报告"),
        "表 4 报告 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的组件控制结果。每个变体均使用三个随机种子，并在训练与检索阶段使用同一种表示。",
    )
    set_plain_text(
        find_paragraph(document, "关系类型提供了最清楚"),
        "关系类型提供了最清楚且方向一致的结构增益。移除关系类型后，DBP15K 和 OpenEA 的 Hits@1 分别下降 3.62 和 2.78 个百分点。拓扑特征也支持结构状态，相对仅使用实体向量分别提高 3.46 和 8.03 个百分点。phrase 视图使两个数据集分别提高 0.81 和 3.95 个百分点。其余对照的变化较小或依赖具体数据集，因此本文将它们视为辅助设计选择，而不是独立贡献。",
    )
    set_plain_text(
        find_paragraph(document, "表6比较完整模型"),
        "表 6 比较完整模型与两种对照。结构查询对照只改变为邻居评分的信号。参数匹配晚期融合对照保持相同的可训练融合参数量，但不读取邻居上下文。",
    )
    set_plain_text(
        find_paragraph(document, "结构查询对照用于分离查询信号"),
        "结构查询对照用于分离查询信号的影响。该对照保留相同的邻域与动态补齐。它也保留 1.5-entmax、门控和联合检索，只将语义查询替换为实体结构表示。晚期融合对照检验更广泛的变化。它保持相同的可训练融合参数量，但不读取邻居上下文，而是在两个分支独立编码后再形成联合表示。",
    )
    set_plain_text(
        find_paragraph(document, "匹配控制结果显示"),
        "匹配比较显示，结构查询的 Hits@1 比语义查询分别低 2.08 和 2.48 个百分点。该结果在当前实验协议下分离出语义引导邻居评分的作用。参数匹配晚期融合对照比完整模型分别低 5.28 和 9.22 个百分点。第二组差距同时反映邻居访问和训练交互路径的变化。由于两个因素一起发生改变，因此该差距不能单独衡量融合时机。",
    )
    set_plain_text(
        find_paragraph(document, "实验支持三项结论"),
        "实验支持三项结论。第一，关系类型在两个受控数据集上均提供方向一致的结构增益。第二，模型将结构上下文与语义共同学习后，结构信息能够增加有效证据。第三，在邻域和可训练融合参数量固定时，语义查询优于结构查询。这些结果共同支持当前实验协议下的关系感知邻居上下文与语义引导选择。",
    )
    set_plain_text(
        find_paragraph(document, "本文使用"),
        "本文使用三个随机种子，因此只报告效应大小和样本波动，不作较强的显著性声明。固定邻居对照同时改变邻居数量和权重函数。参数匹配晚期融合对照同时改变邻居访问和训练交互路径。组件控制只覆盖两个主要数据集。因此，现有证据支持本文测试的设置，而不能直接推广到所有实体对齐场景。",
    )
    set_plain_text(
        find_paragraph(document, "本文研究结构邻居应在何处进入"),
        "本文研究结构邻居应如何进入结构—语义实体对齐。关系类型在两个受控数据集上均提高了结构证据的有效性。可训练联合表示使结构上下文在语义信息之外产生增益。在其他条件一致时，语义查询优于结构查询。完整邻域访问与可训练交互路径的组合优于不读取邻居上下文的晚期融合，但该差距不能单独说明融合时机。实验结果支持当前协议下的关系感知邻居上下文与语义引导选择。",
    )

    global_replacements = [
        ("DBP15K zh-en", "DBP15K ZH–EN"),
        ("DBP15K ja-en", "DBP15K JA–EN"),
        ("DBP15K fr-en", "DBP15K FR–EN"),
        ("OpenEA en-fr", "OpenEA EN–FR-15K-V2"),
        ("OpenEA EN–FR-V2", "OpenEA EN–FR-15K-V2"),
        ("EventEA en-en", "EventEA EN–EN-20K"),
        ("邻域和模型规模相同", "邻域和可训练融合参数量相同"),
        ("训练路径", "训练交互路径"),
    ]
    replace_text_nodes(document._element, global_replacements)

    implementation = find_paragraph(document, "模型使用 AdamW")
    replace_text_nodes(
        implementation._p,
        [
            ("关系感知 GNN 固定为 3 层", "关系感知图神经网络固定为 3 层"),
            ("DBP15K 最多训练 36 个 epoch，OpenEA 和 EventEA 最多训练 50 个 epoch", "DBP15K 最多训练 36 个 epoch，OpenEA 与 EventEA 最多训练 50 个 epoch"),
        ],
    )

    table1_caption = find_caption(document, "表 1")
    intro = table1_caption.insert_paragraph_before("表 1 汇总五个数据集的统计信息与实验划分。", style="Normal")
    intro.paragraph_format.keep_with_next = True

    captions = {
        "表 1": "五个数据集的统计信息与实验划分",
        "表 2": "五个数据集的最终测试结果（均值 ± 样本标准差，n = 3）",
        "表 3": "文献报告结果与本文模型的定位性比较",
        "表 4": "结构组件、邻居聚合与语义视图的消融结果（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
        "表 5": "分支、训练表示与主检索表示对照（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
        "表 6": "邻居访问、查询信号与容量控制结果（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
    }
    for label, body in captions.items():
        rewrite_plain_caption(find_caption(document, label), label, body)

    dataset_names = [
        "DBP15K ZH–EN",
        "DBP15K JA–EN",
        "DBP15K FR–EN",
        "OpenEA EN–FR-15K-V2",
        "EventEA EN–EN-20K",
    ]
    for table_index in (0, 1):
        for row_index, dataset_name in enumerate(dataset_names, start=1):
            family, benchmark = dataset_name.split(" ", 1)
            set_cell_text(document.tables[table_index].cell(row_index, 0), f"{family}\n{benchmark}")
    set_cell_text(document.tables[2].cell(7, 0), "OpenEA\nEN–FR-15K-V2")

    set_cell_text(document.tables[3].cell(0, 1), "DBP15K Hits@1/MRR")
    set_cell_text(document.tables[3].cell(0, 2), "OpenEA Hits@1/MRR")
    set_cell_text(document.tables[4].cell(0, 3), "DBP15K Hits@1/MRR")
    set_cell_text(document.tables[4].cell(0, 4), "OpenEA Hits@1/MRR")
    set_cell_text(document.tables[5].cell(0, 3), "DBP15K Hits@1/MRR")
    set_cell_text(document.tables[5].cell(0, 4), "OpenEA Hits@1/MRR")
    set_cell_text(document.tables[4].cell(4, 0), "无可训练参数的平均融合")
    set_cell_text(document.tables[5].cell(1, 0), "完整模型")
    set_cell_text(document.tables[5].cell(2, 0), "结构查询对照")
    set_cell_text(document.tables[5].cell(3, 0), "参数匹配晚期融合对照")
    set_cell_text(
        document.tables[5].cell(3, 1),
        "不读取邻居上下文；两个分支独立编码后形成并检索晚期融合表示",
    )
    remove_table_row_by_label(document.tables[3], "去除结构监督")

    set_plain_text(find_paragraph(document, "声明"), "声明")
    funding = find_paragraph(document, "基金项目")
    set_plain_text(funding, "基金项目  本研究未获得专项资助。")
    for label in (
        "基金项目",
        "作者贡献",
        "数据可用性",
        "代码可用性",
        "利益冲突",
        "伦理审批",
        "参与同意",
        "发表同意",
    ):
        bold_label(find_paragraph(document, label), label)

    reference_eight = find_paragraph(document, "[8]")
    if "10.14778/3407790.3407828" not in combined_text(reference_eight):
        for run in reference_eight.runs:
            if run.text:
                run.text = run.text.rstrip() + " https://doi.org/10.14778/3407790.3407828"
                break

    replace_figure(document)
    set_math_vectors(document)
    protect_latin_terms(document)
    format_chinese_typography(document)
    format_tables(document)
    add_table_spacing(document)
    set_document_language(document)
    format_reference_language(document)
    remove_line_numbering(document)

    for paragraph in document.paragraphs:
        if paragraph.style.name == "Caption":
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(4)
        if paragraph.style.name == "Formula":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(revise())
