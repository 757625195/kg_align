from __future__ import annotations

import csv
import re
import sys
from copy import deepcopy
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "文献报告值修订版_20260828.docx"
)
OUTPUT = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "方法学控制与直接联合检索修订版_20260829.docx"
)
SUMMARY = ROOT / "outputs/methodology_controls_20260829/summary/direct_joint_aggregate.tsv"
FIGURE = ROOT / "outputs/paper_assets/figure_1_new_main_validated_protocol.png"

sys.path.insert(0, str(ROOT / "scripts"))
from mathify_all_document_symbols import set_markup  # noqa: E402


CITATION_RE = re.compile(r"\[(\d+(?:\s*(?:,|–|-)\s*\d+)*)\]")


def combined_text(paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if combined_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def rewrite(document: Document, prefix: str, markup: str) -> None:
    set_markup(find_paragraph(document, prefix), markup)


def insert_after(paragraph, text: str, style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    new_p.getparent().replace(new_p, new_paragraph._p)
    if style:
        new_paragraph.style = style
    set_markup(new_paragraph, text)
    return new_paragraph


def insert_before(paragraph, text: str, style: str | None = None):
    new_paragraph = paragraph._parent.add_paragraph()
    new_paragraph._p.getparent().remove(new_paragraph._p)
    paragraph._p.addprevious(new_paragraph._p)
    if style:
        new_paragraph.style = style
    set_markup(new_paragraph, text)
    return new_paragraph


def remove_paragraph(paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def read_summary() -> Dict[Tuple[str, str], Dict[str, float]]:
    if not SUMMARY.exists():
        raise FileNotFoundError(f"Run methodology control aggregation first: {SUMMARY}")
    rows: Dict[Tuple[str, str], Dict[str, float]] = {}
    with SUMMARY.open(encoding="utf-8", newline="") as handle:
        for raw in csv.DictReader(handle, delimiter="\t"):
            if int(raw["runs"]) != 3:
                raise RuntimeError(f"Expected three runs for {raw['variant']} {raw['dataset']}")
            rows[(raw["variant"], raw["dataset"])] = {
                key: float(raw[key])
                for key in (
                    "Hits@1_mean",
                    "Hits@1_std",
                    "Hits@10_mean",
                    "Hits@10_std",
                    "MRR_mean",
                    "MRR_std",
                )
            }
    required = {
        (variant, dataset)
        for variant in (
            "main",
            "learned_structure_init",
            "fixed_k8_softmax",
            "no_relation_types",
            "no_structure_supervision",
            "no_layer_selector",
            "no_token_view",
            "no_phrase_view",
            "no_global_view",
            "structure_only",
            "semantic_only",
            "mean_fusion",
            "structural_neighbor_query",
            "late_concat_matched",
            "disjoint_relation_vocab",
            "bidirectional_same_relation",
        )
        for dataset in ("dbp15k_zh_en", "openea_en_fr")
    }
    required.update(
        ("main", dataset)
        for dataset in ("dbp15k_ja_en", "dbp15k_fr_en", "eventea_en_en")
    )
    missing = sorted(required - rows.keys())
    if missing:
        raise RuntimeError(f"Missing completed three-seed controls: {missing}")
    return rows


def result(rows, variant: str, dataset: str) -> Dict[str, float]:
    return rows[(variant, dataset)]


def metric_text(row: Dict[str, float], metric: str) -> str:
    return f"{row[f'{metric}_mean']:.4f} ± {row[f'{metric}_std']:.4f}"


def pair_text(row: Dict[str, float]) -> str:
    return f"{metric_text(row, 'Hits@1')} / {metric_text(row, 'MRR')}"


def pp(main: Dict[str, float], control: Dict[str, float], metric: str = "Hits@1") -> float:
    return (main[f"{metric}_mean"] - control[f"{metric}_mean"]) * 100.0


def set_cell(cell, markup: str) -> None:
    set_markup(cell.paragraphs[0], markup)


def append_row_like(table, template_index: int = -1):
    template = deepcopy(table.rows[template_index]._tr)
    table._tbl.append(template)
    return table.rows[-1]


def insert_column(table, index: int, width_inches: float) -> None:
    table.add_column(Inches(width_inches))
    for row in table.rows:
        cells = list(row._tr.findall(qn("w:tc")))
        new_cell = cells[-1]
        row._tr.remove(new_cell)
        row._tr.insert(index, new_cell)
    grid = table._tbl.tblGrid
    cols = list(grid)
    new_col = cols[-1]
    grid.remove(new_col)
    grid.insert(index, new_col)


def set_table_widths(table, widths: Iterable[int]) -> None:
    widths = list(widths)
    grid = table._tbl.tblGrid
    for grid_col, width in zip(list(grid), widths):
        grid_col.set(qn("w:w"), str(width))
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def rewrite_core_text(document: Document, rows) -> None:
    main_dbp = result(rows, "main", "dbp15k_zh_en")
    main_open = result(rows, "main", "openea_en_fr")
    dis_dbp = result(rows, "disjoint_relation_vocab", "dbp15k_zh_en")
    dis_open = result(rows, "disjoint_relation_vocab", "openea_en_fr")
    bi_dbp = result(rows, "bidirectional_same_relation", "dbp15k_zh_en")
    bi_open = result(rows, "bidirectional_same_relation", "openea_en_fr")
    query_dbp = result(rows, "structural_neighbor_query", "dbp15k_zh_en")
    query_open = result(rows, "structural_neighbor_query", "openea_en_fr")
    late_dbp = result(rows, "late_concat_matched", "dbp15k_zh_en")
    late_open = result(rows, "late_concat_matched", "openea_en_fr")

    main_results = [result(rows, "main", dataset) for dataset in (
        "dbp15k_zh_en", "dbp15k_ja_en", "dbp15k_fr_en", "openea_en_fr", "eventea_en_en"
    )]
    h1_min = min(row["Hits@1_mean"] for row in main_results)
    h1_max = max(row["Hits@1_mean"] for row in main_results)

    rewrite(
        document,
        "跨语言知识图谱实体对齐需要同时处理",
        "跨语言知识图谱实体对齐旨在识别不同语言知识图谱中指向同一现实对象的实体，是知识融合的基础任务。现有方法虽然分别利用图结构和文本语义，"
        "但关系类型、不同传播深度与碎片化文本往往被独立处理，结构邻居在联合表示中的作用也缺少受控验证。为此，本文提出关系感知结构上下文与多尺度语义融合方法。"
        "结构编码器通过共享关系投影进行关系感知消息传递，并以节点级层选择整合不同传播深度；语义编码器从同一实体文本中提取 token、phrase 和 global 三种粒度的表示。"
        "在此基础上，融合模块从完整一跳出邻域中利用 1.5-entmax 选择稀疏结构证据，并将邻居上下文与语义表示联合编码。模型采用双向 InfoNCE 训练，并使用 CSLS 进行跨图检索。"
        f"五个数据集上的三随机种子实验显示，Hits@1 位于 {h1_min:.4f}—{h1_max:.4f}，样本标准差均不超过 0.0053。"
        f"在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2 上，去除关系类型使 Hits@1 分别下降 {pp(main_dbp, result(rows, 'no_relation_types', 'dbp15k_zh_en')):.2f} 和 {pp(main_open, result(rows, 'no_relation_types', 'openea_en_fr')):.2f} 个百分点，"
        f"仅使用语义分支则分别下降 {pp(main_dbp, result(rows, 'semantic_only', 'dbp15k_zh_en')):.2f} 和 {pp(main_open, result(rows, 'semantic_only', 'openea_en_fr')):.2f} 个百分点。"
        f"在保持邻域集合与融合容量不变时，将语义查询替换为结构查询仅下降 {pp(main_dbp, query_dbp):.2f} 和 {pp(main_open, query_open):.2f} 个百分点。"
        "结果表明，关系感知邻居上下文与结构—语义联合建模能够改善实体检索，而查询信号本身的影响相对有限。",
    )

    rewrite(
        document,
        "第一，结构相似不等于实体等价",
        "第一，结构相似不等于实体等价。例如，人物 A 通过“出生地”连接北京，人物 B 通过“工作地点”连接北京；若模型把不同关系退化为无类型连接，就会高估两人的相似性[5,7]。"
        "第二，语义证据往往是碎片化的。例如，单独出现的“Washington”无法区分人物、州和城市，需要结合关系或属性上下文。"
        "第三，人工核验跨图实体对需要逐对判断两侧实体是否指向同一现实对象，公开数据中的种子对齐因而只覆盖实体集合的一部分[4,8]。"
        "例如，训练集中已确认“姚明—Yao Ming”并不意味着其他人物都已标注；有限种子使模型必须从少量对应关系中学习可迁移的结构与语义线索。",
    )
    rewrite(document, "1.2 研究目标与主要贡献", "1.2 研究目标、研究问题与贡献")
    objective = find_paragraph(document, "本研究旨在构建结构与语义协同")
    set_markup(
        objective,
        "本研究旨在评估关系感知结构上下文与多尺度语义联合建模对跨语言实体对齐的作用，并通过受控实验区分关系建模、传播深度、语义粒度、邻居查询信号与邻居上下文的影响。",
    )
    insert_after(
        objective,
        "本研究围绕以下三个问题展开。RQ1：所提出的方法在跨语言、跨知识库来源和事件知识图谱上能否获得稳定的实体检索结果？"
        "RQ2：结构初始化、结构监督、关系类型、跨图关系编号共享、边方向、层选择、邻居稀疏化以及 token、phrase 和 global 语义视图分别产生何种影响？"
        "RQ3：在保持邻居集合和融合容量一致时，使用语义表示计算邻居权重是否优于使用结构表示；不读取邻居上下文的参数匹配晚期融合与主模型有何差异？",
        objective.style.name,
    )
    rewrite(
        document,
        "围绕上述目标，本文的主要贡献如下",
        "围绕上述问题，本文完成了以下工作。",
    )
    rewrite(
        document,
        "设计参数共享的关系感知结构编码器",
        "提出一种关系感知的结构—语义实体对齐模型。结构编码器利用关系类型区分不同边的作用，并结合不同传播深度的节点表示；语义编码器从同一实体文本中提取词级、局部短语和全局上下文信息。两类表示在统一的对齐目标下进行学习。",
    )
    rewrite(
        document,
        "构建多尺度语义与结构上下文融合方法",
        "在联合表示形成之前引入实体的一跳结构邻域。模型根据当前实体与邻居之间的相关性分配稀疏权重，再将邻居摘要、实体自身结构和语义表示合并。为分析性能变化的来源，本文还设置了结构查询和参数量匹配的晚期融合模型作为对照。",
    )
    rewrite(
        document,
        "开展多数据集实验和受控消融分析",
        "在五个实体对齐数据集上进行三次独立实验，并在两个代表性数据集上完成组件消融和控制实验。结果表明，关系类型、多尺度语义和邻居上下文均能改善实体检索；在邻域保持一致时，语义查询与结构查询的结果接近。因此，实验支持在联合表示中保留结构邻域信息，但尚不能说明语义查询本身稳定优于结构查询。",
    )
    for contribution_prefix in (
        "提出一种关系感知的结构—语义实体对齐模型",
        "在联合表示形成之前引入实体的一跳结构邻域",
        "在五个实体对齐数据集上进行三次独立实验",
    ):
        contribution = find_paragraph(document, contribution_prefix)
        for run in contribution.runs:
            run.bold = False

    recent_anchor = find_paragraph(document, "已有文本增强实体对齐研究表明")
    insert_after(
        recent_anchor,
        "近期研究进一步关注噪声标注下的大语言模型监督[16]、真实知识图谱的多源域适应[17]以及基于多步推理代理的实体对齐[18]。这些方法改变了监督来源或推理流程；本文不引入伪标签、外部大语言模型或代理推理，而是在固定种子监督下分析结构邻居与语义表示的交互。",
        recent_anchor.style.name,
    )

    rewrite(
        document,
        "其中，Aseed 称为种子对齐集合",
        "其中，[[\\mathcal A_{\\mathrm{seed}}]] 为种子对齐集合，包含由人工或数据集规则确认的等价实体对，[[e_i\\equiv e_j]] 表示两者指向同一现实对象。"
        "本文采用传导式一对一闭世界评估：训练时可见两张图的全部实体和边，但只使用训练种子作为跨图实体监督；测试查询在另一侧完整候选集合中恰有一个真实对应实体。",
    )
    task_anchor = find_paragraph(document, "目标是学习评分函数")
    insert_after(
        task_anchor,
        "关系预处理是任务输入的一部分。主协议对关系 URI 或名称进行规范化，并把两图同名关系映射到共享编号；DBP15K 还使用数据集提供的 sup_rel_ids 合并已知对应关系。"
        "因此，主协议包含由关系名称和基准文件提供的跨图关系对应信息。第 5.3 节的独立关系词表对照为左右图分配互不共享的关系编号，用于量化该信息对结果的影响。",
        task_anchor.style.name,
    )
    rewrite(
        document,
        "图 1 关系感知结构上下文与多尺度语义融合方法",
        "图 1 关系感知结构上下文与多尺度语义融合方法。结构分支沿传入边执行关系感知传播并选择不同深度，语义分支从同一输入构造 token、phrase 和 global 视图；"
        "融合模块沿原始方向读取全部一跳出邻居并以 1.5-entmax 稀疏加权，所得结构上下文与语义表示形成联合表示。训练和主检索均使用联合表示，验证集仅选择 [[k_{\\mathrm{CSLS}}]]。",
    )

    rewrite(
        document,
        "接着结构编码器对整张知识图谱执行",
        "接着结构编码器在每次前向计算中对整张知识图谱执行有向关系感知消息传递。对三元组 [[(j,r,i)]]，头实体 [[j]] 沿原始方向向尾实体 [[i]] 发送消息。"
        "主协议不自动添加反向边，因此结构编码器只聚合传入边。设 [[\\mathbf h_i^{(l)}]] 为实体 [[i]] 在第 [[l]] 层的状态，[[\\mathbf e_r]] 为关系嵌入，"
        "[[\\mathcal N_{\\mathrm{in}}(i)]] 为指向 [[i]] 的传入边，第 [[l]] 层首先计算：",
    )
    rewrite(
        document,
        "结构编码在整图上执行一次",
        "结构编码在每次前向计算中对整图执行一次，随后按批次实体索引其一跳出邻域内全部邻居的结构表示。这里的融合邻域沿原始边方向读取出邻居，"
        "与第 3.3 节聚合传入边的结构编码方向不同。该设置区分“由哪些事实更新当前节点”和“当前节点指向哪些候选证据”；第 5.3 节另设共享关系类型的双向边对照，"
        "检验这一方向选择是否决定结果。不同大小的邻域仅在批次内动态补齐，填充位置由掩码排除，无出邻居实体保留完全掩蔽的占位位置。",
    )
    semantic_anchor = find_paragraph(document, "对于实体 i，本文将其名称")
    insert_after(
        semantic_anchor,
        "例如，实体“姚明”的共享序列可包含名称“姚明”、关系短语“出生地”和“效力球队”以及属性片段“身高 2.29 米”。token 视图汇总“姚明”“身高”等单项线索，"
        "phrase 视图识别“出生地 上海”或“效力球队 火箭队”等局部组合，global 视图则允许名称、关系和相距较远的属性在完整序列中相互参照。三种视图改变的是编码范围，而不是数据来源。",
        semantic_anchor.style.name,
    )
    entmax_anchor = find_paragraph(document, "式（17）将全部有效邻居")
    insert_after(
        entmax_anchor,
        "仍以“姚明”为例，若一跳出邻居包括“休斯敦火箭队”“上海”和一个低相关候选，语义查询可依据“篮球运动员”“效力球队”等线索提高火箭队邻居的得分；"
        "1.5-entmax 可把低相关候选的权重压为 0，再由式（19）决定保留姚明自身结构还是采用加权邻居摘要。该例只说明计算路径，不预设模型一定选择某个邻居。",
        entmax_anchor.style.name,
    )
    rewrite(
        document,
        "由此，语义信息是在结构邻居聚合之前",
        "由此，语义查询在邻居摘要形成前参与权重计算，更新后的结构上下文再与语义表示生成联合表示。本文将其称为语义查询邻居。"
        "为避免把架构位置直接等同于有效贡献，第 3.5.4 节使用相同邻域的结构查询控制，使语义只在最终联合门进入。",
    )
    rewrite(
        document,
        "为控制训练阶段融合容量并比较邻居查询信号",
        "为分别考察语义进入邻居计算的时机和邻居上下文的可见性，本文设置两种控制。结构查询控制保留相同的一跳出邻域、动态补齐、1.5-entmax 和门控参数，"
        "但以实体自身结构表示计算邻居权重；语义仅在结构上下文形成后进入联合门。参数匹配晚期融合不读取邻居上下文，而在结构分支和语义分支独立编码后形成训练表示：",
    )
    rewrite(
        document,
        "式（22）先分别归一化结构表示与语义表示",
        "式（22）先分别归一化结构表示与语义表示，再通过拼接和非线性映射得到用于 InfoNCE 训练及检索的 [[\\mathbf z_i^{\\mathrm{late}}]]。"
        "该基线与主模型保持相同的可训练融合参数量。结构查询控制单独改变查询信号；晚期基线同时移除邻居可见性并改变训练交互路径，因而只能估计二者的联合影响。",
    )
    hard_negative = find_paragraph(document, "难负候选（hard negative candidate）")
    remove_paragraph(hard_negative)
    rewrite(
        document,
        "因此，结构分支辅助目标在第一个 epoch",
        "因此，结构分支辅助目标在第一个 epoch 的权重为 0.1，并在最后一个 epoch 衰减为 0。该目标只在训练前期约束结构空间；第 5.3 节的消融用于判断其实际影响。",
    )
    rewrite(
        document,
        "训练阶段使用式（25）的联合与结构监督优化编码器",
        "训练阶段使用式（25）优化编码器。主检索直接使用训练所约束的联合表示 [[\\mathbf z_i^{\\mathrm{joint}}]]，不再通过验证集对结构和语义分支进行二次加权。"
        "对联合表示的余弦相似度采用 CSLS 校正[24]：",
    )
    rewrite(
        document,
        "其中，rL 和 rR 分别表示实体到另一侧最近",
        "其中，[[r_L]] 和 [[r_R]] 分别表示查询实体与候选实体到另一侧最近 [[k_{\\mathrm{CSLS}}]] 个实体的平均余弦相似度。"
        "每个训练运行只依据验证集 MRR 从候选集合中选择 [[k_{\\mathrm{CSLS}}]]，测试集仅评估一次。旧的验证集分支加权检索作为第 5.5 节敏感性对照，不参与主结果。",
    )

    rewrite(document, "4.1 数据集与协议分层", "4.1 数据集与实验划分")
    remove_paragraph(find_paragraph(document, "实验围绕三个研究问题展开"))
    rewrite(
        document,
        "实现采用 AdamW",
        "模型使用 AdamW[25] 优化，batch size 为 512，表示维度为 128，关系感知 GNN 固定为 3 层，dropout 为 0.1。语义输入由 300 维映射至 128 维；"
        "global 视图使用两层、四头 Transformer，phrase 视图采用宽度 3 和 5 的卷积。1.5-entmax 温度为 [[T=0.25]]，InfoNCE 温度为 [[\\tau=0.07]]。"
        "随机种子为 42、43 和 44；DBP15K 最多训练 36 个 epoch，OpenEA 与 EventEA 最多训练 50 个 epoch。每 5 个 epoch 依据验证集 MRR 保存检查点，连续 4 次验证未提升时提前停止。"
        "主检索不选择分支融合权重；每个运行从 [[k_{\\mathrm{CSLS}}\\in\\{3,5,7,10,15,20\\}]] 中按验证集 MRR 选择一个值。",
    )

    rewrite(
        document,
        "表 2 汇总完整模型",
        "表 2 汇总直接使用联合表示检索的五数据集结果。每个数据集使用三个随机种子独立训练，DBP15K 从训练对固定留出 10% 作为验证集，OpenEA 与 EventEA 使用官方验证集。",
    )
    rewrite(
        document,
        "表 2 给出的五组结果",
        "五个数据集 Hits@1 的均值位于 "
        f"{h1_min:.4f}—{h1_max:.4f}，三个随机种子的样本标准差均不超过 "
        f"{max(row['Hits@1_std'] for row in main_results):.4f}。这说明当前协议下的随机波动小于不同数据条件之间的结果差异；数据集差异与输入限制在第 6.1 节讨论。",
    )
    rewrite(
        document,
        "表 3 采用既有文献公开报告的结果",
        "表 3 采用既有文献公开报告的结果，而非本地重跑。DBP15K ZH–EN 中，MTransE、JAPE、BootEA 和 RDGCN 取自 RDGCN 的汇总[5]，RREA-text 取自 RREA 的文本增强设置[6]；"
        "OpenEA EN–FR-V2 取自 OpenEA 官方五折平均[8]。本文在 DBP15K 报告三个随机种子均值，在 OpenEA 报告官方第一折的三个随机种子均值。由于划分、输入和重复运行协议不完全一致，该表只用于定位性能范围。",
    )
    rewrite(
        document,
        "在上述定位性比较中",
        f"定位性比较中，本文在 DBP15K ZH–EN 的 Hits@1、Hits@10 和 MRR 分别为 {main_dbp['Hits@1_mean']:.4f}、{main_dbp['Hits@10_mean']:.4f} 和 {main_dbp['MRR_mean']:.4f}；"
        f"OpenEA EN–FR-V2 分别为 {main_open['Hits@1_mean']:.4f}、{main_open['Hits@10_mean']:.4f} 和 {main_open['MRR_mean']:.4f}。跨文献数值不构成严格排名，本文的主要证据来自同一协议的多随机种子消融和控制实验。",
    )
    rewrite(
        document,
        "表 4 在 DBP15K",
        "表 4 在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2 上比较结构初始化、结构监督、邻居选择、关系类型、跨图关系编号共享、边方向、层选择和三种语义视图。"
        "所有行均重新训练三个随机种子，并直接检索各自训练表示。独立关系词表移除同名关系与 sup_rel_ids 的跨图合并，使 DBP15K 和 OpenEA 的关系编号数分别由 2,133/258 增至 3,024/359；"
        "该对照不匹配关系嵌入参数量，估计的是移除跨图对应与扩大关系词表的联合效应。双向边控制为每条边加入反向边但复用原关系类型，因此不增加关系嵌入参数。",
    )
    rewrite(
        document,
        "表 4 结构组件与语义视图消融",
        "表 4 结构组件、关系预处理与语义视图控制（Hits@1/MRR，均值[[\\pm]]样本标准差，[[n=3]]）",
    )
    rewrite(
        document,
        "表 4 显示",
        f"表 4 显示，移除关系类型在 DBP15K 和 OpenEA 上分别降低 Hits@1 {pp(main_dbp, result(rows, 'no_relation_types', 'dbp15k_zh_en')):.2f} 和 {pp(main_open, result(rows, 'no_relation_types', 'openea_en_fr')):.2f} 个百分点。"
        f"改用独立关系词表后分别降低 {pp(main_dbp, dis_dbp):.2f} 和 {pp(main_open, dis_open):.2f} 个百分点，说明关系类型消融的增益同时包含关系语义建模与预处理阶段跨图关系对应信息，不能全部归因于编码器。"
        f"加入共享关系类型的反向边后相对主协议分别变化 {(bi_dbp['Hits@1_mean']-main_dbp['Hits@1_mean'])*100:+.2f} 和 {(bi_open['Hits@1_mean']-main_open['Hits@1_mean'])*100:+.2f} 个百分点。"
        f"去除结构监督仅分别降低 {pp(main_dbp, result(rows, 'no_structure_supervision', 'dbp15k_zh_en')):.2f} 和 {pp(main_open, result(rows, 'no_structure_supervision', 'openea_en_fr')):.2f} 个百分点，因此它只提供小幅辅助。"
        "phrase 对 OpenEA 的影响最大，global 的作用随数据集变化，token 的边际影响较小。",
    )
    rewrite(
        document,
        "表 5 比较完整模型",
        "表 5 将训练表示与主检索表示分列，比较完整模型、结构单分支、语义单分支和无可训练平均融合。所有结果直接检索训练所用表示，不再通过验证集改变结构与语义权重。",
    )
    rewrite(
        document,
        "表 5 分支与融合必要性实验",
        "表 5 分支、训练表示与主检索表示控制（Hits@1/MRR，均值[[\\pm]]样本标准差，[[n=3]]）",
    )
    rewrite(
        document,
        "表 5 表明",
        f"表 5 表明，语义单分支明显强于结构单分支，但仍低于完整模型：DBP15K 与 OpenEA 的 Hits@1 分别低 {pp(main_dbp, result(rows, 'semantic_only', 'dbp15k_zh_en')):.2f} 和 {pp(main_open, result(rows, 'semantic_only', 'openea_en_fr')):.2f} 个百分点。"
        "无可训练平均融合也低于完整模型，说明结构信息的主要价值体现在与语义表示共同形成可训练联合表示，而不是单独检索。",
    )
    rewrite(
        document,
        "表 6 在固定结构与语义编码器",
        "表 6 比较三类控制。完整模型和结构查询控制具有相同的一跳出邻域、融合参数及直接联合检索，二者只改变语义是否在邻居聚合前参与查询；参数匹配晚期融合不读取邻居上下文；"
        "旧分支加权行仅改变完整模型的检索表示，用于检验训练—检索表示不一致的影响。",
    )
    rewrite(
        document,
        "在邻居可见范围和训练容量不变时",
        f"在邻居集合、训练容量和检索表示不变时，结构查询相对语义查询的 Hits@1 在 DBP15K 和 OpenEA 上仅低 {pp(main_dbp, query_dbp):.2f} 和 {pp(main_open, query_open):.2f} 个百分点。"
        "因此，当前证据不支持把语义提前参与查询描述为主要独立贡献。"
        f"参数匹配晚期融合分别低 {pp(main_dbp, late_dbp):.2f} 和 {pp(main_open, late_open):.2f} 个百分点，表明邻居可见性与训练交互路径的联合改变影响较大，但不能把该差值仅解释为融合时机。"
        "直接联合检索与旧分支加权检索的差异不超过 0.20 个百分点，且方向随数据集变化；采用前者主要是为了保证训练与测试表示一致。",
    )
    rewrite(
        document,
        "综合 RQ1–RQ3",
        "综合 RQ1–RQ3，模型在五个评测设置上训练稳定。关系类型和结构初始化提供较稳定的结构证据，但独立关系词表控制表明，主协议的共享关系编号也贡献了部分增益。"
        "结构查询与语义查询在相同邻域下结果接近，说明是否在邻居聚合前使用语义并非当前模型的主要决定因素；不读取邻居的晚期基线明显较弱，但其差值仍包含邻居可见性和训练路径两项变化。",
    )
    rewrite(
        document,
        "内部有效性方面",
        "内部有效性方面，主实验和消融使用相同划分、验证规则和三个随机种子，但 [[n=3]] 仍不足以支持高功效显著性检验。"
        "主检索直接使用联合表示并仅从验证集选择 [[k_{\\mathrm{CSLS}}]]，消除了旧协议中的训练—检索表示不一致。关系词表独立化和双向边控制披露了关系对应信息与边方向的影响。"
        "然而，独立词表对照同时增加关系嵌入数量，参数匹配晚期融合则同时改变邻居可见性和训练交互路径，因而前者不能给出严格容量匹配的关系监督效应，后者不能识别单一融合时机的因果效应。"
        "外部有效性方面，组件控制只覆盖两个代表性数据集，OpenEA 主结果只使用一个官方划分，且尚未验证开放世界和无对应实体场景。",
    )
    rewrite(
        document,
        "本文围绕关系感知结构上下文",
        f"本文围绕关系感知结构上下文与多尺度语义表示开展建模和受控实验。五个数据集的三随机种子 Hits@1 位于 {h1_min:.4f}—{h1_max:.4f}。"
        "直接联合检索保证了训练与测试表示一致；独立关系词表控制表明共享关系对应信息会影响关系组件的增益解释；双向边控制进一步限定了入边传播与出邻居融合的方向假设。"
        f"在相同邻域下，结构查询相对语义查询仅低 {pp(main_dbp, query_dbp):.2f} 和 {pp(main_open, query_open):.2f} 个百分点，因此本文把主要贡献限定为关系感知邻居上下文与多尺度语义的联合建模，而不把语义查询时机或结构辅助损失描述为已获得普遍验证的独立创新。",
    )

    reference_heading = find_paragraph(document, "参考文献")
    declaration = insert_before(reference_heading, "声明", "Heading 1")
    data_p = insert_after(
        declaration,
        "数据可用性  本文使用的 DBP15K、OpenEA 和 EventEA 均为公开基准，来源见相应引用。",
        "Normal",
    )
    code_p = insert_after(
        data_p,
        "代码可用性  复现实验的代码、配置与结果汇总可由通讯作者在合理请求下提供。",
        "Normal",
    )
    insert_after(code_p, "利益冲突  作者声明不存在利益冲突。", "Normal")


def update_tables(document: Document, rows) -> None:
    datasets = [
        ("DBP15K zh-en", "dbp15k_zh_en"),
        ("DBP15K ja-en", "dbp15k_ja_en"),
        ("DBP15K fr-en", "dbp15k_fr_en"),
        ("OpenEA en-fr", "openea_en_fr"),
        ("EventEA en-en", "eventea_en_en"),
    ]
    main_table = document.tables[1]
    for row_index, (_, dataset) in enumerate(datasets, start=1):
        value = result(rows, "main", dataset)
        main_table.cell(row_index, 1).text = metric_text(value, "Hits@1")
        main_table.cell(row_index, 2).text = metric_text(value, "Hits@10")
        main_table.cell(row_index, 3).text = metric_text(value, "MRR")

    comparison = document.tables[2]
    dbp = result(rows, "main", "dbp15k_zh_en")
    openea = result(rows, "main", "openea_en_fr")
    for table_row, value in ((6, dbp), (13, openea)):
        comparison.cell(table_row, 2).text = f"{value['Hits@1_mean']:.4f}"
        comparison.cell(table_row, 3).text = f"{value['Hits@10_mean']:.4f}"
        comparison.cell(table_row, 4).text = f"{value['MRR_mean']:.4f}"

    ablation = document.tables[3]
    variant_rows = [
        ("完整模型", "main"),
        ("可学习实体初始化", "learned_structure_init"),
        ("固定 8 邻居 + softmax", "fixed_k8_softmax"),
        ("去除关系类型", "no_relation_types"),
        ("去除层选择器", "no_layer_selector"),
        ("去除 token 视图", "no_token_view"),
        ("去除 phrase 视图", "no_phrase_view"),
        ("去除 global 视图", "no_global_view"),
        ("去除结构监督", "no_structure_supervision"),
        ("独立关系词表", "disjoint_relation_vocab"),
        ("双向边（共享关系类型）", "bidirectional_same_relation"),
    ]
    while len(ablation.rows) < len(variant_rows) + 1:
        append_row_like(ablation)
    while len(ablation.rows) > len(variant_rows) + 1:
        ablation._tbl.remove(ablation.rows[-1]._tr)
    for row_index, (label, variant) in enumerate(variant_rows, start=1):
        set_cell(ablation.cell(row_index, 0), label)
        set_cell(ablation.cell(row_index, 1), pair_text(result(rows, variant, "dbp15k_zh_en")))
        set_cell(ablation.cell(row_index, 2), pair_text(result(rows, variant, "openea_en_fr")))

    branches = document.tables[4]
    insert_column(branches, 2, 1.3)
    branch_rows = [
        ("完整模型", "联合表示", "联合表示", "main"),
        ("仅结构分支", "结构表示", "结构表示", "structure_only"),
        ("仅语义分支", "语义表示", "语义表示", "semantic_only"),
        ("无可训练平均融合", "结构与语义等权平均", "同一等权平均", "mean_fusion"),
    ]
    headers = ("变体", "训练表示", "主检索表示", "DBP H@1 / MRR", "OpenEA H@1 / MRR")
    for index, header in enumerate(headers):
        set_cell(branches.cell(0, index), header)
    for row_index, (label, train_repr, retrieval_repr, variant) in enumerate(branch_rows, start=1):
        values = (
            label,
            train_repr,
            retrieval_repr,
            pair_text(result(rows, variant, "dbp15k_zh_en")),
            pair_text(result(rows, variant, "openea_en_fr")),
        )
        for col_index, value in enumerate(values):
            set_cell(branches.cell(row_index, col_index), value)
    set_table_widths(branches, [1250, 1500, 1400, 1975, 1975])

    controls = document.tables[5]
    append_row_like(controls)
    set_cell(controls.cell(0, 2), "可训练融合参数")
    control_rows = [
        (
            "完整模型",
            "完整一跳出邻域；语义查询；直接联合检索",
            "197,377",
            pair_text(result(rows, "main", "dbp15k_zh_en")),
            pair_text(result(rows, "main", "openea_en_fr")),
        ),
        (
            "结构查询邻居",
            "相同邻域与门控；结构查询；直接联合检索",
            "197,377",
            pair_text(result(rows, "structural_neighbor_query", "dbp15k_zh_en")),
            pair_text(result(rows, "structural_neighbor_query", "openea_en_fr")),
        ),
        (
            "参数匹配晚期融合",
            "不读取邻居；独立分支后形成并检索 z_late",
            "197,377",
            pair_text(result(rows, "late_concat_matched", "dbp15k_zh_en")),
            pair_text(result(rows, "late_concat_matched", "openea_en_fr")),
        ),
        (
            "旧分支加权检索",
            "完整模型训练；测试改用验证集分支加权",
            "197,377",
            "0.7295 ± 0.0037 / 0.7706 ± 0.0034",
            "0.6324 ± 0.0065 / 0.6978 ± 0.0043",
        ),
    ]
    for row_index, values in enumerate(control_rows, start=1):
        for col_index, value in enumerate(values):
            set_cell(controls.cell(row_index, col_index), value)
    set_table_widths(controls, [1200, 2550, 1150, 1800, 1800])


def remove_highlights(document: Document) -> None:
    for highlight in list(document._element.xpath(".//w:highlight")):
        highlight.getparent().remove(highlight)


def replace_figure(document: Document) -> None:
    if not FIGURE.exists():
        raise FileNotFoundError(f"Generate the revised architecture figure first: {FIGURE}")
    image_parts = [
        relationship.target_part
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/image")
    ]
    if len(image_parts) != 1:
        raise RuntimeError(f"Expected one document image, found {len(image_parts)}")
    image_parts[0]._blob = FIGURE.read_bytes()


def citation_numbers(content: str) -> List[int]:
    result: List[int] = []
    for part in re.split(r"\s*,\s*", content):
        range_match = re.fullmatch(r"(\d+)\s*[–-]\s*(\d+)", part)
        if range_match:
            start, end = map(int, range_match.groups())
            result.extend(range(start, end + 1))
        else:
            result.append(int(part))
    return result


def compress_numbers(numbers: List[int]) -> str:
    if not numbers:
        return ""
    parts: List[str] = []
    start = previous = numbers[0]
    for value in numbers[1:] + [None]:
        if value is not None and value == previous + 1:
            previous = value
            continue
        parts.append(str(start) if start == previous else f"{start}–{previous}")
        if value is not None:
            start = previous = value
    return ",".join(parts)


def replace_text_range(nodes, start: int, end: int, replacement: str) -> None:
    offsets = []
    cursor = 0
    for node in nodes:
        text = node.text or ""
        offsets.append((cursor, cursor + len(text)))
        cursor += len(text)
    first = next(i for i, (_, stop) in enumerate(offsets) if stop > start)
    last = next(i for i, (begin, stop) in enumerate(offsets) if begin < end <= stop)
    first_begin = offsets[first][0]
    last_begin = offsets[last][0]
    prefix = (nodes[first].text or "")[: start - first_begin]
    suffix = (nodes[last].text or "")[end - last_begin :]
    nodes[first].text = prefix + replacement + (suffix if first == last else "")
    if first != last:
        for index in range(first + 1, last):
            nodes[index].text = ""
        nodes[last].text = suffix


def renumber_paragraph_citations(paragraph, mapping: Dict[int, int]) -> None:
    if paragraph.style is not None and paragraph.style.name == "Reference":
        return
    nodes = paragraph._p.xpath(".//w:t|.//m:t")
    full_text = "".join(node.text or "" for node in nodes)
    for match in reversed(list(CITATION_RE.finditer(full_text))):
        old = citation_numbers(match.group(1))
        mapped = [mapping[number] for number in old]
        replace_text_range(nodes, match.start(), match.end(), f"[{compress_numbers(mapped)}]")


def reorder_references(document: Document) -> None:
    # References 21 and 22 were not used by the paper's research question.
    references = [p for p in document.paragraphs if p.style.name == "Reference"]
    by_old: Dict[int, object] = {}
    for paragraph in references:
        match = re.match(r"\[(\d+)\]", combined_text(paragraph))
        if match:
            by_old[int(match.group(1))] = paragraph
    for old_number in (21, 22):
        remove_paragraph(by_old.pop(old_number))

    body_paragraphs = [p for p in document.paragraphs if p.style.name != "Reference"]
    body_paragraphs.extend(
        p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
    )
    order: List[int] = []
    for paragraph in body_paragraphs:
        for match in CITATION_RE.finditer(combined_text(paragraph)):
            for number in citation_numbers(match.group(1)):
                if number not in order:
                    order.append(number)
    expected_old = sorted(by_old)
    if sorted(order) != expected_old:
        raise RuntimeError(f"Cited references {sorted(order)} do not match retained references {expected_old}")
    mapping = {old: new for new, old in enumerate(order, start=1)}
    for paragraph in body_paragraphs:
        renumber_paragraph_citations(paragraph, mapping)

    references = [p for p in document.paragraphs if p.style.name == "Reference"]
    xml_by_old = {int(re.match(r"\[(\d+)\]", combined_text(p)).group(1)): deepcopy(p._p) for p in references}
    anchor = references[0]._p
    for old_number in order:
        xml = xml_by_old[old_number]
        nodes = xml.xpath(".//w:t")
        text = "".join(node.text or "" for node in nodes)
        match = re.match(r"\[(\d+)\]", text)
        replace_text_range(nodes, match.start(), match.end(), f"[{mapping[old_number]}]")
        anchor.addprevious(xml)
    for paragraph in references:
        remove_paragraph(paragraph)


def add_corresponding_email(document: Document) -> None:
    author = find_paragraph(document, "作者")
    insert_after(author, "通讯作者电子邮箱：auroral020417@gmail.com", author.style.name)


def set_cjk_font(document: Document, family: str = "Arial Unicode MS") -> None:
    """Use a CJK font that renders consistently in both Word and LibreOffice."""
    for style in document.styles:
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), family)

    for run in document.element.body.xpath(".//w:r"):
        text = "".join(run.xpath(".//w:t/text()"))
        if not any("\u3400" <= char <= "\u9fff" for char in text):
            continue
        rpr = run.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attribute}"), family)


def validate(document: Document) -> None:
    all_text = "\n".join(combined_text(p) for p in document.paragraphs)
    all_text += "\n" + "\n".join(
        combined_text(p)
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    )
    required = [
        "主检索直接使用训练所约束的联合表示",
        "独立关系词表",
        "双向边（共享关系类型）",
        "去除结构监督",
        "训练表示",
        "主检索表示",
        "利益冲突",
    ]
    missing = [text for text in required if text not in all_text]
    if missing:
        raise RuntimeError(f"Missing required revision text: {missing}")
    forbidden = [
        "最终测试不直接使用训练表示",
        "验证网格为 α",
        "难负候选（hard negative candidate）",
        "参数量 DBP",
    ]
    present = [text for text in forbidden if text in all_text]
    if present:
        raise RuntimeError(f"Outdated text remains: {present}")
    references = [p for p in document.paragraphs if p.style.name == "Reference"]
    numbers = [int(re.match(r"\[(\d+)\]", combined_text(p)).group(1)) for p in references]
    if numbers != list(range(1, len(references) + 1)):
        raise RuntimeError(f"Reference numbering is not sequential: {numbers}")
    if len(document.tables[3].rows) != 12:
        raise RuntimeError("Ablation table does not contain all controls")
    if len(document.tables[4].columns) != 5:
        raise RuntimeError("Branch table did not split training and retrieval representations")


def main() -> None:
    rows = read_summary()
    document = Document(SOURCE)
    add_corresponding_email(document)
    rewrite_core_text(document, rows)
    update_tables(document, rows)
    remove_highlights(document)
    replace_figure(document)
    reorder_references(document)
    set_cjk_font(document)
    validate(document)
    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    validate(reopened)
    print(f"WROTE {OUTPUT}")


if __name__ == "__main__":
    main()
