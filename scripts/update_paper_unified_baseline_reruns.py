from __future__ import annotations

import html.entities
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
MATH_DEPS = ROOT / "tmp" / "docx_math_deps"
sys.path.insert(0, str(MATH_DEPS))

import mathml2omml  # noqa: E402
from latex2mathml.converter import convert as latex_to_mathml  # noqa: E402

SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "Springer引用与逻辑修订版_20260826.docx"
)
OUTPUT = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "统一协议经典模型重跑比较_20260827.docx"
)


def combined_text(paragraph) -> str:
    return "".join(
        paragraph._p.xpath(".//w:t/text()|.//m:t/text()")
    ).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph beginning {prefix!r}, found {len(matches)}"
        )
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    for child in list(paragraph._p):
        if child.tag not in {qn("w:pPr")}:
            paragraph._p.remove(child)
    paragraph.add_run(text)


def latex_to_omml(expression: str):
    mathml = latex_to_mathml(expression.strip(), display="inline")
    omml = mathml2omml.convert(mathml, html.entities.name2codepoint)
    omml = re.sub(
        r"(<m:groupChr><m:groupChrPr>.*?)</m:groupChr>(<m:e>)",
        r"\1</m:groupChrPr>\2",
        omml,
    )
    omml = omml.replace("<m:oMath>", f"<m:oMath {nsdecls('m')}>", 1)
    return parse_xml(omml)


def set_paragraph_with_inline_math(
    paragraph, before: str, expression: str, after: str
) -> None:
    for child in list(paragraph._p):
        if child.tag not in {qn("w:pPr")}:
            paragraph._p.remove(child)
    paragraph.add_run(before)
    paragraph._p.append(latex_to_omml(expression))
    paragraph.add_run(after)


def insert_before_anchor(paragraph, anchor: str, insertion: str) -> None:
    matches = [
        node
        for node in paragraph._p.xpath(".//w:t")
        if node.text and anchor in node.text
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one text node containing {anchor!r}")
    matches[0].text = matches[0].text.replace(anchor, insertion + anchor, 1)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for run in paragraph.runs:
        run.bold = bold
        run.font.size = Pt(9)
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def clear_highlights(document: Document) -> None:
    for highlight in document.element.xpath(".//w:highlight"):
        highlight.getparent().remove(highlight)


def update_comparison_table(document: Document) -> None:
    table = document.tables[2]
    if len(table.rows) != 14 or len(table.columns) != 5:
        raise RuntimeError(
            f"Unexpected comparison table shape: {len(table.rows)}x{len(table.columns)}"
        )

    table._tbl.remove(table.rows[-1]._tr)
    rows = [
        ("MTransE[2]", "0.1537", "0.5474", "0.2814"),
        ("JAPE[3]*", "0.1782", "0.5680", "0.3072"),
        ("BootEA[4]", "0.5439", "0.8041", "0.6324"),
        ("RDGCN[5]†", "0.0876", "0.3831", "0.1803"),
        ("RREA-basic[6]", "0.6512", "0.9095", "0.7451"),
        ("本文模型", "0.7200", "0.8473", "0.7658"),
        ("MTransE[2]", "0.1923", "0.5448", "0.3089"),
        ("JAPE[3]*", "0.2170", "0.5752", "0.3350"),
        ("BootEA[4]", "0.6589", "0.9090", "0.7478"),
        ("RDGCN-random[5]†", "0.2039", "0.6851", "0.3625"),
        ("RREA-basic[6]", "0.7584", "0.9705", "0.8369"),
        ("本文模型", "0.6250", "0.8204", "0.6931"),
    ]

    headers = ("数据集", "方法", "Hits@1", "Hits@10", "MRR")
    for column, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[column], value, bold=True)

    set_cell_text(table.rows[1].cells[0], "DBP15K ZH–EN")
    set_cell_text(table.rows[7].cells[0], "OpenEA EN–FR-V2")
    best_rows = {
        5: {3},
        6: {2, 4},
        11: {2, 3, 4},
    }
    for offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            set_cell_text(
                table.rows[offset].cells[column],
                value,
                bold=column in best_rows.get(offset, set()),
            )

    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))


def update_prose(document: Document) -> None:
    set_paragraph_text(
        find_paragraph(document, "5.2 与经典方法的定位性比较"),
        "5.2 统一协议下与经典方法的重跑比较",
    )
    set_paragraph_with_inline_math(
        find_paragraph(document, "表 3 中 DBP15K"),
        "为获得直接可比的结果，表 3 不再混用不同论文的报告值，而是在 DBP15K ZH–EN "
        "和 OpenEA EN–FR-V2 上重新运行 MTransE、JAPE、BootEA、RDGCN 和 RREA-basic。"
        "所有方法均使用随机种子 42、相同的训练/验证/测试划分以及完整的目标图实体候选集，"
        "并依据验证集 MRR 在每个数据集上独立选择 ",
        r"k_{\mathrm{CSLS}}\in\{3,5,7,10,15,20\}",
        "。"
        "本文模型在此表中同样报告随机种子 42 的单次结果；三随机种子统计仍见表 2。",
    )
    table_caption = find_paragraph(document, "表 3 与经典实体对齐方法的定位性比较")
    set_paragraph_text(
        table_caption,
        "表 3 统一协议下经典实体对齐方法的重跑结果（seed = 42）",
    )
    table_caption.paragraph_format.page_break_before = True
    table_caption.paragraph_format.keep_with_next = True
    blank_after_table = document.paragraphs[139]
    if combined_text(blank_after_table):
        raise RuntimeError("Expected the paragraph after Table 3 to be blank")
    set_paragraph_text(
        blank_after_table,
        "实现边界如下。MTransE、JAPE、BootEA 和 RDGCN 使用 OpenEA 官方源码，"
        "RREA-basic 使用原作者源码且不启用半监督扩充。受计算预算限制，JAPE* 的属性子阶段训练 5 个 epoch，"
        "关系子阶段仍采用官方早停；因此该结果是计算适配版本。RDGCN† 在 DBP15K 上的名称词向量覆盖为 "
        "626/34,460；OpenEA-V2 的实体标识经过匿名化，无法构造原始名称向量，故该数据集使用与实体身份无关的 "
        "Glorot 随机初始化，并记为 RDGCN-random。上述数值表示统一本地协议下的源码重跑结果，"
        "不等同于原论文在其各自设置下报告的结果。",
    )
    set_paragraph_text(
        find_paragraph(document, "在 DBP15K ZH–EN 上，本文模型的 Hits@1"),
        "在 DBP15K ZH–EN 上，本文模型的 Hits@1、Hits@10 和 MRR 分别为 0.7200、0.8473 和 0.7658。"
        "相较表现最强的重跑基线 RREA-basic，其 Hits@1 和 MRR 分别高 6.88 和 2.07 个百分点，"
        "但 Hits@10 低 6.22 个百分点。OpenEA EN–FR-V2 上，本文模型的三项指标为 0.6250、0.8204 和 0.6931，"
        "低于 RREA-basic 的 0.7584、0.9705 和 0.8369，也低于 BootEA 的 0.6589、0.9090 和 0.7478。"
        "因此，本轮直接比较支持本文模型在 DBP15K 顶位检索上的优势，但不支持跨数据集的一致领先结论。",
    )

    limitations = find_paragraph(document, "内部有效性方面")
    addition = (
        "表 3 的源码重跑只覆盖随机种子 42，不能据此估计基线方差；JAPE 和 RDGCN 又包含已明确标注的"
        "计算或输入适配，因此该表只支持本地统一协议下的实现比较。"
    )
    anchor = "外部有效性方面，"
    insert_before_anchor(limitations, anchor, addition)


def main() -> None:
    document = Document(SOURCE)
    update_comparison_table(document)
    update_prose(document)
    clear_highlights(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
