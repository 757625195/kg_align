# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document

from build_interim_new_main_paper import (
    clear_text_highlights,
    formula_text,
    replace_starts,
    set_cell,
)


EXPECTED_VARIANTS = {
    "main",
    "learned_structure_init",
    "no_structure_supervision",
    "fixed_k8_softmax",
    "no_relation_types",
    "no_layer_selector",
    "no_token_view",
    "no_phrase_view",
    "no_global_view",
    "structure_only",
    "semantic_only",
    "mean_fusion",
    "late_concat_matched",
    "structural_neighbor_query",
}

COMPONENT_ROWS = [
    ("完整模型", "main"),
    ("可学习实体初始化", "learned_structure_init"),
    ("去除结构分支监督", "no_structure_supervision"),
    ("固定 8 邻居 + softmax", "fixed_k8_softmax"),
    ("去除关系类型", "no_relation_types"),
    ("去除层选择器", "no_layer_selector"),
    ("去除 token 视图", "no_token_view"),
    ("去除 phrase 视图", "no_phrase_view"),
    ("去除 global 视图", "no_global_view"),
]


def load_results(path: Path) -> tuple[dict, dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    completed = set(payload["completed_variants"])
    missing = sorted(EXPECTED_VARIANTS - completed)
    if missing:
        raise RuntimeError(f"Ablations are incomplete: {', '.join(missing)}")
    summary = {
        (row["variant"], row["dataset"]): row
        for row in payload["summary"]
    }
    deltas = {
        (row["variant"], row["dataset"]): row
        for row in payload["paired_deltas_vs_main"]
    }
    return summary, deltas


def metric_cell(summary: dict, variant: str, dataset: str) -> str:
    row = summary[(variant, dataset)]
    return (
        f'{row["Hits@1_mean"]:.4f} ± {row["Hits@1_std"]:.4f} / '
        f'{row["MRR_mean"]:.4f} ± {row["MRR_std"]:.4f}'
    )


def short_metric(summary: dict, variant: str, dataset: str, metric: str) -> str:
    row = summary[(variant, dataset)]
    return f'{row[f"{metric}_mean"]:.4f} ± {row[f"{metric}_std"]:.4f}'


def delta_pp(deltas: dict, variant: str, dataset: str, metric: str = "Hits@1") -> float:
    return 100.0 * deltas[(variant, dataset)][f"delta_{metric}_mean"]


def ensure_rows(table, count: int) -> None:
    while len(table.rows) < count:
        table.add_row()


def update_tables(document: Document, summary: dict) -> None:
    component = document.tables[4]
    ensure_rows(component, len(COMPONENT_ROWS) + 1)
    for row, (label, variant) in zip(component.rows[1:], COMPONENT_ROWS):
        set_cell(row.cells[0], label, centered=False)
        set_cell(row.cells[1], metric_cell(summary, variant, "dbp15k_zh_en"))
        set_cell(row.cells[2], metric_cell(summary, variant, "openea_en_fr"))

    branches = document.tables[5]
    branch_rows = [
        ("完整模型", "结构邻居上下文 + 语义 + 门控", "main"),
        ("仅结构分支", "结构表示", "structure_only"),
        ("仅语义分支", "语义表示", "semantic_only"),
        ("无可训练融合", "结构与语义等权平均", "mean_fusion"),
    ]
    for row, (label, representation, variant) in zip(branches.rows[1:], branch_rows):
        set_cell(row.cells[0], label, centered=False)
        set_cell(row.cells[1], representation, centered=False)
        set_cell(row.cells[2], metric_cell(summary, variant, "dbp15k_zh_en"))
        set_cell(row.cells[3], metric_cell(summary, variant, "openea_en_fr"))

    fusion = document.tables[7]
    fusion_rows = [
        ("完整模型", "main"),
        ("结构查询邻居", "structural_neighbor_query"),
        ("参数匹配晚期融合", "late_concat_matched"),
    ]
    for row, (label, variant) in zip(fusion.rows[1:], fusion_rows):
        set_cell(row.cells[0], label, centered=False)
        set_cell(row.cells[1], short_metric(summary, variant, "dbp15k_zh_en", "Hits@1"))
        set_cell(row.cells[2], short_metric(summary, variant, "dbp15k_zh_en", "MRR"))
        set_cell(row.cells[3], short_metric(summary, variant, "openea_en_fr", "Hits@1"))
        set_cell(row.cells[4], short_metric(summary, variant, "openea_en_fr", "MRR"))


def rewrite_analysis(document: Document, summary: dict, deltas: dict) -> None:
    learned_dbp = delta_pp(deltas, "learned_structure_init", "dbp15k_zh_en")
    learned_open = delta_pp(deltas, "learned_structure_init", "openea_en_fr")
    phrase_dbp = delta_pp(deltas, "no_phrase_view", "dbp15k_zh_en")
    phrase_open = delta_pp(deltas, "no_phrase_view", "openea_en_fr")
    global_dbp = delta_pp(deltas, "no_global_view", "dbp15k_zh_en")
    global_open = delta_pp(deltas, "no_global_view", "openea_en_fr")
    late_dbp = delta_pp(deltas, "late_concat_matched", "dbp15k_zh_en")
    late_open = delta_pp(deltas, "late_concat_matched", "openea_en_fr")
    query_dbp = delta_pp(deltas, "structural_neighbor_query", "dbp15k_zh_en")
    query_open = delta_pp(deltas, "structural_neighbor_query", "openea_en_fr")

    replace_starts(
        document,
        "跨语言知识图谱实体对齐需要联合利用",
        "跨语言知识图谱实体对齐需要联合利用关系结构与实体语义，但独立实体初始化、固定邻居截断和稠密注意力会削弱结构分支的可迁移性。本文提出一种拓扑初始化的关系感知结构上下文与多尺度语义融合方法。结构分支以八维局部拓扑统计和小比例可学习实体残差初始化，通过关系感知消息传播与节点级层选择融合不同结构深度；语义分支从共享输入构造 token、phrase 和 global 三种视图。融合阶段保留全部一跳出邻居，采用批次内动态补齐和温度为 0.25 的 1.5-entmax 稀疏选择。五个数据集的三随机种子测试平均 Hits@1、Hits@10 和 MRR 分别为 0.7620、0.8858 和 0.8073。统一协议消融表明，共享拓扑初始化和关系类型在两个代表性数据集上均产生稳定增益；phrase 与 global 视图呈现数据集依赖性；参数匹配晚期融合的 Hits@1 分别低 2.18 和 9.44 个百分点，而将语义查询替换为结构查询仅低 0.44 和 0.22 个百分点。结果支持在最终融合前保留结构邻居上下文，但不支持把结构监督衰减或特定查询信号描述为普遍有效的独立贡献。",
    )
    replace_starts(
        document,
        "因此，结构分支在第一个 epoch",
        "因此，结构分支在第一个 epoch 的权重为 0.1，并在最后一个 epoch 衰减为 0。该设计的动机是在训练早期约束结构空间，同时避免后期辅助目标主导联合表示；其是否改善最终检索由第 5.4 节的去除结构监督实验检验，而不作为未经验证的性能声明。",
    )

    replace_starts(
        document,
        "表 5 将报告",
        "表 5 在 DBP15K ZH–EN 与 OpenEA EN–FR 上比较结构初始化、结构监督、邻居选择、关系类型、层选择和三种语义视图。每个变体均重新训练三个随机种子，并独立依据验证集选择融合权重、传播深度和 CSLS 邻域。",
    )
    replace_starts(
        document,
        "本轮统一协议消融正在运行",
        f"将共享拓扑初始化替换为独立可学习实体向量后，Hits@1 在 DBP15K 和 OpenEA 上分别变化 {learned_dbp:+.2f} 与 {learned_open:+.2f} 个百分点；去除关系类型分别下降 3.03 与 3.10 个百分点。这两项在两个数据集上方向一致，构成最稳定的结构证据。固定 8 个邻居并恢复 softmax 分别下降 0.62 与 1.02 个百分点，说明全部一跳邻居与稀疏选择具有较小但一致的联合收益。去除层选择器在 DBP15K 上仅下降 0.10 个百分点，在 OpenEA 上下降 1.71 个百分点。语义视图方面，去除 token 几乎不改变结果；去除 phrase 分别变化 {phrase_dbp:+.2f} 与 {phrase_open:+.2f} 个百分点；去除 global 分别变化 {global_dbp:+.2f} 与 {global_open:+.2f} 个百分点，表明多尺度视图的作用取决于数据集。去除结构监督在 DBP15K 上提高 0.62 个百分点、在 OpenEA 上降低 0.58 个百分点，因此该辅助项未形成跨数据集一致增益。",
    )
    replace_starts(
        document,
        "表 6 将比较",
        "表 6 比较完整模型、结构单分支、语义单分支和无可训练平均融合，以区分两个信息源和可训练融合模块的作用。各变体使用独立验证配置，结构单分支固定 $\\alpha=0$，语义单分支固定 $\\alpha=1$。",
    )
    replace_starts(
        document,
        "上述分支与融合实验正在运行",
        "结构单分支的 Hits@1 在 DBP15K 和 OpenEA 上分别为 0.1684 和 0.1014，语义单分支分别为 0.6424 和 0.5379。这说明语义是当前输入条件下的主要独立检索信号，但结构分支并非没有边际作用：完整模型相对语义单分支分别提高 7.74 和 9.46 个百分点。无可训练平均融合达到 0.7019 和 0.5832，仍比完整模型低 1.78 和 4.93 个百分点，表明可训练的邻居上下文与门控优于结构和语义的固定等权组合。",
    )
    replace_starts(
        document,
        "参数匹配融合控制正在按统一协议运行",
        f"相对完整模型，结构查询邻居的 Hits@1 在 DBP15K 和 OpenEA 上分别变化 {query_dbp:+.2f} 与 {query_open:+.2f} 个百分点，差值较小，说明本轮结果不能把语义查询本身视为主要增益来源。参数匹配晚期融合分别变化 {late_dbp:+.2f} 与 {late_open:+.2f} 个百分点，支持在最终融合前保留并利用结构邻居上下文。需要指出，晚期基线同时改变邻居可见性和交互位置，因此该差值反映二者的联合效应，尚不能在两者之间作完全独立的因果归因。",
    )
    replace_starts(
        document,
        "针对 RQ1",
        "针对 RQ1，五个数据集的三随机种子结果显示，新主配置的数据集等权平均 Hits@1、Hits@10 和 MRR 分别为 0.7620、0.8858 和 0.8073。针对 RQ2，共享拓扑初始化与关系类型在两个代表性数据集上均有稳定贡献；全部邻居稀疏选择和层选择产生较小或数据集相关的收益；phrase 对 OpenEA 更重要，global 在两个数据集上方向相反，token 未显示独立贡献；结构监督衰减也未形成一致增益。针对 RQ3，结构查询与语义查询的差异较小，而参数匹配晚期融合明显较弱，说明主要证据支持在最终融合前利用结构邻居上下文，而非特定查询信号。由于晚期基线同时改变可见信息和交互位置，二者仍不能完全分离。",
    )
    replace_starts(
        document,
        "本文实现了一种拓扑初始化",
        "本文实现了一种拓扑初始化的关系感知结构上下文与多尺度语义融合方法。五个数据集的三随机种子主实验达到平均 Hits@1 0.7620、Hits@10 0.8858 和 MRR 0.8073。两个代表性数据集上的统一协议消融表明，共享拓扑初始化、关系类型以及最终融合前可见的结构邻居上下文构成主要稳定证据；固定平均融合和参数匹配晚期融合均明显较弱。与此同时，结构监督、层选择和三种语义视图的作用并非全部跨数据集一致，其中 token 未显示独立增益，global 的影响方向相反。因此，本文的证据支持关系感知邻居上下文与语义表示的联合建模，但不把每个实现组件都描述为普遍有效的独立创新。",
    )


def verify(document: Document) -> None:
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    tables = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    if "运行中" in body or "运行中" in tables:
        raise RuntimeError("Unfinished experiment markers remain")
    tags = []
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Formula":
            continue
        text = formula_text(paragraph)
        if text.endswith(")") and "(" in text:
            suffix = text.rsplit("(", 1)[-1][:-1]
            if suffix.isdigit():
                tags.append(int(suffix))
    if tags != list(range(1, 28)):
        raise RuntimeError(f"Unexpected equation numbering: {tags}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--results", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    summary, deltas = load_results(args.results)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)
    update_tables(document, summary)
    rewrite_analysis(document, summary, deltas)
    clear_text_highlights(document)
    verify(document)
    document.save(args.output)
    print(f"WROTE {args.output}")


if __name__ == "__main__":
    main()
