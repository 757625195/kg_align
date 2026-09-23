from __future__ import annotations

import hashlib
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "方法学控制与直接联合检索修订版_20260829.docx"
)
OUTPUT = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "实验结论与逻辑一致性修订版_20260830.docx"
)
EXPECTED_SOURCE_SHA256 = "33d4740362d6c8ce9b6fe1ecbac3710d2e3d589f55c1c2497eaa836aead1f53d"

sys.path.insert(0, str(ROOT / "scripts"))
from mathify_all_document_symbols import set_markup  # noqa: E402


def combined_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//w:t/text()|.//m:t/text()")).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if combined_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def rewrite(document: Document, prefix: str, text: str) -> None:
    paragraph = find_paragraph(document, prefix)
    set_markup(paragraph, text)
    for run in paragraph.runs:
        run.bold = False


def insert_after_table(table, text: str) -> Paragraph:
    paragraph_xml = OxmlElement("w:p")
    table._tbl.addnext(paragraph_xml)
    paragraph = Paragraph(paragraph_xml, table._parent)
    paragraph.style = "Normal"
    set_markup(paragraph, text)
    return paragraph


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_markup(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def all_document_text(document: Document) -> str:
    parts = [combined_text(p) for p in document.paragraphs]
    parts.extend(
        combined_text(p)
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    )
    return "\n".join(parts)


def validate(document: Document) -> None:
    text = all_document_text(document)
    required = [
        "0.7077 ± 0.0023 / 0.7497 ± 0.0016",
        "0.6096 ± 0.0043 / 0.6752 ± 0.0019",
        "结构查询相对语义查询的 Hits@1 在 DBP15K 和 OpenEA 上分别低 2.08 和 2.48 个百分点",
        "不同方法的数据划分、文本输入与后处理并不完全一致",
        "当前实验尚未分别控制跨图关系编号共享和边方向",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise RuntimeError(f"Missing required revised text: {missing}")

    forbidden = [
        "低 2.8 和 2.48 个百分点",
        "仅低 0.08 和 0.48 个百分点",
        "旧分支加权行",
        "独立关系词表控制表明",
        "关系词表独立化和双向边控制",
        "双向边控制进一步限定",
        "综合 RQ1–RQ3",
        "本研究围绕以下三个问题展开。本文主要考察三个问题。",
    ]
    present = [item for item in forbidden if item in text]
    if present:
        raise RuntimeError(f"Outdated or contradictory text remains: {present}")

    if len(document.tables) != 6:
        raise RuntimeError(f"Expected six tables, found {len(document.tables)}")
    if len(document.tables[3].rows) != 10:
        raise RuntimeError("Table 4 structure changed unexpectedly")
    if len(document.tables[5].rows) != 4:
        raise RuntimeError("Table 6 must contain the header and three current controls")


def main() -> None:
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise RuntimeError(
            "The manually edited source changed after review; refusing to edit a stale version. "
            f"Expected {EXPECTED_SOURCE_SHA256}, found {source_hash}."
        )

    document = Document(SOURCE)

    rewrite(
        document,
        "跨语言知识图谱实体对齐旨在识别",
        "跨语言知识图谱实体对齐旨在识别不同语言知识图谱中指向同一现实对象的实体，是知识融合的基础任务。"
        "现有方法虽然分别利用图结构和文本语义，但关系类型、不同传播深度与碎片化文本往往被独立处理，结构邻居在联合表示中的作用也缺少受控验证。"
        "为此，本文提出关系感知结构上下文与多尺度语义融合方法。结构编码器通过共享关系投影进行关系感知消息传递，并以节点级层选择整合不同传播深度；"
        "语义编码器从同一实体文本中提取 token、phrase 和 global 三种粒度的表示。融合模块从完整一跳出邻域中利用 1.5-entmax 选择稀疏结构证据，"
        "并将邻居上下文与语义表示联合编码。模型采用双向 InfoNCE 训练，并使用 CSLS 进行跨图检索。五个数据集上的三随机种子实验中，Hits@1 为 0.6344—0.9239，"
        "样本标准差均不超过 0.0053。在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2 上，去除关系类型使 Hits@1 分别下降 3.62 和 2.78 个百分点，"
        "仅使用语义分支则分别下降 9.09 和 9.66 个百分点。在保持邻域集合与融合容量不变时，结构查询相对语义查询下降 2.08 和 2.48 个百分点；"
        "不读取邻居上下文的参数匹配晚期融合则下降 5.28 和 9.22 个百分点。结果支持在联合表示中保留关系感知结构邻域，但尚不能把晚期融合差距单独归因于融合时机。",
    )

    rewrite(
        document,
        "第一，结构相似不等于实体等价",
        "第一，结构相似不等于实体等价。例如，人物 A 通过“出生地”连接北京，人物 B 通过“工作地点”连接北京；若模型把不同关系退化为无类型连接，就会高估两人的相似性[5,7]。"
        "第二，语义证据往往是碎片化的。例如，单独出现的“Washington”无法区分人物、州和城市，需要结合关系或属性上下文。"
        "第三，可用于训练的种子对齐通常只覆盖部分实体[4,8]，模型需要从有限的已知实体对中学习可迁移的结构与语义线索。",
    )
    rewrite(
        document,
        "本研究围绕以下三个问题展开",
        "本文的实验围绕三个问题展开：模型在五个数据集上的检索结果及多次运行的波动如何；关系感知结构编码、多尺度语义编码和邻居聚合分别发挥什么作用；"
        "在邻域和参数量保持一致时，语义查询与结构查询有何差异，以及移除邻居上下文后晚期融合的性能如何变化。",
    )
    rewrite(
        document,
        "在五个实体对齐数据集上进行三次独立实验",
        "在五个实体对齐数据集上进行三次独立运行，并在两个代表性数据集上完成组件消融和控制实验，据此分析结构、语义和邻居机制的实际作用及结论边界。",
    )
    rewrite(
        document,
        "对于实体 i，本文将其名称",
        "对于实体 [[i]]，本文将其名称、相邻关系名称和属性文本分词后拼接为同一 token 序列，并将每个 token 映射为词向量。"
        "phrase 与 global 视图均使用这条共享序列，不额外读取实体描述、外部句子或其他语料。序列经统一维度映射、归一化和位置编码后得到基础表示 [[H_i]]，掩码 [[M_i]] 标记有效 token。",
    )
    rewrite(
        document,
        "为分别考察语义进入邻居计算的时机",
        "为比较邻居查询信号，并观察移除邻居上下文后的性能变化，本文设置两种控制。结构查询控制保留相同的一跳出邻域、动态补齐、1.5-entmax 和门控参数，"
        "但以实体自身结构表示计算邻居权重；语义仅在结构上下文形成后进入联合门。参数匹配晚期融合不读取邻居上下文，而在结构分支和语义分支独立编码后形成训练表示：",
    )
    rewrite(
        document,
        "模型使用 AdamW",
        "模型使用 AdamW[24] 优化，batch size 为 512，表示维度为 128，关系感知 GNN 固定为 3 层，dropout 为 0.1。语义输入由 300 维映射至 128 维，"
        "global 视图使用两层、四头 Transformer，phrase 视图采用宽度 3 和 5 的卷积。1.5-entmax 温度为 [[T=0.25]]，InfoNCE 温度为 [[\\tau=0.07]]。"
        "随机种子为 42、43 和 44；DBP15K 最多训练 36 个 epoch，OpenEA 与 EventEA 最多训练 50 个 epoch。每 5 个 epoch 依据验证集 MRR 保存检查点，"
        "连续 4 次验证未提升时提前停止。主检索不选择分支融合权重，每个运行从 [[k_{\\mathrm{CSLS}}\\in\\{3,5,7,10,15,20\\}]] 中按验证集 MRR 选择一个值。"
        "关系标识按规范化名称统一；DBP15K 另使用数据集提供的 sup_rel_ids 映射已知对应关系。",
    )

    insert_after_table(
        document.tables[2],
        "由于不同方法的数据划分、文本输入与后处理并不完全一致，表 3 仅用于说明本文结果在既有研究中的相对位置，不构成严格的同协议排名。"
        "在 DBP15K ZH–EN 上，本文模型高于表中的早期结构方法，但低于 RREA-text；在 OpenEA EN–FR-V2 上，本文结果低于 BootEA、KDCoE 和 RDGCN。"
        "这些差距说明当前方法尚不能据此声称达到最优性能。",
    )

    rewrite(
        document,
        "表 4 结构组件、关系预处理与语义视图控制",
        "表 4 结构组件、邻居聚合与语义视图消融（Hits@1/MRR，均值[[\\pm]]样本标准差，[[n=3]]）",
    )
    rewrite(
        document,
        "表 4 显示，移除关系类型",
        "表 4 显示，仅使用可学习实体向量而不采用拓扑初始化时，DBP15K 和 OpenEA 的 Hits@1 分别下降 3.46 和 8.03 个百分点；"
        "移除关系类型则分别下降 3.62 和 2.78 个百分点。去除层选择器以及同时改用固定 8 邻居和 softmax 均造成方向一致但幅度较小的下降。"
        "在语义视图中，phrase 在两个数据集上均产生正向贡献，token 的边际影响不超过 0.27 个百分点；去除 global 后，DBP15K 提高 0.49 个百分点而 OpenEA 降低 1.15 个百分点，"
        "说明其作用具有数据集依赖性。去除结构监督仅分别下降 0.07 和 0.17 个百分点，因此该辅助目标不是主要性能来源。",
    )
    rewrite(
        document,
        "表 6 比较三类控制",
        "表 6 比较完整模型、结构查询控制和参数匹配晚期融合。完整模型与结构查询控制具有相同的一跳出邻域、融合参数和直接联合检索，二者仅改变邻居权重的查询信号；"
        "参数匹配晚期融合保持相同的可训练融合参数量，但不读取邻居上下文。",
    )
    rewrite(
        document,
        "在邻居集合、训练容量和检索表示不变时",
        "在邻居集合、训练容量和检索表示不变时，结构查询相对语义查询的 Hits@1 在 DBP15K 和 OpenEA 上分别低 2.08 和 2.48 个百分点。"
        "两个数据集上的变化方向一致，说明语义在邻居聚合前参与查询能够改善当前设置下的检索结果。参数匹配晚期融合分别低 5.28 和 9.22 个百分点，说明移除邻居上下文并改变训练交互路径会明显降低性能；"
        "由于这两个因素同时变化，该差值不能被单独解释为融合时机的作用。",
    )
    rewrite(
        document,
        "综合 RQ1–RQ3",
        "以上实验得到三点主要发现。首先，五个数据集上的随机种子波动较小，但不同数据集的绝对性能差异明显。其次，语义分支单独检索明显强于结构分支，"
        "而完整模型仍优于语义单分支和无可训练平均融合；结合拓扑初始化与关系类型消融可见，结构证据的作用主要体现在可训练的联合表示中。"
        "最后，在保持邻域和融合容量一致时，语义查询在两个代表性数据集上均优于结构查询；不读取邻居的晚期融合进一步下降，说明查询信号和邻居上下文都会影响联合表示。"
        "不同语义视图的贡献并不完全一致，其中 global 视图表现出明显的数据集依赖性。",
    )
    rewrite(
        document,
        "内部有效性方面",
        "内部有效性方面，主实验和消融使用相同划分、验证规则和三个随机种子，但 [[n=3]] 仍不足以支持高功效显著性检验。"
        "固定 8 邻居加 softmax 的对照同时改变邻居预算和归一化函数，参数匹配晚期融合则同时改变邻居可见性和训练交互路径，因而两者均不能识别单一因素的因果效应。"
        "此外，当前实验尚未分别控制跨图关系编号共享和边方向。外部有效性方面，组件控制只覆盖两个代表性数据集，OpenEA 主结果只使用一个官方划分，"
        "且尚未验证开放世界和无对应实体场景。",
    )
    rewrite(
        document,
        "本文围绕关系感知结构上下文",
        "本文提出关系感知结构上下文与多尺度语义联合模型，并在五个实体对齐数据集上进行多随机种子评估。实验表明，拓扑初始化、关系类型和 phrase 视图能够为联合表示提供有效信息；"
        "虽然结构分支单独检索较弱，完整模型仍优于语义单分支和无可训练平均融合。相同邻域下，语义查询在两个代表性数据集上均优于结构查询；不读取邻居的参数匹配晚期融合进一步下降，"
        "但该对照同时改变了邻居可见性和训练交互路径。因此，现有结果支持在结构—语义联合表示中保留关系感知邻居上下文，并支持语义引导的邻居赋权在当前两个数据集上的作用；"
        "其跨数据集普适性以及结构辅助监督的独立价值仍需进一步验证。",
    )

    set_cell_text(document.tables[3].rows[2].cells[0], "仅可学习实体初始化（无拓扑特征）")
    set_cell_text(
        document.tables[5].rows[2].cells[3],
        "0.7077 ± 0.0023 / 0.7497 ± 0.0016",
    )
    set_cell_text(
        document.tables[5].rows[2].cells[4],
        "0.6096 ± 0.0043 / 0.6752 ± 0.0019",
    )

    validate(document)
    document.save(OUTPUT)
    reopened = Document(OUTPUT)
    validate(reopened)
    print(f"WROTE {OUTPUT}")


if __name__ == "__main__":
    main()
