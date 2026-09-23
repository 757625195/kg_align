from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
JOBS = (
    (
        ROOT
        / "outputs"
        / "Relation-Aware Neighbor Context and Semantics-Guided Selection for Knowledge Graph Entity Alignment_格式统一版.docx",
        ROOT
        / "outputs"
        / "Relation-Aware Neighbor Context and Semantics-Guided Selection for Knowledge Graph Entity Alignment_表5表头统一版.docx",
    ),
    (
        ROOT / "outputs" / "1_springer_proceedings_中文版_格式统一版.docx",
        ROOT / "outputs" / "1_springer_proceedings_中文版_表5表头统一版.docx",
    ),
)


def fix_table5_header(source: Path, output: Path) -> None:
    document = Document(source)
    if len(document.tables) < 5:
        raise RuntimeError(f"Expected at least five tables in {source}")

    header_cells = document.tables[4].rows[0].cells
    reference_shading = header_cells[0]._tc.get_or_add_tcPr().find(qn("w:shd"))
    if reference_shading is None:
        reference_shading = OxmlElement("w:shd")
        reference_shading.set(qn("w:val"), "clear")
        reference_shading.set(qn("w:color"), "auto")
        reference_shading.set(qn("w:fill"), "E7E6E6")

    for cell in header_cells:
        cell_properties = cell._tc.get_or_add_tcPr()
        existing = cell_properties.find(qn("w:shd"))
        if existing is not None:
            cell_properties.remove(existing)
        cell_properties.append(deepcopy(reference_shading))

    document.save(output)
    print(output)


if __name__ == "__main__":
    for input_path, output_path in JOBS:
        fix_table5_header(input_path, output_path)
