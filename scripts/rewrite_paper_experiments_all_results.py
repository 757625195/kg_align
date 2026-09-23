from __future__ import annotations

import argparse
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DATASET_ROWS = [
    ("DBP15K zh-en", "38,960", "165,556", "4,050/450/10,500", "38,960×14×300", "2,133"),
    ("DBP15K ja-en", "39,594", "170,698", "4,050/450/10,500", "39,594×40×300", "1,870"),
    ("DBP15K fr-en", "39,654", "221,720", "4,050/450/10,500", "39,654×13×300", "2,036"),
    ("OpenEA en-fr", "30,000", "176,430", "3,000/1,500/10,500", "30,000×32×300", "258"),
    ("OpenEA en-de", "30,000", "177,499", "3,000/1,500/10,500", "30,000×32×300", "194"),
    ("OpenEA D-W", "30,000", "157,348", "3,000/1,500/10,500", "30,000×32×300", "288"),
    ("OpenEA D-Y", "30,000", "129,033", "3,000/1,500/10,500", "30,000×32×300", "93"),
    ("EventEA en-en", "58,719", "80,146", "5,770/2,885/20,196", "58,719×32×300", "175"),
]


REPLACEMENTS = {
    "4 实验设置": "4 实验设计",
    "4.1 数据集与划分": "4.1 数据集与协议分层",
    "DBP15K 使用固定 raw_split/0_3": (
        "实验采用三类互补协议。机制验证协议用于表 3、表 5、表 6 和图 2：DBP15K ZH-EN 使用固定 "
        "raw_split/0_3 的 4,500 个训练对并按固定轮数训练，OpenEA EN-FR 使用官方 721_5fold/1 "
        "训练、验证和测试划分。探索性诊断协议用于表 7 和表 8，目前仅运行随机种子 42。跨数据集扩展协议用于表 9："
        "DBP15K 三个语言对均从原训练对固定留出 10% 作为验证集，OpenEA 与 EventEA 使用官方验证集。所有协议均在完整"
        "目标实体集合上排名，不以测试真值缩小候选集。由于训练对数量、验证方式和简化配置不同，表 3 与表 9 不合并为"
        "同一统计总体。数据来源和基准规范见 JAPE[2]、OpenEA[16] 与 EventEA[20]。"
    ),
    "4.2 评价指标与比较协议": "4.2 评价指标与统计原则",
    "4.3 实现细节": "4.3 实现细节与后处理",
    "公共设置为": (
        "公共设置为：batch size 128，结构表示和联合表示维度均为 128，关系感知 GNN 为 3 层；语义编码器使用 2 层、"
        "4 个注意力头，dropout 为 0.1，邻居预算为 8。训练目标采用温度 0.07 的 InfoNCE、间隔 0.2 的排序约束，"
        "每个正对配置 1 个随机负例和 1 个批内难负例，不使用全局难负样本库。八数据集扩展实验采用消融结果支持的"
        "简化配置：关闭 global 语义视图；DBP15K 保留关系感知邻居采样，OpenEA 与 EventEA 使用均匀邻居采样。"
        "重排序器在基础模型的 top-20 候选上训练 12 轮，固定融合权重为 0.35；随后在 k 属于 {5, 10, 20}、blend 属于 "
        "{0.50, 0.75, 1.00} 的网格中按验证 MRR 选择 CSLS 参数。"
    ),
    "5 结果": "5 实验结果与分析",
    "5.1 主结果": "5.1 整体性能与训练稳定性",
    "DBP15K 的标准差很小": (
        "在机制验证协议下，基础模型在 DBP15K ZH-EN 和 OpenEA EN-FR 上的 Hits@1 分别为 "
        "0.6768±0.0005 和 0.4885±0.0112。DBP15K 的方差较小，而 OpenEA 的标准差高出一个数量级，"
        "说明后者对初始化、官方验证流程和属性 token 序列更敏感。仅加入重排序器时，两个数据集的原始余弦 Hits@1 "
        "分别增加 0.61 和 0.27 个百分点；使用 CSLS 后，基础模型的 Hits@1 分别增加 1.63 和 2.16 个百分点。"
        "因此，表 3 支持后处理能够改善检索，但重排序器本身的贡献有限。表 3 使用原始机制实验协议，其绝对数值不与"
        "采用留出验证集和简化配置的表 9 直接比较。"
    ),
    "5.2 与已有方法的比较": "5.2 与经典方法的定位性比较",
    "下表中的 MTransE": (
        "表 4 的 MTransE、JAPE、BootEA 和 RDGCN 数值取自 RDGCN 在 DBP15K ZH-EN 上的汇总[4]，"
        "RREA-text 取自 RREA 原论文的文本增强设置[5]。RREA 表中的第二指标被标为 MRR，但 RDGCN 的 0.846 与其 "
        "Hits@10 一致，且在 Hits@1 为 0.822 时 MRR 不可能达到 0.964；本文据数值语义将 0.964 记为 Hits@10，"
        "并将其 MRR 标为未可靠报告。由于文献方法与本文在名称翻译、文本初始化、数据划分和多次运行统计上不一致，"
        "该表仅用于定位性能区间，而非严格复现排名。"
    ),
    "当前模型明显优于早期": (
        "当前主模型的 Hits@1 高于 MTransE、JAPE 和 BootEA，但仍低于 RDGCN 和 RREA-text。即使 CSLS 后的 "
        "0.6931 接近 RDGCN 的 0.7075，也不能据此声称达到同等水平，因为二者的输入语义与评估协议不同。"
        "这一结果表明，本文方法已经形成可用基线，但其主要证据应来自协议一致的内部对照、机制消融和跨数据集复验，"
        "而不是非严格文献表中的名次。"
    ),
    "5.3 三随机种子核心消融": "5.3 核心组件贡献",
    "跨模态增强是当前最稳定": (
        "跨模态增强是最稳定的必要组件。关闭该模块后，DBP15K 与 OpenEA 的 Hits@1 分别下降 5.62 和 6.69 个百分点，"
        "MRR 分别下降 4.94 和 6.29 个百分点，说明在结构聚合前利用语义状态筛选上下文优于末端无参数平均。"
        "移除多尺度语义编码器后，DBP15K Hits@1 仅下降 0.69 个百分点，但 OpenEA 下降 17.47 个百分点；"
        "这一差异与 OpenEA 由关系名、属性名和值构成的较长语义序列一致，表明多粒度语义建模的收益具有明显的数据依赖性。"
    ),
    "多尺度语义编码器在 DBP15K": (
        "损失项消融呈现不同结论。仅保留 InfoNCE、去除 hard negative 或去除 ranking loss 时，DBP15K Hits@1 "
        "相对主模型最多下降 0.07 个百分点；OpenEA 的变化也不超过 0.22 个百分点，其中去除 ranking loss 后 Hits@1 "
        "还略高 0.01 个百分点，仅 InfoNCE 的 MRR 略高 0.05 个百分点。因此，现有结果不能把复合负采样与排序目标"
        "描述为已经验证的主要创新，其超参数、负样本数量和采样策略仍需独立优化。"
    ),
    "5.4 早期交互、晚期拼接与简单平均": "5.4 早期交互与晚期融合对照",
    "为直接回答": (
        "为检验交互位置，三种融合方式使用相同的数据划分、编码器、损失、训练轮数、随机种子和完整候选检索流程。"
        "早期交互在结构聚合前以语义状态选择邻域证据；晚期拼接让两个分支独立编码后使用两层 MLP 投影；简单平均不增加"
        "融合参数。三种设置均运行随机种子 42、43 和 44，表 6 报告均值与样本标准差。"
    ),
    "早期交互相对晚期拼接": (
        "早期交互相对晚期拼接在 DBP15K 和 OpenEA 上分别提高 Hits@1 9.86 和 27.32 个百分点，且三个随机种子"
        "方向一致；相对简单平均则分别提高 5.62 和 6.69 个百分点。晚期拼接还低于简单平均 4.24 和 20.63 个百分点，"
        "其 OpenEA Hits@1 标准差为 0.0368，明显高于早期交互的 0.0112。结果说明，在本文实现与参数规模下，"
        "把语义条件引入邻域聚合比末端重新混合两个已成形空间更有效、更稳定。"
    ),
    "该对照支持": (
        "该对照只能支持早期交互优于本文设置的两个晚期基线，不能把全部差值唯一归因于交互时机。早期模块使用 "
        "197,377 个融合参数并访问固定预算邻居，晚期 MLP 使用 49,408 个参数且不访问邻居，简单平均没有融合参数；"
        "信息可见范围、函数形式与容量同时发生变化。因而仍需参数匹配的晚期 MLP、无语义邻居注意力和随机邻居注意力"
        "基线，才能进一步分离交互位置的因果效应。"
    ),
    "5.5 补充消融分析": "5.5 细粒度结构与语义诊断",
    "更细的组件实验": (
        "表 7 的细粒度组件实验目前只完成随机种子 42，因此用于生成后续配置假设，不用于稳健排名或显著性结论。"
        "各变体应与所属结构组或语义组 base 比较，而不与不同训练批次的主表数值直接相减。"
    ),
    "关系感知消息和层融合": (
        "结构分支中，去除关系感知消息后，DBP15K 和 OpenEA Hits@1 分别下降 5.26 和 2.68 个百分点；"
        "去除层融合后分别下降 1.41 和 3.19 个百分点，二者方向跨数据集一致。关系感知邻居采样则不稳定："
        "DBP15K 下降 0.25 个百分点，OpenEA 反而提高 0.89 个百分点。该观察促使扩展实验在 DBP15K 保留关系采样，"
        "而在 OpenEA 与 EventEA 中采用均匀采样，但这一数据集特定选择仍需三随机种子复验。"
    ),
    "语义视图结果对最初假设": (
        "语义分支中，only phrase view 相对三视图 base 在 DBP15K 和 OpenEA 上分别提高 0.80 和 2.91 个百分点；"
        "关闭 global view 分别提高 0.79 和 3.44 个百分点。only token view 在 OpenEA 上下降 8.47 个百分点，"
        "说明局部短语信息不可缺少，但当前 Transformer global 分支未形成稳定互补。扩展实验据此关闭 global view；"
        "这一简化提高了跨数据集验证的可解释性，但仍需通过门控权重、视图相关性和过拟合分析确认原因。"
    ),
    "5.6 主动学习结果": "5.6 主动学习的样本效率",
    "主动学习使用": (
        "主动学习实验采用 20% 初始种子、5 轮查询和每轮 90 对的预算，仅运行随机种子 42。same_budget_no_query "
        "使用与初始种子加 450 次查询相同数量的随机已知标签，但不执行主动选择；full_train_no_al 使用完整训练集，"
        "只作为监督上限参考。三种低预算设置保持同一测试集，表 8 报告完整候选排名结果。"
    ),
    "主动查询相对于": (
        "相对 20% seed 基线，主动查询使 DBP15K 和 OpenEA Hits@1 分别提高 0.75 和 10.58 个百分点，说明查询流程"
        "能够利用新增标签。然而，在更严格的同预算比较中，主动查询在 OpenEA 上高 1.87 个百分点，在 DBP15K 上反而低 "
        "1.47 个百分点，因而不能得出普遍提高样本效率的结论。该模块还缺少三随机种子、相同训练轮数以及正负人工反馈"
        "统一计价的实验；完整训练集结果仅说明低预算方案与完整监督之间仍存在明显差距。"
    ),
    "5.7 多数据集扩展验证与验证集选择后处理": "5.7 跨数据集稳健性与验证集选择后处理",
    "为检验最新简化配置": (
        "表 9 检验由前述诊断得到的简化配置及后处理流程能否跨数据集复现。实验覆盖 DBP15K zh-en、ja-en、fr-en，"
        "OpenEA en-fr、en-de、D-W、D-Y，以及 EventEA en-en 共八个数据集；其中 D-W 和 D-Y 为同语种异源知识库，"
        "EventEA 用于检验事件实体和复杂属性异质性。每个数据集运行随机种子 42、43 和 44，并依次比较基础模型、"
        "top-20 MLP 重排序器以及在重排序分数上由验证 MRR 选择参数的 CSLS。DBP15K 从训练对中固定留出 10%，"
        "OpenEA 与 EventEA 使用官方验证集；测试集只用于最终报告。"
    ),
    "三个新增数据集的最终组合": (
        "最终组合在 24/24 次独立运行中均超过对应基础模型，八数据集平均 Hits@1 绝对增益为 1.75 个百分点；"
        "按数据集聚合的增益范围为 0.55 至 3.26 个百分点。OpenEA en-fr 与 en-de 的提升最大，而高基线的 DBP15K "
        "fr-en 提升最小。D-W 和 D-Y 排除了语言差异后仍分别提高 2.15 和 1.16 个百分点，EventEA 提高 1.17 个百分点，"
        "说明局部密度校正的收益不仅存在于翻译型跨语言任务，也出现在异源知识库和事件型异构场景。"
    ),
    "三个新增数据集上，Reranker": (
        "重排序器单独作用时，24 次运行中 16 次提高、7 次下降、1 次持平，平均 Hits@1 仅增加 0.08 个百分点；"
        "在八个数据集上的平均增益远低于最终组合。因此，Reranker 不应被单独表述为主要性能来源。更准确的结论是，"
        "它总体未显著破坏基础候选排序，而稳定增量主要出现在加入验证集选择的 CSLS 之后。表 9 的最终数值应作为"
        " Reranker 与 CSLS 的组合效果解释。"
    ),
    "按随机种子 42": (
        "验证集选择结果进一步表明，CSLS 参数不能跨数据集统一写死。OpenEA en-fr 与 DBP15K ja-en 的三个种子均偏好 "
        "k=5，OpenEA D-W 均选择 (k, blend)=(10, 1.00)，OpenEA en-de 的三个种子均选择 blend=1.00；"
        "但 DBP15K zh-en、fr-en、OpenEA D-Y 和 EventEA 的最优 k 或 blend 随种子变化。因而，固定网格并按验证 MRR "
        "独立选择参数，比根据测试集事后指定统一配置更符合可复现实验原则。"
    ),
    "需要指出，表 9": (
        "表 9 评价的是当前代码下简化配置与后处理的内部稳健性，而非与外部方法的严格 SOTA 比较。DBP15K 使用"
        "训练对留出验证集，OpenEA 与 EventEA 使用官方划分；EventEA 的候选集合还包含两侧知识图谱中的上下文实体，"
        "并使用官方实体名称文件，其绝对数值不宜与弱化名称信息的 OpenEA V2 直接比较。综合全部实验，可以支持的结论是："
        "关系感知消息、层融合、多尺度局部语义和早期交互具有较一致的内部证据，验证集选择 CSLS 具有最稳定的跨数据集增益；"
        "global 视图、复合损失、关系采样和主动学习仍需要更严格的复验。"
    ),
}


CAPTIONS = {
    "表 2": "表 2  八数据集统计与扩展实验划分",
    "表 3": "表 3  双数据集基础模型与后处理结果（机制验证协议，均值±样本标准差，n=3）",
    "表 4": "表 4  与经典方法的非严格对照（DBP15K ZH-EN）",
    "表 5": "表 5  核心组件消融结果（均值±样本标准差，n=3）",
    "表 6": "表 6  早期交互、晚期拼接与简单平均对照（均值±样本标准差，n=3）",
    "表 7": "表 7  单随机种子细粒度组件诊断",
    "表 8": "表 8  主动学习样本效率对照（单随机种子）",
    "表 9": "表 9  八数据集简化配置与验证集选择后处理结果（均值±样本标准差，n=3）",
}


def set_run_fonts(run, east_asia: str = "STSong", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:cs"), latin)


def replace_paragraph_text(paragraph, text: str) -> None:
    paragraph.text = text
    for run in paragraph.runs:
        set_run_fonts(run)


def replace_text_preserving_embedded_objects(paragraph, old: str, new: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    combined = "".join(node.text or "" for node in text_nodes)
    start = combined.find(old)
    if start < 0:
        raise RuntimeError(f"Could not locate paragraph text to replace: {old}")
    end = start + len(old)

    positions = []
    cursor = 0
    for node in text_nodes:
        text = node.text or ""
        positions.append((node, cursor, cursor + len(text)))
        cursor += len(text)

    overlapping = [item for item in positions if item[1] < end and item[2] > start]
    if not overlapping:
        raise RuntimeError("Replacement span does not intersect any Word text node")

    first_node, first_start, _ = overlapping[0]
    last_node, last_start, _ = overlapping[-1]
    before = (first_node.text or "")[: start - first_start]
    after = (last_node.text or "")[end - last_start :]
    first_node.text = before + new + (after if first_node is last_node else "")
    for node, _, _ in overlapping[1:-1]:
        node.text = ""
    if last_node is not first_node:
        last_node.text = after


def find_paragraph(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def find_caption(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Caption" and paragraph.text.strip().startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one caption starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def update_metric_paragraph(document: Document) -> None:
    paragraph = find_paragraph(document, "评价采用 Hits@1")
    old_prefix = "评价采用 Hits@1、Hits@10 和平均倒数排名（MRR）。主结果基于原始余弦相似度，CSLS（"
    new_prefix = (
        "本文以 Hits@1、Hits@10 和平均倒数排名（MRR）衡量检索质量，所有指标均基于完整目标实体集合的排序。"
        "基础模型以原始余弦相似度为主，CSLS（"
    )
    replace_text_preserving_embedded_objects(paragraph, old_prefix, new_prefix)
    old_suffix = (
        "）仅作为补充后处理结果单独报告。主模型、核心消融及融合位置对照均运行随机种子 42、43 和 44，"
        "表中给出均值与样本标准差。经典方法的结果仅在数据集、输入特征和评价设置不完全一致的条件下用于研究定位，"
        "不作为严格的复现排名。"
    )
    new_suffix = (
        "）作为检索后处理单独报告。除明确标注为单随机种子的探索性实验外，内部比较均采用随机种子 42、43 和 44，"
        "报告均值与样本标准差。模型检查点、重排序器和 CSLS 参数只使用训练集或验证集选择，测试集仅用于最终报告。"
        "外部方法仅在协议不完全一致的条件下用于定位，不进行显著性或最优性声明。"
    )
    replace_text_preserving_embedded_objects(paragraph, old_suffix, new_suffix)


def update_dataset_table(document: Document) -> None:
    table = next(
        table
        for table in document.tables
        if [cell.text.strip() for cell in table.rows[0].cells]
        == ["数据集", "实体数", "关系三元组", "训练/验证/测试对齐", "序列形状", "关系类型"]
    )
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    template_row = deepcopy(table.rows[1]._tr)
    for row in list(table.rows[1:]):
        table._tbl.remove(row._tr)
    for _ in DATASET_ROWS:
        table._tbl.append(deepcopy(template_row))

    headers = ["数据集", "实体数（合计）", "三元组（合计）", "训练/验证/测试", "语义序列形状", "关系类型"]
    for row_index, values in enumerate([headers, *DATASET_ROWS]):
        row = table.rows[row_index]
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for col_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                set_run_fonts(run)
                run.font.size = Pt(8.0)
                run.bold = row_index == 0

    header_pr = table.rows[0]._tr.get_or_add_trPr()
    if header_pr.find(qn("w:tblHeader")) is None:
        header_pr.append(OxmlElement("w:tblHeader"))


def update_document(document: Document) -> None:
    overview = document.paragraphs[136]
    if overview.text.strip():
        raise RuntimeError("Expected the paragraph after Table 2 to be blank")
    replace_paragraph_text(
        overview,
        "实验围绕五个问题展开：整体模型是否有效且稳定；哪些结构、语义和损失组件产生可重复贡献；早期交互是否优于"
        "晚期融合；主动查询是否在相同标注预算下提高样本效率；简化配置与验证集选择后处理能否跨语言、同语种异源"
        "知识库和事件型知识图谱保持增益。",
    )

    for prefix, replacement in REPLACEMENTS.items():
        replace_paragraph_text(find_paragraph(document, prefix), replacement)

    for prefix, replacement in CAPTIONS.items():
        replace_paragraph_text(find_caption(document, prefix), replacement)

    update_metric_paragraph(document)
    update_dataset_table(document)

    body_indent = find_paragraph(document, "在机制验证协议下").paragraph_format.first_line_indent
    section_57_prefixes = [
        "表 9 检验由前述诊断",
        "最终组合在 24/24",
        "重排序器单独作用时",
        "验证集选择结果进一步表明",
        "表 9 评价的是当前代码下",
    ]
    for prefix in section_57_prefixes:
        find_paragraph(document, prefix).paragraph_format.first_line_indent = body_indent
    find_paragraph(document, "最终组合在 24/24").paragraph_format.space_before = Pt(8)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.source.resolve() == args.output.resolve():
        raise ValueError("Output must differ from source")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source, args.output)
    document = Document(args.output)
    update_document(document)
    document.save(args.output)


if __name__ == "__main__":
    main()
