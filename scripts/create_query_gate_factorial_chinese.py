from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "1_二因素审稿实验标黄修订版_20260831.docx"
OUTPUT = ROOT / "outputs" / "1_二因素审稿实验标黄修订版_中文版_20260831.docx"
FIGURE = ROOT / "outputs" / "paper_assets" / "figure_1_springer_submission_zh_600dpi.png"

TITLE = "知识图谱实体对齐中的关系感知邻居上下文与语义引导选择"
PLACEHOLDER_RE = re.compile(r"\{\{EQ(\d+)\}\}")


PLAIN = {
    0: TITLE,
    1: "林心怡",
    2: "摘要",
    3: (
        "知识图谱实体对齐旨在识别独立构建的图谱中指向同一对象的实体。本文提出一种关系感知邻居上下文模型。"
        "结构编码器保留关系类型，并组合不同传播深度的图状态。语义编码器从同一实体文本序列提取 token、phrase 和 global 三种视图。"
        "融合模块为完整的一跳出邻域评分，使用 1.5-entmax 抑制弱相关邻居，并将所得结构上下文与实体语义结合。"
        "本文在五个数据集上采用三个随机种子进行实验。完整模型在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的 Hits@1 分别为 0.7285 和 0.6344，"
        "相对语义分支分别提高 9.09 和 9.66 个百分点。二因素对照分别改变点积查询来源和兼容性门控来源。"
        "在两种门控设置上取平均后，语义查询使 Hits@1 分别变化 +0.07 和 +0.13 个百分点；在两种查询设置上取平均后，语义门控使其分别变化 +0.03 和 -0.18 个百分点。"
        "结果表明，关系类型和结构邻居上下文是性能提升的主要来源，查询来源与门控来源的影响较小。"
    ),
    4: (
        "关键词：知识图谱实体对齐；跨语言知识图谱；事件知识图谱；关系感知图神经网络；"
        "结构—语义融合；稀疏邻居选择"
    ),
    5: "1 引言",
    6: "1.1 问题与证据缺口",
    8: (
        "现有方法已经证明图结构和文本信息均有价值。MTransE 和 JAPE 将图谱映射到共享或相互关联的向量空间[2,3]，"
        "BootEA 将高置信实体对加入训练集[4]，RDGCN 和 RREA 则保留关系与邻域信息[5,6]。"
        "这些研究留下了一个尚待回答的设计问题：结构邻居应在何时进入联合表示？"
    ),
    9: (
        "邻居选择需要同时考虑关系类型与相关性，而且并非每个邻居都有助于识别当前实体。"
        "因此，本文保留关系类型，使用实体语义为结构邻居加权，并将筛选后的邻居上下文与语义表示结合。"
    ),
    10: "1.2 研究问题与贡献",
    11: (
        "本文围绕三个问题展开。关系感知邻居上下文能否提供语义表示之外的对齐证据？"
        "当邻域和融合容量固定时，点积查询来源与兼容性门控来源分别产生何种影响？"
        "关系类型、邻居上下文和多尺度语义编码如何影响联合模型？"
    ),
    12: "本文的主要贡献如下。",
    13: (
        "本文提出一种关系感知结构编码器。共享变换控制参数规模，关系嵌入保留边类型差异，"
        "节点级层选择器组合不同传播深度的图状态。"
    ),
    14: (
        "本文设计语义引导的结构邻居选择方法。语义编码器从同一实体文本提取 token、phrase 和 global 三种视图，"
        "融合模块使用所得表示为完整的一跳出邻域评分。1.5-entmax 将弱相关邻居的权重置零，"
        "筛选后的邻居上下文与实体语义共同形成联合表示。"
    ),
    15: (
        "本文在五个数据集上使用三个随机种子，检验完整模型在跨语言和事件知识图谱设置中的表现。"
        "组件对照评估关系类型、多尺度语义和邻居访问。二因素实验独立改变点积查询来源与兼容性门控来源，"
        "同时固定邻域、可训练融合容量、上下文门控和检索表示。"
    ),
    16: "2 相关工作",
    17: "2.1 共享表示方法",
    18: (
        "MTransE 学习不同语言嵌入空间之间的映射[2]，JAPE 在共享结构空间中加入属性关联[3]。"
        "这类方法通过显式几何映射连接图谱，但难以直接利用特定关系对应的局部邻域。后续研究因此转向关系感知结构编码器。"
    ),
    19: "2.2 关系感知结构编码",
    20: (
        "GraphSAGE 通过聚合局部邻居学习节点表示[9]。R-GCN 在消息变换中保留关系类型，"
        "并通过共享基或分块结构限制参数增长[7]。这些研究表明，图编码器需要同时表示邻接实体和边类型。"
    ),
    21: (
        "RDGCN 通过实体图与关系对偶图引入关系证据[5]，RREA 使用关系反射变换保留关系差异[6]。"
        "RPR-RHGT 使用图 Transformer 编码筛选后的多步关系路径[10]。本文采用关系感知消息传递，"
        "并考察一跳结构邻居如何进入联合表示。"
    ),
    22: "2.3 结构—语义交互",
    23: (
        "Transformer 能够建立序列中不同位置之间的联系[11]，层次注意力网络则表明不同文本层次能够提供不同证据[12]。"
        "本文的语义编码器遵循这一多尺度思想，从同一实体文本提取 token、phrase 和 global 三种视图，"
        "并使用融合后的语义表示参与邻居选择。"
    ),
    24: (
        "文本增强实体对齐研究表明，语义信息能够补充图结构。RREA 分别报告文本增强和纯结构设置[6]，"
        "MCLEA 通过对比目标对齐多种数据类型[13]。其他研究发现，简单拼接或注意力可能混合尚未对齐的表示空间[14]。"
        "这些结果促使本文对邻居评分信号开展受控检验。本文在保持邻域和可训练融合参数量不变时，分别控制查询来源与兼容性门控来源。"
    ),
    25: (
        "近期方法还使用大语言模型[15]、域适应[16]和逐步推理代理[17]增强跨图监督或推理。"
        "本文研究固定实体对齐框架内局部结构邻居与实体语义之间的交互。"
    ),
    26: "3 问题定义与方法",
    27: "3.1 实体对齐任务",
    28: "本文将两个知识图谱定义如下。",
    34: "3.2 模型概述",
    35: (
        "模型包含结构编码、语义编码、邻居上下文融合和对比学习四个部分。结构编码器保留关系类型，并组合不同传播深度。"
        "语义编码器从同一实体文本构建 token、phrase 和 global 三种视图。融合模块使用语义信息为全部一跳出邻居评分，"
        "形成稀疏结构上下文，再将该上下文与实体语义结合。双向 InfoNCE 同时约束联合表示和结构表示，"
        "检索阶段使用 CSLS 校正训练后联合表示之间的相似度。图 1 展示完整信息流。"
    ),
    38: "3.3 关系感知结构编码器",
    39: (
        "局部度和邻居统计可以描述节点的结构角色[18,19]，结构身份方法也根据节点在图中的角色进行比较。"
        "受此启发，模型使用八维局部拓扑向量初始化结构分支。"
    ),
    46: (
        "结构编码与上下文选择使用不同的边方向。消息传递通过传入边更新目标节点，从而聚合指向该实体的事实。"
        "融合模块读取一跳出邻居，并将该实体发出的事实视为候选上下文。这两种方向分别服务于节点更新和上下文选择。"
    ),
    56: "每层传播均采用 LayerNorm、ReLU 和 dropout，以稳定不同深度的特征分布并降低过拟合风险。",
    57: "堆叠传播使实体能够汇集更远邻居的证据。",
    62: "3.4 多尺度语义编码器",
    63: (
        "三个语义视图使用同一实体文本，但编码范围不同。token 视图汇总词级线索，phrase 视图提取局部组合，"
        "global 视图建模不同位置之间的依赖。"
    ),
    65: (
        "token 视图直接汇总基础序列。均值池化保留全部有效 token 的信息，注意力池化提高重要 token 的权重。"
        "模型对两种输出取平均，形成词级表示。"
    ),
    67: (
        "掩码排除填充位置。均值池化汇总全部有效 token，注意力池化提高区分性词项的贡献，二者的平均结果构成 token 视图。"
    ),
    68: "phrase 视图组合基础序列中相邻的 token，从名称、关系短语和属性文本中提取局部模式。",
    71: (
        "global 视图使用 Transformer 编码器连接相距较远的 token，使名称、关系和属性线索能够在完整实体文本中相互影响。"
    ),
    74: "不同实体可能依赖不同文本粒度。因此，模型根据三个视图的表示估计各视图权重。",
    77: "3.5 语义引导的结构上下文",
    78: (
        "融合模块从完整的一跳出邻域中构造实体特定的结构上下文。语义查询首先为结构邻居评分，1.5-entmax 将弱相关候选的权重置零。"
        "模型随后将加权邻居摘要与实体自身结构状态结合，并通过最终门控将该上下文与实体语义融合。"
    ),
    79: "3.5.1 查询来源与兼容性门控来源",
    83: "点积可能遗漏匹配向量各维度之间的细粒度关系。因此，模型定义如下交互特征。",
    90: "3.5.2 结构上下文构造",
    92: (
        "模型随后使用实体自身结构表示、邻居摘要和语义表示，为每个维度计算上下文门控。"
        "在所有二因素对照中，式（18）的语义状态保持不变，因此实验只改变式（14）的查询来源和式（16）的兼容性门控来源。"
    ),
    96: "3.5.3 最终联合表示",
    101: (
        "在完整模型中，语义信息同时通过点积查询和兼容性门控进入邻居加权。所得结构上下文再与原始语义状态结合。"
        "第 3.5.4 节在保持上下文构造和最终融合不变的条件下，分别检验这两条条件化路径。"
    ),
    102: "3.5.4 二因素匹配对照",
    106: "3.6 训练目标",
    115: "结构损失在第一个 epoch 的权重为 0.1，并在最后一个 epoch 降至零。",
    116: "3.7 使用 CSLS 检索",
    120: "4 实验设计",
    121: "4.1 数据集与实验划分",
    122: (
        "表 1 汇总五个数据集的统计信息和实验划分。三个 DBP15K 数据对与 OpenEA EN–FR-15K-V2 采用跨语言设置，"
        "EventEA EN–EN-20K 采用同语言事件知识图谱设置。"
    ),
    123: "表 1 五个数据集的统计信息与实验划分",
    124: "4.2 评价指标与统计报告",
    125: (
        "评估时，每个实体都在完整候选集合中参与排序。本文报告 Hits@1、Hits@10 和平均倒数排名（MRR）。"
        "主实验与消融实验均给出三个随机种子的均值和样本标准差。"
    ),
    126: "4.3 实现细节与超参数选择",
    128: "5 结果与分析",
    129: "5.1 多数据集结果",
    130: (
        "表 2 报告四个跨语言数据集和一个同语言事件知识图谱数据集上的联合表示检索结果。每个数据集均使用三个随机种子。"
        "DBP15K 从训练实体对中留出 10% 作为验证集，OpenEA 和 EventEA 使用官方验证集。"
    ),
    131: "表 2 五个数据集的最终测试结果（均值 ± 样本标准差，n = 3）",
    133: (
        "五个数据集的多随机种子结果波动较小。因此，同一完整配置可以用于跨语言图谱对和同语言事件知识图谱对，"
        "并为后续匹配对照提供稳定的实验基础。"
    ),
    134: "5.2 与已发表结果的相对位置",
    135: (
        "表 3 列出早期研究公开报告的结果。MTransE、JAPE、BootEA 和 RDGCN 的 DBP15K 数值取自 RDGCN[5]，"
        "RREA-text 数值取自 RREA[6]，OpenEA EN–FR-15K-V2 数值采用 OpenEA 官方五折平均结果[8]。"
    ),
    136: "表 3 已发表结果与本文模型的定位性比较",
    139: "5.3 组件对照证据",
    140: (
        "表 4 报告 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的组件对照。"
        "每个变体均使用三个随机种子，并在训练与检索阶段使用同一种表示。"
    ),
    141: "表 4 结构组件、邻居聚合与语义视图的消融结果（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
    142: (
        "关系类型是两个数据集上最稳定的结构因素。移除关系类型后，DBP15K 和 OpenEA 的 Hits@1 分别下降 3.62 和 2.78 个百分点。"
        "拓扑特征相对仅使用可学习实体向量分别提高 3.46 和 8.03 个百分点。移除 phrase 视图后，Hits@1 分别下降 0.81 和 3.95 个百分点。"
        "其余变体的影响较小或随数据集变化。"
    ),
    143: "5.4 结构上下文在联合表示中的价值",
    144: (
        "表 5 比较完整模型、两个单分支模型和无可训练参数的平均融合。每个变体在训练与检索阶段使用同一种表示，"
        "验证过程不调整分支权重。"
    ),
    145: "表 5 分支、训练表示与主检索表示对照（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
    146: (
        "完整模型相对语义分支在 DBP15K 和 OpenEA 上的 Hits@1 分别提高 9.09 和 9.66 个百分点，并且优于无可训练参数的平均融合。"
        "这些结果表明，学习型交互能够将结构上下文转化为语义表示之外的对齐证据。结构单分支的较低结果进一步说明，"
        "这部分增益依赖结构与语义的联合学习。"
    ),
    147: "5.5 查询与兼容性门控的二因素对照",
    148: (
        "表 6 报告四种查询—门控组合在统一协议下的重跑结果。各组合使用相同的完整一跳出邻域、197,377 个可训练融合参数、"
        "1.5-entmax、温度、上下文门控、最终融合、训练目标和直接联合检索。实验只改变式（14）和式（16）的信息来源。"
        "晚期融合保留相同参数量，并作为不读取邻居的已有参照。"
    ),
    149: "表 6 查询—门控二因素对照与晚期融合参照（Hits@1/MRR，均值 ± 样本标准差，n = 3）",
    150: (
        "在使用语义门控时，语义查询效应定义为 SS−TS；在使用结构门控时，该效应定义为 ST−TT。"
        "DBP15K 上配对 Hits@1 效应分别为 +0.07 ± 0.06 和 +0.07 ± 0.20 个百分点，OpenEA 上分别为 +0.29 ± 1.07 和 -0.04 ± 0.82 个百分点。"
        "在使用语义查询时，语义门控效应定义为 SS−ST；在使用结构查询时，该效应定义为 TS−TT。"
        "DBP15K 上相应效应分别为 +0.03 ± 0.07 和 +0.03 ± 0.27 个百分点，OpenEA 上分别为 -0.01 ± 0.34 和 -0.34 ± 0.52 个百分点。"
    ),
    151: (
        "对另一因素取平均后，语义查询对 Hits@1 的主效应分别为 +0.07 和 +0.13 个百分点，语义门控的主效应分别为 +0.03 和 -0.18 个百分点。"
        "查询—门控交互效应分别为 +0.00 和 +0.33 个百分点。这些效应小于完整模型相对参数匹配晚期融合的 5.28 和 9.22 个百分点差距。"
        "因此，直接机制证据更有力地支持邻居访问与学习型交互路径，而不是其中任一条语义条件化路径。"
    ),
    152: "6 讨论",
    153: "6.1 实验结论",
    154: (
        "组件对照将关系类型和结构邻居上下文确定为性能提升的主要来源。完整模型在两个数据集上均优于语义分支，"
        "说明筛选后的结构上下文提供了文本之外的证据。二因素实验进一步表明，当邻域、容量、上下文门控和检索表示固定后，"
        "点积查询来源与兼容性门控来源引起的变化较小。"
    ),
    155: (
        "增益大小取决于输入数据。DBP15K 提供较有用的名称和局部图线索，EventEA 包含多样的事件关系和属性[25]，"
        "OpenEA 使用编码后的实体标识以降低名称偏置[8]。在本文预处理中，GloVe[26] 覆盖约 52.18% 的 OpenEA token。"
        "当文本能够引导邻居选择，且关系类型能够区分相似局部结构时，本方法表现更好。"
    ),
    156: "6.2 证据范围",
    157: (
        "主实验覆盖四个跨语言图谱对和同语言 EventEA EN–EN 图谱对。组件对照集中在 DBP15K ZH–EN 与 OpenEA EN–FR-15K-V2。"
        "表 6 的二因素结果来自独立完成的三随机种子重跑，支持这两个数据集上的机制结论。固定八邻居 softmax 对照同时改变邻居数量和加权函数。"
        "参数匹配晚期融合同时移除邻居访问并改变训练交互路径，因此其差异表示整体设计效应。"
    ),
    158: "6.3 适用条件与后续工作",
    159: (
        "后续研究可以从输入语义、图规模和任务设置三个方面扩展。子词编码器可以替代未知 token 的随机向量，"
        "跨语言数据对可以使用多语言编码器[27,28]。更大规模图谱可用于检验完整邻域选择的效率，"
        "开放世界数据集可用于评估不存在匹配实体的情况。这些研究将进一步检验语义引导邻居选择在文本线索较弱和候选空间较大时的表现。"
    ),
    160: "7 结论",
    161: (
        "本文研究关系感知结构邻居在结构—语义联合表示中的作用。关系类型在两个受控数据集上增强结构证据，完整模型也优于语义分支。"
        "完整邻域访问与学习型交互路径的组合进一步优于参数匹配晚期融合。二因素实验将点积查询与兼容性门控分开，"
        "并表明二者的信息来源选择对结果的影响小于邻居访问。五个数据集的结果表明，该完整配置可用于跨语言图谱对和同语言事件知识图谱对。"
    ),
    162: "附录 A 查询—门控二因素结果",
    163: (
        "每个组合的第一个字母表示点积查询来源，第二个字母表示兼容性门控来源。S 表示语义状态，T 表示结构状态。"
        "每个单元格报告 Hits@1/MRR，以及验证集选择的 CSLS k。"
    ),
    164: "表 A1 查询—门控二因素对照的逐随机种子结果",
    165: "参考文献",
}


MARKUP = {
    7: (
        "知识图谱通常以三元组 {{EQ0}} 表示事实。在三元组中，{{EQ1}}、{{EQ2}} 和 {{EQ3}} 分别表示头实体、关系和尾实体。"
        "该三元组表示关系 {{EQ4}} 将头实体 {{EQ5}} 与尾实体 {{EQ6}} 相连[1]。不同系统通常独立构建知识图谱。"
        "因此，同一对象可能具有不同的名称、标识符、属性和值格式。实体对齐需要在实体集合 {{EQ7}} 与 {{EQ8}} 之间找到对应实体。"
    ),
    30: "其中，{{EQ0}}、{{EQ1}} 和 {{EQ2}} 分别表示 {{EQ3}} 侧的实体集合、关系集合和三元组集合。任务还提供少量种子对齐集合。",
    32: (
        "集合 {{EQ0}} 包含由人工标注或数据集规则确认的对应实体对。符号 {{EQ1}} 表示两个实体指向同一个现实对象。"
    ),
    33: (
        "模型学习评分函数 {{EQ0}}，使真实对应实体的得分高于其他候选实体。评估采用一对一闭世界设置。"
    ),
    37: (
        "图 1 所提出模型的信息流。结构分支保留关系类型并选择有效传播深度，语义分支从同一输入构建三种视图。"
        "语义引导的查询为全部一跳出邻居评分，1.5-entmax 去除弱相关证据，模型随后将筛选后的结构上下文与语义信息结合。"
        "训练和检索均使用联合表示，验证集仅选择 {{EQ0}}。"
    ),
    40: (
        "模型在两张图中分别计算八维拓扑向量，该向量不依赖跨图关系标识。对于实体 {{EQ0}}，变量 {{EQ1}}、{{EQ2}} 和 {{EQ3}} 分别表示入度、出度和总度；"
        "{{EQ4}} 与 {{EQ5}} 分别表示传入和传出关系类型的数量；{{EQ6}} 与 {{EQ7}} 分别表示传入源实体和传出目标实体的 {{EQ8}} 均值。"
        "节点没有相应邻居时，模型将该均值记为零。变量 {{EQ9}} 衡量入边与出边的平衡。原始特征向量定义如下。"
    ),
    42: "模型分别在左右图内部对各拓扑特征进行标准化，得到 {{EQ0}}。初始结构状态由共享拓扑投影与较小的实体特定残差组成。",
    44: (
        "矩阵 {{EQ0}} 是左右图共享的线性投影，向量 {{EQ1}} 是可训练实体向量。拓扑项为相似局部角色提供可比较的起点。"
        "残差系数设为 0.1，以避免局部统计相似的不同实体获得相同表示。"
    ),
    45: (
        "在每次前向计算中，结构编码器在整张图上执行有向关系感知消息传递。对于三元组 {{EQ0}}，头实体 {{EQ1}} 向尾实体 {{EQ2}} 发送消息。"
        "消息沿原始边方向传播，编码器只使用传入边。变量 {{EQ3}} 表示实体 {{EQ4}} 在第 {{EQ5}} 层的状态，{{EQ6}} 表示关系嵌入，"
        "集合 {{EQ7}} 包含实体 {{EQ8}} 的传入边。第 {{EQ9}} 层首先计算如下消息。"
    ),
    48: (
        "矩阵 {{EQ0}} 与 {{EQ1}} 分别投影源实体和关系嵌入。同一层中的全部关系共享这两个矩阵，关系差异由嵌入 {{EQ2}} 保留，"
        "因此模型无需为每种关系配置独立的完整矩阵。向量 {{EQ3}} 表示边 {{EQ4}} 发送给目标实体 {{EQ5}} 的消息。"
        "模型对所有到达实体 {{EQ6}} 的消息取平均。"
    ),
    50: "模型先将实体自身状态投影为 {{EQ0}}，再根据自身状态和平均邻居消息为每个维度计算门控。",
    53: "式（3）中的 {{EQ0}} 为每个表示维度给出 0 到 1 之间的权重。式（4）保留投影后的自身状态，并通过该门控加入邻居证据。",
    54: "式（4）给出第 {{EQ1}} 层的原始输出 {{EQ0}}。除最后一层外，模型在进入下一层前对该输出进行变换。",
    58: (
        "最深层可能丢失实体自身或直接邻居的信息。因此，编码器保留投影后的输入和每层原始输出 {{EQ0}}。"
        "状态 {{EQ1}} 不含邻居信息，状态 {{EQ2}} 包含至多 {{EQ3}} 跳信息。模型先计算平均上下文 {{EQ4}}，"
        "再使用两层多层感知机（MLP）结合该上下文为每个状态评分。{{EQ5}} 将这些分数转换为深度权重 {{EQ6}}。"
    ),
    61: (
        "层选择 MLP 的输入维度为 {{EQ0}}，隐藏维度为 {{EQ1}}。该网络使用高斯误差线性单元（GELU）和 dropout，并输出一个分数。"
        "式（7）为 {{EQ2}} 到 {{EQ3}} 的每个深度分配独立权重。LayerNorm 和 L2 归一化随后得到结构表示 {{EQ4}}。"
        "因此，每个实体可以使用不同的传播深度组合。"
    ),
    64: (
        "对于实体 {{EQ0}}，模型将名称、相邻关系名称和属性文本拼接为一个 token 序列，并把每个 token 映射为词向量。"
        "phrase 与 global 视图使用同一序列，因此三个视图的信息均来自该序列中的名称、关系名称和属性文本。"
        "共享投影、归一化和位置编码得到基础表示 {{EQ1}}，掩码 {{EQ2}} 标记有效 token。"
    ),
    76: (
        "模型对三个视图向量加权后进行拼接，并由输出 MLP 处理拼接结果。随后，模型将该输出与三个加权视图之和的投影相加。"
        "L2 归一化得到语义表示 {{EQ0}}。"
    ),
    80: (
        "对于实体 {{EQ0}}，令 {{EQ1}} 表示其语义状态，{{EQ2}} 表示实体自身的结构状态，{{EQ3}} 表示出邻居 {{EQ4}} 的结构状态。"
        "点积查询使用来源 {{EQ5}}，兼容性门控使用来源 {{EQ6}}。每个来源均可取 {{EQ7}} 或 {{EQ8}}。"
        "完整模型将两个来源都设为 {{EQ9}}，键和值由 {{EQ10}} 投影得到。"
    ),
    82: "式（14）衡量查询来源 {{EQ0}} 与结构邻居 {{EQ1}} 的兼容性。缩放因子 {{EQ2}} 控制点积的数值范围，其中 {{EQ3}}。",
    85: (
        "其中，拼接项 {{EQ0}} 保留两个输入向量，{{EQ1}} 描述对应维度的共同激活，{{EQ2}} 记录逐维绝对差。"
        "该特征组合沿用自然语言推理中的匹配特征设计[20]，并作为可学习兼容性门控的输入。可训练映射与 sigmoid 函数随后产生邻居门值。"
    ),
    87: (
        "标量 {{EQ0}} 衡量邻居 {{EQ1}} 与门控来源 {{EQ2}} 的兼容性。模型在归一化前将该值与点积分数组合。"
        "将 {{EQ3}} 与 {{EQ4}} 分开后，两条语义条件化路径可以独立检验。"
    ),
    89: (
        "式（17）中，{{EQ0}} 和 {{EQ1}} 分别包含批次补齐后全部邻居位置的分数与门值。有效邻居的掩码偏置 {{EQ2}} 为零，"
        "填充位置为 {{EQ3}}。掩码在温度缩放后加入。温度 {{EQ4}} 控制权重集中程度，1.5-entmax 可以将低分有效邻居的权重置为精确的零[21]。"
        "有效权重之和为 1；实体没有有效邻居时，邻居摘要为全零向量。"
    ),
    91: (
        "模型使用式（17）的稀疏权重对邻居值向量加权平均，得到实体 {{EQ1}} 的结构邻居摘要 {{EQ0}}。"
        "完整模型在加权前保留全部一跳出邻居，因此摘要可以利用完整局部邻域，同时用稀疏权重降低无关邻居的影响。"
    ),
    95: (
        "门向量 {{EQ0}} 控制每个维度中实体自身结构与邻居摘要的比例。门值接近 1 时，式（19）偏向 {{EQ1}}；"
        "门值接近 0 时，式（19）偏向 {{EQ2}}。若实体没有有效邻居，模型令 {{EQ3}}，从而避免填充位置改变实体表示。"
    ),
    97: "模型得到结构上下文 {{EQ0}} 后计算联合门控，以确定每个维度中结构证据与语义证据的比例。",
    100: (
        "式（21）首先得到门控向量 {{EQ0}}。联合门值接近 1 时更偏向结构上下文，接近 0 时更偏向语义证据。"
        "完整模型随后计算等权基础向量 {{EQ1}}，并输出 {{EQ2}}。该固定残差限制门控造成的大幅偏移，同时不增加可训练参数。"
    ),
    103: (
        "二因素对照将查询来源 {{EQ0}} 与兼容性门控来源 {{EQ1}} 交叉组合，得到语义查询+语义门控（SS）、结构查询+语义门控（TS）、"
        "语义查询+结构门控（ST）和结构查询+结构门控（TT）。四种配置使用相同的完整一跳出邻域、动态补齐、1.5-entmax、温度、"
        "可训练层、InfoNCE 目标、式（18）的上下文门控、最终联合门控和联合检索。参数匹配晚期融合保留相同的融合参数量，但不使用邻居上下文。"
    ),
    105: (
        "式（22）分别归一化结构表示和语义表示，再将二者拼接并进行非线性映射，得到用于 InfoNCE 训练和检索的 {{EQ0}}。"
        "该基线与完整模型具有相同数量的可训练融合参数。"
    ),
    107: (
        "模型在单阶段训练中使用双向 InfoNCE[22]。每个小批次从种子对齐集合采样 {{EQ0}} 个实体对，"
        "并使用同一批实体对计算联合损失和结构损失。两项损失在一次前向与反向传播中共同优化。"
        "左图实体 {{EQ1}} 与右图实体 {{EQ2}} 之间的温度缩放余弦相似度为"
    ),
    109: (
        "变量 {{EQ0}} 表示温度。正确实体对位于相似度矩阵的对角线上。模型在两个检索方向上计算交叉熵，并取平均得到联合损失 {{EQ1}}。"
    ),
    111: (
        "式（24）直接训练联合表示。基于拓扑的结构分支也需要在两张图之间建立共享空间，因此模型在训练早期对 {{EQ0}} 使用同形式的双向 InfoNCE，"
        "得到结构损失 {{EQ1}}。总目标定义如下。"
    ),
    113: (
        "模型只在训练早期使用结构监督，其权重随训练线性下降。变量 {{EQ0}} 表示当前 epoch，{{EQ1}} 表示总 epoch 数。"
        "训练进度和结构损失权重定义如下。"
    ),
    117: "式（25）用于训练编码器。主检索直接使用训练后的联合表示 {{EQ0}}。CSLS[23] 校正联合表示之间的余弦相似度。",
    119: (
        "{{EQ0}} 与 {{EQ1}} 表示平均余弦相似度，分别使用查询实体和候选实体在另一张图中的 {{EQ2}} 个最近实体。"
        "每次运行均依据验证集 MRR 从候选集合中选择 {{EQ3}}，随后只在测试集上评估一次。"
    ),
    127: (
        "模型使用 AdamW[24]，批大小为 512，表示维度为 128。结构编码器包含三层关系感知图神经网络，所有组件的 dropout 均为 0.1。"
        "语义输入由 300 维投影至 128 维。global 视图采用两层、四头 Transformer，phrase 视图采用宽度为 3 和 5 的卷积。"
        "1.5-entmax 温度为 {{EQ0}}，InfoNCE 温度为 {{EQ1}}。实验使用随机种子 42、43 和 44。DBP15K 最多训练 36 个 epoch，"
        "OpenEA 和 EventEA 最多训练 50 个 epoch。模型每五个 epoch 检查一次验证集 MRR，连续四次没有提升时停止训练。"
        "主检索直接使用训练后的联合表示。每次运行均依据验证集 MRR 从 {{EQ2}} 中选择一个值。"
    ),
}


TABLE_REPLACEMENTS = {
    0: {
        "Dataset": "数据集",
        "Entities (total)": "实体总数",
        "Triples (total)": "三元组总数",
        "Train/validation/test": "训练/验证/测试",
        "Semantic sequence shape": "语义序列形状",
        "Relation types": "关系类型",
    },
    1: {"Dataset": "数据集"},
    2: {
        "Dataset": "数据集",
        "Method": "方法",
        "Not reported": "未报告",
        "Proposed model": "本文模型",
    },
    3: {
        "Variant": "变体",
        "Full model": "完整模型",
        "Learnable entity initialization only (no topology features)": "仅使用可学习实体初始化（无拓扑特征）",
        "Fixed 8-neighbor budget + softmax": "固定 8 个邻居 + softmax",
        "Without relation types": "去除关系类型",
        "Without layer selector": "去除层选择器",
        "Without token view": "去除 token 视图",
        "Without phrase view": "去除 phrase 视图",
        "Without global view": "去除 global 视图",
    },
    4: {
        "Variant": "变体",
        "Training representation": "训练表示",
        "Primary retrieval representation": "主检索表示",
        "Full model": "完整模型",
        "Joint representation": "联合表示",
        "Structural branch only": "仅结构分支",
        "Structural representation": "结构表示",
        "Semantic branch only": "仅语义分支",
        "Semantic representation": "语义表示",
        "Parameter-free mean fusion": "无可训练参数的平均融合",
        "Equal mean of structure and semantics": "结构与语义等权平均",
        "Same equal-weight mean": "相同的等权平均",
    },
    5: {
        "Cell": "组合",
        "Query source": "查询来源",
        "Gate source": "门控来源",
        "Neighbor access": "邻居访问",
        "Semantic": "语义",
        "Structure": "结构",
        "Complete one-hop": "完整一跳邻域",
        "Late fusion": "晚期融合",
        "None": "无",
    },
    6: {
        "Dataset": "数据集",
        "Cell": "组合",
        "Seed 42": "随机种子 42",
        "Seed 43": "随机种子 43",
        "Seed 44": "随机种子 44",
    },
}


def equation_children(paragraph):
    return [
        copy.deepcopy(child)
        for child in paragraph._p.iterchildren()
        if child.tag in {qn("m:oMath"), qn("m:oMathPara")}
    ]


def set_fonts(run, east_asia: str = "Arial Unicode MS", size: float | None = None) -> None:
    latin_font = "Times New Roman" if east_asia == "Times New Roman" else east_asia
    run.font.name = latin_font
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin_font)
    fonts.set(qn("w:hAnsi"), latin_font)
    fonts.set(qn("w:cs"), latin_font)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:hint"), "eastAsia")
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attribute}")
        if key in fonts.attrib:
            del fonts.attrib[key]


def append_text(paragraph, text: str) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    set_fonts(run)


def clear_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text: str) -> None:
    clear_content(paragraph)
    append_text(paragraph, text)


def set_markup(paragraph, markup: str) -> None:
    equations = equation_children(paragraph)
    clear_content(paragraph)
    cursor = 0
    used: set[int] = set()
    for match in PLACEHOLDER_RE.finditer(markup):
        append_text(paragraph, markup[cursor:match.start()])
        index = int(match.group(1))
        if index >= len(equations):
            raise RuntimeError(f"Paragraph has {len(equations)} equations, but markup requests EQ{index}")
        paragraph._p.append(copy.deepcopy(equations[index]))
        used.add(index)
        cursor = match.end()
    append_text(paragraph, markup[cursor:])
    if len(used) != len(equations):
        raise RuntimeError(f"Markup retained {len(used)} of {len(equations)} equations: {markup}")


def has_yellow(paragraph) -> bool:
    return any(run._r.xpath("./w:rPr/w:highlight[@w:val='yellow']") for run in paragraph.runs)


def highlight_runs(paragraph) -> None:
    for run in paragraph.runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def translate_tables(document: Document) -> None:
    for table_index, replacements in TABLE_REPLACEMENTS.items():
        table = document.tables[table_index]
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                original = cell.text.strip()
                if original in replacements:
                    cell.text = replacements[original]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(1)
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    if row_index == 0 or len(cell.text) < 28:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        set_fonts(run, size=8.5)
                        if row_index == 0:
                            run.bold = True


def replace_figure(document: Document) -> None:
    if not FIGURE.exists():
        raise FileNotFoundError(FIGURE)
    shape = document.inline_shapes[0]
    relation_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    document.part.related_parts[relation_id]._blob = FIGURE.read_bytes()
    with Image.open(FIGURE) as image:
        width_px, height_px = image.size
    shape.width = Inches(6.35)
    shape.height = Inches(6.35 * height_px / width_px)
    properties = shape._inline.docPr
    properties.set("name", "图 1 所提出模型的信息流")
    properties.set(
        "descr",
        "关系感知结构编码、多尺度语义编码、完整一跳出邻域稀疏加权、InfoNCE训练与联合表示检索的信息流。",
    )


def normalize_typography(document: Document) -> None:
    style_specs = {
        "Normal": (10.5, "Arial Unicode MS", False),
        "List Number": (10.5, "Arial Unicode MS", False),
        "Heading 1": (12.0, "Arial Unicode MS", True),
        "Heading 2": (11.0, "Arial Unicode MS", True),
        "Heading 3": (10.5, "Arial Unicode MS", True),
        "Caption": (9.0, "Arial Unicode MS", False),
        "Reference": (9.0, "Times New Roman", False),
    }
    for name, (size, east_asia, bold) in style_specs.items():
        if name not in document.styles:
            continue
        style = document.styles[name]
        style.font.name = east_asia
        style.font.size = Pt(size)
        style.font.bold = bold
        rpr = style._element.get_or_add_rPr()
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:ascii"), east_asia)
        fonts.set(qn("w:hAnsi"), east_asia)
        fonts.set(qn("w:cs"), east_asia)
        fonts.set(qn("w:eastAsia"), east_asia)

    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        east_asia = "Arial Unicode MS"
        if style_name == "Reference":
            east_asia = "Times New Roman"
        for run in paragraph.runs:
            set_fonts(run, east_asia=east_asia)
        if style_name in {"Normal", "List Number"} and paragraph.text.strip():
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.5
        if style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if style_name == "Caption":
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True

    title = document.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_fonts(run, "Arial Unicode MS", 16)
        run.bold = True
    document.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in document.paragraphs[1].runs:
        set_fonts(run, "Arial Unicode MS", 11)

    for section in document.sections:
        if section.header.paragraphs:
            header = section.header.paragraphs[0]
            set_plain(header, "知识图谱实体对齐中的关系感知邻居上下文")
            header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in header.runs:
                set_fonts(run, "Arial Unicode MS", 8.5)


def set_language(document: Document) -> None:
    document.core_properties.language = "zh-CN"
    for part in document.part.package.parts:
        element = getattr(part, "_element", None)
        if element is None:
            continue
        for lang in element.xpath(".//w:lang"):
            lang.set(qn("w:val"), "zh-CN")
            lang.set(qn("w:eastAsia"), "zh-CN")


def equation_count(document: Document) -> int:
    return len(document._element.xpath(".//m:oMath"))


def build() -> Path:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    if len(document.paragraphs) != 194 or len(document.tables) != 7:
        raise RuntimeError("Unexpected source structure")
    original_equations = equation_count(document)
    highlighted = {index for index, paragraph in enumerate(document.paragraphs) if has_yellow(paragraph)}

    for index, text in PLAIN.items():
        if document.paragraphs[index]._p.xpath(".//m:oMath|.//m:oMathPara"):
            raise RuntimeError(f"Paragraph {index} contains math and requires markup")
        set_plain(document.paragraphs[index], text)
    for index, markup in MARKUP.items():
        set_markup(document.paragraphs[index], markup)

    for index in highlighted:
        highlight_runs(document.paragraphs[index])

    translate_tables(document)
    replace_figure(document)
    normalize_typography(document)
    set_language(document)

    document.core_properties.title = TITLE
    document.core_properties.author = "林心怡"
    document.core_properties.last_modified_by = "林心怡"
    document.core_properties.subject = "知识图谱实体对齐"
    document.core_properties.keywords = (
        "知识图谱实体对齐；跨语言知识图谱；事件知识图谱；关系感知图神经网络；结构—语义融合；稀疏邻居选择"
    )
    document.save(OUTPUT)

    checked = Document(OUTPUT)
    if len(checked.paragraphs) != 194 or len(checked.tables) != 7:
        raise RuntimeError("Paragraph or table count changed")
    if equation_count(checked) != original_equations:
        raise RuntimeError(f"Equation count changed: {original_equations} -> {equation_count(checked)}")
    if len(checked.inline_shapes) != 1:
        raise RuntimeError("Figure count changed")
    text = "\n".join(paragraph.text for paragraph in checked.paragraphs[:166])
    required = [
        TITLE,
        "二因素对照分别改变点积查询来源和兼容性门控来源",
        "语义查询对 Hits@1 的主效应分别为 +0.07 和 +0.13 个百分点",
        "附录 A 查询—门控二因素结果",
    ]
    missing = [item for item in required if item not in text]
    if missing:
        raise RuntimeError(f"Required Chinese content missing: {missing}")
    return OUTPUT


if __name__ == "__main__":
    print(build())
