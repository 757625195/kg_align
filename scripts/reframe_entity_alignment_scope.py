from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "outputs"

ZH_SOURCE = OUTPUTS / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_中文学术润色终稿_20260830_图1重绘版.docx"
ZH_OUTPUT = OUTPUTS / "知识图谱实体对齐中的关系感知邻居上下文与语义引导选择_任务范围重写版_20260830.docx"

EN_SOURCE = OUTPUTS / "Relation_Aware_Neighbor_Context_Bilingual_Polished_English_20260830_Figure1_Redesigned.docx"
EN_OUTPUT = OUTPUTS / "Relation_Aware_Neighbor_Context_Knowledge_Graph_Entity_Alignment_Reframed_20260830.docx"


def replace_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def append_text(paragraph, text: str) -> None:
    if text.strip() in paragraph.text:
        return
    run = paragraph.add_run(text)
    if len(paragraph.runs) > 1:
        previous = paragraph.runs[-2]
        if previous._r.rPr is not None:
            run._r.insert(0, deepcopy(previous._r.rPr))


def rewrite_chinese() -> None:
    document = Document(ZH_SOURCE)
    paragraphs = document.paragraphs

    title = "知识图谱实体对齐中的关系感知邻居上下文与语义引导选择"
    abstract = (
        "知识图谱实体对齐用于识别独立构建图谱中指向同一对象的实体。本文提出一种语义引导的关系感知邻居上下文模型，"
        "使实体语义在联合编码前参与结构证据选择。结构编码器在消息传递中保留关系类型并组合不同传播深度，语义编码器从同一实体文本序列提取 "
        "token、phrase 和 global 三种视图。融合模块使用 1.5-entmax 为完整的一跳出邻域分配稀疏权重，抑制弱相关邻居，再将结构上下文与实体语义编码为联合表示。"
        "实验在四个跨语言数据集和一个同语言事件图谱数据集上采用三个随机种子。完整模型在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2 上的 Hits@1 "
        "分别为 0.7285 和 0.6344，较语义单分支提高 9.09 和 9.66 个百分点。在邻域和融合参数量一致时，语义查询较结构查询提高 2.08 和 2.48 个百分点。"
        "结果表明，关系类型、结构邻居上下文和语义引导选择能够改善不同图谱设置下的实体对齐。"
    )
    keywords = (
        "关键词：知识图谱实体对齐；跨语言知识图谱；事件知识图谱；关系感知图神经网络；结构—语义融合；稀疏邻居选择"
    )

    replace_text(paragraphs[0], title)
    replace_text(paragraphs[3], abstract)
    replace_text(paragraphs[4], keywords)
    append_text(
        paragraphs[7],
        "本文将实体对齐作为一般的跨图检索任务，实验覆盖四个跨语言图谱对和一个同语言事件图谱对。",
    )
    replace_text(
        paragraphs[15],
        "在四个跨语言数据集和一个同语言事件图谱数据集上使用三个随机种子，检验完整模型在不同图谱设置下的表现。匹配对照分别检验关系类型、多尺度语义、邻居访问和查询信号，并在固定邻域与融合参数量的条件下比较语义查询和结构查询。",
    )
    replace_text(
        paragraphs[124],
        "表 1 汇总五个数据集的统计信息与实验划分。DBP15K 的三个图谱对和 OpenEA EN–FR-15K-V2 属于跨语言设置，EventEA EN–EN-20K 属于同语言事件图谱设置。该组合同时考察语言差异和事件结构差异。",
    )
    replace_text(
        paragraphs[132],
        "表 2 报告四个跨语言数据集和一个同语言事件图谱数据集上的联合表示检索结果。每个数据集均使用三个随机种子。DBP15K 从训练对中留出 10% 作为验证集，OpenEA 和 EventEA 使用官方验证集。",
    )
    replace_text(
        paragraphs[135],
        "模型在五个数据集上的结果较为稳定，Hits@1 的最大样本标准差为 0.0053。该结果表明同一完整配置能够用于跨语言图谱对和同语言事件图谱对，并为后续匹配对照提供稳定基础。",
    )
    replace_text(
        paragraphs[159],
        "主实验覆盖四个跨语言图谱对和 EventEA EN–EN 同语言事件图谱对，组件控制集中在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2。组件结论来自两个代表性跨语言场景，EventEA 结果用于检验完整配置在事件图谱上的适用性。三个随机种子用于估计均值与样本波动。固定 8 邻居加 softmax 的对照同时改变邻居数量和权重函数，其差异反映两者的共同作用。参数匹配晚期融合同时移除邻居访问并改变训练交互路径，其结果衡量整体设计差异。",
    )
    replace_text(
        paragraphs[161],
        "后续研究可以从输入语义、图规模和任务设定三个方面扩展。子词编码器可以替代未知 token 的随机向量，其中跨语言图谱可使用多语言编码器[27–28]。更大规模图谱可以检验完整邻域选择的效率，开放世界数据集可以评估不存在对应实体的情况。这些实验将进一步检验语义引导邻居选择在弱文本线索和大候选空间下的稳健性。",
    )
    replace_text(
        paragraphs[163],
        "本文研究关系感知结构邻居在结构—语义联合表示中的作用。关系类型在两个受控数据集上稳定增强结构证据，结构邻居上下文使完整模型超过语义单分支，语义查询也在匹配条件下优于结构查询。完整邻域访问与学习型交互路径的组合进一步超过参数匹配晚期融合。这些结果支持关系感知邻居上下文与语义引导选择在实体对齐中的价值。五数据集主实验进一步表明，完整配置能够用于跨语言图谱对和同语言事件图谱对。",
    )

    document.core_properties.title = title
    document.save(ZH_OUTPUT)


def rewrite_english() -> None:
    document = Document(EN_SOURCE)
    paragraphs = document.paragraphs

    title = "Relation-Aware Neighbor Context and Semantics-Guided Selection for Knowledge Graph Entity Alignment"
    abstract = (
        "Knowledge graph entity alignment identifies entities that refer to the same object across independently constructed graphs. "
        "This paper presents a relation-aware neighbor-context model in which entity semantics guides structural evidence selection before joint encoding. "
        "The structural encoder preserves relation types during message passing and combines graph states from different propagation depths. "
        "The semantic encoder extracts token, phrase, and global views from one entity-text sequence. "
        "The fusion module scores the complete one-hop outgoing neighborhood, applies 1.5-entmax to suppress weak neighbors, and combines the resulting structural context with entity semantics. "
        "Experiments use three random seeds on four cross-lingual datasets and one same-language event-centric dataset. "
        "The full model achieves Hits@1 scores of 0.7285 on DBP15K ZH–EN and 0.6344 on OpenEA EN–FR-15K-V2. "
        "It improves over the semantic branch by 9.09 and 9.66 percentage points. "
        "With the neighborhood and trainable fusion parameter count fixed, semantic queries improve over structural queries by 2.08 and 2.48 points. "
        "The results show that relation types, structural neighbor context, and semantics-guided selection improve entity alignment across different graph settings."
    )
    keywords = (
        "  knowledge graph entity alignment; cross-lingual knowledge graphs; event-centric knowledge graphs; "
        "relation-aware graph neural networks; structural–semantic fusion; sparse neighbor selection"
    )

    replace_text(paragraphs[0], title)
    replace_text(paragraphs[5], abstract)
    paragraphs[6].runs[0].text = "Keywords:"
    if len(paragraphs[6].runs) == 1:
        paragraphs[6].add_run(keywords)
    else:
        paragraphs[6].runs[1].text = keywords
        for run in paragraphs[6].runs[2:]:
            run.text = ""
    append_text(
        paragraphs[9],
        " This study treats entity alignment as a general cross-graph retrieval task. The benchmark suite covers four cross-lingual graph pairs and one same-language event-centric pair.",
    )
    replace_text(
        paragraphs[17],
        "Experiments with three random seeds on four cross-lingual datasets and one same-language event-centric dataset test the full model across different graph settings. Matched controls examine relation types, multi-scale semantics, neighbor access, and the query signal. The semantic-query comparison fixes both the neighborhood and the number of trainable fusion parameters.",
    )
    replace_text(
        paragraphs[126],
        "Table 1 summarizes the statistics and experimental splits of five datasets. The three DBP15K pairs and OpenEA EN–FR-15K-V2 use cross-lingual settings. EventEA EN–EN-20K uses a same-language, event-centric setting. This collection tests both language differences and event-structure differences.",
    )
    replace_text(
        paragraphs[134],
        "Table 2 reports joint-representation retrieval on four cross-lingual datasets and one same-language event-centric dataset. Each dataset uses three random seeds. DBP15K reserves 10% of its training pairs for validation. OpenEA and EventEA use their official validation sets.",
    )
    replace_text(
        paragraphs[137],
        "The results are stable across the five datasets. The largest sample standard deviation of Hits@1 is 0.0053. The same full configuration therefore applies to cross-lingual graph pairs and a same-language event-centric graph pair. This stability supports the matched comparisons in the following sections.",
    )
    replace_text(
        paragraphs[161],
        "The main experiments cover four cross-lingual graph pairs and the same-language EventEA EN–EN pair. Component controls focus on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. The component conclusions therefore come from two representative cross-lingual settings, while EventEA tests the complete configuration on an event-centric graph pair. Three random seeds estimate means and sample variation. The fixed 8-neighbor plus softmax control changes both the neighbor count and the weighting function, so its difference reflects both factors. The parameter-matched late-fusion control removes neighbor access and changes the training interaction path. Its result measures the overall design difference.",
    )
    replace_text(
        paragraphs[163],
        "Future work can extend the input semantics, graph scale, and task setting. Subword encoders can replace random vectors for unknown tokens, with multilingual encoders used for cross-lingual pairs [27, 28]. Larger graphs can test the efficiency of complete-neighborhood selection. Open-world datasets can evaluate entities without matches. These studies can further test semantics-guided neighbor selection with weaker text and larger candidate spaces.",
    )
    replace_text(
        paragraphs[165],
        "This paper examines the role of relation-aware structural neighbors in a structural–semantic joint representation. Relation types strengthen structural evidence on both controlled datasets. Structural neighbor context improves the full model over the semantic branch, and semantic queries outperform structural queries under matched conditions. Complete-neighborhood access with a learned interaction path also exceeds parameter-matched late fusion. These results support relation-aware neighbor context and semantics-guided selection for entity alignment. The five-dataset experiments also show that the complete configuration applies to cross-lingual graph pairs and a same-language event-centric graph pair.",
    )

    document.core_properties.title = title
    document.save(EN_OUTPUT)


def main() -> None:
    rewrite_chinese()
    rewrite_english()
    print(ZH_OUTPUT)
    print(EN_OUTPUT)


if __name__ == "__main__":
    main()
