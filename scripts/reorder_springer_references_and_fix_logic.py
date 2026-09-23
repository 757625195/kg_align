from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
import sys

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "Springer投稿删减版_均值符号修正_20260826.docx"
)
OUTPUT = ROOT / "outputs" / (
    "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐_"
    "Springer引用与逻辑修订版_20260826.docx"
)

sys.path.insert(0, str(ROOT / "scripts"))
from mathify_all_document_symbols import set_markup  # noqa: E402


# Springer numeric references are numbered by first appearance in the text.
# Values are old reference numbers; their positions are the new numbers.
REFERENCE_ORDER = [20, 1, 2, 3, 4, 5, 7, 14, 6, 8, 9, 10, 11, 12, 13, 19, 17, 16, 18, 15]
OLD_TO_NEW = {old: new for new, old in enumerate(REFERENCE_ORDER, start=1)}
CITATION_RE = re.compile(r"\[(\d+(?:\s*(?:,|–|-)\s*\d+)*)\]")


def combined_text(paragraph) -> str:
    return "".join(
        paragraph._p.xpath(".//w:t/text()|.//m:t/text()")
    ).replace("\u00a0", " ").strip()


def find_paragraph(document: Document, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}, found {len(matches)}")
    return matches[0]


def find_formula(document: Document, number: int):
    suffix = f"({number})"
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if combined_text(paragraph).endswith(suffix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one formula ending {suffix}, found {len(matches)}")
    return matches[0]


def rewrite(document: Document, prefix: str, markup: str) -> None:
    # Recover LaTeX commands that Python interprets as control-character
    # escapes inside long prose literals.
    markup = (
        markup.replace("\x07lpha", r"\alpha")
        .replace("\x08ar", r"\bar")
        .replace("\rho", r"\rho")
        .replace("\rmathrm", r"\mathrm")
        .replace("\right", r"\right")
        .replace("\t" + "imes", r"\times")
    )
    set_markup(find_paragraph(document, prefix), markup)


def replace_formula(document: Document, number: int, latex: str) -> None:
    set_markup(find_formula(document, number), f"[[{latex}]].\t({number})")


def rewrite_logic(document: Document) -> None:
    rewrite(
        document,
        "本研究旨在检验结构信息",
        "本研究旨在检验结构信息与语义信息的组合方式如何影响跨语言实体检索，"
        "重点考察三个问题：关系类型和传播深度如何影响结构表示；不同语义粒度分别提供何种信息；"
        "在控制训练阶段融合参数量后，最终检索表示形成前引入一跳结构邻居是否优于仅使用两个独立分支，"
        "以及邻居查询使用语义信号或结构信号是否产生差异。具体研究问题与评价协议见第 4.1 节。",
    )
    rewrite(
        document,
        "本文使用的 GraphSAGE",
        "本文使用的 GraphSAGE 邻域聚合[6]、R-GCN 关系建模[7]、Transformer 自注意力[8]和 InfoNCE 对比目标[9]均来自已有研究。"
        "本文关注的科学问题是：在控制训练阶段融合容量后，最终检索表示形成前可见的一跳结构邻居是否带来额外信息，"
        "以及语义查询是否比结构查询更适合构造该邻居上下文。本文不把晚期基线与主模型的差值单独归因于融合时机，"
        "因为该基线同时改变了邻居可见范围。本文的主要贡献如下：",
    )
    rewrite(
        document,
        "开展多数据集实验和受控消融分析",
        "开展多数据集实验和受控消融分析。 本文在五个数据集上进行三随机种子主实验，并在两个代表性数据集上分别检验"
        "拓扑初始化、早期结构监督、邻居选择、关系类型、层选择及三种语义视图。参数匹配的晚期训练融合基线控制训练容量，"
        "结构查询基线则在邻居可见范围不变时替换查询信号，从而分别提供关于邻居上下文和查询来源的受控证据。",
    )
    rewrite(
        document,
        "已有文本增强实体对齐研究表明",
        "已有文本增强实体对齐研究表明，名称、描述和其他模态能够补充结构证据。RREA 将文本增强结果与纯结构设置分开报告[5]；"
        "MCLEA 通过模态内对比和模态间对齐联合学习多模态表示[12]；多模态知识图谱 Transformer 研究还指出，异质信息的直接组合"
        "可能形成未对齐的表示空间[13]。这些工作支持同时使用结构与语义，但其架构和训练目标往往同时变化，不能据此单独判断交互位置的作用。"
        "本文因此设置容量匹配和查询信号控制，在相同数据划分与检索协议下检验最终表示形成前引入结构邻居上下文的效果。",
    )
    rewrite(
        document,
        "图 1 关系感知结构上下文",
        "图 1 关系感知结构上下文与多尺度语义融合方法。（a）局部拓扑统计与可学习实体残差形成初始结构状态，三层关系消息传播和节点级层选择得到结构表示；"
        "三个语义视图形成语义表示；全部一跳出邻居经动态补齐和 1.5-entmax 稀疏选择形成结构上下文。（b）联合与结构分支均采用双向 InfoNCE，"
        "结构权重从 0.1 线性衰减至 0；验证集独立选择 [[\alpha]]、[[L]] 和 [[k_{\mathrm{CSLS}}]]。结构编码、语义编码、上下文融合、"
        "控制基线、训练目标和检索分别对应式（1）—（7）、式（8）—（13）、式（14）—（21）、式（22）、式（23）—（26）和式（27）。",
    )
    rewrite(
        document,
        "本文先在每张图内部计算",
        "本文先在每张图内部计算关系编号不变的局部拓扑特征。对实体 [[i]]，令 [[d_i^{\mathrm{in}}]]、[[d_i^{\mathrm{out}}]] 和 [[d_i]] 分别表示入度、出度和总度，"
        "[[\rho_i^{\mathrm{in}}]] 与 [[\rho_i^{\mathrm{out}}]] 表示入边和出边的关系类型数。[[\bar{\delta}_i^{\mathrm{in}}]] 与 [[\bar{\delta}_i^{\mathrm{out}}]] "
        "分别表示传入源实体和传出目标实体的 [[\log(1+d_j)]] 均值；无相应邻居时该均值记为 0。方向平衡定义为 "
        "[[b_i=(d_i^{\mathrm{out}}-d_i^{\mathrm{in}})/\max(1,d_i)]]。原始特征写为：",
    )
    rewrite(
        document,
        "加权后的三个视图以拼接形式",
        "加权后的三个视图以拼接形式进入输出 MLP；输出 MLP 的结果再与三个加权视图之和的残差投影相加，最后进行 L2 归一化，得到语义表示 "
        "[[\mathbf z_i^{\mathrm{sem}}]]。式（13）的视图门控网络学习实体级权重，输出 MLP 学习加权视图之间的非线性交互。"
        "当前主配置由 token、phrase 和 global 三个视图共同构成残差，不设置仅由 global 视图提供的专用残差。",
    )
    replace_formula(
        document,
        21,
        r"\widehat{\mathbf z}_i^{\mathrm{joint}}=\operatorname{Norm}\!\left(\mathbf g_i^{\mathrm{joint}}\odot\mathbf c_i^{\mathrm{str}}+(1-\mathbf g_i^{\mathrm{joint}})\odot\mathbf s_i\right)",
    )
    rewrite(
        document,
        "式（21）中，联合门",
        "式（21）先得到门控联合向量 [[\widehat{\mathbf z}_i^{\mathrm{joint}}]]：联合门接近 1 的维度更多使用结构上下文，接近 0 的维度更多使用语义表示。"
        "与代码一致，主配置随后计算等权基础向量 [[\mathbf b_i=\operatorname{Norm}((\mathbf c_i^{\mathrm{str}}+\mathbf s_i)/2)]]，并输出 "
        "[[\mathbf z_i^{\mathrm{joint}}=\operatorname{Norm}(0.9\widehat{\mathbf z}_i^{\mathrm{joint}}+0.1\mathbf b_i)]]。该固定残差用于限制门控输出的偏移，"
        "不引入额外可训练参数。",
    )
    rewrite(
        document,
        "为区分性能变化来自结构邻居可见性",
        "为控制训练阶段融合容量并比较邻居查询信号，本文设置两种基线。结构查询基线保留相同的全部一跳邻居、动态补齐、1.5-entmax 和门控过程，"
        "但以实体自身结构表示而非语义表示计算邻居权重。参数匹配晚期训练融合不读取邻居结构上下文，而在结构分支和语义分支独立编码后形成训练表示：",
    )
    rewrite(
        document,
        "式（22）先分别归一化",
        "式（22）先分别归一化结构表示与语义表示，再通过拼接和非线性映射得到用于 InfoNCE 训练的联合向量。该基线与主模型保持相同的可训练融合参数量。"
        "最终测试仍按第 3.8 节对分支输出进行统一的验证集加权；因此，主模型与该基线的差异同时包含训练交互路径和最终结构邻居可见范围的变化，"
        "不能把结果单独解释为融合时机的因果效应。",
    )
    rewrite(
        document,
        "因此，结构分支在第一个 epoch",
        "因此，结构分支在第一个 epoch 的权重为 0.1，并在最后一个 epoch 衰减为 0。第 5.3 节的消融显示，去除该监督后 DBP15K 的 Hits@1 提高 0.62 个百分点，"
        "而 OpenEA 下降 0.58 个百分点，方向并不一致。因此，本文将其作为当前训练协议的一部分，而不把它描述为已获得稳定验证的独立贡献。",
    )
    rewrite(
        document,
        "训练阶段使用式（25）",
        "训练阶段使用式（25）的联合与结构监督优化编码器。为使所有变体采用同一检索规则，最终测试不直接使用训练表示 [[\mathbf z_i^{\mathrm{joint}}]] 或 "
        "[[\mathbf z_i^{\mathrm{late}}]]，而是组合语义输出与结构输出：[[\mathbf x_i=\operatorname{Norm}(\alpha\mathbf z_{i,\mathrm{enh}}^{\mathrm{sem}}+(1-\alpha)"
        "\mathbf z_{i,\mathrm{enh}}^{\mathrm{str}})]]。在主模型中，结构输出包含加权邻居上下文；在晚期基线中，它退化为实体自身结构表示。随后对 [[\mathbf x_i]] 使用 CSLS 校正[17]：",
    )
    rewrite(
        document,
        "实验围绕三个研究问题展开",
        "实验围绕三个研究问题展开。RQ1：新主配置在跨语言、跨知识库来源与事件知识图谱上能否获得稳定结果？"
        "RQ2：拓扑初始化、早期结构监督、全部邻居稀疏选择、关系类型、节点级层选择以及三种语义视图分别产生何种贡献？"
        "RQ3：在训练阶段融合参数量相同的条件下，最终检索表示包含一跳结构邻居时与仅使用独立分支时有何差异，且邻居查询是否必须使用语义信息？",
    )
    rewrite(
        document,
        "表 3 中 DBP15K",
        "表 3 中 DBP15K ZH–EN 的 MTransE、JAPE、BootEA 和 RDGCN 数值取自 RDGCN 的汇总[4]，RREA-text 取自原论文[5]；"
        "OpenEA 数值取自官方基准的五折平均[14]。本文 OpenEA 结果只来自官方第一折，因此该部分仅用于定位结果范围，不能视为严格同协议排名。",
    )
    rewrite(
        document,
        "将共享拓扑初始化替换为",
        "将共享拓扑初始化替换为独立可学习实体向量后，Hits@1 在 DBP15K 和 OpenEA 上分别下降 3.50 与 7.45 个百分点；去除关系类型分别下降 3.03 与 3.10 个百分点，"
        "这两项在两个数据集上方向一致。去除早期结构监督后，DBP15K 提高 0.62 个百分点而 OpenEA 下降 0.58 个百分点，说明其作用依赖数据集。"
        "固定 8 个邻居并恢复 softmax 分别下降 0.62 与 1.02 个百分点，表明全部一跳邻居与稀疏选择具有较小但一致的联合收益。"
        "去除层选择器在 DBP15K 上下降 0.10 个百分点，在 OpenEA 上下降 1.71 个百分点。语义视图方面，去除 token 几乎不改变结果；"
        "去除 phrase 分别下降 0.29 与 4.87 个百分点；去除 global 后 DBP15K 提高 1.30 个百分点而 OpenEA 下降 0.84 个百分点，表明多尺度视图的作用取决于数据集。",
    )
    rewrite(
        document,
        "表 6 固定结构与语义编码器",
        "表 6 固定结构与语义编码器，并使三种训练融合路径具有相同的可训练参数量。完整模型与结构查询基线均读取全部一跳出邻居，二者只改变查询来源；"
        "参数匹配晚期训练融合不读取邻居上下文。该设计控制训练容量并单独检验查询来源，但晚期基线同时改变训练交互路径和最终邻居可见范围。",
    )
    rewrite(
        document,
        "相对完整模型，结构查询邻居",
        "相对完整模型，结构查询邻居的 Hits@1 在 DBP15K 和 OpenEA 上分别下降 0.44 与 0.22 个百分点，差值较小，说明本轮结果不能把语义查询本身视为主要增益来源。"
        "参数匹配晚期训练融合分别下降 2.18 与 9.44 个百分点，表明在当前统一检索规则下，保留一跳结构邻居上下文与更高结果相关。"
        "由于晚期基线同时移除了邻居上下文并改变训练交互路径，该差值是两项变化的联合效应，不能单独证明提前融合优于晚期融合。",
    )
    rewrite(
        document,
        "RQ1 的结果表明",
        "RQ1 的结果表明，模型能够在五个数据集上稳定训练，但绝对性能随语言、知识库来源和事件属性异构性而变化。"
        "RQ2 的消融显示，拓扑初始化和关系类型在两个代表性数据集上均提供方向一致的增益；早期结构监督、层选择和语义视图则具有不同程度的数据集依赖性。"
        "RQ3 的控制实验表明，语义查询与结构查询差异很小；包含一跳结构邻居的主模型优于不读取邻居的参数匹配基线，但现有设计不能把邻居可见性与训练交互位置完全分离。",
    )
    rewrite(
        document,
        "内部有效性方面",
        "内部有效性方面，主实验和消融均使用相同数据划分、验证规则和三个随机种子，但 [[n=3]] 仍不足以支持高功效显著性检验。"
        "各数据集和变体独立选择 [[\alpha]]、[[L]] 与 [[k_{\mathrm{CSLS}}]]，降低了沿用主模型超参数造成的偏置，也引入了有限验证集上的网格方差。"
        "此外，统一检索阶段使用分支加权而非直接使用训练联合向量；参数匹配晚期基线又同时改变邻居可见范围和训练交互路径，因此第 5.5 节只能支持联合效应，不能识别单一融合时机的因果贡献。"
        "外部有效性方面，组件消融只覆盖 DBP15K ZH–EN 和 OpenEA EN–FR，OpenEA 主结果只运行一个官方划分，尚未验证十万实体以上、开放世界和存在无对应实体的场景。",
    )
    rewrite(
        document,
        "本文围绕关系感知结构上下文",
        "本文围绕关系感知结构上下文与多尺度语义表示的组合方式开展建模和受控实验。结果表明，拓扑初始化和关系类型在两个代表性数据集上产生方向一致的增益，"
        "语义查询与结构查询的差异较小；在容量匹配的训练条件和统一分支检索规则下，包含一跳结构邻居上下文的模型优于不读取邻居的基线。"
        "由于该基线同时改变邻居可见范围和训练交互路径，本文不把差值单独归因于提前或晚期融合。该结论限定了当前实验能够支持的主张。",
    )

    set_markup(find_paragraph(document, "3.6.4 融合位置的控制基线"), "3.6.4 训练融合与邻居可见性的控制基线")
    set_markup(find_paragraph(document, "5.5 融合位置、邻居可见范围与容量控制"), "5.5 邻居上下文、查询信号与容量控制")


def set_cell_markup(cell, markup: str) -> None:
    set_markup(cell.paragraphs[0], markup)


def fix_result_tables(document: Document) -> None:
    comparison = document.tables[2]
    current_dbp = comparison.rows[6].cells
    current_dbp[2].text = "0.7198"
    current_dbp[3].text = "0.8470"
    current_dbp[4].text = "0.7661"

    ablation = document.tables[3]
    set_cell_markup(ablation.rows[1].cells[1], r"[[0.7198\pm0.0007]] / [[0.7661\pm0.0008]]")
    set_cell_markup(ablation.rows[1].cells[2], r"[[0.6324\pm0.0065]] / [[0.6978\pm0.0043]]")

    template_row = deepcopy(ablation.rows[2]._tr)
    ablation.rows[1]._tr.addnext(template_row)
    supervision_row = ablation.rows[2]
    set_cell_markup(supervision_row.cells[0], "去除早期结构监督")
    set_cell_markup(supervision_row.cells[1], r"[[0.7259\pm0.0037]] / [[0.7706\pm0.0034]]")
    set_cell_markup(supervision_row.cells[2], r"[[0.6266\pm0.0063]] / [[0.6921\pm0.0052]]")

    design = document.tables[5]
    design.rows[0].cells[3].text = "训练融合路径 / 可训练参数"
    design.rows[1].cells[3].text = "邻居聚合 + 联合门控 / 197,377"
    design.rows[2].cells[3].text = "邻居聚合 + 联合门控 / 197,377"
    design.rows[3].cells[3].text = "独立分支后形成训练向量 / 197,377"

    set_markup(
        find_paragraph(document, "表 6 参数匹配融合控制的设计差异"),
        "表 6 参数匹配训练融合与邻居可见性控制的设计差异",
    )
    set_markup(
        find_paragraph(document, "表 7 参数匹配融合控制的测试结果"),
        "表 7 邻居可见性与查询信号控制的测试结果（均值[[\pm]]样本标准差，[[n=3]]）",
    )


def mapped_citation_content(content: str) -> str:
    return re.sub(r"\d+", lambda match: str(OLD_TO_NEW[int(match.group())]), content)


def replace_text_range(nodes, start: int, end: int, replacement: str) -> None:
    offsets = []
    cursor = 0
    for node in nodes:
        text = node.text or ""
        offsets.append((cursor, cursor + len(text)))
        cursor += len(text)

    first_index = next(index for index, (_, stop) in enumerate(offsets) if stop > start)
    last_index = next(index for index, (begin, _) in enumerate(offsets) if begin < end <= offsets[index][1])
    first_begin, _ = offsets[first_index]
    last_begin, _ = offsets[last_index]
    first_text = nodes[first_index].text or ""
    last_text = nodes[last_index].text or ""
    prefix = first_text[: start - first_begin]
    suffix = last_text[end - last_begin :]

    nodes[first_index].text = prefix + replacement + (suffix if first_index == last_index else "")
    if first_index != last_index:
        for index in range(first_index + 1, last_index):
            nodes[index].text = ""
        nodes[last_index].text = suffix


def renumber_citations_in_paragraph(paragraph) -> None:
    if paragraph.style is not None and paragraph.style.name == "Reference":
        return
    nodes = paragraph._p.xpath(".//w:t")
    full_text = "".join(node.text or "" for node in nodes)
    matches = list(CITATION_RE.finditer(full_text))
    for match in reversed(matches):
        replacement = "[" + mapped_citation_content(match.group(1)) + "]"
        replace_text_range(nodes, match.start(), match.end(), replacement)


def renumber_body_citations(document: Document) -> None:
    paragraphs = list(document.paragraphs)
    paragraphs.extend(
        paragraph
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    for paragraph in paragraphs:
        renumber_citations_in_paragraph(paragraph)


def replace_reference_number(reference_xml, new_number: int) -> None:
    text_nodes = reference_xml.xpath(".//w:t")
    full_text = "".join(node.text or "" for node in text_nodes)
    match = re.match(r"\[(\d+)\]", full_text)
    if match is None:
        raise RuntimeError(f"Reference entry has no leading number: {full_text[:80]}")
    replace_text_range(text_nodes, match.start(), match.end(), f"[{new_number}]")


def reorder_reference_list(document: Document) -> None:
    references = [p for p in document.paragraphs if p.style.name == "Reference"]
    if len(references) != len(REFERENCE_ORDER):
        raise RuntimeError(f"Expected {len(REFERENCE_ORDER)} references, found {len(references)}")

    by_old_number = {}
    for paragraph in references:
        match = re.match(r"\[(\d+)\]", combined_text(paragraph))
        if match is None:
            raise RuntimeError(f"Could not identify reference number: {combined_text(paragraph)[:80]}")
        by_old_number[int(match.group(1))] = deepcopy(paragraph._p)

    anchor = references[0]._p
    for new_number, old_number in enumerate(REFERENCE_ORDER, start=1):
        reference_xml = by_old_number[old_number]
        replace_reference_number(reference_xml, new_number)
        anchor.addprevious(reference_xml)

    for paragraph in references:
        paragraph._p.getparent().remove(paragraph._p)


def convert_detached_overbars(document: Document) -> int:
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    converted = 0
    for group in list(document._element.xpath(".//m:groupChr")):
        char = group.find(f"./{{{math_ns}}}groupChrPr/{{{math_ns}}}chr")
        if char is None or char.get(f"{{{math_ns}}}val") != "¯":
            continue
        expression = group.find(f"./{{{math_ns}}}e")
        if expression is None:
            raise RuntimeError("Detached overbar has no expression")
        bar = group.makeelement(qn("m:bar"))
        bar_pr = bar.makeelement(qn("m:barPr"))
        position = bar.makeelement(qn("m:pos"), {qn("m:val"): "top"})
        bar_pr.append(position)
        bar.append(bar_pr)
        bar.append(deepcopy(expression))
        group.getparent().replace(group, bar)
        converted += 1
    return converted


def citation_first_appearance(document: Document) -> list[int]:
    seen = []
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Reference":
            continue
        for match in CITATION_RE.finditer(combined_text(paragraph)):
            for number in re.findall(r"\d+", match.group(1)):
                value = int(number)
                if value not in seen:
                    seen.append(value)
    # Table citations occur after the narrative has introduced their sources,
    # but include them in the completeness check.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for match in CITATION_RE.finditer(cell.text):
                    for number in re.findall(r"\d+", match.group(1)):
                        value = int(number)
                        if value not in seen:
                            seen.append(value)
    return seen


def validate(document: Document) -> None:
    first_order = citation_first_appearance(document)
    expected = list(range(1, len(REFERENCE_ORDER) + 1))
    if first_order != expected:
        raise RuntimeError(f"Citation first-appearance order is {first_order}, expected {expected}")

    references = [combined_text(p) for p in document.paragraphs if p.style.name == "Reference"]
    numbers = [int(re.match(r"\[(\d+)\]", entry).group(1)) for entry in references]
    if numbers != expected:
        raise RuntimeError(f"Reference list order is {numbers}")

    required_text = [
        "去除早期结构监督",
        "不能把结果单独解释为融合时机的因果效应",
        "统一检索阶段使用分支加权",
        "0.7198",
    ]
    body = "\n".join(combined_text(p) for p in document.paragraphs)
    body += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    missing = [text for text in required_text if text not in body]
    if missing:
        raise RuntimeError(f"Missing required logical revisions: {missing}")

    if "MTransE[2]" not in document.tables[2].cell(1, 1).text:
        raise RuntimeError("Comparison-table citations were not renumbered")
    if document.tables[2].cell(6, 2).text != "0.7198":
        raise RuntimeError("DBP15K current-model comparison value is inconsistent")
    if "0.7198" not in combined_text(document.tables[3].cell(1, 1).paragraphs[0]):
        raise RuntimeError("Ablation full-model row is inconsistent")


def main() -> None:
    document = Document(SOURCE)
    rewrite_logic(document)
    fix_result_tables(document)
    renumber_body_citations(document)
    reorder_reference_list(document)
    converted = convert_detached_overbars(document)
    validate(document)
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    validate(reopened)
    print(f"WROTE {OUTPUT}")
    print(f"citations=20 references=20 converted_mean_bars={converted}")


if __name__ == "__main__":
    main()
