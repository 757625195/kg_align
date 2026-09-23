# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import html.entities
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
MATH_DEPS = ROOT / "tmp" / "docx_math_deps"
sys.path.insert(0, str(MATH_DEPS))

import mathml2omml  # noqa: E402
from latex2mathml.converter import convert as latex_to_mathml  # noqa: E402


MAIN_RESULTS = {
    "DBP15K zh-en": (0.7197777778, 0.0007397575, 0.8469523810, 0.0007438333, 0.7661196391, 0.0007843247),
    "DBP15K ja-en": (0.7844126984, 0.0062250560, 0.8864126984, 0.0037499307, 0.8213347197, 0.0052459260),
    "DBP15K fr-en": (0.9183174603, 0.0004294524, 0.9655238095, 0.0010605265, 0.9358180761, 0.0002387150),
    "OpenEA en-fr": (0.6324444444, 0.0065224770, 0.8195555556, 0.0053810226, 0.6978041132, 0.0043051869),
    "EventEA en-en": (0.7551660395, 0.0014791022, 0.9106093616, 0.0014540262, 0.8156018058, 0.0012451485),
}


def normalize(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def latex_to_omml(expression: str):
    # Normal Python string literals interpret a few LaTeX prefixes as control
    # escapes. Repair only those recognized escapes before MathML conversion.
    expression = (
        expression.replace("\x07lpha", r"\alpha")
        .replace("\x08ar", r"\bar")
        .replace("\x08eta", r"\beta")
        .replace("\x08oldsymbol", r"\mathbf")
        .replace("\tau", r"\tau")
        .replace("\times", r"\times")
        .replace("\rho", r"\rho")
        .replace("\frac", r"\frac")
    )
    mathml = latex_to_mathml(expression.strip(), display="inline")
    omml = mathml2omml.convert(mathml, html.entities.name2codepoint)
    omml = re.sub(
        r"(<m:groupChr><m:groupChrPr>.*?)</m:groupChr>(<m:e>)",
        r"\1</m:groupChrPr>\2",
        omml,
    )
    omml = omml.replace("<m:oMath>", f"<m:oMath {nsdecls('m')}>", 1)
    return parse_xml(omml)


def clear_paragraph(paragraph: Paragraph) -> None:
    properties = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)


def clear_text_highlights(document: Document) -> None:
    for highlight in list(document.element.body.iter(qn("w:highlight"))):
        highlight.getparent().remove(highlight)


def add_mixed_text(paragraph: Paragraph, text: str) -> None:
    token_re = re.compile(r"\$([^$]+)\$")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        paragraph._p.append(latex_to_omml(match.group(1)))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def set_mixed_paragraph(paragraph: Paragraph, text: str) -> Paragraph:
    clear_paragraph(paragraph)
    add_mixed_text(paragraph, text)
    return paragraph


def find_exact(document: Document, text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if normalize(p.text) == normalize(text)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph {text!r}, found {len(matches)}")
    return matches[0]


def find_starts(document: Document, prefix: str) -> Paragraph:
    matches = [p for p in document.paragraphs if normalize(p.text).startswith(normalize(prefix))]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting {prefix!r}, found {len(matches)}")
    return matches[0]


def replace_starts(document: Document, prefix: str, text: str) -> Paragraph:
    return set_mixed_paragraph(find_starts(document, prefix), text)


def insert_before(document: Document, anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    paragraph = document.add_paragraph(style=style)
    add_mixed_text(paragraph, text)
    anchor._p.addprevious(paragraph._p)
    return paragraph


def insert_after(document: Document, anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    paragraph = document.add_paragraph(style=style)
    add_mixed_text(paragraph, text)
    anchor._p.addnext(paragraph._p)
    return paragraph


def display_math(expression: str, tag: str | None = None):
    math_paragraph = OxmlElement("m:oMathPara")
    properties = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "centerGroup")
    properties.append(justification)
    math_paragraph.append(properties)
    math_paragraph.append(latex_to_omml(expression))
    if tag is not None:
        math_paragraph.append(latex_to_omml(rf"\qquad\qquad ({tag})"))
    return math_paragraph


def insert_formula_before(document: Document, anchor: Paragraph, expression: str, tag: str | None = None) -> Paragraph:
    paragraph = document.add_paragraph(style="Formula")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph._p.append(display_math(expression, tag))
    anchor._p.addprevious(paragraph._p)
    return paragraph


def insert_formula_after(document: Document, anchor: Paragraph, expression: str, tag: str) -> Paragraph:
    paragraph = document.add_paragraph(style="Formula")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph._p.append(display_math(expression, tag))
    anchor._p.addnext(paragraph._p)
    return paragraph


def formula_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.xpath(".//*[local-name()='t']/text()"))


def find_formula(document: Document, tag: int) -> Paragraph:
    suffix = f"({tag})"
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Formula" and formula_text(paragraph).endswith(suffix)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one equation {suffix}, found {len(matches)}")
    return matches[0]


def replace_formula(document: Document, tag: int, expression: str) -> Paragraph:
    paragraph = find_formula(document, tag)
    clear_paragraph(paragraph)
    paragraph._p.append(display_math(expression, str(tag)))
    return paragraph


def renumber_csls(document: Document) -> None:
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Formula"
        and formula_text(paragraph).startswith("CSLS")
        and formula_text(paragraph).endswith("(25)")
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one CSLS equation (25), found {len(matches)}")
    paragraph = matches[0]
    for node in paragraph._p.xpath(".//*[local-name()='t']"):
        if node.text == "(25)":
            node.text = "(27)"
            return
    raise RuntimeError("Could not renumber the CSLS equation")


def set_cell(cell, text: str, header: bool = False, centered: bool = True) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(8.5)
            run.bold = header
            run.font.name = "Heiti SC" if header else "Songti SC"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), run.font.name)


def metric_text(mean: float, std: float) -> str:
    return f"{mean:.4f} ± {std:.4f}"


def update_tables(document: Document) -> None:
    main_table = document.tables[1]
    for row in main_table.rows:
        dataset = row.cells[0].text.strip()
        if dataset not in MAIN_RESULTS:
            continue
        h1, h1s, h10, h10s, mrr, mrrs = MAIN_RESULTS[dataset]
        set_cell(row.cells[1], metric_text(h1, h1s))
        set_cell(row.cells[2], metric_text(h10, h10s))
        set_cell(row.cells[3], metric_text(mrr, mrrs))

    comparison = document.tables[2]
    for row in comparison.rows:
        if row.cells[1].text.strip() != "本文模型":
            continue
        dataset = row.cells[0].text.strip()
        if dataset == "DBP15K ZH–EN" or not dataset:
            previous = comparison.rows[comparison.rows.index(row) - 1].cells[0].text.strip() if False else ""
        if row._tr is comparison.rows[6]._tr:
            values = MAIN_RESULTS["DBP15K zh-en"]
        else:
            values = MAIN_RESULTS["OpenEA en-fr"]
        set_cell(row.cells[2], f"{values[0]:.4f}")
        set_cell(row.cells[3], f"{values[2]:.4f}")
        set_cell(row.cells[4], f"{values[4]:.4f}")

    category = document.tables[3]
    headers = ["数据集类别", "数量", "Hits@1", "Hits@10", "MRR", "有效深度", "验证选择"]
    for index, text in enumerate(headers):
        set_cell(category.rows[0].cells[index], text, header=True)
    rows = [
        ["DBP15K", "3", "0.8075", "0.8996", "0.8411", "3", "各语言对独立选择 α 与 k"],
        ["OpenEA EN–FR", "1", "0.6324", "0.8196", "0.6978", "3", "α=0.5，k=3"],
        ["EventEA", "1", "0.7552", "0.9106", "0.8156", "3", "α=0.6，k=5"],
        ["全部数据集", "5", "0.7620", "0.8858", "0.8073", "—", "数据集等权平均"],
    ]
    for row, values in zip(category.rows[1:], rows):
        for index, text in enumerate(values):
            set_cell(row.cells[index], text, centered=index != 6)

    component = document.tables[4]
    component_labels = [
        "完整模型",
        "可学习实体初始化",
        "去除结构分支监督",
        "固定 8 邻居 + softmax",
        "去除关系类型",
        "去除层选择器",
        "去除 token 视图",
    ]
    for row, label in zip(component.rows[1:], component_labels):
        set_cell(row.cells[0], label, centered=False)
        set_cell(row.cells[1], "运行中")
        set_cell(row.cells[2], "运行中")

    branches = document.tables[5]
    for row in branches.rows[1:]:
        set_cell(row.cells[2], "运行中")
        set_cell(row.cells[3], "运行中")

    design = document.tables[6]
    design_rows = [
        ["完整模型", "全部一跳出邻居", "语义表示", "邻居聚合前 + 最终门控 / 197,377"],
        ["结构查询邻居", "全部一跳出邻居", "自身结构表示", "邻居聚合前 + 最终门控 / 197,377"],
        ["参数匹配晚期融合", "不可见", "无邻居查询", "独立编码后末端融合 / 197,377"],
    ]
    for row, values in zip(design.rows[1:], design_rows):
        for index, text in enumerate(values):
            set_cell(row.cells[index], text, centered=index != 3)

    fusion_results = document.tables[7]
    for row in fusion_results.rows[1:]:
        for cell in row.cells[1:]:
            set_cell(cell, "运行中")


def replace_images(document: Document, output: Path, figure1: Path, figure2: Path) -> None:
    if len(document.inline_shapes) < 2:
        raise RuntimeError("Expected two embedded figures")
    replacements = {}
    descriptions = [
        "拓扑初始化、关系感知传播、多尺度语义、全一跳邻居稀疏选择以及联合和结构监督的信息流",
        "五个数据集主模型 Hits@1、Hits@10 和 MRR 的三随机种子均值及样本标准差",
    ]
    for shape, figure, description in zip(list(document.inline_shapes)[:2], (figure1, figure2), descriptions):
        blips = shape._inline.xpath(".//*[local-name()='blip']")
        rel_id = blips[0].get(qn("r:embed"))
        relationship = document.part.rels[rel_id]
        replacements[f"word/{relationship.target_ref}"] = figure.read_bytes()
        props = shape._inline.xpath(".//*[local-name()='docPr']")
        if props:
            props[0].set("descr", description)

    document.save(output)
    temp = output.with_suffix(".tmp.docx")
    with zipfile.ZipFile(output, "r") as source_zip:
        with zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
            for item in source_zip.infolist():
                target_zip.writestr(item, replacements.get(item.filename, source_zip.read(item.filename)))
    os.replace(temp, output)


def rewrite_document(document: Document) -> None:
    replace_starts(
        document,
        "跨语言知识图谱实体对齐需要同时处理",
        "跨语言知识图谱实体对齐需要联合利用关系结构与实体语义，但独立实体初始化、固定邻居截断和稠密注意力会削弱结构分支的可迁移性。本文提出一种关系感知结构上下文与多尺度语义融合方法。结构分支以八维局部拓扑统计初始化实体状态，并叠加 0.1 倍可学习实体残差；关系感知消息传播和节点级层选择进一步融合不同结构深度。语义分支从共享输入构造 token、phrase 和 global 三种视图。融合阶段保留全部一跳出邻居，采用批次内动态补齐和温度为 0.25 的 1.5-entmax 稀疏选择，不增加邻居选择参数。模型以联合表示上的双向 InfoNCE 为主目标，并加入初始权重为 0.1、随训练线性衰减至 0 的结构分支 InfoNCE。五个数据集的三随机种子测试结果达到平均 Hits@1 0.7620、Hits@10 0.8858 和 MRR 0.8073，其中 DBP15K FR–EN 的 Hits@1 为 0.9183，OpenEA EN–FR 为 0.6324。结果表明该统一配置能够在跨语言、跨来源和事件知识图谱上保持稳定检索性能。",
    )
    replace_starts(
        document,
        "关键词：",
        "关键词：知识图谱实体对齐；跨语言知识图谱；拓扑初始化；关系感知图神经网络；稀疏邻居选择；多尺度语义编码",
    )

    replace_starts(
        document,
        "实现并验证一种轻量级关系感知结构编码器",
        "实现一种拓扑初始化的轻量关系感知结构编码器。八维局部拓扑统计为左右图提供关系编号不变的初始结构坐标，较小的可学习实体残差保留个体差异；共享关系投影和节点级层选择控制参数量并融合不同传播深度。",
    )
    replace_starts(
        document,
        "构建语义编码与结构编码早期交互模块",
        "构建基于全部一跳出邻居的结构—语义交互模块。变长邻域仅在当前批次内动态补齐，1.5-entmax 在不增加可训练参数的条件下产生稀疏邻居权重，并以参数匹配晚期融合和结构查询邻居作为控制基线。",
    )
    replace_starts(
        document,
        "在五个数据集上完成三随机种子主实验",
        "在五个数据集上完成三随机种子主实验，并为每个数据集独立使用验证集选择融合权重、有效传播深度和 CSLS 邻域。组件消融仅在 DBP15K ZH–EN 与 OpenEA EN–FR 上按同一协议运行。",
    )

    replace_starts(
        document,
        "模型由结构编码、语义编码、结构邻居上下文和对比学习四部分组成",
        "模型由拓扑初始化的结构编码、多尺度语义编码、全一跳邻居上下文和对比学习四部分组成。结构编码器在两张知识图谱上执行关系感知消息传递，并融合不同传播深度；语义编码器从名称、关系和属性序列构造 token、phrase 和 global 三个视图。对当前批次实体，模型保留全部一跳出邻居并按批次最大邻域宽度动态补齐，再以 1.5-entmax 形成稀疏结构上下文。训练同时约束联合表示与结构表示，验证阶段按数据集选择融合权重、有效传播深度和 CSLS 邻域。",
    )
    replace_starts(
        document,
        "图 1 关系感知结构上下文",
        "图 1 拓扑初始化的关系感知结构上下文与多尺度语义融合方法。（a）八维图内拓扑统计与可学习实体残差形成初始结构状态，三层关系消息传播和节点级层选择得到结构表示；三个语义视图形成语义表示；全部一跳出邻居经动态补齐和 1.5-entmax 稀疏选择形成结构上下文。（b）联合与结构分支均采用双向 InfoNCE，结构权重从 0.1 线性衰减至 0；验证集独立选择 $\alpha$、$L$ 和 $k_{\mathrm{CSLS}}$。",
    )

    structure_heading = find_exact(document, "3.3 关系感知轻量结构编码器")
    structure_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph._p is structure_heading._p
    )
    first_structure = document.paragraphs[structure_index + 1]
    intro = insert_before(
        document,
        first_structure,
        "为避免左右知识图谱完全从相互独立的实体编号向量开始，本文先在每张图内部计算关系编号不变的局部拓扑特征。对实体 $i$，令 $d_i^{\mathrm{in}}$、$d_i^{\mathrm{out}}$ 和 $d_i$ 分别表示入度、出度和总度，$\rho_i^{\mathrm{in}}$ 与 $\rho_i^{\mathrm{out}}$ 表示入边和出边的关系类型数，$\bar{\delta}_i^{\mathrm{in}}$ 与 $\bar{\delta}_i^{\mathrm{out}}$ 表示相邻实体对数总度的均值，$b_i$ 表示出入度方向平衡。原始特征写为：",
    )
    topo_formula = insert_formula_before(
        document,
        first_structure,
        r"\mathbf{f}_i=[\log(1+d_i^{\mathrm{in}});\log(1+d_i^{\mathrm{out}});\log(1+d_i);\log(1+\rho_i^{\mathrm{in}});\log(1+\rho_i^{\mathrm{out}});\bar{\delta}_i^{\mathrm{in}};\bar{\delta}_i^{\mathrm{out}};b_i]",
    )
    topo_explanation = insert_before(
        document,
        first_structure,
        "八个维度分别在左、右知识图谱内部进行标准化，得到 $\widetilde{\mathbf f}_i$。初始结构状态由共享拓扑投影与较小的实体残差共同给出：",
    )
    insert_formula_before(
        document,
        first_structure,
        r"\mathbf{h}_i^{(0)}=\mathbf{W}_{\mathrm{top}}\widetilde{\mathbf{f}}_i+0.1\mathbf{e}_i",
    )
    insert_before(
        document,
        first_structure,
        "其中，$\mathbf W_{\mathrm{top}}\in\mathbb R^{128\times 8}$ 为两侧知识图谱共享的线性投影，$\mathbf e_i\in\mathbb R^{128}$ 为可学习实体向量。拓扑项为具有相近局部角色的实体提供可比较起点，0.1 倍残差则避免仅凭八项统计把不同实体压缩为同一表示。",
    )
    set_mixed_paragraph(
        first_structure,
        "在上述初始状态上，结构编码器对整张知识图谱执行有向关系感知消息传递。对于三元组 $(j,r,i)$，头实体 $j$ 沿该边向尾实体 $i$ 发送消息；当前数据处理不自动添加反向边。设 $\mathbf h_i^{(l)}$ 为实体 $i$ 在第 $l$ 层的状态，$\mathbf e_r$ 为关系 $r$ 的嵌入，$\mathcal N_{\mathrm{in}}(i)$ 为所有指向 $i$ 的传入边。第 $l$ 层首先计算：",
    )

    set_mixed_paragraph(find_exact(document, "3.5 固定预算邻居提取"), "3.5 全部一跳出邻居与批次内动态补齐")
    replace_starts(
        document,
        "整图结构编码每次前向计算一次",
        "整图结构编码在每次前向计算中执行一次，随后按当前批次实体编号从结构表示矩阵索引邻居。邻接表沿三元组的头实体到尾实体方向构造；主配置对实体 $i$ 保留集合 $\mathcal N_{\mathrm{out}}(i)$ 中的全部一跳出邻居，不预先截断为固定数量。由于不同实体的邻居数不同，系统仅将当前批次补齐到该批次的最大邻域宽度；填充位置使用实体自身编号占位，并由掩码排除。无出邻居实体的邻域宽度至少保留 1 个占位位置，但其掩码全部为 0。该整理过程不增加可训练参数，也不会复制有效邻居。",
    )

    replace_starts(
        document,
        "本节的目标是从固定预算邻居中提取",
        "本节的目标是从全部一跳出邻居中提取与当前实体相关的结构证据，并将其与实体语义表示合成为联合表示。计算分为三个阶段：语义查询首先评估各结构邻居的相关性，1.5-entmax 将低相关候选赋予精确零权重；随后，加权邻居摘要与实体自身结构状态形成结构上下文；最后，结构上下文与语义表示通过逐维门控得到联合表示。",
    )
    replace_starts(
        document,
        "对于实体 ，融合模块接收",
        "对于实体 $i$，融合模块接收语义表示 $\mathbf s_i=\mathbf z_i^{\mathrm{sem}}$、自身结构表示 $\mathbf t_i=\mathbf z_i^{\mathrm{str}}$、全部一跳出邻居结构表示集合及有效掩码。语义表示经查询投影得到 $\mathbf q_i$；第 $j$ 个结构邻居分别经键和值投影得到 $\mathbf k_j$ 与 $\mathbf v_j$。查询与键的缩放点积给出初始相关性分数：",
    )
    replace_starts(
        document,
        "式（14）衡量实体语义",
        "式（14）衡量实体语义与第 $j$ 个结构邻居的匹配程度。点积越大，表示该邻居与当前实体的语义越一致；除以 $\sqrt d$ 用于控制高维点积的数值范围，避免稀疏权重在训练早期过度集中。",
    )
    replace_formula(
        document,
        17,
        r"\boldsymbol{\alpha}_i=\operatorname{entmax}_{1.5}\left(\frac{\mathbf a_i+\log(\mathbf p_i+\epsilon)+\mathbf m_i}{T}\right),\quad T=0.25.",
    )
    replace_starts(
        document,
        "式（17）中的掩码",
        "式（17）将全部有效邻居的分数向量记为 $\mathbf a_i$，逐维门值记为 $\mathbf p_i$，掩码偏置 $\mathbf m_i$ 将填充位置排除。与 softmax 对所有有效位置分配正权重不同，1.5-entmax 可以把低分邻居的权重压到精确的 0；温度 $T=0.25$ 控制分布集中程度。有效权重仍归一化为和 1，无有效邻居时输出全零向量。",
    )
    replace_starts(
        document,
        "根据式（17）得到的权重",
        "根据式（17）得到的稀疏权重，对邻居值向量加权求和，形成实体 $i$ 的结构邻居摘要 $\bar{\mathbf c}_i$。由于主配置不截断一跳出邻居，该摘要可以利用完整局部邻域，同时由稀疏权重抑制与当前实体不相关的证据。",
    )
    replace_starts(
        document,
        "为区分性能变化来自结构邻居可见性",
        "为区分性能变化来自结构邻居可见性、邻居查询信号还是融合位置，本文设置两种参数匹配基线。结构查询基线保留相同的全部一跳邻居、动态补齐、1.5-entmax 和门控过程，但以实体自身结构表示而非语义表示计算邻居权重。参数匹配晚期融合不读取邻居结构上下文，而在结构分支和语义分支独立编码后进行融合：",
    )

    set_mixed_paragraph(find_exact(document, "3.7 单阶段双向 InfoNCE 对齐训练"), "3.7 联合与结构分支的双向 InfoNCE 训练")
    replace_starts(
        document,
        "模型在联合表示上直接进行实体对齐训练",
        "模型采用单阶段训练，不设置独立预热阶段。对于包含 $B$ 个已对齐实体对的批次，首先计算左图第 $i$ 个实体与右图第 $j$ 个实体之间的温度缩放余弦相似度：",
    )
    replace_starts(
        document,
        "其中， 为温度参数",
        "其中，$\tau=0.07$ 为温度参数，正确实体对位于相似度矩阵对角线。模型分别计算左图到右图和右图到左图的交叉熵并取平均，得到联合表示损失 $\mathcal L_{\mathrm{joint}}$：",
    )
    objective_intro = replace_starts(
        document,
        "式（24）构成全部主实验的训练目标",
        "式（24）直接约束联合表示。为使拓扑初始化的结构分支在训练早期建立跨图坐标，本文还对 $\mathbf z^{\mathrm{str}}$ 使用同形式的双向 InfoNCE，记为 $\mathcal L_{\mathrm{str}}$。总目标为：",
    )
    objective = insert_formula_after(
        document,
        objective_intro,
        r"\mathcal L(p)=\mathcal L_{\mathrm{joint}}+\lambda_{\mathrm{str}}(p)\mathcal L_{\mathrm{str}}.",
        "25",
    )
    schedule_intro = insert_after(
        document,
        objective,
        "结构监督只用于训练前期，并随训练进度线性减弱。设当前 epoch 为 $e$，总 epoch 数为 $E$，归一化进度与权重定义为：",
    )
    schedule = insert_formula_after(
        document,
        schedule_intro,
        r"p=\frac{e-1}{E-1},\qquad \lambda_{\mathrm{str}}(p)=0.1(1-p).",
        "26",
    )
    insert_after(
        document,
        schedule,
        "因此，结构分支在第一个 epoch 的权重为 0.1，并在最后一个 epoch 衰减为 0。该辅助项不新增编码模块；其作用是先稳定结构空间，再让后期优化集中于最终联合表示。",
    )
    replace_starts(
        document,
        "模型采用 AdamW 优化器",
        "模型采用 AdamW 优化器[16]。DBP15K 的学习率为 $5\times10^{-4}$，最多训练 36 个 epoch；OpenEA 与 EventEA 的学习率为 $3\times10^{-4}$，最多训练 50 个 epoch。所有数据集每 5 个 epoch 在验证集上评估 MRR 并保存最佳检查点，连续 4 次验证未提升时提前停止。主实验不使用结构预热、伪标签、hard-negative 间隔损失或 ranking loss。",
    )
    renumber_csls(document)
    replace_starts(
        document,
        "训练阶段直接使用联合表示完成双向 InfoNCE 对齐",
        "训练阶段使用式（25）的联合与结构监督优化编码器。最终检索在同一检查点上进一步组合语义表示与融合模块输出的结构上下文：语义权重为 $\alpha$，结构权重为 $1-\alpha$，融合向量经 L2 归一化后使用 CSLS 校正[17]：",
    )

    replace_starts(
        document,
        "实验围绕三个研究问题展开",
        "实验围绕三个研究问题展开。RQ1：新主配置在跨语言、跨知识库来源与事件知识图谱上能否获得稳定结果？RQ2：拓扑初始化、早期结构监督、全部邻居稀疏选择、关系类型、节点级层选择以及三种语义视图分别产生何种贡献？RQ3：在融合参数量相同的条件下，早期读取结构邻居与两个分支独立编码后的晚期融合有何差异，且邻居查询是否必须使用语义信息？",
    )
    replace_starts(
        document,
        "公共设置为 batch size 512",
        "公共设置为 batch size 512，结构与联合表示维度 128，关系感知 GNN 最大深度 3，语义编码器包含 2 层和 4 个注意力头，dropout 为 0.1，InfoNCE 温度为 0.07。结构状态由八维拓扑特征投影与 0.1 倍实体残差初始化；结构 InfoNCE 权重从 0.1 线性衰减至 0。主配置保留全部一跳出邻居，按批次最大邻居数动态补齐，并采用 $T=0.25$ 的 1.5-entmax。验证网格为 $\alpha\in\{0.2,\ldots,0.8\}$、$L\in\{1,2,3\}$、$k_{\mathrm{CSLS}}\in\{3,5,7,10,15,20\}$，CSLS 强度固定为 1.0。每个数据集和消融变体均依据三个随机种子的平均验证 MRR 独立选择唯一配置。",
    )

    replace_starts(
        document,
        "表 2 汇总完整模型",
        "表 2 汇总新主配置在五个数据集上的最终检索结果。每个数据集使用随机种子 42、43 和 44 独立训练，并报告 Hits@1、Hits@10 和 MRR 的均值与样本标准差。DBP15K 从训练对中固定留出 10% 作为验证集；OpenEA 与 EventEA 使用官方验证集。",
    )
    replace_starts(
        document,
        "DBP15K FR–EN 的 Hits@1",
        "DBP15K FR–EN 的 Hits@1 和 MRR 最高，分别为 0.9183 和 0.9358；EventEA 的 Hits@1 为 0.7552；OpenEA EN–FR 的 Hits@1 最低，为 0.6324。五个数据集的 Hits@1 样本标准差均不超过 0.0066，说明随机种子波动小于数据集间差异。",
    )
    replace_starts(
        document,
        "在 DBP15K ZH–EN 上，本文模型",
        "在 DBP15K ZH–EN 上，本文模型的 Hits@1 为 0.7198，高于 MTransE、JAPE、BootEA 和 RDGCN 表中数值，但低于 RREA-text。OpenEA EN–FR 的 Hits@1 为 0.6324，高于官方 MTransE、JAPE 和 GCN-Align，低于 BootEA、KDCoE 和 RDGCN。由于 OpenEA 文献值为官方五折平均，而本文只使用官方第一折，该比较仅用于方法定位。",
    )
    replace_starts(
        document,
        "为比较不同任务类型",
        "为比较不同任务类型，表 4 将五个数据集划分为 DBP15K、OpenEA EN–FR 和 EventEA 三类，并对 DBP15K 的三个语言对作等权平均。表中的传播深度、融合权重和 CSLS 邻域均由各数据集验证集独立选择。",
    )
    replace_starts(
        document,
        "DBP15K 的选择后平均 Hits@1",
        "DBP15K 三个语言对的平均 Hits@1 为 0.8075，EventEA 为 0.7552，OpenEA EN–FR 为 0.6324。DBP15K 的较高结果表明名称与结构线索相对充分；OpenEA 的编码 URI 与有限跨语言词向量覆盖仍构成更困难的语义条件。",
    )
    replace_starts(
        document,
        "数据集等权平均 Hits@1",
        "五个数据集等权平均 Hits@1 为 0.7620，Hits@10 为 0.8858，MRR 为 0.8073。图 2 同时展示三项指标及跨种子标准差，避免只用单一 Hits@1 描述模型性能。",
    )
    replace_starts(
        document,
        "图 2 五数据集原始检查点",
        "图 2 新主配置在五个数据集上的测试结果（三随机种子均值，误差线为样本标准差）",
    )

    replace_starts(
        document,
        "表 5 报告关系类型",
        "表 5 将报告拓扑初始化、结构分支监督、全部邻居稀疏选择、关系类型、层选择和语义视图消融。所有配置只关闭目标机制，并在两个代表性数据集上重新训练三个随机种子；当前中间版保留实验位置，最终数值将在全部运行和验证选择结束后统一填入。",
    )
    replace_starts(
        document,
        "去除关系类型后",
        "本轮统一协议消融正在运行。为避免把旧配置结果错误归因于新主模型，本中间版不沿用历史下降幅度；最终版将报告均值、样本标准差、相对主模型差值以及跨种子配对差值。",
    )
    replace_starts(
        document,
        "表 6 比较完整模型",
        "表 6 将比较完整模型、结构单分支、语义单分支和无可训练平均融合，以区分两个信息源和可训练融合模块的作用。各变体使用独立验证配置，结构单分支固定 $\alpha=0$，语义单分支固定 $\alpha=1$。",
    )
    replace_starts(
        document,
        "相对完整模型，结构单分支",
        "上述分支与融合实验正在运行。最终分析将以三随机种子结果判断结构分支是否在新拓扑初始化和早期结构监督下获得更充分作用，并避免根据单个种子作结论。",
    )
    replace_starts(
        document,
        "表 7 固定结构与语义编码器",
        "表 7 固定结构与语义编码器，并使三种融合路径具有相同的融合参数量。完整模型与结构查询基线均读取全部一跳出邻居，二者只改变查询来源；参数匹配晚期融合不读取邻居上下文，用于分离查询时机、结构信息可见范围和模型容量。",
    )
    replace_starts(
        document,
        "与完整模型相比，参数匹配晚期融合",
        "参数匹配融合控制正在按统一协议运行。最终版将分别比较完整模型与结构查询邻居、参数匹配晚期融合的配对差值；只有超过跨种子波动且在两个数据集方向一致的变化才作为主要证据。",
    )

    replace_starts(
        document,
        "针对 RQ1，五个数据集",
        "针对 RQ1，五个数据集的三随机种子结果显示，新主配置在 DBP15K、OpenEA 和 EventEA 上均能稳定训练，数据集等权平均 Hits@1、Hits@10 和 MRR 分别达到 0.7620、0.8858 和 0.8073。DBP15K FR–EN 取得最高 Hits@1，OpenEA EN–FR 仍是最困难的设置。RQ2 和 RQ3 涉及组件归因，必须以正在运行的统一协议消融为依据；本中间版不复用旧模型结论。",
    )
    replace_starts(
        document,
        "OpenEA EN–FR-15K-V2 对实体 URI",
        "OpenEA EN–FR-15K-V2 对实体 URI 进行编码以降低名称偏置[14]，当前输入使用英文 GloVe 覆盖词表内 token，并以基于 MD5 种子的确定性随机向量表示词表外 token[15]。本地统计显示，该数据集 token 的 GloVe 覆盖率约为 52.18%，大量专名和编码标识缺少可迁移的跨语言语义。新主配置将 Hits@1 提高到 0.6324，但仍低于 DBP15K 三语言对平均值 0.8075。这一差异与输入语义覆盖和跨图结构差异一致，但其因果贡献仍需专门实验验证。",
    )
    replace_starts(
        document,
        "内部有效性方面，全部消融",
        "内部有效性方面，主实验和消融均使用相同数据划分、验证规则和三个随机种子，但 $n=3$ 仍不足以支持高功效显著性检验。各数据集和变体独立选择 $\alpha$、$L$ 与 $k_{\mathrm{CSLS}}$，降低了沿用主模型超参数造成的偏置，也引入了有限验证集上的网格方差。外部有效性方面，组件消融只覆盖 DBP15K ZH–EN 和 OpenEA EN–FR，OpenEA 主结果只运行一个官方划分，尚未验证十万实体以上、开放世界和存在无对应实体的场景。",
    )
    replace_starts(
        document,
        "后续研究应优先解决三个问题",
        "后续研究应优先解决三个问题。第一，以跨语言子词模型或可解释名称编码替代随机词表外向量，并在 OpenEA 全部官方划分上复现。第二，继续检验拓扑初始化、结构监督衰减和稀疏邻居选择在更大图上的效率与稳定性。第三，在开放世界和存在无对应实体的场景中评估检索校准与泛化能力。",
    )
    replace_starts(
        document,
        "本文提出并实现一种结合轻量关系消息",
        "本文实现了一种拓扑初始化的关系感知结构上下文与多尺度语义融合方法。模型以八维图内拓扑统计和小比例实体残差建立初始结构状态，保留全部一跳出邻居，并通过批次内动态补齐与 1.5-entmax 构造稀疏结构上下文；联合表示与结构表示分别接受双向 InfoNCE 监督，结构权重从 0.1 线性衰减至 0。五个数据集的三随机种子主实验达到平均 Hits@1 0.7620、Hits@10 0.8858 和 MRR 0.8073。关于各组件的独立贡献、提前与晚期融合差异以及查询信号的结论，将以同一协议下正在完成的多随机种子消融为准。",
    )

    update_tables(document)
    document.core_properties.subject = "新主配置、五数据集主结果与统一协议消融中间版"


def verify(output: Path) -> None:
    document = Document(output)
    tags = []
    unnumbered = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Formula":
            continue
        text = formula_text(paragraph)
        match = re.search(r"\((\d+)\)$", text)
        if match:
            tags.append(int(match.group(1)))
        else:
            unnumbered += 1
    if tags != list(range(1, 28)):
        raise RuntimeError(f"Unexpected native equation tags: {tags}")
    if unnumbered < 2:
        raise RuntimeError("Topology initialization equations were not preserved as native equations")
    xml = document._element.xml
    if xml.count("<m:oMath") < 100:
        raise RuntimeError("Native Word equations were unexpectedly lost")
    stale = ["固定预算邻居提取", "0.7262", "0.7721", "下降 3.40"]
    body_text = "\n".join(p.text for p in document.paragraphs)
    for fragment in stale:
        if fragment in body_text:
            raise RuntimeError(f"Stale content remains: {fragment}")
    for required in ("0.7620", "1.5-entmax", "运行中"):
        if required not in body_text and required not in "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells):
            raise RuntimeError(f"Required content missing: {required}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--figure1", required=True)
    parser.add_argument("--figure2", required=True)
    args = parser.parse_args()

    source = Path(args.source)
    output = Path(args.output)
    figure1 = Path(args.figure1)
    figure2 = Path(args.figure2)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    document = Document(output)
    rewrite_document(document)
    clear_text_highlights(document)
    replace_images(document, output, figure1, figure2)
    verify(output)
    print(f"WROTE {output}")


if __name__ == "__main__":
    main()
