import argparse
import json
import shutil
from copy import deepcopy
from pathlib import Path
from typing import Dict, Iterable, List

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def replace_paragraph(document: Document, starts_with: str, text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(starts_with)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {starts_with!r}, got {len(matches)}")
    paragraph = matches[0]
    paragraph.clear()
    run = paragraph.add_run(text)
    set_east_asia_font(run, "Songti SC")
    return paragraph


def replace_heading(document: Document, old: str, new: str) -> Paragraph:
    paragraph = replace_paragraph(document, old, new)
    for run in paragraph.runs:
        set_east_asia_font(run, "Heiti SC")
        run.bold = True
    return paragraph


def find_paragraph(document: Document, exact_text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if p.text.strip() == exact_text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph {exact_text!r}, got {len(matches)}")
    return matches[0]


def insert_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = style
    run = new_paragraph.add_run(text)
    set_east_asia_font(run, "Heiti SC" if style.startswith("Heading") else "Songti SC")
    if style.startswith("Heading"):
        run.bold = True
    return new_paragraph


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths: Iterable[int], indent: int = 120) -> None:
    widths = list(widths)
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        prevent_row_split(row)
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def clear_text_highlights(document: Document) -> None:
    """Remove inherited review highlighting while preserving table-cell shading."""
    for highlight in list(document.element.body.iter(qn("w:highlight"))):
        highlight.getparent().remove(highlight)


def format_table(table, widths: List[int], first_column_left: bool = True) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                shade_cell(cell, "F4F6F9")
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if first_column_left and column_index == 0
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_east_asia_font(run, "Heiti SC" if row_index == 0 else "Songti SC")
                    run.font.size = Pt(8.5)
                    run.bold = row_index == 0


def add_caption_before(target: Paragraph, text: str) -> Paragraph:
    paragraph = target.insert_paragraph_before(style="Caption")
    run = paragraph.add_run(text)
    set_east_asia_font(run, "Heiti SC")
    run.bold = True
    return paragraph


def add_table_before(document: Document, target: Paragraph, headers: List[str], rows: List[List[str]], widths: List[int]):
    table = document.add_table(rows=1, cols=len(headers))
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, widths)
    target._p.addprevious(table._tbl)
    return table


def add_paragraph_before(target: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    paragraph = target.insert_paragraph_before(style=style)
    run = paragraph.add_run(text)
    set_east_asia_font(run, "Heiti SC" if style.startswith("Heading") else "Songti SC")
    if style.startswith("Heading"):
        run.bold = True
    return paragraph


def metric(row: Dict, name: str) -> str:
    return f"{row[f'{name}_mean']:.4f}±{row[f'{name}_std']:.4f}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument(
        "--results",
        default="outputs/controlled_ablation_20260820/ablation_summary.json",
    )
    args = parser.parse_args()

    source = Path(args.source)
    output = Path(args.output)
    results_path = Path(args.results)
    if not results_path.is_absolute():
        results_path = Path(__file__).resolve().parents[1] / results_path
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, output)
    document = Document(output)

    results = json.loads(results_path.read_text(encoding="utf-8"))
    indexed = {(row["dataset"], row["variant"]): row for row in results}

    title = document.paragraphs[0]
    title.clear()
    run = title.add_run("关系感知结构上下文与多尺度语义融合的\n跨语言知识图谱实体对齐方法")
    set_east_asia_font(run, "Heiti SC")
    run.font.size = Pt(17)
    run.bold = True

    replace_paragraph(
        document,
        "跨语言知识图谱实体对齐需要同时处理",
        "跨语言知识图谱实体对齐需要同时处理关系结构不一致、实体文本碎片化和监督样本有限等问题。本文构建一种关系感知结构上下文与多尺度语义融合方法：关系嵌入和共享投影形成轻量关系消息，节点级层选择器融合不同传播深度，token、phrase 和 global 三个视图编码实体文本，固定预算结构邻居用于形成局部上下文。模型采用单阶段双向 InfoNCE 训练，并仅依据验证集按数据集选择语义权重、有效传播深度和 CSLS k。五个数据集的三随机种子主结果达到平均 Hits@1 0.7262、MRR 0.7721。进一步在 DBP15K ZH–EN 和 OpenEA EN–FR 上完成 42 次同协议控制实验：去除关系类型使 Hits@1 分别下降 3.40 和 2.35 个百分点，去除层选择器分别下降 2.15 和 1.90 个百分点；参数量相同的晚期融合分别下降 1.32 和 7.52 个百分点。相反，将语义查询改为结构查询后 Hits@1 仅变化 +0.23 和 +0.02 个百分点，说明现有证据支持关系感知、多深度选择和结构邻居可见性，但尚不能证明语义必须在邻居聚合前参与查询。",
    )
    replace_paragraph(
        document,
        "关键词：",
        "关键词：知识图谱实体对齐；跨语言知识图谱；关系感知图神经网络；结构邻居上下文；多尺度语义编码；受控消融",
    )
    replace_paragraph(
        document,
        "第一，结构相似不等于实体等价",
        "第一，结构相似不等于实体等价。例如，人物 A 通过“出生地”连接北京，人物 B 通过“工作地点”连接北京；若模型只看到两者都连接北京而忽略关系类型，就会高估其相似性。第二，语义证据往往是碎片化的。例如，名称 token“Washington”本身无法区分人物、州和城市；局部片段“born in Washington”和完整属性序列提供的证据也不相同。第三，监督信号有限。种子对齐只覆盖实体集合的一部分；例如早期模型把行星“Mercury”错配为同名汽车品牌并将其作为伪标签后，错误实体的关系和属性会继续拉近，进而改变后续候选排序。本文不使用伪标签，该例仅说明监督稀缺条件下避免额外错误监督的必要性。",
    )
    replace_paragraph(
        document,
        "本研究的核心目标是检验",
        "本研究围绕四个可检验问题展开：关系类型和节点级多深度选择是否提高结构区分性；token、phrase 和 global 三个语义视图是否分别产生稳定增益；在融合容量相同的条件下，性能差异主要来自语义查询时机还是结构邻居可见范围；按数据集进行验证选择是否稳定改善检索。由此，本文不预设每个组件都有效，而以多随机种子控制实验区分得到支持、数据集依赖和未获支持的假设。",
    )
    replace_paragraph(
        document,
        "设计一种轻量级的关系感知结构编码器",
        "实现并验证一种轻量级关系感知结构编码器。关系类型通过小维度嵌入进入共享消息投影，节点级层选择器融合 0 至 3 跳状态；在两个代表性数据集上移除关系类型或层选择器均导致 Hits@1 下降。",
    )
    replace_paragraph(
        document,
        "构建语义编码与结构编码早期交互模块",
        "构建结构邻居上下文与多尺度语义融合路径，并用参数量相同的晚期融合及无语义查询邻居基线拆分容量、邻居可见性和查询时机。结果表明结构邻居可见性具有稳定价值，但现有语义查询方式未显示独立增益。",
    )
    replace_paragraph(
        document,
        "在五个数据集上完成三个随机种子的重新训练",
        "在五个数据集上完成三随机种子主实验，并在 DBP15K ZH–EN 与 OpenEA EN–FR 上新增 42 次同协议消融训练。所有配置只依据验证集平均 MRR选择，论文同时报告正结果与未获支持的组件。",
    )
    replace_paragraph(
        document,
        "已有文本增强实体对齐方法表明",
        "已有文本增强实体对齐方法表明，名称和描述能够改善跨语言匹配。RREA 的文本版本将文本增强与纯结构设置分开比较[5]。MCLEA 分别编码多种模态，并用模态内对比与模态间对齐目标共同构造实体表示[12]；Li 等人的多模态图谱 Transformer 工作指出，直接拼接或注意力组合异质信息可能形成未对齐的信息空间[13]。这些研究说明语义与结构均有价值，但不能预先证明某一交互位置或每个语义尺度都必要。本文因此在相同数据划分、训练目标和验证协议下，对关系类型、层选择、三个语义视图、参数匹配晚期融合和邻居查询信号分别进行控制。",
    )
    replace_paragraph(
        document,
        "早期错误（early error）",
        "早期错误（early error）指模型尚未形成稳定表示时产生的高置信错误。例如，同名实体“Mercury”可能在早期被错配到错误类别；若该实体对被直接加入监督集合，其邻居和属性会进一步围绕错误对应关系收缩。朴素伪标签（naive pseudo-labeling）是指只根据模型分数接受预测，而不进行互惠一致性、类型约束或可靠性过滤。当前方法不生成伪标签。",
    )
    replace_paragraph(
        document,
        "本文将“协同”定义为",
        "前向计算包含两个潜在交互点：语义或结构状态首先用于构造固定预算邻居上下文，随后该上下文与语义状态通过逐维门控形成联合表示。为避免仅凭架构图把两者统称为有效“协同”，第 5.8 节在相同参数量下分别控制邻居是否可见、邻居查询是否使用语义以及融合发生的位置。",
    )
    replace_paragraph(
        document,
        "DBP15K 的语义序列使用已有预处理特征",
        "DBP15K 的语义序列使用已有预处理特征，其长度与形状见表 2。OpenEA 与 EventEA 的序列按“实体名称—关系名称—属性名称—属性值”的顺序分词，最大长度为 32；例如城市实体可形成“Washington country United States population ...”这样的输入，其中 token 视图保留单词证据，phrase 视图捕获相邻片段，global 视图建模远距离组合。词表内 token 是能够在 GloVe 词表中查到的词，直接使用 300 维预训练向量[15]；词表外 token 是 GloVe 中不存在的名称或符号，使用基于 MD5 种子的确定性随机向量，使不同运行得到相同初始化。三类数据的名称可用性、序列长度和词表外比例不同，因此跨数据集差异不能只归因于模型结构。",
    )
    p58 = replace_paragraph(
        document,
        "中间层进一步使用 LayerNorm",
        "中间层进一步使用 LayerNorm、ReLU 和 dropout。当前主配置堆叠 3 层：第一层吸收直接邻居，第二、三层通过前一层状态继续获得两跳和三跳邻域证据。",
    )
    insert_after(
        p58,
        "本文所称消息路径是沿有向边逐层传播的计算路线，而不是显式生成的关系路径。例如“姚明—出生地→上海—所在国家→中国”需要两层传播才能让姚明间接获得中国的信息。RPR-RHGT 会生成并筛选可靠关系路径[10]；当前代码不执行路径生成、筛选或路径编码，因此本文只主张多跳邻域聚合，不主张路径级推理。",
    )
    replace_paragraph(
        document,
        "输入语义序列为",
        "输入语义序列为由名称、关系和属性 token 组成的矩阵。每一行对应一个 300 维 token 向量，非零行形成有效位置掩码。输入先经过线性投影、LayerNorm、正弦位置编码和 dropout 得到共享序列状态；token、phrase 和 global 三个视图都从同一序列及同一掩码构造，不额外引入外部句子。",
    )
    replace_paragraph(
        document,
        "加权后的三个向量仍以拼接形式进入",
        "加权后的三个向量以拼接形式进入输出 MLP，并加上 global 视图残差；若 global 视图关闭，则使用其余启用视图的均值作为残差。输出经 L2 归一化得到语义表示。MLP 用于学习实体级非线性权重，而非预设固定比例。第 5.7 节的控制实验显示，该机制的总体可用性不等于每个视图都必要：phrase 视图在 OpenEA 上有明显贡献，而 token 与 global 视图未显示独立增益。",
    )
    replace_heading(document, "3.9 语义引导的结构上下文与联合表示", "3.9 邻居上下文、查询信号与联合表示")
    replace_paragraph(
        document,
        "其中  和  保留独立证据",
        "其中两个原始向量保留各自证据，逐元素乘积表示同一维度上的共同激活，绝对差表示逐维不一致。这种“原向量—差异—乘积”的局部匹配表示也用于 ESIM 的推理匹配层[19]；本文将其迁移到语义—结构邻居匹配，并通过消融检验其所在交互位置，而不把该特征组合本身视为原创。邻居门控随后计算为：",
    )
    replace_paragraph(
        document,
        "本文仅将公式（13）—（20）",
        "公式（13）—（20）给出主模型的邻居上下文与联合表示路径。主模型使用语义表示作为邻居查询，但第 5.8 节进一步以结构查询替换语义查询，从而检验这一步是否产生独立贡献。",
    )
    replace_paragraph(
        document,
        "本文另外设置两个不使用结构邻居的控制组",
        "本文设置两个参数受控基线。第一，参数匹配晚期融合不读取结构邻居，只在结构与语义分支独立编码后进行末端融合。第二，无语义查询邻居基线保留同一邻居集合和全部门控参数，但使用实体自身结构表示查询邻居，语义仅在最终联合门控中出现。",
    )
    replace_paragraph(
        document,
        "晚期拼接基线先独立得到并归一化",
        "参数匹配晚期融合先独立得到并归一化结构表示和语义表示，其核心映射仍采用式（21）的拼接 MLP：",
    )
    replace_paragraph(
        document,
        "其中，，。",
        "其中第一层将 2d=256 维输入映射到 512 维，第二层映射回 d=128 维。该 MLP 含 197,248 个参数；另加入 128 维输出尺度和 1 个残差混合系数，使有效可训练参数总数达到 197,377，与主模型融合模块完全一致。",
    )
    replace_paragraph(
        document,
        "另一个无参数平均基线为",
        "无语义查询邻居基线与主模型共享公式（13）—（20）的参数和邻居输入，仅将公式（13）、（15）和（17）中的查询条件由语义表示替换为实体自身结构表示；最终式（19）—（20）仍融合语义与结构上下文。",
    )
    replace_paragraph(
        document,
        "按实际前向路径统计",
        "按有效前向路径统计，主模型、参数匹配晚期融合和无语义查询邻居基线的融合模块均含 197,377 个可训练参数。三者分别对应“邻居可见且语义查询”“邻居不可见且末端融合”和“邻居可见但结构查询”，因此可以依次分离模型容量、结构邻居可见范围和语义查询时机。",
    )
    replace_heading(document, "3.10 双向 InfoNCE 对齐训练", "3.10 单阶段双向 InfoNCE 对齐训练")
    replace_paragraph(
        document,
        "该式即完整训练目标",
        "式（23）即当前代码和全部主实验的完整训练目标。相似度矩阵每行、每列的非对角实体构成批内对比负例。难负候选（hard negative candidate）是当前表示下得分很高但并非真实对应的实体；本文不额外挖掘这类候选，也不使用排序损失、伪标签或独立 warm-up 阶段。因此不存在“第一阶段损失”和“第二阶段损失”的差异，所有 epoch 使用同一训练对、同一 InfoNCE 目标，只有学习率随调度器变化。",
    )
    replace_paragraph(
        document,
        "模型采用 AdamW 优化器",
        "模型采用 AdamW 优化器[16]。DBP15K 的学习率为 5×10⁻⁴，最多训练 36 个 epoch；OpenEA 与 EventEA 的学习率为 3×10⁻⁴，最多训练 50 个 epoch。所有数据集每 5 个 epoch 在验证集上评估 MRR并保存最佳检查点，连续 4 次验证未提升时允许提前停止。主实验与新增 42 次消融均不进行结构预热、联合阶段切换或末轮权重平均。",
    )
    replace_paragraph(
        document,
        "实验围绕三个问题展开",
        "实验回答四个研究问题：RQ1，完整模型在跨语言和事件型图谱上是否稳定；RQ2，关系类型、层选择以及 token、phrase、global 视图分别产生何种贡献；RQ3，在融合参数量相同的条件下，邻居可见范围与语义查询时机如何影响结果；RQ4，按数据集独立选择语义权重、有效传播深度与 CSLS k 是否改善检索。",
    )
    replace_paragraph(
        document,
        "所有主结果采用统一协议",
        "表 3、表 5–8 和图 2 来自五个数据集、随机种子 42、43、44 的 15 次主模型训练。表 9–11 在 DBP15K ZH–EN 与 OpenEA EN–FR 上复用同协议主模型结果，并新增 7 个变体×2 个数据集×3 个种子，共 42 次训练。所有运行采用相同数据划分、InfoNCE 目标、验证频率和提前停止规则；每个变体分别依据三个种子的平均验证 MRR选择 α、L、k。表 4 的文献数值仅用于定位性能区间。",
    )
    replace_paragraph(
        document,
        "本文以 Hits@1",
        "本文以 Hits@1、Hits@10 和平均倒数排名（MRR）衡量完整候选集合上的检索质量。五数据集主结果及七项消融均报告随机种子 42、43、44 的均值与样本标准差。每个数据集—变体组合只使用训练集或验证集选择检查点和 α、L、k，测试集仅用于唯一配置确定后的最终报告。由于 n=3 不足以支持可靠的正态性与显著性检验，本文报告绝对差值和跨种子离散程度，不使用夸大的显著性措辞。文献方法因文本输入、划分和后处理不同，仅作定位性比较。",
    )
    replace_paragraph(
        document,
        "公共设置为：batch size 512",
        "公共设置为 batch size 512，结构与联合表示维度 128，关系感知 GNN 最大深度 3，语义编码器 2 层、4 个注意力头，dropout 0.1，邻居预算 8，InfoNCE 温度 0.07。验证网格为 α∈{0.20,…,0.80}、L∈{1,2,3}、k∈{3,5,7,10,15,20}，CSLS 校正强度固定为 1.0。关系、层和语义视图消融仅关闭目标信息源；无关系类型变体以全部关系嵌入的均值替代边类型，保持关系通道参数量不变。主模型、结构查询基线和参数匹配晚期融合的融合参数均为 197,377。",
    )
    replace_paragraph(
        document,
        "最优 α 位于",
        "最优 α 位于 0.50–0.70，k 覆盖 3、5、15 和 20，说明语义占比与 CSLS 局部尺度需要按数据集校准。五个主数据集均选择 L=3。新增消融也分别重新选择 α、L、k，避免把主模型的后处理参数强加给结构或语义已改变的变体。",
    )
    replace_paragraph(
        document,
        "本轮 15 次训练均保留",
        "五数据集主模型始终保留 token、phrase、global 三个语义视图，global 分支及其两层 Transformer 未作结构改动。第 5.7 节仅在两个代表性数据集上关闭单一视图，因此用于判断独立贡献，而不改变五数据集主配置。",
    )

    # Keep the classic-method table synchronized with the revised contribution language.
    comparison_table = document.tables[3]
    comparison_table.cell(6, 1).text = "关系感知 GNN + 多尺度语义 + 结构邻居上下文"
    format_table(comparison_table, [1691640, 1062990, 1062990, 1062990, 1062990])

    section6 = find_paragraph(document, "6 讨论")
    add_paragraph_before(section6, "5.7 关系、层选择与语义视图消融", "Heading 2")
    add_paragraph_before(
        section6,
        "表 9 在两个代表性数据集上比较完整模型与五个单组件消融。每个变体均重新训练三个随机种子，并独立选择 α、L、k。",
    )
    add_caption_before(section6, "表 9 同协议组件消融结果（Hits@1/MRR，均值±样本标准差，n=3）")
    component_variants = [
        "main",
        "no_relation_types",
        "no_layer_selector",
        "no_token_view",
        "no_phrase_view",
        "no_global_view",
    ]
    component_rows = []
    for variant in component_variants:
        dbp = indexed[("dbp15k_zh_en", variant)]
        open_ea = indexed[("openea_en_fr", variant)]
        component_rows.append([
            dbp["label_zh"],
            f"{dbp['semantic_weight']:.2f}/{dbp['depth']}/{dbp['csls_k']}",
            f"{metric(dbp, 'Hits@1')} / {metric(dbp, 'MRR')}",
            f"{open_ea['semantic_weight']:.2f}/{open_ea['depth']}/{open_ea['csls_k']}",
            f"{metric(open_ea, 'Hits@1')} / {metric(open_ea, 'MRR')}",
        ])
    add_table_before(
        document,
        section6,
        ["变体", "DBP α/L/k", "DBP H@1 / MRR", "OpenEA α/L/k", "OpenEA H@1 / MRR"],
        component_rows,
        [1800, 1200, 2400, 1200, 2520],
    )
    add_paragraph_before(
        section6,
        "去除关系类型后，DBP15K 与 OpenEA 的 Hits@1 分别下降 3.40 和 2.35 个百分点，MRR 分别下降 3.87 和 2.67 个百分点；去除层选择器的 Hits@1 分别下降 2.15 和 1.90 个百分点。两项结构组件在两个数据集上的方向一致。语义视图呈现不同结论：去除 phrase 视图在 OpenEA 上使 Hits@1 下降 4.24 个百分点，但在 DBP15K 上提高 0.68 个百分点；去除 token 或 global 视图未产生一致下降。因此，实验支持 phrase 信息的场景依赖价值，但不支持“三个视图均为必要组件”的强主张。",
    )
    add_paragraph_before(section6, "5.8 交互位置、邻居可见范围与容量控制", "Heading 2")
    add_paragraph_before(
        section6,
        "为解释晚期融合差异，表 10 固定结构与语义编码器，并使三种融合路径具有相同的 197,377 个有效参数。",
    )
    add_caption_before(section6, "表 10 参数匹配融合控制的设计差异")
    add_table_before(
        document,
        section6,
        ["变体", "结构邻居可见", "邻居查询信号", "语义进入位置 / 融合参数"],
        [
            ["完整模型", "是", "语义表示", "邻居聚合前 + 最终门控 / 197,377"],
            ["无语义查询邻居", "是", "自身结构表示", "仅最终门控 / 197,377"],
            ["参数匹配晚期融合", "否", "无邻居查询", "独立编码后末端融合 / 197,377"],
        ],
        [1800, 1600, 2200, 3520],
    )
    add_caption_before(section6, "表 11 参数匹配融合控制的测试结果（均值±样本标准差，n=3）")
    control_variants = ["main", "structural_neighbor_query", "late_concat_matched"]
    control_rows = []
    for variant in control_variants:
        dbp = indexed[("dbp15k_zh_en", variant)]
        open_ea = indexed[("openea_en_fr", variant)]
        control_rows.append([
            dbp["label_zh"],
            metric(dbp, "Hits@1"),
            metric(dbp, "MRR"),
            metric(open_ea, "Hits@1"),
            metric(open_ea, "MRR"),
        ])
    add_table_before(
        document,
        section6,
        ["变体", "DBP H@1", "DBP MRR", "OpenEA H@1", "OpenEA MRR"],
        control_rows,
        [1800, 1830, 1830, 1830, 1830],
    )
    add_paragraph_before(
        section6,
        "与完整模型相比，参数匹配晚期融合在 DBP15K 和 OpenEA 上的 Hits@1 分别下降 1.32 和 7.52 个百分点，表明增加末端 MLP 容量不能替代结构邻居上下文。无语义查询邻居基线的 Hits@1 则分别变化 +0.23 和 +0.02 个百分点，MRR 分别变化 +0.17 和 −0.02 个百分点，均小于跨种子标准差。由于该基线保留邻居信息但取消语义查询，现有结果把主要收益定位于结构邻居可见性，而非语义必须在聚合前参与查询。",
    )

    replace_paragraph(
        document,
        "三随机种子结果支持三点",
        "主实验与受控消融共同支持四点。第一，关系类型和节点级层选择在两个代表性数据集上均产生方向一致的增益。第二，结构邻居上下文比参数匹配的末端融合更有效，尤其在 OpenEA 上差距较大。第三，语义提前查询邻居并未显示独立优势，因而不能将“早期语义引导”列为已验证贡献。第四，语义视图贡献具有数据集依赖性，只有 phrase 视图在 OpenEA 上显示明显必要性；按数据集选择 α 和 CSLS k 仍然重要。",
    )
    replace_paragraph(
        document,
        "五数据集的 15 次训练均保留",
        "消融结果对最初假设形成了区分。关系类型与多深度选择假设得到两个数据集的一致支持；“三个语义视图分别必要”的假设未得到支持；“语义必须在结构上下文形成前参与查询”的假设也未得到支持。完整模型优于参数匹配晚期融合，不能单独证明语义查询时机，因为晚期基线同时看不到结构邻居；加入结构查询基线后，完整模型与其结果几乎相同。由此，较稳健的解释是关系感知结构编码和邻居上下文可见性构成主要有效因素。",
    )
    replace_paragraph(
        document,
        "内部有效性方面",
        "内部有效性方面，所有新增消融使用相同数据划分、训练目标和三随机种子，但 n=3 仍不足以支持高功效显著性检验；每个变体独立选择 α、L、k 虽保证公平，却也引入验证网格方差。消融只覆盖 DBP15K ZH–EN 与 OpenEA EN–FR，不能直接推广到其余三个主数据集。外部有效性方面，尚未验证 100K 以上规模、开放世界和无对应实体场景。词表外随机向量及不同数据集的文本预处理也可能限制跨语言语义质量。最后，表 4 的文献方法输入与划分不同，定位性比较不能替代统一复现。",
    )
    replace_paragraph(
        document,
        "后续研究应优先解决三个问题",
        "后续研究应优先解决三个问题。第一，重新设计语义查询或采用对比约束，使语义在邻居选择阶段产生可测量而非冗余的作用。第二，采用累积视图实验（token、token+phrase、token+phrase+global）和更多随机种子，进一步确认 phrase 的场景依赖性以及 token/global 的冗余来源。第三，在更大规模、开放世界、无对应实体和跨知识库来源场景中评估效率、校准稳定性与泛化性。",
    )
    replace_paragraph(
        document,
        "本文构建了一套结合关系感知结构建模",
        "本文构建了一套结合轻量关系消息、节点级多深度选择、固定预算结构邻居上下文和多尺度语义编码的实体对齐方法。模型采用单阶段双向 InfoNCE 训练，并在验证集上按数据集选择语义权重、有效传播深度和 CSLS k。五数据集三随机种子主实验表明，验证选择后的平均 Hits@1 和 MRR 分别为 0.7262 和 0.7721。",
    )
    replace_paragraph(
        document,
        "本文不主张所有语义视图",
        "两个代表性数据集上的 42 次新增训练进一步限定了方法贡献：关系类型和层选择均得到一致支持；参数匹配晚期融合明显弱于保留邻居上下文的模型；phrase 视图只在 OpenEA 上显示明显贡献，而 token、global 及语义提前查询邻居均未显示稳定独立增益。",
    )
    replace_paragraph(
        document,
        "最终五数据集平均 Hits@1",
        "因此，本文的科学结论不是“所有模块都有效”，而是通过容量和信息可见范围受控的实验，将主要证据定位到关系感知、多深度结构选择和结构邻居上下文。语义查询时机与部分语义视图仍需重新设计和扩大验证范围。该结论比仅比较完整模型与文献基线更可复核，也为后续简化模型提供了明确方向。",
    )

    reference = document.add_paragraph(style="Reference")
    reference_run = reference.add_run(
        "[19] Chen Q, Zhu X, Ling ZH, Wei S, Jiang H, Inkpen D (2017) Enhanced LSTM for natural language inference. In: Proceedings of the 55th Annual Meeting of the Association for Computational Linguistics, vol 1, pp 1657–1668. https://doi.org/10.18653/v1/P17-1152"
    )
    set_east_asia_font(reference_run, "Songti SC")

    document.core_properties.title = "关系感知结构上下文与多尺度语义融合的跨语言知识图谱实体对齐方法"
    document.core_properties.subject = "统一协议多随机种子消融修订版"
    clear_text_highlights(document)
    document.save(output)
    print(f"WROTE {output}")


if __name__ == "__main__":
    main()
