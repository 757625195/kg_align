from __future__ import annotations

from pathlib import Path

from docx import Document

from create_latest_chinese_translation import (
    format_chinese_typography,
    format_reference_language,
    protect_latin_terms,
    set_document_language,
)


ROOT = Path(__file__).resolve().parents[1]
EN_SOURCE = ROOT / "outputs" / "Relation_Aware_Neighbor_Context_Springer_English_Formatted_20260830.docx"
ZH_SOURCE = ROOT / "outputs" / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_Springer中文修订稿_20260830.docx"
EN_OUTPUT = ROOT / "outputs" / "Relation_Aware_Neighbor_Context_Springer_Anti_Defensive_English_20260830.docx"
ZH_OUTPUT = ROOT / "outputs" / "跨语言实体对齐中的关系感知邻居上下文与语义引导选择_Springer去防御性中文译稿_20260830.docx"


EN_REWRITES = {
    5: (
        "Cross-lingual knowledge graph entity alignment identifies matching entities across graphs written in "
        "different languages. This paper introduces a relation-aware neighbor-context model that lets entity "
        "semantics select structural evidence before joint encoding. The structural encoder preserves relation "
        "types and combines several propagation depths. The semantic encoder builds token, phrase, and global views "
        "from one entity-text sequence. A semantic query scores the complete one-hop outgoing neighborhood, and "
        "1.5-entmax removes weak evidence. The model then combines the selected structural context with entity "
        "semantics. Experiments use three random seeds on five datasets. The full model reaches Hits@1 scores of "
        "0.7285 on DBP15K ZH–EN and 0.6344 on OpenEA EN–FR-15K-V2. With the neighborhood and trainable fusion "
        "parameter count fixed, semantic queries improve Hits@1 over structural queries by 2.08 and 2.48 percentage "
        "points. The full model also improves over semantic-only retrieval by 9.09 and 9.66 points. Removing relation "
        "types lowers Hits@1 by 3.62 and 2.78 points. The full model exceeds a parameter-matched late-fusion control "
        "by 5.28 and 9.22 points. This comparison measures the combined contribution of neighbor access and the "
        "learned interaction path. These results establish relation-aware neighbor context and semantics-guided "
        "selection as effective components of the evaluated alignment framework."
    ),
    10: (
        "Existing methods establish the value of both graph structure and text. MTransE and JAPE map graphs into "
        "shared or linked vector spaces [2, 3]. BootEA adds high-confidence pairs to the training set [4]. RDGCN and "
        "RREA preserve relation and neighborhood information [5, 6]. Their results leave one design question open: "
        "when should structural neighbors enter a joint representation?"
    ),
    15: (
        "The structural encoder preserves relation types during message passing. Shared transformations control "
        "parameter growth, while relation embeddings retain edge-type differences. A node-level layer selector lets "
        "each entity combine shallow and deep graph states."
    ),
    16: (
        "The semantic encoder builds token, phrase, and global views from one entity-text sequence. The resulting "
        "semantic state scores the complete one-hop outgoing neighborhood, and 1.5-entmax removes weak neighbors "
        "before joint encoding."
    ),
    17: (
        "Experiments use three random seeds on five datasets that cover different languages, graph sources, and event "
        "data. Matched controls isolate relation types, semantic views, neighbor access, and the query signal. The "
        "query comparison fixes the neighborhood and trainable fusion parameter count."
    ),
    20: (
        "MTransE learns mappings between language-specific embedding spaces [2]. JAPE adds attribute links to a "
        "shared structural space [3]. These methods give graph integration a clear geometric form. Their mappings "
        "offer limited access to relation-specific neighborhood structure, which motivates relation-aware encoders."
    ),
    23: (
        "RDGCN adds relation evidence through an entity graph and a relation dual graph [5]. RREA uses relation "
        "reflection to preserve relation differences [6]. RPR-RHGT encodes selected multi-step relation paths with a "
        "graph Transformer [10]. These models establish the value of typed and multi-hop structure. This study "
        "complements them with a compact relation-aware message-passing encoder and a controlled test of how one-hop "
        "structural context enters the joint representation."
    ),
    25: (
        "Transformers connect positions across a sequence [11]. Hierarchical attention networks show that different "
        "levels of text provide distinct evidence [12]. The semantic encoder applies this multi-scale principle "
        "through token, phrase, and global views. These views jointly guide neighbor selection."
    ),
    26: (
        "Text-based entity alignment studies show that semantics can support graph structure. RREA reports text-based "
        "and structure-only settings [6]. MCLEA aligns several data types through contrastive objectives [13]. Other "
        "work shows that simple fusion may mix spaces that are not aligned [14]. Together, these findings motivate a "
        "controlled comparison of neighbor-query signals. This study fixes the neighborhood and trainable fusion "
        "parameter count while changing the query source."
    ),
    27: (
        "Recent studies use large language models [15], domain adaptation [16], or step-by-step agents [17]. These "
        "approaches change the supervision source or reasoning process. This study targets a complementary question: "
        "how do local structural neighbors and entity semantics interact within a fixed alignment framework?"
    ),
    37: (
        "The model combines four components. The structural encoder preserves relation types and combines several "
        "propagation depths. The semantic encoder builds token, phrase, and global views from the same entity text. "
        "The fusion module uses semantics to score every one-hop outgoing neighbor. It forms a sparse structural "
        "context and combines that context with entity semantics. Bidirectional InfoNCE trains the joint and structural "
        "representations. CSLS then adjusts similarities between the trained joint representations. Figure 1 shows "
        "the complete information flow."
    ),
    41: (
        "Local degree and neighbor statistics describe a node's structural role [18, 19]. Structural-identity methods "
        "also compare nodes by their graph roles. These observations motivate an eight-dimensional topology vector "
        "that initializes the structural branch. The following definition specifies this feature set."
    ),
    48: (
        "The structural encoder and context module assign complementary roles to edge direction. Incoming edges update "
        "the current node during graph propagation. Outgoing neighbors represent facts emitted by the current entity "
        "and serve as candidate context during fusion. This division aligns graph propagation with incoming evidence "
        "and context selection with outgoing facts."
    ),
    70: (
        "The validity mask restricts attention to observed tokens. Mean pooling preserves broad lexical evidence, and "
        "attention pooling highlights the most informative terms. Their average balances coverage and selectivity."
    ),
    105: (
        "The semantic query shapes neighbor weights before the model builds structural context. The final gate then "
        "combines this context with entity semantics. Section 3.5.4 evaluates this design through matched controls."
    ),
    107: (
        "The structural-query control isolates the query signal. It retains the same neighborhood, dynamic padding, "
        "1.5-entmax, gates, and joint retrieval. It substitutes the entity's structural representation for the "
        "semantic query. The late-fusion control measures a broader design change. It matches the trainable fusion "
        "parameter count and forms the joint representation after separate branch encoding, with neighbor context "
        "excluded from fusion."
    ),
    142: (
        "Table 3 positions the model against published results. Differences in data splits, text inputs, and "
        "post-processing make these values contextual. Tables 4–6 provide the same-protocol evidence for the paper's "
        "claims."
    ),
    147: (
        "Relation types provide the clearest consistent structural gain. Their removal lowers Hits@1 by 3.62 points on "
        "DBP15K and 2.78 points on OpenEA. Topology features improve the structural state by 3.46 and 8.03 points over "
        "entity vectors alone. The phrase view improves both datasets by 0.81 and 3.95 points. The remaining controls "
        "show smaller or dataset-dependent effects and define the model's supporting design choices."
    ),
    149: (
        "Table 5 compares the full model with two single branches and parameter-free mean fusion. Each comparison uses "
        "the same representation for training and retrieval and keeps branch weights fixed during validation."
    ),
    151: (
        "The full model improves Hits@1 over semantic-only retrieval by 9.09 points on DBP15K and 9.66 points on "
        "OpenEA. It also outperforms parameter-free mean fusion. These gains show that learned interaction converts "
        "structural context into useful alignment evidence. The low structural-only scores further show that joint "
        "learning with semantics creates this gain."
    ),
    153: (
        "Table 6 tests the query signal and context-aware interaction with matched controls. The structural-query "
        "control shares the full neighborhood, gates, and fusion parameters but scores neighbors with structure. The "
        "parameter-matched late-fusion control uses the same fusion parameter count after independent branch encoding."
    ),
    155: (
        "Semantic queries improve Hits@1 over structural queries by 2.08 points on DBP15K and 2.48 points on OpenEA. "
        "This controlled gap measures the contribution of semantics-guided neighbor scoring. The full model improves "
        "over parameter-matched late fusion by 5.28 and 9.22 points. This second gap measures the combined contribution "
        "of neighbor access and the learned interaction path."
    ),
    158: (
        "The experiments establish three findings. First, relation types produce a consistent structural gain on both "
        "controlled datasets. Second, joint learning turns structural context into useful evidence alongside semantics. "
        "Third, semantic queries outperform structural queries when the neighborhood and trainable fusion parameter "
        "count are fixed. Together, these findings identify relation-aware neighbor context and semantics-guided "
        "selection as the main sources of the model's gains."
    ),
    161: (
        "The evidence covers five datasets, with component controls on DBP15K ZH–EN and OpenEA EN–FR-15K-V2. Three "
        "random seeds quantify effect size and sample variation. The fixed-neighbor control jointly changes neighbor "
        "count and the weighting function. The parameter-matched late-fusion control jointly changes neighbor access "
        "and the training interaction path. These design properties define the causal scope of the reported "
        "comparisons."
    ),
    163: (
        "Three extensions will test the method under broader conditions. Multilingual subword encoders can replace "
        "random vectors for unknown tokens [27, 28]. Larger graphs can test the efficiency of complete-neighborhood "
        "selection. Open-world datasets can test entities that lack matches. These experiments will measure the "
        "robustness of semantics-guided neighbor selection with weaker text and larger candidate spaces."
    ),
    165: (
        "This study shows how structural neighbors can enter structural–semantic entity alignment. Relation types "
        "strengthen structural evidence on both controlled datasets. A learned joint representation turns structural "
        "context into gains beyond semantic-only retrieval. Semantic queries also outperform structural queries under "
        "matched conditions. Complete neighbor access with a learned interaction path outperforms parameter-matched "
        "late fusion. Together, these results establish relation-aware neighbor context and semantics-guided selection "
        "as effective design choices for cross-lingual entity alignment."
    ),
}


ZH_REWRITES = {
    5: (
        "跨语言知识图谱实体对齐用于识别不同语言图谱中的对应实体。本文提出一种采用语义引导选择的关系感知邻居上下文模型，使实体语义能够在联合编码之前筛选结构证据。"
        "结构编码器保留关系类型并组合多个传播深度。语义编码器从同一实体文本序列构建 token、phrase 和 global 三种视图。语义查询为完整的一跳出邻域评分，"
        "1.5-entmax 去除弱相关证据，模型随后将筛选后的结构上下文与实体语义结合。实验在五个数据集上使用三个随机种子。完整模型在 DBP15K ZH–EN 和 "
        "OpenEA EN–FR-15K-V2 上的 Hits@1 分别达到 0.7285 和 0.6344。在邻域和可训练融合参数量固定时，语义查询相对结构查询分别提高 2.08 和 2.48 个百分点。"
        "与仅语义检索相比，Hits@1 在 OpenEA 上提高 9.66 个百分点，在 DBP15K 上提高 9.09 个百分点。去除关系类型后，Hits@1 分别下降 3.62 和 2.78 个百分点。完整模型相对参数匹配晚期融合对照分别提高 "
        "5.28 和 9.22 个百分点，该比较衡量邻居访问与学习型交互路径的联合贡献。结果表明，关系感知邻居上下文和语义引导选择是该实体对齐框架中的有效组成部分。"
    ),
    10: (
        "现有方法已经确立图结构和文本信息的价值。MTransE 和 JAPE 将图谱映射到共享或相互关联的向量空间[2–3]。BootEA 将高置信实体对加入训练集[4]。"
        "RDGCN 和 RREA 保留关系与邻域信息[5–6]。这些研究留下了一个关键设计问题：结构邻居应在何时进入联合表示？"
    ),
    15: (
        "结构编码器在消息传递中保留关系类型。共享变换控制参数增长，关系嵌入保留边类型差异。节点级层选择器使每个实体能够组合浅层和深层图状态。"
    ),
    16: (
        "语义编码器从同一实体文本序列构建 token、phrase 和 global 三种视图。所得语义状态为完整的一跳出邻域评分，1.5-entmax 在联合编码之前去除弱相关邻居。"
    ),
    17: (
        "实验在五个数据集上使用三个随机种子，这些数据集覆盖不同语言、图谱来源与事件数据。匹配对照分别检验关系类型、语义视图、邻居访问和查询信号。查询比较固定邻域和可训练融合参数量。"
    ),
    20: (
        "MTransE 学习不同语言嵌入空间之间的映射[2]。JAPE 在共享结构空间中加入属性关联[3]。这些方法为图谱整合提供清晰的几何形式，但其映射对关系特定邻域结构的表达能力有限，"
        "由此推动了关系感知编码器的发展。"
    ),
    23: (
        "RDGCN 通过实体图与关系对偶图引入关系证据[5]。RREA 使用关系反射变换保留关系差异[6]。RPR-RHGT 使用图 Transformer 编码筛选后的多步关系路径[10]。"
        "这些模型确立了带类型和多跳结构的价值。本文以紧凑的关系感知消息传递编码器补充现有研究，并受控检验一跳结构上下文如何进入联合表示。"
    ),
    25: (
        "Transformer 可以建立序列中不同位置之间的联系[11]。层次注意力网络表明，不同文本粒度能够提供不同证据[12]。语义编码器通过 token、phrase 和 global 三种视图实现这一多尺度原则，"
        "并使用三种视图共同引导邻居选择。"
    ),
    26: (
        "文本增强实体对齐研究表明，语义信息能够补充图结构。RREA 分别报告文本增强和纯结构设置[6]。MCLEA 通过对比目标对齐多种数据类型[13]。其他研究指出，简单融合可能混合尚未对齐的表示空间[14]。"
        "这些结果共同支持对邻居查询信号进行受控比较。本文固定邻域和可训练融合参数量，仅改变查询来源。"
    ),
    27: (
        "近期研究使用大语言模型[15]、域适应[16]或逐步推理代理[17]，并由此改变监督来源或推理过程。本文研究一个互补问题：在固定的实体对齐框架内，局部结构邻居如何与实体语义交互？"
    ),
    37: (
        "模型由四个部分构成。结构编码器保留关系类型并组合多个传播深度。语义编码器从同一实体文本构建 token、phrase 和 global 三种视图。融合模块使用语义信息为全部一跳出邻居评分，"
        "形成稀疏结构上下文，并将该上下文与实体语义结合。双向 InfoNCE 训练联合表示和结构表示，CSLS 随后校正训练后联合表示之间的相似度。图 1 展示完整的信息流。"
    ),
    41: (
        "局部度统计和邻居统计能够描述节点的结构角色[18–19]，结构身份方法也根据图中角色比较节点。这些研究启发本文构造一个八维拓扑向量来初始化结构分支。下文给出该特征集合的具体定义。"
    ),
    48: (
        "结构编码器与上下文模块为边方向分配互补作用。图传播通过传入边更新当前节点；融合阶段将当前实体发出的事实所指向的出邻居作为候选上下文。该划分使图传播对应传入证据，使上下文选择对应实体发出的事实。"
    ),
    70: (
        "有效性掩码将注意力限制在实际 token 上。均值池化保留广泛的词汇证据，注意力池化突出信息量更高的词项。两者的平均在覆盖范围与选择性之间取得平衡。"
    ),
    93: (
        "仍以“姚明”为例，若一跳出邻居包括“休斯敦火箭队”“上海”和一个低相关候选，语义查询可依据“篮球运动员”“效力球队”等线索提高火箭队邻居的得分。"
        "随后，1.5-entmax 将低相关候选的权重压为 0，再由式（19）决定保留姚明自身结构还是采用加权邻居摘要。"
    ),
    105: (
        "语义查询在模型构造结构上下文之前改变邻居权重，最终门控再将该上下文与实体语义结合。第 3.5.4 节通过匹配对照评估这一设计。"
    ),
    107: (
        "结构查询对照用于分离查询信号。该对照保留相同的邻域、动态补齐、1.5-entmax、门控和联合检索，并以实体结构表示替代语义查询。晚期融合对照检验范围更广的设计变化。"
        "该对照匹配可训练融合参数量，在两个分支独立编码后形成联合表示，并将邻居上下文排除在融合过程之外。"
    ),
    142: (
        "表 3 用于确定本文模型在已发表结果中的相对位置。数据划分、文本输入和后处理差异使这些数值具有背景参考性质。表 4–6 提供相同实验协议下支撑本文主张的证据。"
    ),
    147: (
        "关系类型提供了最清楚且方向一致的结构增益。移除关系类型后，DBP15K 和 OpenEA 的 Hits@1 分别下降 3.62 和 2.78 个百分点。拓扑特征相对仅使用实体向量分别提高 3.46 和 8.03 个百分点。"
        "phrase 视图使两个数据集分别提高 0.81 和 3.95 个百分点。其余对照呈现较小或依赖数据集的效应，并由此界定模型中的辅助设计选择。"
    ),
    149: (
        "表 5 比较完整模型、两个单分支模型和无可训练参数的平均融合。每项比较在训练与检索时使用相同表示，并在验证阶段保持分支权重固定。"
    ),
    151: (
        "完整模型相对仅语义检索在 DBP15K 和 OpenEA 上的 Hits@1 分别提高 9.09 和 9.66 个百分点，也优于无可训练参数的平均融合。"
        "这些增益表明，学习型交互能够将结构上下文转化为有效的对齐证据。结构单分支的较低结果进一步说明，该增益来自结构与语义的联合学习。"
    ),
    153: (
        "表 6 通过匹配对照检验查询信号和上下文交互。结构查询对照共享完整邻域、门控和融合参数，但使用结构表示为邻居评分。参数匹配晚期融合对照在两个分支独立编码后使用相同数量的融合参数。"
    ),
    155: (
        "语义查询相对结构查询在 DBP15K 和 OpenEA 上的 Hits@1 分别提高 2.08 和 2.48 个百分点。该受控差距衡量语义引导邻居评分的贡献。完整模型相对参数匹配晚期融合分别提高 5.28 和 9.22 个百分点。"
        "第二组差距衡量邻居访问与学习型交互路径的联合贡献。"
    ),
    158: (
        "实验确立了三项发现。第一，关系类型在两个受控数据集上均产生方向一致的结构增益。第二，联合学习使结构上下文能够与语义共同形成有效证据。第三，在邻域和可训练融合参数量固定时，语义查询优于结构查询。"
        "这些发现表明，关系感知邻居上下文和语义引导选择是模型增益的主要来源。"
    ),
    161: (
        "现有证据覆盖五个数据集，组件对照集中在 DBP15K ZH–EN 和 OpenEA EN–FR-15K-V2。三个随机种子用于量化效应大小和样本波动。固定邻居对照同时改变邻居数量和权重函数，"
        "参数匹配晚期融合对照同时改变邻居访问和训练交互路径。这些设计特征界定了各项比较的因果解释范围。"
    ),
    163: (
        "三类扩展实验可以检验更广泛的条件。多语言子词编码器可以替代未知 token 的随机向量[27–28]。更大规模的图谱可以检验完整邻域选择的效率。开放世界数据集可以检验不存在对应实体的情况。"
        "这些实验将衡量语义引导邻居选择在文本线索更弱和候选空间更大时的稳健性。"
    ),
    165: (
        "本文说明了结构邻居如何进入结构—语义实体对齐。关系类型在两个受控数据集上增强了结构证据。学习型联合表示使结构上下文产生超过仅语义检索的增益。"
        "在匹配条件下，语义查询也优于结构查询。完整邻域访问与学习型交互路径的组合优于参数匹配晚期融合。结果表明，关系感知邻居上下文和语义引导选择是跨语言实体对齐中的有效设计。"
    ),
}


EN_NODE_REWRITES = {
    42: [
        (
            "The local topology vector has eight values. The model computes this vector separately in each graph. The computation does not require matching relation identifiers across graphs.",
            "The model computes an eight-value topology vector separately in each graph and independently of cross-graph relation identifiers.",
        )
    ],
    66: [
        (
            "They do not read extra descriptions, outside sentences, or other text sources.",
            "All three views therefore draw exclusively from the name, relation names, and attribute text in this sequence.",
        )
    ],
    88: [
        (
            "The pattern is not the final similarity score. It gives the model information for judging whether a neighbor fits the current entity.",
            "This pattern supplies the learned gate with evidence for judging whether a neighbor fits the current entity. The dot-product score and gate then determine the final neighbor weight.",
        )
    ],
    104: [
        (
            "This fixed residual limits large changes from the gate. It does not add trainable parameters.",
            "This fixed residual limits large changes from the gate while keeping the trainable parameter count unchanged.",
        )
    ],
    111: [
        (
            "The model uses single-stage bidirectional InfoNCE training [22]. It does not use a separate warm-up stage.",
            "The model optimizes bidirectional InfoNCE objectives in a single training stage [22].",
        )
    ],
    121: [
        (
            " directly. It does not use validation data to change the structural or semantic weight. CSLS",
            " directly. CSLS",
        )
    ],
    131: [
        (
            "The main retrieval step does not tune a weight between the two branches. Instead, validation MRR selects one ",
            "The main retrieval step keeps the learned joint representation unchanged. Validation MRR selects one ",
        )
    ],
}


ZH_NODE_REWRITES = {
    42: [
        (
            "具体而言，本文构造的局部拓扑特征包含八个维度，分别在每张图内部计算，且不依赖跨图关系编号的一致性。",
            "本文在每张图内部独立计算八维局部拓扑特征，并使其独立于跨图关系编号。",
        )
    ],
    66: [
        (
            "phrase 与 global 视图均使用这条共享序列，不额外读取实体描述、外部句子或其他语料。",
            "三个视图均取自该序列中的名称、关系名称和属性文本。",
        )
    ],
    88: [
        (
            "它本身不是最终相似度，而是供后续可学习映射判断邻居兼容性的特征表示。",
            "该特征组合为可学习门控提供判断邻居兼容性的证据，点积分数与门值随后共同决定最终邻居权重。",
        )
    ],
    104: [
        (
            "该固定残差用于限制门控输出的偏移，不引入额外可训练参数。",
            "该固定残差限制门控输出的偏移，并保持可训练参数量不变。",
        )
    ],
    111: [
        ("模型采用单阶段的双向", "模型在单一训练阶段优化双向"),
        ("，不设置独立预热阶段。", "。"),
    ],
    121: [
        ("，不再通过验证集对结构和语义分支进行二次加权。", "。"),
    ],
    131: [
        (
            "主检索不选择分支融合权重，每个运行从 ",
            "主检索保持训练所得联合表示不变。每个运行从 ",
        )
    ],
}


def set_plain_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def replace_text_nodes(paragraph, replacements: list[tuple[str, str]]) -> None:
    found = {source: False for source, _ in replacements}
    for node in paragraph._p.xpath(".//w:t"):
        value = node.text or ""
        for source, target in replacements:
            if source in value:
                value = value.replace(source, target)
                found[source] = True
        node.text = value
    missing = [source for source, present in found.items() if not present]
    if missing:
        raise RuntimeError(f"Missing text replacements in paragraph: {missing}")


def rewrite_document(
    source: Path,
    output: Path,
    paragraph_rewrites: dict[int, str],
    node_rewrites: dict[int, list[tuple[str, str]]],
    chinese: bool,
) -> None:
    document = Document(source)
    if len(document.paragraphs) != 204:
        raise RuntimeError(f"Unexpected paragraph count in {source}: {len(document.paragraphs)}")

    for index, text in paragraph_rewrites.items():
        if document.paragraphs[index]._p.xpath(".//m:oMath|.//m:oMathPara"):
            raise RuntimeError(f"Paragraph {index} contains native math and cannot be replaced wholesale")
        set_plain_text(document.paragraphs[index], text)

    for index, replacements in node_rewrites.items():
        replace_text_nodes(document.paragraphs[index], replacements)

    document.core_properties.subject = (
        "Cross-lingual knowledge graph entity alignment; anti-defensive academic revision"
        if not chinese
        else "跨语言知识图谱实体对齐；去防御性学术写作修订"
    )

    if chinese:
        protect_latin_terms(document)
        # WORD JOINER prevents LibreOffice/Word from splitting decimal terms such as
        # "9.66" and "1.5-entmax" at a line boundary.
        for node in document._element.xpath(".//w:t"):
            node.text = (node.text or "").replace("\ufeff", "\u2060")
        format_chinese_typography(document)
        set_document_language(document)
        format_reference_language(document)

    document.save(output)


def main() -> None:
    rewrite_document(EN_SOURCE, EN_OUTPUT, EN_REWRITES, EN_NODE_REWRITES, chinese=False)
    rewrite_document(ZH_SOURCE, ZH_OUTPUT, ZH_REWRITES, ZH_NODE_REWRITES, chinese=True)
    print(EN_OUTPUT)
    print(ZH_OUTPUT)


if __name__ == "__main__":
    main()
